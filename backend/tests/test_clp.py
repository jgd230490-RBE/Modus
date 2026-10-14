"""
CLP monitoring, 14 Oct 2026 (bundle G: W8 data fields, W7 charts) — backend/clp.py, GET /api/clp/monitoring,
and the W8 write paths (the site CLP cap, the peak bands and the Euro class on Config).

What is asserted: clp.py reads and never writes, names no team slot id, imports neither main nor network,
and every SQL literal carries the tenant predicate; each chart takes its figure from the module that owns
it (chart 1 = bulletin.route_cap_rows = the Conditions panel's exceeded list; chart 1b =
gate_arrivals.location_visits; chart 4 = costlines' trips per month and vehicle; chart 6 = areas.area_load;
chart 10's target = carbon_target.status's target_t; chart 16 = the scenario's delta_trips); an empty chart
names the missing field and carries no zero; team NAMES only; a team code's scope; the endpoint's edges;
the word rule over every string in the reply.

WHAT THIS DOES NOT PROVE
------------------------
  * The HTTP layer is stubbed (http_smoke.py drives the real one).
  * Nothing in a browser (parse_frontend / render_frontend / dom_frontend read the fixture written last).
  * SQLite only; hand-made polylines, not HERE geometry.

Run (PowerShell):  $env:PYTHONUTF8 = "1"; .venv-tests\\Scripts\\python backend/tests/test_clp.py
"""
import copy
import json
import math
import os
import re
import sys
import tempfile
import types

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, as test_gate_arrivals.py ----------------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)

TMP = tempfile.mkdtemp(prefix="rbe_clp_")
from portable import remove_scratch_db, offline   # noqa: E402  tests/portable.py
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import weeks  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import costing  # noqa: E402
import config  # noqa: E402
import costlines  # noqa: E402
import gates  # noqa: E402
import gate_arrivals  # noqa: E402
import gate_hours  # noqa: E402
import network  # noqa: E402
import taxonomy  # noqa: E402
import areas  # noqa: E402
import bulletin  # noqa: E402
import carbon_target  # noqa: E402
import clp  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


def hav(a, b):
    lon1, lat1, lon2, lat2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 2 * 6371.0088 * math.asin(math.sqrt(h))


def plen(pts):
    return sum(hav(pts[i - 1], pts[i]) for i in range(1, len(pts)))


def box(x1, y1, x2, y2):
    return [[x1, y1], [x2, y1], [x2, y2], [x1, y2], [x1, y1]]


def poly(*rings):
    return {"type": "Polygon", "coordinates": [list(r) for r in rings]}


def strings(o):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for v in o.values():
            yield from strings(v)
    elif isinstance(o, list):
        for v in o:
            yield from strings(v)


ART = network.DEFAULT_PROFILE          # "Artic Tipper (44t)", 29 t a load
RIGID = "Rigid 8-wheeler (32t)"        # 20 t a load
BAL = "Large aggregate / ballast"
ATTR = ["Source: Office for National Statistics licensed under the Open Government Licence v.3.0",
        "Contains OS data © Crown copyright and database right 2026"]

# =========================================================================== #
#  0. A small scheme by hand: two pits, two yards, three routes, two vehicles    #
# =========================================================================== #
remove_scratch_db(db, TMP)
db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db(); db.init_gates_db()
db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db(); db.init_costing_db()
db.init_tenant()
taxonomy.seed_taxonomy()
config.invalidate()
config.seed_from_file(conversions)
config.invalidate()
costing.invalidate()
carbon_target.invalidate()
access.set_current("planner123")
q1 = network.create_location("Pit West", "origin", lat=52.55, lon=-0.95, loc_type="Quarry", supplies=[BAL])["id"]
d1 = network.create_location("Yard East", "destination", lat=52.55, lon=-0.85, loc_type="Compound", receives=[BAL])["id"]
q2 = network.create_location("Pit North", "origin", lat=52.58, lon=-0.98, loc_type="Quarry", supplies=[BAL])["id"]
d2 = network.create_location("Yard North", "destination", lat=52.58, lon=-0.92, loc_type="Compound", receives=[BAL])["id"]
R1 = network.create_route(q1, d1, material_category=BAL, route_id="R001", ipt="IPT1")["id"]
R2 = network.create_route(q2, d2, material_category=BAL, route_id="R002", ipt="IPT2")["id"]
R3 = network.create_route(q1, d2, material_category=BAL, route_id="R003", ipt="IPT1")["id"]   # never routed
L1 = [[round(-0.95 + 0.01 * i, 6), 52.55] for i in range(11)]          # W -> E across -0.90
L2 = [[round(-0.98 + 0.01 * i, 6), 52.58] for i in range(7)]           # wholly in W
L2_RET = [[-0.92, 52.58], [-0.92, 52.45], [-0.98, 52.45], [-0.98, 52.58]]   # back via S


def add_geom(rid, pts, leg="loaded", profile=ART, hr=0.5):
    db.execute(
        "INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
        "duration_hr, computed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (db.current_tenant(), rid, profile, leg, 0, json.dumps(pts), round(plen(pts), 3), hr, "2026-10-14T00:00:00Z"))


for _p in (ART, RIGID):
    add_geom(R1, L1, profile=_p); add_geom(R1, list(reversed(L1)), leg="return", profile=_p)
    add_geom(R2, L2, profile=_p); add_geom(R2, L2_RET, leg="return", profile=_p)


def _line(rid, sect, month, veh, ipt):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline="earthworks", section_id=sect, material_type=BAL, material_description=None,
        vehicle_type=veh, submitted_by="tester", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=q) for m, q in month], ipt=ipt))
    main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline="earthworks", section_id=sect)


_line(R1, "WS1", [(10, 2900.0), (11, 1450.0)], ART, "IPT1")     # 100 loads in Oct, 50 in Nov
_line(R2, "WS2", [(10, 2000.0)], RIGID, "IPT2")                 # 100 loads at 20 t
_line(R3, "WS3", [(10, 580.0)], ART, "IPT1")                    # 20 loads, not routed
config.invalidate()
factors = conversions.load_factors()
WD = 22.0
acc_planner = access.current()
M = clp.monitoring(10, acc_planner, factors=factors, working_days=WD)
C = M["charts"]

# =========================================================================== #
#  1. Source level                                                              #
# =========================================================================== #
src = open(os.path.join(BACKEND, "clp.py"), encoding="utf-8").read()
code_only = re.sub(r'"""[\s\S]*?"""', "", src)
ok("⭐ clp.py is read-only: no db.execute, no INSERT / UPDATE / DELETE",
   "db.execute" not in code_only and not re.search(r"\b(INSERT|UPDATE|DELETE)\b", code_only))
ok("⭐ clp.py imports neither main nor network and names no team slot id",
   not re.search(r"^\s*(import|from)\s+(network|main)\b", src, re.M) and re.search(r"\bIPT ?[1-6]\b", src) is None)
_stmts = re.findall(r'db\.query\(\s*((?:"[^"]*"\s*)+)', code_only)
ok("🔴 tenant isolation: both SQL literals in clp.py constrain tenant_id",
   len(_stmts) == 2 and all("tenant_id = ?" in s for s in _stmts), str(_stmts))
ok("the reply's keys and the chart order are the contract's",
   set(M) == {"month", "label", "working_days", "distance_unit", "months", "order", "charts", "with_data", "peak_bands",
              "peak_words", "scope", "scope_line", "limit", "demo_notice"}
   and M["order"] == ["route_caps", "site_caps", "vehicle_types", "euro_class", "councils", "carbon_rate", "trips_avoided"]
   and list(C) == M["order"] and M["label"] == "Oct 2026" and M["working_days"] == 22.0)
ok("every chart carries key, number, title, the CLP item it serves, state and missing",
   all({"key", "number", "title", "serves", "state", "missing"} <= set(c) and c["key"] == k for k, c in C.items())
   and [C[k]["number"] for k in M["order"]] == ["1", "1", "4", "5", "6", "10", "16"])

# =========================================================================== #
#  2. Empty states before anything is typed: each names its field, never a 0     #
# =========================================================================== #
ok("⭐ chart 1 with no cap typed: empty, names the planning block's field, still lists the routes (no cap, no figure 0)",
   C["route_caps"]["state"] == "empty"
   and C["route_caps"]["missing"] == ("No planning cap typed on any route with volume this month (Data → Routes → a route's "
                                      "planning block: Max vehicles / day).")
   and all(r["cap"] is None and r["pct"] is None for r in C["route_caps"]["rows"]), str(C["route_caps"])[:300])
ok("⭐ chart 1b with no site cap: empty, names Data → Locations → CLP cap",
   C["site_caps"]["state"] == "empty" and "Data → Locations → a location: CLP cap, lorry visits a day" in C["site_caps"]["missing"])
ok("⭐ chart 5 with no Euro class typed: empty, names Config → Vehicles → Euro class, and the vehicles in use",
   C["euro_class"]["state"] == "empty"
   and C["euro_class"]["missing"] == "No Euro class typed for the vehicles the plan uses (Config → Vehicles → Euro class)."
   and sorted(C["euro_class"]["vehicles_untyped"]) == sorted([clp._veh_label(ART, factors), clp._veh_label(RIGID, factors)]))
ok("⭐ chart 6 with no council boundaries: empty, says an admin imports them",
   C["councils"]["state"] == "empty" and C["councils"]["missing"] == "No council boundaries are loaded for this scheme (an admin imports them).")
ok("⭐ chart 16 with no target: empty, 'No decisions agreed yet' and where they are set",
   C["trips_avoided"]["state"] == "empty"
   and C["trips_avoided"]["missing"] == "No decisions agreed yet: a planner sets them with the carbon target in the Carbon planner.")
ok("⭐ chart 10 with no target: still drawn (the plan's own figures), the target line None and the 'no target' sentence, never 0",
   C["carbon_rate"]["state"] == "set" and C["carbon_rate"]["target_kg_per_tonne"] is None and "target_kg_per_load" not in C["carbon_rate"]
   and C["carbon_rate"]["target_note"] == "No carbon target set yet. A planner sets it in the Carbon planner.")
ok("with_data lists exactly the charts in state 'set'", M["with_data"] == ["vehicle_types", "carbon_rate"], str(M["with_data"]))

# =========================================================================== #
#  3. Chart 4: loads by vehicle, the costlines trips rule                          #
# =========================================================================== #
VT = C["vehicle_types"]
_cl = costlines.lines(acc_planner, status="Approved", factors=factors)["lines"]
_by = {}
for _l in _cl:
    _by.setdefault((_l["vehicle_type"], _l["month_index"]), 0)
    _by[(_l["vehicle_type"], _l["month_index"])] += _l["trips"]
ok("the over-time window is every month with approved lines: Oct and Nov 2026",
   [m["label"] for m in M["months"]] == ["Oct 2026", "Nov 2026"] and [m["month_index"] for m in VT["months"]] == [10, 11])
ok("⭐ chart 4 equals costlines' trips per vehicle and month (the one trips rule): artic 120 + 50, rigid 100 + 0",
   {(s["group"], m): s["loads"][i] for s in VT["series"] for i, m in enumerate((10, 11))}
   == {(ART, 10): _by.get((ART, 10), 0), (ART, 11): _by.get((ART, 11), 0), (RIGID, 10): _by.get((RIGID, 10), 0), (RIGID, 11): 0}
   and _by[(ART, 10)] == 120 and _by[(RIGID, 10)] == 100 and VT["totals"] == [220, 50], str(VT["series"])[:300])
ok("⭐ shares are the server's: Oct artic round(120 × 100 ÷ 220, 1) = 54.5 %, rigid 45.5 %; Nov artic 100 %",
   VT["series"][0]["group"] == ART and VT["series"][0]["pct"] == [54.5, 100.0] and VT["series"][1]["pct"] == [45.5, 0.0])
ok("chart 4 headline names the biggest vehicle and its share of every load in the window (170 of 270 = 63 %)",
   VT["headline"] == f"{clp._veh_label(ART, factors)} carries 63 % of the loads", VT["headline"])

# =========================================================================== #
#  4. Chart 10: all-in carbon per load and per tonne                               #
# =========================================================================== #
CR = C["carbon_rate"]
import carbon  # noqa: E402
_sc = carbon.scenario(acc_planner, status="Approved", factors=factors)
_m10 = [l for l in _sc["lines"] if l["month_index"] == 10]
_t10 = sum(l["total_t"] for l in _m10)
ok("⭐ chart 10 per month = the scenario's all-in total ÷ its loads and ÷ its tonnes (Oct: 220 loads, 5 480 t)",
   CR["months"][0]["loads"] == 220 and CR["months"][0]["tonnes"] == 5480.0
   and CR["months"][0]["kg_per_load"] == round(_t10 * 1000 / 220, 1)
   and CR["months"][0]["kg_per_tonne"] == round(_t10 * 1000 / 5480, 2), str(CR["months"][0]))
ok("chart 10 says All-in and names the one forecast month with no road route (R003)",
   CR["unit_note"] == "All-in" and CR["lines_unrouted"] == 1 and "1 forecast month with no road route yet carries no transport carbon here." in CR["basis"])
ok("...no decisions line without a target", all(r["kg_per_load_decisions"] is None for r in CR["months"]))

# =========================================================================== #
#  5. Typing the fields (W8): caps, Euro class, areas, a target                    #
# =========================================================================== #
# the route cap through its real write path; the site cap through the location PUT
_rp = network.set_route_planning(R1, {"max_vehicles_per_day"}, max_vehicles_per_day=4)
ok("the route cap is written through network.set_route_planning", not (_rp or {}).get("error"), str(_rp))
ok("🔴 the site cap refuses 0, a negative and a word, each in words; blank is no cap",
   network.clp_cap_problem(0) == "the CLP cap must be above 0 (clear it to mean no cap)"
   and network.clp_cap_problem(-3) == "the CLP cap must be above 0 (clear it to mean no cap)"
   and network.clp_cap_problem("lots") == "the CLP cap must be a number of lorry visits a day"
   and network.clp_cap_problem("") is None and network.clp_cap_problem(None) is None
   and network.set_clp_cap("L999", 4).get("error") == "not found")
os.environ["ADMIN_TOKEN"] = "adm"
_loc = main.update_location(d1, main.LocationIn(name="Yard East", role="destination", receives=[BAL], lat=52.55, lon=-0.85,
                                                loc_type="Compound", clp_cap_per_day=4), token="adm")
ok("⭐ the location PUT with clp_cap_per_day stores it; GeoJSON properties carry it",
   not (_loc or {}).get("error") and db.query("SELECT clp_cap_per_day FROM locations WHERE tenant_id = ? AND id = ?",
                                               (db.current_tenant(), d1))[0]["clp_cap_per_day"] == 4.0
   and {f["properties"]["id"]: f["properties"].get("clp_cap_per_day") for f in network.locations_geojson()["features"]}.get(d1) == 4.0,
   str(_loc))
main.update_location(d1, main.LocationIn(name="Yard East", role="destination", receives=[BAL], lat=52.55, lon=-0.85,
                                         loc_type="Compound"), token="adm")
ok("🔴 a PUT that does not send clp_cap_per_day leaves it alone (an older client cannot blank it)",
   db.query("SELECT clp_cap_per_day FROM locations WHERE tenant_id = ? AND id = ?", (db.current_tenant(), d1))[0]["clp_cap_per_day"] == 4.0)
_e400 = _raises(lambda: main.update_location(d1, main.LocationIn(name="Yard East", role="destination", receives=[BAL], lat=52.55,
                                                                 lon=-0.85, clp_cap_per_day=0), token="adm"))
ok("🔴 a PUT with clp_cap_per_day 0 is a 400 with the sentence, and nothing changes",
   getattr(_e400, "status_code", None) == 400 and "above 0" in getattr(_e400, "detail", "")
   and db.query("SELECT clp_cap_per_day FROM locations WHERE tenant_id = ? AND id = ?", (db.current_tenant(), d1))[0]["clp_cap_per_day"] == 4.0)
os.environ.pop("ADMIN_TOKEN", None)
# Config: the Euro class and the peak bands go through config.validate
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_bad = copy.deepcopy(_doc)
_bad["vehicles"][ART]["euro_class"] = "Euro 6"
_bad["planning"]["peak_bands"] = [{"from": "07:00", "to": "10:00"}, {"from": "09:30", "to": "11:00"}]
_pr = config.validate(_bad)
ok("🔴 Config refuses an unknown Euro class and overlapping peak bands, each in words",
   f"vehicle '{ART}': euro_class must be one of Euro IV, Euro V, Euro VI, Zero emission" in _pr
   and "planning.peak_bands: peak band 2 overlaps another band" in _pr, str(_pr))
_bad2 = copy.deepcopy(_doc)
_bad2["planning"]["peak_bands"] = "07:00-10:00"
ok("...and peak bands that are not a list", "planning.peak_bands must be a list of {from, to} windows" in config.validate(_bad2),
   str(config.validate(_bad2)))
_doc["vehicles"][ART]["euro_class"] = "Euro VI"
_doc["planning"]["peak_bands"] = [{"from": "07:00", "to": "10:00"}, {"from": "16:00", "to": "19:00"}]
ok("a valid Euro class and two peak bands pass", config.validate(_doc) == [], str(config.validate(_doc)))
config.save(_doc, by="tester")
config.invalidate()
factors = conversions.load_factors()
# the council boundaries (the test_areas boxes)
FC = {"type": "FeatureCollection", "attribution": ATTR, "features": [
    {"type": "Feature", "properties": {"area_id": "W", "name": "Westshire", "tier": "district"}, "geometry": poly(box(-1.0, 52.5, -0.9, 52.6))},
    {"type": "Feature", "properties": {"area_id": "E", "name": "Eastshire", "tier": "district"}, "geometry": poly(box(-0.9, 52.5, -0.8, 52.6))},
    {"type": "Feature", "properties": {"area_id": "S", "name": "Southshire", "tier": "district"}, "geometry": poly(box(-1.0, 52.4, -0.8, 52.5))},
    {"type": "Feature", "properties": {"area_id": "C", "name": "Testshire", "tier": "county"}, "geometry": poly(box(-1.0, 52.4, -0.8, 52.6))},
]}
_imp = areas.import_collection(FC)
ok("the four council areas import and the routed legs are clipped", _imp.get("imported") == 4, str(_imp)[:200])
# the target: the rigid on R002 swapped for the artic (2 000 t at 29 t = 69 loads, not 100)
DECISIONS = [{"driver": "vehicle", "scope": {"route_id": R2}, "vehicle": ART}]
_t = main.put_carbon_target(main.CarbonTargetIn(mode="pct", levers=DECISIONS, ui_state={}, updated_by="Sam (planner)"))
ok("the carbon target is set from one decision (a lorry swap on R002)", (_t.get("target") or {}).get("levers") == DECISIONS, str(_t)[:300])
M2 = clp.monitoring(10, acc_planner, factors=factors, working_days=WD)
C2 = M2["charts"]

# chart 1
RC = C2["route_caps"]
_cond = bulletin.conditions_of_month(10, factors=factors)
ok("⭐ chart 1: R001 carries round(100 ÷ 22, 3) = 4.545 movements a day against a cap of 4: over, 114 %, over by 0.545",
   RC["state"] == "set" and RC["rows"][0]["route_id"] == R1 and RC["rows"][0]["movements_per_day"] == round(100 / 22, 3)
   and RC["rows"][0]["cap"] == 4 and RC["rows"][0]["over"] and RC["rows"][0]["pct"] == round(100 / 22 * 100 / 4)
   and RC["rows"][0]["over_by"] == round(100 / 22 - 4, 3) and RC["rows"][0]["route"] == "Pit West → Yard East", str(RC)[:300])
ok("🔴 chart 1 is the Conditions panel's rule: its over rows are exactly conditions_of_month's exceeded list",
   [(r["route_id"], r["cap"], r["movements_per_day"], r["over_by"]) for r in RC["rows"] if r["over"]]
   == [(x["route_id"], x["cap"], x["two_way_a_day"], x["over_by"]) for x in _cond["route_caps"]["exceeded"]]
   and _cond["route_caps"]["routes_with_cap"] == RC["routes_with_cap"] == 1, str(_cond["route_caps"]))
ok("chart 1 lists the routes with no cap after the capped ones, never compared; the headline counts the over route",
   [r["route_id"] for r in RC["rows"]][1:] and all(r["cap"] is None and not r["over"] for r in RC["rows"][1:])
   and RC["routes_without_cap"] == 2 and RC["headline"] == "1 route over its cap" and clp.CAP_DIFF in RC["basis"],
   str(RC.get("headline")))
# chart 1b
SC = C2["site_caps"]
_lv = gate_arrivals.location_visits(10, factors=factors, working_days=WD)
ok("⭐ chart 1b: Yard East has 100 ÷ 22 = 4.5 visits a day against a CLP cap of 4: over, 114 %; the figures are location_visits'",
   SC["state"] == "set" and SC["rows"][0]["site"] == "Yard East" and SC["rows"][0]["visits_per_day"] == _lv[d1]["visits_per_day"] == 4.5
   and SC["rows"][0]["cap"] == 4.0 and SC["rows"][0]["over"] and SC["rows"][0]["pct"] == 114
   and SC["headline"] == "1 site over its CLP cap", str(SC)[:300])
ok("chart 1b: Pit West counts the lorries coming to load (R001 100 + R003 20 = 120 ÷ 22 = 5.5), no cap, not compared",
   {r["site"]: r for r in SC["rows"]}["Pit West"]["visits_per_day"] == round(120 / 22, 1)
   and {r["site"]: r for r in SC["rows"]}["Pit West"]["cap"] is None and not {r["site"]: r for r in SC["rows"]}["Pit West"]["over"])
# chart 5
EC = C2["euro_class"]
ok("⭐ chart 5 with the artic typed Euro VI: set; the rigid's loads are 'Not typed' (never guessed), drawn last",
   EC["state"] == "set" and [s["group"] for s in EC["series"]] == ["Euro VI", "Not typed"]
   and EC["series"][0]["loads"] == [120, 50] and EC["series"][1]["loads"] == [100, 0]
   and EC["vehicles_untyped"] == [clp._veh_label(RIGID, factors)] and EC["headline"] == "Euro VI: 63 % of the loads", str(EC)[:300])
# chart 6
CO = C2["councils"]
_al = areas.area_load(10, working_days=WD, geometry=False)
_ald = {a["name"]: a for a in _al["areas"] if a["tier"] == "district"}
ok("⭐ chart 6 is area_load's districts (the county row left out), sorted by distance; movements and laden km are its own",
   CO["state"] == "set" and CO["tier"] == "district" and {r["council"] for r in CO["rows"]} <= set(_ald)
   and all(r["movements_per_day"] == round(_ald[r["council"]]["movements_per_day"], 1)
           and r["laden_distance"] == round(_ald[r["council"]]["laden_lorry_km"], 1) for r in CO["rows"])
   and [r["distance"] for r in CO["rows"]] == sorted([r["distance"] for r in CO["rows"]], reverse=True), str(CO)[:400])
ok("chart 6: laden + empty = the area's lorry-km (to the rounding), the double-count sentence and the licence lines ride along",
   all(abs((r["laden_distance"] or 0) + (r["empty_distance"] or 0) - round(_ald[r["council"]]["lorry_km"], 1)) <= 0.11 for r in CO["rows"])
   and CO["double_count"] == areas.DOUBLE_COUNT and CO["attribution"] == ATTR and CO["distance_word"] == "lorry-km")
ok("chart 6: Southshire carries R002's empty return only — distance, no movements",
   {r["council"]: r for r in CO["rows"]}.get("Southshire", {}).get("movements_per_day") == 0.0
   and ({r["council"]: r for r in CO["rows"]}.get("Southshire", {}).get("empty_distance") or 0) > 0, str(CO["rows"]))
# chart 10 with the target
CR2 = C2["carbon_rate"]
_st = carbon_target.status(acc_planner, carbon_target.get(), factors=factors)
_base = carbon.scenario(acc_planner, status="Approved", factors=factors)["baseline"]["totals"]
ok("⭐ chart 10's target is carbon_target.status's target_t shared out over the plan's tonnes",
   CR2["target_t"] == _st["target_t"]
   and CR2["target_kg_per_tonne"] == round(_st["target_t"] * 1000 / _base["tonnes"], 2) and CR2["target_note"] is None, str(CR2)[:300])
ok("🔴 per load carries NO target line, and says why: the lorry swap raises Oct's kg a load (fewer, heavier loads) while the total falls",
   "target_kg_per_load" not in CR2 and CR2["per_load_note"] == clp.PER_LOAD_NOTE and clp.PER_LOAD_NOTE in CR2["basis"]
   and CR2["months"][0]["kg_per_load_decisions"] > CR2["months"][0]["kg_per_load"]
   and CR2["months"][0]["kg_per_tonne_decisions"] < CR2["months"][0]["kg_per_tonne"], str(CR2["months"][0]))
ok("chart 10 with decisions: Oct per load with the decisions is the scenario's (fewer loads, so more kg a load); the headline names both figures",
   CR2["months"][0]["kg_per_load_decisions"] is not None and CR2["months"][0]["kg_per_tonne_decisions"] is not None
   and CR2["headline"] == f"{CR2['plan_kg_per_tonne']:g} kg CO₂e a tonne planned; the target allows {CR2['target_kg_per_tonne']:g}")
# chart 16
TA = C2["trips_avoided"]
ok("⭐ chart 16: the swap takes R002 from 100 loads to ceil(2000 ÷ 29) = 69 in Oct — 31 lorry trips avoided, said in the headline",
   TA["state"] == "set" and TA["avoided"] == 100 - math.ceil(2000 / 29) == 31
   and TA["months"][0] == {"month_index": 10, "label": "Oct 2026", "loads_planned": 220, "loads_with_decisions": 189, "avoided": 31}
   and TA["months"][1]["avoided"] == 0 and TA["headline"] == "31 lorry trips avoided", str(TA)[:300])
ok("chart 16 names the decision in words and the lorry-km it saves (a figure ≥ 0)",
   TA["decisions"] == (carbon_target.get() or {}).get("lever_words") and len(TA["decisions"]) == 1
   and TA["distance_saved"] is not None and TA["distance_saved"] >= 0 and TA["distance_word"] == "lorry-km")
ok("⭐ with everything typed every chart has data", M2["with_data"] == M2["order"], str(M2["with_data"]))
ok("the peak bands ride on the reply in words",
   M2["peak_bands"] == [{"from": "07:00", "to": "10:00"}, {"from": "16:00", "to": "19:00"}]
   and M2["peak_words"].startswith("No deliveries 07:00–10:00 and 16:00–19:00"))

# =========================================================================== #
#  6. Miles, a team code, the endpoint, the words                                  #
# =========================================================================== #
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc.setdefault("tenant", {})["distance_unit"] = "mi"
config.save(_doc, by="tester"); config.invalidate()
Mmi = clp.monitoring(10, acc_planner, factors=conversions.load_factors(), working_days=WD)
_w = {r["council"]: r for r in Mmi["charts"]["councils"]["rows"]}
ok("⭐ a miles tenant reads lorry-miles: each distance is the km ÷ 1.609344",
   Mmi["distance_unit"] == "mi" and Mmi["charts"]["councils"]["distance_word"] == "lorry-miles"
   and all(abs(_w[k]["distance"] - round(_ald[k]["lorry_km"] / config.KM_PER_MILE, 1)) <= 0.051 for k in _w)
   and Mmi["charts"]["trips_avoided"]["distance_word"] == "lorry-miles", str(_w)[:200])
_doc["tenant"].pop("distance_unit", None)
config.save(_doc, by="tester"); config.invalidate()
_txt = json.dumps(M2, ensure_ascii=False)
ok("🔴 team NAMES only: no team slot id anywhere in the reply", re.search(r"\bIPT ?[1-6]\b", _txt) is None)
_banned = re.compile(r"\b(bake[ds]?|baking|line-month|cached|geometry|server|endpoint|truck|trucks|real-time|predict\w*|optimi[sz]\w*|telematics)\b|\btoken\b", re.I)
_hits = [s for s in strings(M2) if _banned.search(s)]
ok("🔴 the word rule over every string in the reply (no bake / line-month / cached / geometry / server / endpoint / token / truck …)",
   not _hits, str(_hits[:3]))
os.environ["IPT2_CODE"] = "ipt2-secret"
try:
    access.set_current("ipt2-secret")
    acc_t2 = access.current()
    MT = clp.monitoring(10, acc_t2, factors=conversions.load_factors(), working_days=WD)
finally:
    os.environ.pop("IPT2_CODE", None)
    access.set_current("planner123")
ok("⭐ a team code: scope team; the vehicle chart holds only its lines (100 rigid loads); the caps and councils are the scheme's",
   MT["scope"] == "team" and MT["charts"]["vehicle_types"]["totals"] == [100]
   and MT["charts"]["route_caps"]["rows"] == C2["route_caps"]["rows"] and MT["charts"]["councils"]["rows"] == C2["councils"]["rows"]
   and "the whole scheme for the caps and councils" in MT["scope_line"], str(MT["charts"]["vehicle_types"])[:200])
ok("...and its target is its own share (target_scope team)", MT["charts"]["carbon_rate"]["target_scope"] == "team")
ok("the endpoint: main.clp_monitoring(10) returns the same dict as clp.monitoring(10, acc, …)",
   main.clp_monitoring(10) == clp.monitoring(10, access.current(), factors=conversions.load_factors(), working_days=main._working_days(conversions.load_factors())))
_e = [_raises(lambda: main.clp_monitoring(0)), _raises(lambda: main.clp_monitoring(main.MONTH_COUNT + 1))]
ok("month 0 and 61 are a 400 with the sentence", all(getattr(x, "status_code", None) == 400 and "month must be 1.." in x.detail for x in _e))
access.set_current("")
_e401 = _raises(lambda: main.clp_monitoring(10))
access.set_current("planner123")
ok("no code -> 401", getattr(_e401, "status_code", None) == 401, repr(_e401))
ok("an empty month (Jan 2026): every chart empty, each with a sentence, no figure",
   all(c["state"] == "empty" and c["missing"] for c in clp.monitoring(1, acc_planner, factors=conversions.load_factors(), working_days=WD)["charts"].values()
       if c["key"] in ("route_caps", "site_caps", "councils")))

# =========================================================================== #
#  6b. W10: the CLP pack — exactly the charts with data, the demo notice on every page and sheet  #
# =========================================================================== #
import io as _io  # noqa: E402
import export as _export  # noqa: E402
import clp_pack  # noqa: E402
import openpyxl as _openpyxl  # noqa: E402
from portable import pdf_text as _pdf_text  # noqa: E402
offline(_export)
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc.setdefault("tenant", {})["demo_notice"] = "Fictional demo scheme · nothing here is real"
config.save(_doc, by="tester"); config.invalidate()
MP = clp.monitoring(10, acc_planner, factors=conversions.load_factors(), working_days=WD)
ok("the pack's month carries the demo notice from the tenant block", MP["demo_notice"] == "Fictional demo scheme · nothing here is real")
_pdf = clp_pack.build_pdf(MP, generated="14 Oct 2026")
_pt = re.sub(r"\s+", " ", _pdf_text(_pdf) or "")
_pages = len(re.findall(r"page \d+", _pt))
ok("⭐ W10 PDF: titled for the month, the demo notice on EVERY page, the stamp 'planned, not observed'",
   "Wayscope · CLP monitoring pack · Oct 2026" in _pt and _pages >= 2
   and _pt.count("Fictional demo scheme · nothing here is real") >= _pages
   and "Planned figures from the approved plan, not observed: Wayscope holds no gate log or arrival times." in _pt, f"{_pages} pages")
ok("⭐ W10 PDF: every chart with data, in the section's order, each with its headline, how it is worked out and its source",
   all(MP["charts"][k]["title"] in _pt and MP["charts"][k]["headline"] in _pt for k in MP["with_data"])
   and [_pt.index(MP["charts"][k]["title"]) for k in MP["with_data"]] == sorted(_pt.index(MP["charts"][k]["title"]) for k in MP["with_data"])
   and _pt.count("How it is worked out: ") == len(MP["with_data"]) and _pt.count("Source: Wayscope") == len(MP["with_data"]), _pt[:300])
ok("W10 PDF: the figures are printed as a table too (Pit West → Yard East 4.5 against 4, 114 %; 31 trips avoided)",
   "Pit West → Yard East 4.5 4 114 % Over the cap" in _pt and "Oct 2026 220 189 31" in _pt, _pt[:200])
ok("🔴 W10 PDF: no team id, none of the banned words",
   re.search(r"\bIPT ?[1-6]\b", _pt) is None
   and not re.search(r"\b(bake[ds]?|line-month|cached|geometry|server|endpoint|token|trucks?|real-time|predict\w*|optimi[sz]\w*)\b", _pt, re.I))
# a month with charts missing: they are named with what they need, never drawn as empty
_ME = copy.deepcopy(MP)
for _k in ("euro_class", "trips_avoided"):
    _ME["charts"][_k] = {**{x: _ME["charts"][_k][x] for x in ("key", "number", "title", "serves")}, "state": "empty",
                         "missing": clp.monitoring(10, acc_planner, factors=factors, working_days=WD)["charts"][_k].get("missing") or "no data"}
_ME["with_data"] = [k for k in _ME["order"] if _ME["charts"][k]["state"] == "set"]
_pe = re.sub(r"\s+", " ", _pdf_text(clp_pack.build_pdf(_ME, generated="14 Oct 2026")) or "")
ok("⭐ W10 PDF: a chart with no data is not drawn — it is listed under 'Not drawn this month' with what it needs",
   "Not drawn this month" in _pe and f"{_ME['charts']['euro_class']['title']}: " in _pe
   and _pe.count(_ME["charts"]["euro_class"]["title"]) == 1 and _pe.count("How it is worked out: ") == len(_ME["with_data"]))
_wb = _openpyxl.load_workbook(_io.BytesIO(clp_pack.build_xlsx(MP, generated="2026-10-14T12:00:00Z")))
ok("⭐ W10 XLSX: About first, then one sheet a chart with data, named for it",
   _wb.sheetnames == ["About"] + [clp_pack.SHEET[k] for k in MP["with_data"]], str(_wb.sheetnames))
_ab = [[c.value for c in r] for r in _wb["About"].iter_rows()]
ok("W10 XLSX: the demo notice heads every sheet in the demo bar's colours; About names the month, the scope and the basis",
   all(_wb[n]["A1"].value == "Fictional demo scheme · nothing here is real" and (_wb[n]["A1"].font.color.rgb or "").endswith(_export.DEMO_FG_HEX)
       for n in _wb.sheetnames)
   and ["Month", "Oct 2026"] == [x for x in _ab if x[0] == "Month"][0][:2] and any(x[0] == "Scope" and x[1] == MP["scope_line"] for x in _ab))
_rc = [[c.value for c in r] for r in _wb[clp_pack.SHEET["route_caps"]].iter_rows()]
ok("W10 XLSX: a chart's sheet holds its table and says how it is worked out and its source",
   ["Pit West → Yard East", "4.5", "4", "114 %", "Over the cap"] in [r[:5] for r in _rc]
   and any(r[0] == "How it is worked out" and r[1] == MP["charts"]["route_caps"]["basis"] for r in _rc)
   and any(r[0] == "Source" for r in _rc))
_wbe = _openpyxl.load_workbook(_io.BytesIO(clp_pack.build_xlsx(_ME)))
ok("W10 XLSX: a chart with no data has no sheet; About lists it as not drawn with what it needs",
   clp_pack.SHEET["euro_class"] not in _wbe.sheetnames
   and any(x[0] == "Not drawn" and str(x[1]).startswith(_ME["charts"]["euro_class"]["title"] + ": ") for x in [[c.value for c in r] for r in _wbe["About"].iter_rows()]))
ok("W10: the file name is the month's: wayscope-clp-pack-2026-10.pdf / .xlsx",
   clp_pack.file_name(MP, "pdf") == "wayscope-clp-pack-2026-10.pdf" and clp_pack.file_name({"month": 13}, "xlsx") == "wayscope-clp-pack-2027-01.xlsx")
_rp = main.clp_pack_pdf(10)
_rx = main.clp_pack_xlsx(10)
ok("⭐ W10 endpoints: the PDF and the XLSX for a month, as attachments under the month's file name",
   _rp.media_type == "application/pdf" and _rp.headers["Content-Disposition"] == 'attachment; filename="wayscope-clp-pack-2026-10.pdf"'
   and (_rp.body or b"")[:4] == b"%PDF"
   and _rx.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
   and _rx.headers["Content-Disposition"] == 'attachment; filename="wayscope-clp-pack-2026-10.xlsx"' and (_rx.body or b"")[:2] == b"PK")
_e = [_raises(lambda: main.clp_pack_pdf(0)), _raises(lambda: main.clp_pack_xlsx(None))]
ok("W10 endpoints: no month or month 0 is a 400", all(getattr(x, "status_code", None) == 400 for x in _e))
access.set_current("")
_e401 = _raises(lambda: main.clp_pack_pdf(10))
access.set_current("planner123")
ok("🔴 W10 endpoints: no code -> 401", getattr(_e401, "status_code", None) == 401)
_doc["tenant"]["demo_notice"] = None
config.save(_doc, by="tester"); config.invalidate()

# =========================================================================== #
#  7. The fixture for the frontend harnesses (planner, Oct 2026, every chart set)  #
# =========================================================================== #
_fix = os.path.join(HERE, "fixtures", "clp_monitoring.json")
if not FAIL:
    with open(_fix, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(M2, indent=1, ensure_ascii=False))
ok("the fixture backend/tests/fixtures/clp_monitoring.json is written from a clean run, LF on disk",
   os.path.exists(_fix) and b"\r\n" not in open(_fix, "rb").read())

print(f"\n{PASS} passed, {len(FAIL)} failed")
for f in FAIL:
    print("  FAIL:", f)
sys.exit(1 if FAIL else 0)
