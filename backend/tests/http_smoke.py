"""
http_smoke.py — the REAL HTTP layer, end to end. Wayscope (G2, 16 Sep 2026; rebrand 29 Sep).

    pip install -r backend/requirements.txt httpx      # a venv, or CI
    python3 backend/tests/http_smoke.py

Every other harness stubs FastAPI and calls endpoint bodies directly, so none of them
proves that a route is registered, that a query parameter or a JSON body reaches the
function, that the gate middleware reads the cookie it set, or that a Response comes
back with the headers the browser needs. This one boots `main.app` with FastAPI's
TestClient (lifespan and middleware included) against a scratch SQLite database and
drives it the way the browser and an admin would:

    gate → sign-in → empty tenant → import the demo package → read it back through
    /api/meta, the map overlay, the Look-ahead and both exports → export the tenant →
    refuse bad input.

When FastAPI is not installed (the stubbed harnesses' environment) it prints SKIPPED
and exits 0 with "0 passed, 0 failed", so the suite runner stays green there.

NOT covered: PostgreSQL (SQLite only), HERE (no key — routes stay UNBAKED), Mapbox
(the PDF falls back to its schematic map when the static-map service is unreachable),
and any browser rendering.
"""
import json
import math
import os
import re
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

try:
    import fastapi  # noqa: F401
    from fastapi.testclient import TestClient
    import httpx  # noqa: F401
except Exception as e:  # the stubbed sandbox
    print(f"http_smoke.py: SKIPPED — real FastAPI/httpx not importable ({e.__class__.__name__})")
    print("0 passed, 0 failed")
    sys.exit(0)

# --- environment: a deployment with every secret set, no external keys -----------------
TMP = tempfile.mkdtemp(prefix="modus_http_")
for v in ("DATABASE_URL", "TENANT_ID", "ALLOW_DEMO_CODES", "HERE_API_KEY", "MAPBOX_TOKEN",
          "MAPBOX_PUBLIC_TOKEN", "GOOGLE_MAPS_API_KEY", "GOOGLE_STREETVIEW_KEY", "MAP_GATE",
          "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE"):
    os.environ.pop(v, None)
os.environ.update({
    "ADMIN_TOKEN": "adm-token-xyz",
    "PLANNER_CODE": "plan-code-xyz",
    "ADMIN_CODE": "admin-code-xyz",
    "IPT1_CODE": "north-code-xyz",
    "MAP_PASSWORD": "map-pass-xyz",
    "GATE_SECRET": "gate-secret-for-tests-only",
    # 22 Oct 2026 (read_cache): no background warm-up here — every build below is one this file asked for
    "WAYSCOPE_WARM_CACHE": "0",
})

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "http.db")
import main  # noqa: E402
import gate  # noqa: E402
# 17 Oct 2026 (demo redesign R2, D104): the harness pins its own date — the shipped demo's as-of date, Thu 26 Nov 2026
# (Nov W4) — so the commit week, the day sheet and every "this week" figure read the same week whatever the calendar says
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from portable import pin_today, DEMO_AS_OF  # noqa: E402
_PINNED = pin_today(DEMO_AS_OF)

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


TOKEN = {"token": "adm-token-xyz"}
PLAN = {"X-Access-Code": "plan-code-xyz"}
NORTH = {"X-Access-Code": "north-code-xyz"}
DEMO = json.load(open(os.path.join(ROOT, "demo", "uk-corridor.package.json"), encoding="utf-8"))

# https: the gate cookie is Secure (as it must be on Render), so a plain-http client would
# never send it back — which is also what a browser does.
with TestClient(main.app, base_url="https://testserver") as c:
    # ---------------------------------------------------------------- 1. boot + gate
    r = c.get("/api/health")
    ok("health answers 200 on SQLite", r.status_code == 200 and r.json().get("backend") == "sqlite", r.text[:120])
    r = c.get("/")
    ok("the app page is served, named Wayscope, with no-cache", r.status_code == 200 and "Wayscope" in r.text
       and r.headers.get("cache-control") == "no-cache", str(r.status_code))
    ok("...and its <head> carries the favicon links and the manifest",
       '<link rel="icon" href="/favicon.ico" sizes="32x32">' in r.text and '<link rel="manifest" href="/site.webmanifest">' in r.text
       and '<link rel="icon" href="/favicon.svg" type="image/svg+xml">' in r.text)
    # 29 Sep: the brand files are served without any cookie — a refused visitor sees the logo.
    for path, ctype in (("/favicon.ico", "image/x-icon"), ("/favicon.svg", "image/svg+xml"),
                        ("/site.webmanifest", "application/manifest+json"),
                        ("/apple-touch-icon.png", "image/png"), ("/icon-192.png", "image/png"),
                        ("/brand/logo-header-on-dark.svg", "image/svg+xml"),
                        ("/brand/logo-stacked-on-dark.svg", "image/svg+xml"),
                        ("/brand/theme-tokens.css", "text/css")):
        r = c.get(path)
        ok(f"{path} is served without a cookie, {ctype}, no-cache",
           r.status_code == 200 and r.headers.get("content-type", "").startswith(ctype)
           and r.headers.get("cache-control") == "no-cache" and len(r.content) > 0,
           f"{r.status_code} {r.headers.get('content-type')}")
    r = c.get("/favicon.svg")
    ok("the SVG favicon is a real SVG", r.status_code == 200 and b"<svg" in r.content[:400])
    r = c.get("/brand/../backend/factors.json")
    ok("🔴 the brand mount cannot be walked out of", r.status_code in (404, 400), str(r.status_code))
    r = c.get("/site.webmanifest")
    _man = r.json()
    ok("the manifest names Wayscope and its icons all answer 200 at the root",
       "Wayscope" in _man.get("name", "") and _man.get("icons")
       and all(c.get(i["src"]).status_code == 200 for i in _man["icons"]), str(_man.get("icons")))
    r = c.get("/map/")
    ok("🔴 the map is closed without a credential — the password page, 401",
       r.status_code == 401 and "password" in r.text.lower(), str(r.status_code))
    ok("...and the password page is branded: Wayscope title, the stacked logo, no RBE-era wording",
       "<title>Wayscope — map access</title>" in r.text and 'src="/brand/logo-stacked-on-dark.svg"' in r.text
       and "alliance" not in r.text.lower() and "Modus" not in r.text)
    # 11 Oct 2026 (the demo bar): the app's iframes ask for /map/?embed=1&v=…; the gate reads the path only
    r = c.get("/map/?embed=1&v=1")
    ok("🔴 the embedded map's address (/map/?embed=1) is closed the same way — the password page, 401",
       r.status_code == 401 and "<title>Wayscope — map access</title>" in r.text, str(r.status_code))
    r = c.get("/help/")
    ok("🔴 the guide is closed without a staff sign-in — 401", r.status_code == 401, str(r.status_code))
    # 13 Oct 2026 (W4): the technical guide sits under the same mount and the same staff gate
    r = c.get("/help/technical/")
    ok("🔴 W4: the technical guide is closed without a staff sign-in — 401, like the guide", r.status_code == 401
       and "Technical guide" not in r.text, str(r.status_code))
    # 12 Oct 2026 (C4): the monthly bulletin's page is staff-only like the guide, with its own refusal page
    r = c.get("/bulletin/?month=10")
    ok("🔴 the monthly bulletin page is closed without a staff sign-in — 401, the bulletin's own sign-in page (the middleware's)",
       r.status_code == 401 and "The monthly bulletin is for signed-in staff" in r.text and "user guide" not in r.text.lower(),
       f"{r.status_code} {r.text[:120]}")
    r = c.get("/api/bulletin/month.pdf?month=10")
    ok("🔴 ...and so is its PDF — 401", r.status_code == 401, str(r.status_code))
    # 14 Oct 2026 (backlog item 2): the route map export is staff only too
    r = c.get("/api/routes/map-export.pdf?month=10")
    ok("🔴 ROUTE MAP EXPORT: without a staff sign-in — 401", r.status_code == 401, str(r.status_code))
    ok("🔴 ...and so are its PNG and the card's list of places (14 Oct, evening) — 401",
       c.get("/api/routes/map-export.png?month=10").status_code == 401
       and c.get("/api/routes/map-export/places?month=10").status_code == 401)
    # 16 Oct 2026 (CLOCS v5 C3): the route sheet is staff only too, and writing its words is an admin's
    r = c.get("/api/routes/R016/route-sheet.pdf")
    ok("🔴 ROUTE SHEET: without a staff sign-in — 401", r.status_code == 401, str(r.status_code))
    r = c.put("/api/admin/routes/R016/sheet-text", json={"route_reason": "x"}, headers={"X-Wayscope": "1"})
    ok("🔴 ...and its words cannot be written without the admin token — 403, nothing written", r.status_code == 403, str(r.status_code))
    # 14 Oct 2026 (night): the community considerations layer is staff only, and fetching it is a planner's
    ok("🔴 COMMUNITY CONSIDERATIONS: the layer, its status and a fetch without a credential — 401 each",
       c.get("/api/community").status_code == 401 and c.get("/api/community/status").status_code == 401
       and c.post("/api/community/refresh", headers={"X-Wayscope": "1"}).status_code == 401)
    # 16 Oct 2026 (CLOCS v5 C1): the collision history and each route's evidence are staff only; fetching is a planner's
    ok("🔴 COLLISION HISTORY: the layer, its status, a fetch and a route's evidence without a credential — 401 each",
       c.get("/api/collisions").status_code == 401 and c.get("/api/collisions/status").status_code == 401
       and c.post("/api/collisions/refresh", headers={"X-Wayscope": "1"}).status_code == 401
       and c.get("/api/routes/R016/collision-evidence.pdf").status_code == 401
       and c.get("/api/routes/R016/collision-evidence.xlsx").status_code == 401)
    # 13 Oct 2026 (D43): the one-page leaflet, page and PDF, the bulletin's staff rule
    r = c.get("/bulletin/leaflet/?month=10")
    ok("🔴 the leaflet page is closed without a staff sign-in — 401, the bulletin's sign-in page",
       r.status_code == 401 and "The monthly bulletin is for signed-in staff" in r.text, f"{r.status_code} {r.text[:120]}")
    ok("🔴 ...and so is the leaflet PDF — 401", c.get("/api/bulletin/leaflet.pdf?month=10").status_code == 401)
    r = c.get("/api/public/alignment")
    ok("🔴 the overlay API is gated like the map — 401 JSON", r.status_code == 401
       and r.headers.get("content-type", "").startswith("application/json"), str(r.status_code))

    r = c.get("/api/meta")
    m = r.json()
    ok("meta is readable before sign-in (the sign-in screen needs it)", r.status_code == 200)
    # NARROWED 12 Oct 2026 (D31): before sign-in the answer is the PUBLIC shape and says so
    ok("a fresh tenant is EMPTY: no routes, teams or sections — and before sign-in the shape is `public`",
       m["routes"] == [] and m["ipts"] == [] and m["work_sections"] == [] and m.get("scope") == "public",
       str({k: len(m[k]) for k in ("routes", "ipts", "work_sections")}) + " " + str(m.get("scope")))
    ok("...and carries the default tenant words: Wayscope, Team, EUR, no country",
       m["tenant"]["name"] == "Wayscope" and m["tenant"]["team_label"] == "Team"
       and m["tenant"]["currency"] == "EUR" and m["tenant"]["country"] is None, str(m["tenant"]))
    ok("D31: the public shape carries what the sign-in screen and the public map read — units, vehicles, labels, months, factors, seasonal windows",
       m["units"] and m["vehicles"] and m["vehicle_labels"].get("langs") and m["months"]["count"] == 60
       and isinstance(m["factors"], dict) and isinstance(m["seasonal_restrictions"], list) and m["disciplines"])

    r = c.post("/api/auth", json={"code": "wrong"})
    ok("a wrong access code is refused — 401", r.status_code == 401)
    r = c.post("/api/auth", json={"code": "submitter123"})
    ok("🔴 the old demo codes do not work when real codes are set", r.status_code == 401)
    r = c.post("/api/auth", json={"code": "plan-code-xyz"})
    ok("the planner code signs in", r.status_code == 200 and r.json()["role"] == "planner", r.text[:160])
    ok("...and the middleware sets the signed gate cookie", gate.COOKIE in r.cookies or gate.COOKIE in c.cookies,
       str(dict(r.headers)))
    ok("...and a planner is offered all six slots while the tenant names none",
       r.json().get("ipts") == ["IPT1", "IPT2", "IPT3", "IPT4", "IPT5", "IPT6"], str(r.json().get("ipts")))

    r = c.get("/map/")
    ok("with the staff cookie the map opens and loads the overlay module",
       r.status_code == 200 and "overlay.js" in r.text and "Wayscope" in r.text, str(r.status_code))
    _map_doc = r.text
    r = c.get("/map/?embed=1&v=1")
    ok("...and /map/?embed=1 (the app's iframe) is the same document, which reads the flag itself (11 Oct, the demo bar)",
       r.status_code == 200 and r.text == _map_doc and "get('embed') === '1'" in r.text, str(r.status_code))
    r = c.get("/map/overlay.js")
    ok("the overlay module is served", r.status_code == 200 and "applyOverlayPackage" in r.text)
    r = c.get("/map/data/alignment.js")
    ok("🔴 no static alignment file is served — 404", r.status_code == 404, str(r.status_code))
    r = c.get("/help/")
    ok("with the staff cookie the guide opens, as Wayscope", r.status_code == 200 and "Wayscope — User guide" in r.text,
       str(r.status_code))
    r = c.get("/help/technical/")
    ok("⭐ W4: ...and with it the technical guide opens at /help/technical/ (the mount serves its index.html), no-cache like the guide",
       r.status_code == 200 and "<title>Wayscope — Technical guide</title>" in r.text and "GET /api/costing/lines" in r.text
       and "no-cache" in (r.headers.get("cache-control") or ""), f"{r.status_code} {r.headers.get('cache-control')}")

    r = c.get("/api/public/alignment")
    b = r.json()
    ok("an empty tenant's overlay is {} plus its words and no teams",
       r.status_code == 200 and set(b) == {"tenant", "teams"} and b["teams"] == [], str(b)[:160])
    etag0 = r.headers.get("etag") or ""
    ok("the overlay carries an ETag", bool(etag0), str(dict(r.headers)))
    r = c.get("/api/public/alignment", headers={"If-None-Match": etag0})
    ok("⭐ a matching If-None-Match gets 304 with no body", r.status_code == 304 and not r.content, str(r.status_code))

    r = c.get("/api/restrictions/layers")
    ok("no country → no restriction provider, no layers", r.status_code == 200
       and r.json().get("provider") is None and r.json().get("layers") == [], r.text[:160])

    # ------------------------------------------------------------ 2. the tenant package
    r = c.get("/api/admin/tenant/status")
    ok("🔴 tenant status without the admin token — 403", r.status_code == 403, str(r.status_code))
    r = c.get("/api/admin/tenant/status", params=TOKEN)
    ok("tenant status with the token says EMPTY", r.status_code == 200 and r.json()["empty"] is True, r.text[:160])
    # 12 Oct 2026 (C11 closed): the bake and every other admin call refuse without the token, with a wrong one, and —
    # the part that was open until now — with ADMIN_TOKEN UNSET even when the right token is sent
    r = c.post("/api/admin/bake-routes")
    ok("🔴 C11: POST /api/admin/bake-routes without the admin token — 403, no HERE call", r.status_code == 403, str(r.status_code))
    r = c.post("/api/admin/bake-routes", params={"token": "wrong"})
    ok("🔴 C11: ...with a wrong token — 403", r.status_code == 403, str(r.status_code))
    _saved_admin = os.environ.pop("ADMIN_TOKEN")
    try:
        r_bake = c.post("/api/admin/bake-routes", params=TOKEN)
        r_stat = c.get("/api/admin/tenant/status", params=TOKEN)
        r_exp = c.get("/api/admin/tenant/export", params=TOKEN)
    finally:
        os.environ["ADMIN_TOKEN"] = _saved_admin
    ok("🔴 C11 CLOSED: with ADMIN_TOKEN unset the admin surface is refused (403) even with the token that would match — bake, status, export",
       r_bake.status_code == 403 and r_stat.status_code == 403 and r_exp.status_code == 403
       and "not configured" in r_bake.text, f"{r_bake.status_code} {r_stat.status_code} {r_exp.status_code} {r_bake.text[:80]}")
    r = c.get("/api/admin/tenant/status", params=TOKEN)
    ok("...and with it set again the token opens the surface as before", r.status_code == 200, str(r.status_code))

    r = c.post("/api/admin/tenant/import", params=TOKEN, json={"modus_package": 99, "tables": {}})
    ok("a package of the wrong format is refused — 400 with the reason",
       r.status_code == 400 and "modus_package" in r.text, r.text[:160])
    r = c.post("/api/admin/tenant/import", json=DEMO)
    ok("🔴 import without the admin token — 403", r.status_code == 403, str(r.status_code))
    r = c.post("/api/admin/tenant/import", params=TOKEN, json=DEMO)
    res = r.json()
    ok("⭐ the demo package imports over HTTP", r.status_code == 200 and res.get("ok") is True, r.text[:200])
    # NARROWED 5 Oct 2026 (demo v2): 74 forecasts (72 approved + Pending + Rejected), 14 gates,
    # 5 zones, 6 haul-road links travel with the package too
    ok("...with every table's rows inserted",
       # NARROWED 11 Oct 2026 (carbon phase 4): + 2 routes, 4 forecasts, 18 weeks, 2 gates (the muck-away choice)
       # NARROWED 14 Oct 2026 (civils materials, D13): + 2 routes, 10 forecasts, 46 weeks, 1 gate (the concrete pour)
       # NARROWED 16 Oct 2026 (an access road to every gate, the user): + 17 zones (Z006-Z022)
       # RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): every table's count read from the package file itself (exact), for
       # every table the package carries rows in, so a regeneration never needs this line touched again
       all(res.get("inserted", {}).get(t) == len(rows) for t, rows in DEMO["tables"].items() if rows)
       and res["inserted"].get("routes") == 43 and res["inserted"].get("forecasts") == 907, str(res.get("inserted")))
    ok("...and the GB diesel price applied", res.get("fuel_index_applied") == ["GB"], str(res.get("fuel_index_applied")))
    r = c.post("/api/admin/tenant/import", params=TOKEN, json=DEMO)
    ok("🔴 a second import into a non-empty tenant is refused — 400",
       r.status_code == 400 and "already holds data" in r.text, r.text[:160])

    # 11 Oct 2026 (the demo bar, review): with the demo tenant in, a visitor with no cookie. Inside the app (the
    # iframe's /map/?embed=1, sent as a frame) the bar above carries the notice, so the password page leaves its pill
    # out; /map/ on its own, and /map/?embed=1 typed into the tab (Sec-Fetch-Dest: document), keep it.
    _jar = c.cookies
    c.cookies = type(_jar)()
    try:
        r_emb = c.get("/map/?embed=1&v=1", headers={"Sec-Fetch-Dest": "iframe"})
        r_emb_nohdr = c.get("/map/?embed=1&v=1")
        r_alone = c.get("/map/")
        r_typed = c.get("/map/?embed=1&v=1", headers={"Sec-Fetch-Dest": "document"})
        # 12 Oct 2026 (D31 / D18): the same visitor — no cookie, no code — reads /api/meta and asks for the network
        r_meta_out = c.get("/api/meta")
        r_meta_hdr = c.get("/api/meta", headers=PLAN)
        r_gates_out = c.get("/api/gates")
        r_who_out = c.get("/api/public/whoami")
    finally:
        c.cookies = _jar
    _mo = r_meta_out.json()
    ok("🔴 D31: with the demo tenant in, /api/meta to a visitor with NO cookie and NO code is the PUBLIC shape: the tenant's words, "
       "the vehicles and factors, but NO routes, NO teams, NO sections",
       r_meta_out.status_code == 200 and _mo["scope"] == "public" and _mo["tenant"]["name"].startswith("Wolds Link")
       and _mo["routes"] == [] and _mo["ipts"] == [] and _mo["work_sections"] == [] and len(_mo["vehicles"]) > 0
       and isinstance(_mo["factors"].get("vehicle_payload_t"), dict) and isinstance(_mo["seasonal_restrictions"], list),
       str({k: _mo.get(k) for k in ("scope", "routes", "ipts", "work_sections")})[:200])
    ok("⭐ D31: ...and the SAME visitor with the planner code in the header (the staff app's fetch) gets the staff shape: forty-three routes (RE-PINNED 17 Oct: the four-year programme), three teams, eight sections",
       r_meta_hdr.status_code == 200 and r_meta_hdr.json()["scope"] == "staff" and len(r_meta_hdr.json()["routes"]) == 43
       and [i["label"] for i in r_meta_hdr.json()["ipts"]] == ["North Team", "Central Team", "South Team"]
       and len(r_meta_hdr.json()["work_sections"]) == 8, str(r_meta_hdr.json().get("scope")))
    ok("🔴 D18: /api/gates to a visitor with no credential is the map's 401 JSON, not the gate list",
       r_gates_out.status_code == 401 and r_gates_out.json().get("detail") == "the map is password protected", r_gates_out.text[:120])
    ok("🔴 whoami: a visitor with no credential gets the gate's 401, and learns nothing", r_who_out.status_code == 401, str(r_who_out.status_code))
    ok("⭐ demo tenant, no cookie: the app's iframe (/map/?embed=1) gets the password page, 401, WITHOUT the notice pill (the bar has it)",
       r_emb.status_code == 401 and 'id="p" name="p" type="password"' in r_emb.text and 'class="demo-notice"' not in r_emb.text
       and r_emb_nohdr.status_code == 401 and 'class="demo-notice"' not in r_emb_nohdr.text, f"{r_emb.status_code} {r_emb_nohdr.status_code}")
    # NARROWED 14 Oct 2026 (the user took the recommendation): the demo package's notice is "the scheme and its plan are
    # invented" (was "nothing here is real", which the bulletin's real government figures contradicted); same three places
    ok("⭐ .../map/ on its own gets the password page WITH the notice pill (no bar above it)",
       r_alone.status_code == 401 and '<div class="demo-notice">Fictional scheme</div>' in r_alone.text,   # NARROWED 15 Oct 2026 (the user)
       str(r_alone.status_code))
    ok("🔴 ...and /map/?embed=1 typed into the tab (a top-level page) keeps the pill too",
       r_typed.status_code == 401 and 'class="demo-notice"' in r_typed.text, str(r_typed.status_code))

    m = c.get("/api/meta").json()
    ok("meta now reads the demo tenant: its name, £, GB",
       m["tenant"]["name"].startswith("Wolds Link") and m["tenant"]["currency_symbol"] == "£"
       and m["tenant"]["country"] == "GB", str(m["tenant"]))
    # NARROWED 12 Oct 2026 (D31): this client holds the STAFF cookie (from /api/auth above) and sends no header — the
    # app's iframe — and that alone earns the staff shape
    ok("⭐ ...three named teams, eight sections, forty-three routes (RE-PINNED 17 Oct) — the staff shape, on the staff COOKIE alone (no header)",
       [i["label"] for i in m["ipts"]] == ["North Team", "Central Team", "South Team"]
       and len(m["work_sections"]) == 8 and len(m["routes"]) == 43 and m["scope"] == "staff",
       str([i["label"] for i in m["ipts"]]) + " " + str(m.get("scope")))
    # RE-PINNED 17 Oct 2026 (R2, §3a, §1.8): the Burrough lane community agreement, 8 a day on each route (was 20 and 30)
    ok("§2.5: every meta route carries max_vehicles_per_day — 8 on R016 and R018 (the demo's village caps), None elsewhere",
       {r["route_id"]: r["max_vehicles_per_day"] for r in m["routes"] if r["max_vehicles_per_day"] is not None} == {"R016": 8, "R018": 8}
       and all("max_vehicles_per_day" in r for r in m["routes"]))

    r = c.post("/api/auth", json={"code": "north-code-xyz"})
    ok("⭐ a team code signs in under the team's NAME", r.status_code == 200
       and r.json()["label"] == "North Team" and r.json()["ipt"] == "IPT1", r.text[:160])
    ok("...and the planner picker offers the tenant's three slots, not six",
       c.post("/api/auth", json={"code": "plan-code-xyz"}).json().get("ipts") == ["IPT1", "IPT2", "IPT3"])

    r = c.get("/api/public/alignment")
    b = r.json()
    ok("⭐ the map overlay arrives: version, six bands, five boundaries, the view",
       b.get("version") == "wolds-link-demo-2" and len(b.get("bands", [])) == 6
       and len(b.get("boundaries", [])) == 5 and b.get("view", {}).get("zoom"), str(list(b))[:160])
    ok("...with the team names the map shows",
       {t["id"]: t["label"] for t in b["teams"]} == {"IPT1": "North Team", "IPT2": "Central Team", "IPT3": "South Team"},
       str(b.get("teams")))
    ok("...and a NEW ETag, so a browser holding the empty overlay refetches", r.headers.get("etag") != etag0)

    # 13 Oct 2026: the ETag covers the tenant block and the team names, not only the overlay row's updated_at.
    # Before, a Config save or a team rename left the tag unchanged, so the map kept its 304 and the old words.
    etag1 = r.headers.get("etag") or ""
    _du0 = b["tenant"]["distance_unit"]
    r = c.get("/api/public/alignment", headers={"If-None-Match": etag1})
    ok("⭐ ETag: with nothing changed, the demo overlay's own ETag gets 304 with no body", r.status_code == 304 and not r.content
       and r.headers.get("etag") == etag1, f"{r.status_code} {r.headers.get('etag')} vs {etag1}")
    _fdoc = c.get("/api/config/factors", headers=PLAN).json()["doc"]
    _fnew = json.loads(json.dumps(_fdoc))
    _fnew.setdefault("tenant", {})["distance_unit"] = "km" if _du0 == "mi" else "mi"
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _fnew, "updated_by": "http_smoke"})
    ok("ETag: the tenant block's distance_unit is changed on Config (PUT /api/admin/config/factors, token)",
       r.status_code == 200 and r.json().get("ok") is True, r.text[:200])
    r = c.get("/api/public/alignment", headers={"If-None-Match": etag1})
    ok("🔴 ETag: after a Config save of the tenant block, the OLD If-None-Match gets 200 (not 304) with the NEW distance unit",
       r.status_code == 200 and r.json()["tenant"]["distance_unit"] == _fnew["tenant"]["distance_unit"],
       f"{r.status_code} {r.content[:80]!r}")
    etag2 = r.headers.get("etag") or ""
    ok("...and a new ETag, which itself answers 304 next time",
       etag2 and etag2 != etag1 and c.get("/api/public/alignment", headers={"If-None-Match": etag2}).status_code == 304,
       f"{etag1} {etag2}")
    ok("...and the overlay itself is unchanged in that 200 (same version, bands)",
       r.status_code == 200 and r.json().get("version") == b.get("version") and r.json().get("bands") == b.get("bands"))
    # a team rename: there is no HTTP route for it (the names arrive with the package), so straight into the table
    _lab0 = db.query("SELECT label FROM ipts WHERE tenant_id = ? AND id = ?", (db.current_tenant(), "IPT1"))[0]["label"]
    db.execute("UPDATE ipts SET label = ? WHERE tenant_id = ? AND id = ?", ("Northern Team", db.current_tenant(), "IPT1"))
    try:
        r = c.get("/api/public/alignment", headers={"If-None-Match": etag2})
    finally:
        db.execute("UPDATE ipts SET label = ? WHERE tenant_id = ? AND id = ?", (_lab0, db.current_tenant(), "IPT1"))
    ok("🔴 ETag: after a team rename, the old If-None-Match gets 200 with the new team name",
       r.status_code == 200 and {t["id"]: t["label"] for t in r.json()["teams"]}.get("IPT1") == "Northern Team",
       f"{r.status_code} {r.content[:80]!r}")
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _fdoc, "updated_by": "http_smoke"})
    ok("ETag: the factors document is put back", r.status_code == 200 and r.json().get("ok") is True, r.text[:200])
    r = c.get("/api/public/alignment", headers={"If-None-Match": etag1})
    ok("...and with the document and the team name back, the FIRST ETag answers 304 again (the tag is the content, not a counter)",
       r.status_code == 304, str(r.status_code))
    ok("🔴 ETag: the 304 path does not read the overlay value — main.public_alignment asks get_overlay_stamp() before any get_overlay_row()",
       (lambda s: "get_overlay_stamp()" in s and s.index("get_overlay_stamp()") < s.index("status_code=304")
        < s.index("get_overlay_row()"))(
           open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read().split("def public_alignment(", 1)[1].split("\n@app.", 1)[0]))

    bad = json.loads(json.dumps({k: v for k, v in b.items() if k not in ("tenant", "teams")}))
    bad["bands"][-1]["chain_to"] = 1
    r = c.put("/api/admin/overlay", params=TOKEN, json=bad)
    ok("🔴 a backwards band is refused over HTTP — 400 naming it", r.status_code == 400 and "runs backwards" in r.text,
       r.text[:160])
    bad["bands"][-1]["chain_to"] = b["bands"][-1]["chain_to"]
    bad["boundaries"][0]["chain_m"] += 3
    r = c.put("/api/admin/overlay", params=TOKEN, json=bad)
    ok("🔴 a boundary off every band edge is refused — 400", r.status_code == 400 and "not a band edge" in r.text,
       r.text[:160])
    r = c.put("/api/admin/overlay", json=b)
    ok("🔴 overlay upload without the token — 403", r.status_code == 403)

    # REVERSED (H1, 29 Sep 2026): a GB tenant now HAS a restriction provider (National
    # Highways planned closures) and an automatic diesel series (DESNZ). The demo's typed
    # placeholder is the index until the first DESNZ fetch replaces it. lazy=0: this must
    # never start a background fetch of gov.uk from the harness.
    r = c.get("/api/restrictions/layers")
    ok("GB (H1) → the restriction provider is National Highways, one layer, 100 m",
       r.json().get("provider", {}).get("key") == "nh_closures" and r.json().get("match_m") == 100.0
       and [l["key"] for l in r.json().get("layers", [])] == ["nh_closures"], r.text[:200])
    r = c.get("/api/fuel-index?lazy=0", headers=PLAN)
    fi = r.json() if r.status_code == 200 else {}
    idx = fi.get("index") or {}
    ok("GB (H1) → the automatic series is DESNZ, and the demo's typed GB placeholder is the index until it is fetched",
       r.status_code == 200 and idx.get("country") == "GB" and idx.get("auto_available") is True
       and (idx.get("provider") or {}).get("key") == "desnz" and idx.get("eur_per_l") == 1.955
       and idx.get("source") == "manual", str(idx)[:200])
    r = c.get("/api/forecast-weeks/restrictions?bucket=commit", headers=PLAN)
    ok("GB (H1) → the look-ahead's restriction status names the provider and the ROADWORKS word",
       r.status_code == 200 and r.json().get("provider") == "nh_closures" and r.json().get("flag_label") == "ROADWORKS", r.text[:200])
    r = c.get("/api/forecast-weeks/tark-tee?bucket=commit", headers=PLAN)
    ok("...and the pre-H1 path still answers (an alias, hidden from the schema)", r.status_code == 200)
    r = c.get("/api/meta")
    # NARROWED 11 Oct 2026 (EU vehicles removed): the four EU names left the product, so the GB demo hides
    # Rigid 7.5t alone — asserted exactly, since "N3 or Rigid 7.5t" would now pass on the one name and prove nothing.
    ok("GB (H1) → /api/meta carries hidden_vehicles (Rigid 7.5t alone, no EU name) and the tenant's distance unit (miles)",
       r.json().get("hidden_vehicles") == ["Rigid 7.5t"]
       and not any(v.startswith("N3 ") for v in r.json().get("vehicles") or [])
       and r.json()["tenant"].get("distance_unit") == "mi", str(r.json().get("hidden_vehicles"))[:200])
    ok("11 Oct 2026: GB → /api/meta offers English vehicle names only, and leads with the UK trade names",
       r.json()["vehicle_labels"]["langs"] == ["en"] and (r.json().get("planning_vehicles") or [None])[0] == "Rigid 8-wheeler (32t)",
       str(r.json().get("vehicle_labels", {}).get("langs")))

    # ------------------------------------------------------- 3. the data, as people read it
    r = c.get("/api/forecasts", headers=PLAN)
    rows = r.json() if r.status_code == 200 else []
    # NARROWED 14 Oct 2026: + the ten concrete-pour line-months. RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): the four-year
    # programme, read from the package file itself
    ok("a planner sees all 907 forecast rows (905 approved, 1 pending, 1 rejected — the four-year programme)",
       r.status_code == 200 and len(rows) == len(DEMO["tables"]["forecasts"]) == 907, f"{r.status_code} {len(rows)}")
    r = c.get("/api/forecasts", headers=NORTH)
    nrows = r.json() if r.status_code == 200 else []
    ok("⭐ the North Team code sees ONLY North Team's rows",
       r.status_code == 200 and nrows and all(x.get("ipt") == "IPT1" for x in nrows) and len(nrows) < len(rows),
       f"{r.status_code} {len(nrows)}")
    r = c.get("/api/forecasts")
    ok("🔴 forecasts without a code — 401", r.status_code == 401, str(r.status_code))

    # ------------------------------------------- 3b. the carbon planner (10 Oct 2026, phase 2)
    # The demo package ships NO geometry, so every line is unbaked: transport is excluded and
    # says so. One route is then baked BY HAND (no HERE) with the illustrative distances of
    # notes/carbon-why-different-1010.md §3 (R001 Kilby Ridge → Harborough, 18.4 + 19.1 km), so
    # the figures can be checked through the real HTTP layer against the note and the Dashboard.
    r = c.post("/api/carbon/scenario", json={})
    ok("🔴 the carbon scenario without a code — 401", r.status_code == 401, str(r.status_code))
    r = c.post("/api/carbon/scenario", headers=PLAN, json={})
    cs = r.json() if r.status_code == 200 else {}
    bt = (cs.get("baseline") or {}).get("totals") or {}
    ok("⭐ the carbon scenario answers a planner: 905 approved line-months (RE-PINNED 17 Oct: the four-year programme), unbaked network EXCLUDED from transport",
       r.status_code == 200 and len(cs.get("lines") or []) == 905 and bt.get("baked_lines") == 0
       and bt.get("excludes_unbaked") is True and bt.get("transport_t") == 0 and bt.get("material_t", 0) > 0,
       f"{r.status_code} {r.text[:200]}")
    ok("...with the known limits, the Planned items and every factor's source, year and status",
       len(cs.get("limits") or []) >= 7 and len(cs.get("planned") or []) >= 4
       and all(f.get("status") for f in cs.get("factors") or []) and cs.get("unit") == "km")
    r = c.post("/api/carbon/scenario", headers=NORTH, json={})
    # NARROWED 17 Oct 2026 (R2, §3a): an invariant read in the same test, was "< 72" (the old demo's line count)
    ok("⭐ the North Team code's scenario holds only North Team's lines",
       r.status_code == 200 and r.json()["lines"] and all(l["ipt"] == "IPT1" for l in r.json()["lines"])
       and 0 < len(r.json()["lines"]) < len(cs.get("lines") or []), f"{r.status_code} {r.text[:160]}")
    r = c.post("/api/carbon/scenario", headers=PLAN,
               json={"allocation": {"Small aggregate": {"Fencing": 70, "Drainage": 40}}})
    ok("🔴 an allocation over 100 % is refused — 400 naming the sum", r.status_code == 400 and "110" in r.text, r.text[:160])
    r = c.post("/api/carbon/scenario", headers=PLAN,
               json={"alternatives": [{"name": "x", "levers": [{"driver": "material", "scope": {"category": "Precast / concrete"}, "epd": {"kg_per_t": 90}}]}]})
    ok("🔴 an EPD without its source is refused — 400", r.status_code == 400 and "source" in r.text, r.text[:160])
    r = c.post("/api/carbon/scenario", headers=PLAN, json={"status": "Everything"})
    ok("an unknown status is refused — 400", r.status_code == 400)
    # phase 4: the demo's landfill-vs-reuse choice. RE-PINNED 17 Oct 2026 (R2, §3a): R019's approved tonnes are read from the
    # package (135,000 t: the Langton surplus Apr 2027 – Sep 2028 and the lifted platform stone Jan–May 2030); the two
    # factors stay literal (the independent check): landfill × 19.55376 kg/t, the recycling hub × 1.01398 kg/t
    _t19 = sum(f["quantity"] for f in DEMO["tables"]["forecasts"] if f["route_id"] == "R019" and f["status"] == "Approved")
    _n19 = sum(1 for f in DEMO["tables"]["forecasts"] if f["route_id"] == "R019" and f["status"] == "Approved")
    _land19, _cut19 = _t19 * 19.55376 / 1000.0, _t19 * (19.55376 - 1.01398) / 1000.0
    r = c.post("/api/carbon/scenario", headers=PLAN, json={"alternatives": [
        {"name": "Reuse the cut", "levers": [{"driver": "disposal", "scope": {"route_id": "R019"}, "to_route": "R020"}]}]})
    cs = r.json() if r.status_code == 200 else {}
    _l19 = [l for l in cs.get("lines") or [] if l["route_id"] == "R019"]
    _a = (cs.get("alternatives") or [{}])[0]
    ok("⭐ demo (phase 4): the surplus cut to the fictional landfill is a Waste line at 19.55376 kg/t (R019's tonnes × the factor) …",
       r.status_code == 200 and _t19 == 135000 and len(_l19) == _n19 == 23
       and all(l["disposal_route"] == "landfill" and l["sheet"] == "Waste" for l in _l19)
       and abs(sum(l["disposal_t"] for l in _l19) - _land19) < 0.002, str(_l19)[:300])
    ok("⭐ …and sending it to the fictional soil-recycling hub moves R019's tonnes × (19.55376 − 1.01398) kg/t, ranked as the disposal decision",
       (_a.get("drivers") or [{}])[0].get("driver") == "disposal" and abs(_a.get("delta_t", 0) + _cut19) < 0.002, str(_a.get("drivers"))[:200])

    # B4 (12 Oct 2026): with NO geometry the map draws nothing, so the bulletin's roll-up is EMPTY and names every
    # line as not drawn — it never presents zeros as figures. The baked case is 3c below.
    r = c.get("/api/bulletin/month?month=10", headers=PLAN)
    _pe = r.json() if r.status_code == 200 else {}
    # NARROWED 14 Oct 2026 (civils materials): 22 October lines on 21 routes (the pour's two Melton lines share a route)
    # RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): October 2026 of the four-year programme is 18 lines on 13 routes
    ok("⭐ the bulletin on the unbaked demo: empty (four headline figures null), all 18 October lines named as on routes the map does not draw",
       r.status_code == 200 and _pe["rollup"]["empty"] is True and _pe["rollup"]["headline"]["two_way_a_day"] is None
       and _pe["rollup"]["headline"]["routes_in_use"] == 0 and _pe["rollup"]["routes_not_drawn"]["lines"] == 18
       and len(_pe["rollup"]["routes_not_drawn"]["routes"]) == 13
       # NARROWED 12 Oct 2026 (the review and fix pass): plain words, "with no stored road route yet", not "not yet baked"
       # NARROWED 22 Oct 2026 (words sweep): "not routed yet", as the map says it
       and _pe["rollup"]["note"] == ("No approved forecast on a route the map draws in this month: 18 lines on 13 routes "
                                     "not routed yet, listed below and not counted.")
       and _pe["conditions"]["route_caps"]["wording"] == "planning flag", f"{r.status_code} {str(_pe)[:300]}")
    # NARROWED 12 Oct 2026 (the tests-and-regressions review): the carbon figures on that empty roll-up are
    # BOTH None (co2_t_a_day used to read 0.0), and there is no co2_error to carry
    ok("…its carbon block: both figures None, not_baked False (no counted line at all), no error",
       _pe["rollup"]["carbon"]["co2_t_a_day"] is None and _pe["rollup"]["carbon"]["co2_t_in_month"] is None
       and _pe["rollup"]["carbon"]["not_baked"] is False and _pe["rollup"]["carbon"]["error"] is None, str(_pe["rollup"].get("carbon")))

    _r001 = [x for x in rows if x["route_id"] == "R001" and x["status"] == "Approved"]
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
               "duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", ("default", "R001", "Artic Tipper (44t)", "loaded", 0, "[[-1.2,52.5],[-1.1,52.6]]", 18.4, 0.42))
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
               "duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", ("default", "R001", "Artic Tipper (44t)", "return", 0, "[[-1.1,52.6],[-1.2,52.5]]", 19.1, 0.44))
    try:
        r = c.post("/api/carbon/scenario", headers=PLAN, json={"alternatives": [
            {"name": "Electric on R001", "levers": [{"driver": "fuel", "scope": {"route_id": "R001"}, "fuel": "electric"}]}]})
        cs = r.json() if r.status_code == 200 else {}
        bt = (cs.get("baseline") or {}).get("totals") or {}
        cl = c.get("/api/costing/lines?status=Approved", headers=PLAN).json()
        cl_r001 = sum(x["co2_t"] for x in cl["lines"] if x["route_id"] == "R001" and x["co2_t"] is not None)
        # RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): was the note's worked example (62,500 t over four months, 75.99 t).
        # R001 now carries the Harborough platform stone (Oct–Dec 2026) and the WS1 sub-base (Nov 2028 – Apr 2029), so the
        # figures are COMPUTED here from R001's own rows; the factors stay literal (they are the independent check):
        # trips = ceil(t ÷ 29 t) per line-month (the app's rule), vehicle-km = trips × (18.4 + 19.1) km, × 0.93939 kg/km
        _q001 = [float(x["quantity"]) for x in _r001]
        _trips001 = sum(math.ceil(q / 29.0 - 1e-9) for q in _q001)
        _vkm001 = _trips001 * (18.4 + 19.1)
        _tr001 = _vkm001 * 0.93939 / 1000.0
        ok("⭐ R001 baked by hand: its 9 approved line-months (90,500 t), transport = trips × both legs × 0.93939 kg/km, over HTTP",
           len(_r001) == 9 and abs(sum(_q001) - 90500) < 1e-6 and bt.get("baked_lines") == 9
           and abs(bt.get("transport_t", 0) - _tr001) < 0.01, f"{bt.get('baked_lines')} {bt.get('transport_t')} vs {_tr001}")
        ok("⭐ ...the SAME figure the Dashboard reads from /api/costing/lines (one number)", abs(bt.get("transport_t", 0) - cl_r001) < 0.002,
           f"{bt.get('transport_t')} vs {cl_r001}")
        # 12 Oct 2026 (B2): the public map's two feeds carry that figure too, per line per month, with no co2_error
        r = c.get("/api/public/forecast-matrix", params={"from": 1, "to": 60, "unit": "t"})
        fm = r.json() if r.status_code == 200 else {}
        _fm001 = [x for x in fm.get("routes") or [] if x["route_id"] == "R001"]
        _fm_sum = sum(v for x in _fm001 for v in (x.get("co2_t") or {}).values() if v is not None)
        ok("⭐ B2: /api/public/forecast-matrix carries co2_t per line per month for the hand-baked R001, keyed like monthly, summing to the costlines figure",
           r.status_code == 200 and _fm001 and all(set(x.get("co2_t") or {}) == set(x["monthly"]) for x in _fm001)
           and all(v is not None for x in _fm001 for v in x["co2_t"].values()) and abs(_fm_sum - cl_r001) < 0.002,
           f"{r.status_code} {_fm_sum} vs {cl_r001} {[x.get('co2_t') for x in _fm001][:2]}")
        ok("B2: ...an UNBAKED route's months read None, never a zero (the rest of the demo is not baked here)",
           any(all(v is None for v in x["co2_t"].values()) for x in fm.get("routes") or [] if x["route_id"] != "R001"))
        ok("B2: ...and co2_error is absent on the happy path", "co2_error" not in fm and "working_days" in fm)
        # 12 Oct 2026 (Track C review): sections=1 adds the largest section-month (the map's section ramp); absent without it
        r = c.get("/api/public/forecast-matrix", params={"from": 1, "to": 60, "unit": "t", "sections": 1})
        ok("Track C review: forecast-matrix with sections=1 answers 200 and carries section_co2_max (None or a positive figure); without it the key is absent",
           r.status_code == 200 and "section_co2_max" in r.json()
           and (r.json()["section_co2_max"] is None or r.json()["section_co2_max"] > 0) and "section_co2_max" not in fm, r.text[:200])
        _m001 = sorted({x["month_index"] for x in _r001})[0]
        r = c.get("/api/public/month-kpis", params={"month": _m001, "unit": "t"})
        mk = r.json() if r.status_code == 200 else {}
        ok("B2: /api/public/month-kpis for R001's first month: co2_t present on its line, co2_error absent",
           r.status_code == 200 and "co2_error" not in mk
           and any(l["route_id"] == "R001" and l["co2_t"] is not None for l in mk.get("lines") or []), f"{r.status_code} {str(mk)[:160]}")
        # 21 Oct 2026 (live OOM kills, branch claude/map-read-memory): every month in one read, through the real app
        r_all = c.get("/api/public/month-kpis/months", params={"unit": "t"})
        _all = r_all.json() if r_all.status_code == 200 else {}
        ok("⭐ 21 Oct: /api/public/month-kpis/months answers 200 through the real app: all 60 months, R001's first month exactly "
           "the month read's answer",
           r_all.status_code == 200 and sorted(int(k) for k in _all.get("months", {})) == list(range(1, 61))
           and json.dumps(_all["months"][str(_m001)]) == json.dumps(mk), f"{r_all.status_code} {r_all.text[:160]}")
        ok("21 Oct: ...a unit it does not know is a 400",
           c.get("/api/public/month-kpis/months", params={"unit": "furlongs"}).status_code == 400)
        _anon = TestClient(main.app, base_url="https://testserver")
        ok("21 Oct: ...and with no cookie and no code it is behind the gate like the month read (401)",
           _anon.get("/api/public/month-kpis/months").status_code == 401
           and _anon.get("/api/public/month-kpis", params={"month": 1}).status_code == 401)
        ok("21 Oct: /api/health still answers 200 (now async: answered while the worker threads are busy)",
           _anon.get("/api/health").status_code == 200)
        # RE-PINNED 17 Oct 2026 (R2): the Carbon Tool basis computed from R001's tonnes × the 18.4 km one-way leg × the tool's
        # 2.27669e-4 t per t per one-way km (laden + unladen, 2024 factors with well-to-tank)
        _nh001 = sum(_q001) * 18.4 * 0.00022766866666666666
        ok("...and the NH Carbon Tool basis beside it: tonnes × one-way km × the tool's factor, labelled, not added in",
           abs(bt.get("transport_nh_t", 0) - _nh001) < 0.01 and "NH Carbon Tool" in (cs.get("methods") or {}).get("nh_tool", ""),
           str(bt.get("transport_nh_t")))
        alt = (cs.get("alternatives") or [{}])[0]
        fc = cs.get("fleet_compare") or {}
        ok("⭐ electric on R001: the ranked driver is diesel vs electric, carbon down, more trips at the lower payload",
           (alt.get("drivers") or [{}])[0].get("driver") == "fuel" and alt.get("delta_t", 0) < -50
           and fc.get("lines") == 9 and fc["electric"]["trips"] > fc["diesel"]["trips"], str(alt.get("drivers"))[:200])   # RE-PINNED 17 Oct 2026: R001's 9 lines
        # after the show (11 Oct 2026, items 12 and 14): well-to-wheel and the modules, over HTTP.
        # RE-PINNED 17 Oct 2026 (R2): R001's artic-km × the 0.22058 kg/km well-to-tank row, computed from its own rows
        _wtt001 = _vkm001 * 0.22058 / 1000.0
        ok("⭐ well-to-wheel over HTTP: R001's tank-to-wheel + well-to-tank (artic-km × 0.22058), beside the total (not in it)",
           abs(bt.get("transport_wtt_t", 0) - _wtt001) < 0.01 and abs(bt.get("transport_wtw_t", 0) - (_tr001 + _wtt001)) < 0.02
           and abs(bt.get("total_t", 0) - (bt.get("material_t", 0) + bt.get("transport_t", 0) + bt.get("disposal_t", 0))) < 0.01,
           f"{bt.get('transport_wtt_t')} {bt.get('transport_wtw_t')}")
        st = (cs.get("baseline") or {}).get("stages") or {}
        ok("...the life-cycle modules add up to the total over HTTP", st and abs(sum(st.values()) - bt.get("total_t", 0)) < 0.01, str(st))
        # item 11: the summary PDF through the real HTTP layer — the demo notice and the stamp on it
        r = c.post("/api/carbon/summary.pdf", json={})
        ok("🔴 the carbon summary without a code — 401", r.status_code == 401, str(r.status_code))
        r = c.post("/api/carbon/summary.pdf", headers=PLAN, json={"alternatives": [
            {"name": "Electric on R001", "levers": [{"driver": "fuel", "scope": {"route_id": "R001"}, "fuel": "electric"}]}],
            "commentary": "Smoke check & <nothing> else"})
        from portable import pdf_text
        _st = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
        ok("⭐ the carbon summary over HTTP: a one-page PDF, stamped, with the fictional-scheme notice and the commentary as typed",
           r.status_code == 200 and r.headers.get("content-type", "").startswith("application/pdf") and r.content[:5] == b"%PDF-"
           and "Wayscope planning estimate, DESNZ 2026 basis, not Carbon Tool figures" in _st
           and "DEMO Fictional scheme" in _st and "Smoke check & <nothing> else" in _st   # NARROWED 15 Oct 2026 (the user)
           and "Battery electric lorries, where Config has an electric profile for that lorry — route R001" in _st,
           f"{r.status_code} {_st[:300]}")
        r = c.post("/api/carbon/summary.pdf", headers=PLAN, json={"commentary": "x" * 1501})
        ok("🔴 a commentary over 1,500 characters is refused — 400", r.status_code == 400 and "1500" in r.text, r.text[:160])

        # --------------------------------------- 3c. the monthly bulletin (B4, 12 Oct 2026), while R001 is baked
        r = c.get("/api/bulletin/month?month=10")
        ok("🔴 the bulletin without a code — 401", r.status_code == 401, str(r.status_code))
        ok("🔴 a month outside the horizon is refused — 422", c.get("/api/bulletin/month?month=0", headers=PLAN).status_code == 422)
        r = c.get("/api/bulletin/month?month=10", headers=PLAN)
        pg = r.json() if r.status_code == 200 else {}
        hl = (pg.get("rollup") or {}).get("headline") or {}
        _wd = pg.get("working_days")
        ok("⭐ the bulletin answers a planner for Oct 2026: the four headline figures present, working days stated, the demo notice on it",
           r.status_code == 200 and pg["month"] == {"index": 10, "label": "Oct 2026", "label_long": "October 2026",
                                                     "from": "2026-10-01", "to": "2026-10-31"}
           and all(hl.get(k) is not None for k in ("two_way_a_day", "movements_in_month", "routes_in_use", "tonnes"))
           and _wd and hl["working_days"] == _wd and pg["demo_notice"] == "Fictional scheme" and "nothing here is real" not in r.text   # NARROWED 15 Oct 2026
           and pg["plan_only"] is True, f"{r.status_code} {str(pg)[:300]}")
        _kp = c.get("/api/public/month-kpis?month=10", headers=PLAN).json()
        _r001 = [l for l in _kp["lines"] if l["route_id"] == "R001"]
        # RE-PINNED 17 Oct 2026 (R2, §3a): October 2026 is now 18 lines on 13 routes; R001's October is 2,800 t of platform
        # stone. Read from the package, so the arithmetic stays the map bar's
        _oct = [f for f in DEMO["tables"]["forecasts"] if f["month_index"] == 10 and f["status"] == "Approved"]
        _t001 = sum(f["quantity"] for f in _oct if f["route_id"] == "R001")
        ok("⭐ …the figures are the map bar's rules over the map's own payload: R001 is the ONE drawn route — its October tonnes ÷ 29 t "
           "= loads ÷ working days; vehicles = ceil per line from the same trips_per_vehicle_day; the 12 undrawn routes named",
           _t001 == 2800 and len(_oct) == 18 and len({f["route_id"] for f in _oct}) == 13
           and hl["routes_in_use"] == 1 and abs(hl["tonnes"] - _t001) < 1e-6 and abs(hl["movements_in_month"] - _t001 / 29) < 1e-3
           and abs(hl["two_way_a_day"] - _t001 / 29 / _wd) < 1e-3 and len(_r001) == 1 and _r001[0]["trips_per_vehicle_day"]
           and pg["rollup"]["vehicles"]["fleet"] == -(-(_t001 / 29 / _wd) // _r001[0]["trips_per_vehicle_day"])
           # NARROWED 14 Oct 2026: 21 lines on 20 undrawn routes (the concrete pour). RE-PINNED 17 Oct 2026: 17 on 12
           and pg["rollup"]["routes_not_drawn"]["lines"] == 17 and len(pg["rollup"]["routes_not_drawn"]["routes"]) == 12
           and "R001" not in pg["rollup"]["routes_not_drawn"]["routes"], f"{hl} {pg['rollup'].get('vehicles')} {_r001}")
        _cl10 = sum(x["co2_t"] for x in cl["lines"] if x["route_id"] == "R001" and x["month_index"] == 10 and x["co2_t"] is not None)
        ok("⭐ …carbon is the costlines figure the Dashboard reads (R001, October), stamped as the Carbon page stamps it",
           abs(pg["rollup"]["carbon"]["co2_t_in_month"] - _cl10) < 0.002 and _cl10 > 0
           and pg["rollup"]["carbon"]["stamp"] == "Wayscope planning estimate, DESNZ 2026 basis, not Carbon Tool figures",
           f"{pg['rollup']['carbon']} vs {_cl10}")
        ok("…and the block is a figure, not a gap: not_baked False, no error (12 Oct 2026 review)",
           pg["rollup"]["carbon"]["not_baked"] is False and pg["rollup"]["carbon"]["error"] is None)
        _b = pg["rollup"]["busiest_routes"]
        ok("…the busiest route is R001 by its places, with \"no planning cap typed\"; the team row reads the tenant's name, never an id",
           len(_b) == 1 and _b[0]["label"] == "Kilby Ridge Quarry → Harborough Compound" and _b[0]["cap_wording"] == "no planning cap typed"
           and [t["key"] for t in pg["rollup"]["breakdown"]["team"]] == ["South Team"] and "IPT" not in r.text, str(_b)[:200])
        _cn = pg["conditions"]
        # NARROWED 14 Oct 2026: 21 routes carry October volume (the concrete pour's two). RE-PINNED 17 Oct 2026 (R2): 13
        ok("⭐ …conditions of the month from what is STORED: no stored check, so \"never checked\" on all 13 routes (National Highways provider)",
           _cn["restrictions"]["status"] == "never checked" and len(_cn["restrictions"]["never_checked_routes"]) == 13
           and _cn["restrictions"]["hits"] == [] and _cn["restrictions"]["provider"] == "nh_closures", str(_cn["restrictions"])[:200])
        # REVERSED 14 Oct 2026 (the user took the recommendation): the public bulletin no longer carries the stockpile over
        # capacity (a site matter). The sentence and the figures stay pinned on bulletin.conditions_of_month in
        # test_bulletin.py, which still computes them. NARROWED 17 Oct 2026 (R2, §3a): the four-year demo has no stockpile
        # over capacity at its shipped date at all, so: no key, and NO stockpile named on the public page
        ok("⭐ …the public page carries NO stockpile: no stockpiles_over key, no stockpile named in the body, and the count leaves it out",
           "stockpiles_over" not in _cn and not any(n in r.text for n in ("Stockpile Tilton", "Stockpile North", "Stockpile South"))
           and _cn["count"] == len(_cn["restrictions"]["hits"]) + len(_cn["seasonal"]) + len(_cn["route_caps"]["exceeded"]),
           str(sorted(_cn)))
        ok("…and the bulletin never says \"the plan puts\" (D23)", "plan puts" not in r.text)
        _cap = _cn["route_caps"]
        # RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): the Burrough caps are 8 a day each, and October 2026 exceeds neither,
        # so the exceeded case is read in May 2027 (month 17), where R018 runs 11.3 loads a day against its cap of 8
        ok("…October 2026: both village caps typed, neither exceeded, still the \"planning flag\" wording",
           _cap["wording"] == "planning flag" and _cap["routes_with_cap"] == 2 and _cap["exceeded"] == [], str(_cap)[:300])
        r17 = c.get("/api/bulletin/month?month=17", headers=PLAN)
        _cap17 = ((r17.json() if r17.status_code == 200 else {}).get("conditions") or {}).get("route_caps") or {}
        ok("⭐ …the typed cap on R018 (8 a day) is exceeded by May 2027's loaded trips a day and reads \"planning flag\", never a block",
           r17.status_code == 200 and _cap17.get("wording") == "planning flag" and _cap17.get("routes_with_cap") == 2
           and [e["route_id"] for e in _cap17.get("exceeded") or []] == ["R018"]
           and _cap17["exceeded"][0]["cap"] == 8 and _cap17["exceeded"][0]["two_way_a_day"] > 8
           and "never a block" in _cap17["note"] and "planning flag" in r17.text, str(_cap17)[:300])
        ok("…no seasonal window on the demo (the seed has none)", _cn["seasonal"] == [])
        r2 = c.get("/api/bulletin/month?month=10", headers=NORTH)
        pn = r2.json() if r2.status_code == 200 else {}
        # REVERSED 12 Oct 2026 (D37): a team code gets every team's row too — here South Team's, the one drawn route
        ok("⭐ the North Team code gets the SCHEME's figures (the public map's), every team's row included (D37)",
           r2.status_code == 200 and pn["rollup"]["headline"] == hl and pn["rollup"]["busiest_routes"] == _b
           and pn["rollup"]["breakdown"]["team"] == pg["rollup"]["breakdown"]["team"] and len(pn["rollup"]["breakdown"]["team"]) == 1
           and pn["scope"]["team"] == "North Team"
           and pn["rollup"]["breakdown"]["discipline"] == pg["rollup"]["breakdown"]["discipline"], f"{r2.status_code} {str(pn.get('scope'))}")
        # 12 Oct 2026 (C4): the bulletin as people read it, over HTTP — the PDF and the page, from the same roll-up.
        # Mapbox is never called: portable.offline() refuses the static-image fetch, so the map is the schematic.
        # 14 Oct 2026: the fetch is now the basemap alone (basemap.py), refused the same way
        import export
        from portable import offline
        _real_open = offline(export)
        try:
            # ---- 14 Oct 2026 (backlog item 2): the route map export over HTTP (real reportlab, the demo package) ----
            # 14 Oct 2026, evening (the CLP plans): just the map, as a regional / local / site plan, PDF or PNG
            r = c.get("/api/routes/map-export.pdf", params={"month": 10}, headers=PLAN)
            _rt = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            # NARROWED 14 Oct 2026, evening: one page, just the map (no demo bar, no table, no scheme's block); the tag on the map
            ok("⭐ ROUTE MAP EXPORT: the planner downloads the map as filtered — ONE A4 landscape page, a regional plan, the demo tag on the map",
               r.status_code == 200 and r.headers.get("content-type") == "application/pdf" and b"/MediaBox [ 0 0 841.8898 595.2756 ]" in r.content
               and r.content.count(b"/Type /Page\n") == 1
               and re.search(r'filename="route-map-2026-10-regional(-1-\d+)?\.pdf"', r.headers.get("content-disposition", "")) is not None
               and r.headers.get("cache-control") == "no-store"
               and "Regional plan" in _rt and "Demo - fictional scheme" in _rt and "Two-way movements a day" in _rt   # NARROWED 15 Oct 2026
               and "not an approved CLP" in _rt, f"{r.status_code} {r.headers.get('content-disposition')} {_rt[:300]}")
            ok("🔴 ...no team id, no 'truck', no 'baked' in it; no table, no scheme's block, no Carbon Tool rows (just the map)",
               not re.search(r"\bIPT ?[1-6]\b", _rt) and not re.search(r"\btrucks?\b", _rt, re.I)
               and not re.search(r"\b(un)?baked?\b", _rt, re.I) and "PLANNING FLAG" not in _rt
               and "To be completed by the scheme" not in _rt and "Carbon Tool" not in _rt, _rt[:300])
            # REVERSED 14 Oct 2026, evening: the old options are ignored — the export is only ever the map
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "carbon_tool": 1, "map_only": 1}, headers=PLAN)
            _rc = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            ok("🔴 ...carbon_tool=1 and map_only=1 change nothing: no Carbon Tool rows, the same one page",
               r.status_code == 200 and "Carbon Tool" not in _rc and r.content.count(b"/Type /Page\n") == 1, _rc[-300:])
            r = c.get("/api/routes/map-export/places", params={"month": 10}, headers=PLAN)
            _pl = r.json() if r.status_code == 200 else {}
            ok("⭐ ...the card's list: the three plans with the CLP template's scales, and the places to centre on (names, never ids)",
               r.status_code == 200 and [p["kind"] for p in _pl.get("plans", [])] == ["regional", "local", "site"]
               and _pl["plans"][1]["scales"] == [2000, 2500, 3000] and _pl["plans"][2]["scales"] == [500, 1000]
               and isinstance(_pl.get("places"), list) and not any(re.match(r"^L\d+$", str(n)) for n in _pl.get("places", [])), str(_pl)[:300])
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "plan": "site", "scale": 2500}, headers=PLAN)
            ok("🔴 ...a scale the plan does not offer: 422 with the sentence", r.status_code == 422
               and r.json().get("detail") == "the site plan is drawn at 1:500, 1:1,000", r.text[:200])
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "plan": "local", "at": "Nowhere"}, headers=PLAN)
            ok("🔴 ...a place not at the end of a route: 422", r.status_code == 422 and "Nowhere is not at the end of a route" in r.text, r.text[:200])
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "plan": "local"}, headers=PLAN)
            _rl = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            ok("...a local plan over HTTP: its name, and a scale only when there is a place to centre on (the demo package carries no road routes)",
               r.status_code == 200 and "Local plan" in _rl and (("Scale 1:2,500 at A4" in _rl) == bool(_pl.get("places"))),
               f"{r.status_code} {_rl[:200]}")
            r = c.get("/api/routes/map-export.png", params={"month": 10}, headers=PLAN)
            ok("⭐ ...the PNG for a Word document: an image, named for the month and plan, not cached",
               r.status_code == 200 and r.headers.get("content-type") == "image/png" and r.content[:8] == b"\x89PNG\r\n\x1a\n"
               and re.search(r'filename="route-map-2026-10-regional(-1-\d+)?\.png"', r.headers.get("content-disposition", "")) is not None
               and r.headers.get("cache-control") == "no-store", f"{r.status_code} {r.headers.get('content-type')} {r.text[:120] if r.status_code != 200 else ''}")
            # ---- 16 Oct 2026 (CLOCS v5 C2 / C3): the route sheet for drivers over HTTP (the demo package) ----
            r = c.put("/api/admin/routes/R016/sheet-text", params=TOKEN,
                      json={"route_reason": "  Agreed in the outline CLP.  ", "permitted_deviations": "The signed diversion only."})
            ok("⭐ ROUTE SHEET: the admin types why this route and the permitted diversions (stored trimmed)",
               r.status_code == 200 and r.json() == {"id": "R016", "route_reason": "Agreed in the outline CLP.",
                                                     "permitted_deviations": "The signed diversion only."}, r.text[:200])
            r = c.put("/api/admin/routes/R016/sheet-text", params=TOKEN, json={"route_reason": "x" * 401})
            ok("🔴 ...401 characters: 400 Bad Request with the sentence, and the reason is unchanged",
               r.status_code == 400 and r.json().get("detail") == "why this route is 401 characters: at most 400 fit on the route sheet"
               and next(x for x in c.get("/api/routes/status").json() if x["id"] == "R016")["route_reason"] == "Agreed in the outline CLP.",
               r.text[:200])
            ok("🔴 ...an unknown route: 404", c.put("/api/admin/routes/R999/sheet-text", params=TOKEN, json={"route_reason": "a"}).status_code == 404)
            r = c.get("/api/routes/R016/route-sheet.pdf", headers=PLAN)
            _st = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            ok("⭐ ...the planner downloads it: ONE A4 PORTRAIT page named for the route, not cached, the typed words on it, the demo "
               "heading; the demo package carries no road routes, so the map says so",
               r.status_code == 200 and r.headers.get("content-type") == "application/pdf"
               and b"/MediaBox [ 0 0 595.2756 841.8898 ]" in r.content and r.content.count(b"/Type /Page\n") == 1
               and 'filename="route-sheet-R016.pdf"' in r.headers.get("content-disposition", "")
               and r.headers.get("cache-control") == "no-store" and "Route sheet" in _st
               and "Agreed in the outline CLP." in _st and "The signed diversion only." in _st and "DEMO · " in _st
               and "No road route has been worked out for this route yet." in _st, f"{r.status_code} {_st[:300]}")
            ok("🔴 ...no team id, no 'truck', no 'baked' on it", not re.search(r"\bIPT ?[1-6]\b", _st)
               and not re.search(r"\btrucks?\b", _st, re.I) and not re.search(r"\b(un)?baked?\b", _st, re.I), _st[:300])
            ok("🔴 ...a team code gets its OWN team's route (R016, North) and a 404 for another team's (R001, South)",
               c.get("/api/routes/R016/route-sheet.pdf", headers=NORTH).status_code == 200
               and c.get("/api/routes/R001/route-sheet.pdf", headers=NORTH).status_code == 404)
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "team": "IPT2"}, headers=NORTH)
            _rn = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            ok("🔴 ...a team code asking for another team gets its OWN team's routes",
               r.status_code == 200 and " · North Team · " in _rn and "Central Team" not in _rn, _rn[:240])
            r = c.get("/api/routes/map-export.pdf", params={"month": 10, "team": "Central Team"}, headers=PLAN)
            _rp = " ".join(pdf_text(r.content).split()) if r.status_code == 200 else ""
            ok("...the planner can filter by the team's NAME (the map's deep link) as well as its slot",
               r.status_code == 200 and " · Central Team · " in _rp, _rp[:240])
            ok("🔴 ...a month outside the horizon is refused — 422",
               c.get("/api/routes/map-export.pdf", params={"month": 61}, headers=PLAN).status_code == 422)
            r = c.get("/api/bulletin/month.pdf?month=10", headers=PLAN)
            _bt = pdf_text(r.content) if r.status_code == 200 else ""
            _bp = [" ".join(x.split()) for x in _bt.split("\f") if x.strip()]
            _bj = " ".join(_bt.split())
            r_pdf_bytes = r.content if r.status_code == 200 else b""
            # NARROWED 12 Oct 2026 (the user, evening): movements in the month print whole ("583", was "582.8")
            _figs = (f"{hl['two_way_a_day']:,.0f}", f"{hl['movements_in_month']:,.0f}", f"{hl['routes_in_use']:,}", f"{hl['tonnes']:,.0f} t")
            ok("⭐ the bulletin PDF over HTTP: a PDF named for the month, not cached, \"fictional demo scheme\" on every page",
               r.status_code == 200 and r.headers.get("content-type", "").startswith("application/pdf") and r.content[:5] == b"%PDF-"
               and r.headers.get("content-disposition") == 'attachment; filename="bulletin-2026-10.pdf"'
               and r.headers.get("cache-control") == "no-store" and len(_bp) >= 2
               and all("demo fictional scheme" in " ".join(x.split()).lower() for x in _bp), f"{r.status_code} {len(_bp)} pages")   # NARROWED 15 Oct 2026: "DEMO Fictional scheme" on every page
            ok("⭐ ...page 1 carries the four headline figures equal to /api/bulletin/month's roll-up",
               bool(_bp) and all(f in _bp[0] for f in _figs), f"{_figs} | {_bp[0][:400] if _bp else ''}")
            # REVERSED 14 Oct 2026 (the user took the recommendation): the stockpile over capacity is a site matter and is
            # no longer in the public bulletin — its absence is asserted where its presence was
            ok("...the busiest route by its places, NO stockpile, the never-checked routes, the carbon stamp, the planning flag",
               # NARROWED 14 Oct 2026 (the publishing check): plain words, the flagged route by its places (no "Route R016"),
               # the stockpile as its balance against its capacity; the stamp is in Sources (two columns: read in drawing order)
               "Kilby Ridge Quarry to Harborough Compound" in _bj and "40,000 t capacity" not in _bj
               and "Stockpile Tilton" not in _bj and "Over capacity" not in _bj and "Not checked against published road works" in _bj
               and "Wayscope planning estimate, DESNZ 2026 basis, not Carbon Tool figures" in " ".join(pdf_text(r_pdf_bytes, raw=True).split())
               # RE-PINNED 17 Oct 2026 (R2): October exceeds no cap, so no flag sentence; May 2027's PDF carries R018's (below)
               and "Route R016" not in _bj and "above its planning flag of" not in _bj, _bj[:600])
            r17p = c.get("/api/bulletin/month.pdf?month=17", headers=PLAN)
            _b17 = " ".join(pdf_text(r17p.content).split()) if r17p.status_code == 200 else ""
            # (the flagged route reads by its places only when the map draws it; undrawn here, it reads "One route" — never its id)
            ok("⭐ ...May 2027's bulletin PDF carries the flag sentence, never the route's id: above its planning flag of 8 a day",
               r17p.status_code == 200 and "R018" not in _b17 and "above its planning flag of 8 a day." in _b17, _b17[:600])
            ok("🔴 ...no team id, none of the banned words",
               not re.search(r"\bIPT\s*[-_ ]?\s*[1-6]\b", _bj) and not re.search(r"\b(secure|real[- ]time|telematics|optimi[sz]e|predict|live tracking)\b", _bj, re.I)
               and not re.search(r"\bAI\b", _bj))
            r = c.get("/bulletin/?month=10", headers=PLAN)
            ok("⭐ the bulletin page over HTTP with a code: HTML, not cached, the same four figures, the PDF alongside",
               r.status_code == 200 and r.headers.get("content-type", "").startswith("text/html")
               and r.headers.get("cache-control") == "no-store" and all(f in r.text for f in _figs)
               and 'href="/api/bulletin/month.pdf?month=10">Download the PDF</a>' in r.text
               and "Fictional scheme" in r.text, f"{r.status_code} {r.text[:200]}")   # NARROWED 15 Oct 2026 (the user)
            r = c.get("/api/bulletin/leaflet.pdf?month=10", headers=PLAN)
            _lp = [" ".join(x.split()) for x in pdf_text(r.content).split("\f") if x.strip()] if r.status_code == 200 else []
            ok("⭐ D43 the leaflet PDF over HTTP: one page, named for the month, not cached, the demo bar, the four figures, the brief",
               r.status_code == 200 and r.headers.get("content-type", "").startswith("application/pdf")
               and r.headers.get("content-disposition") == 'attachment; filename="leaflet-2026-10.pdf"'
               and r.headers.get("cache-control") == "no-store" and len(_lp) == 1
               and "demo fictional scheme" in " ".join(_lp[0].split()).lower() and all(f in _lp[0] for f in _figs) and "Key points" in _lp[0],     # NARROWED 14 Oct 2026; and 15 Oct (the wording)
               f"{r.status_code} {len(_lp)}")
            ok("⭐ D43 the newsletter PDF over HTTP is four pages, the brand band on each",
               len(_bp) == 4 and all("Lorry movements bulletin · October 2026" in x for x in _bp), f"{len(_bp)}")
            r = c.get("/bulletin/leaflet/?month=10", headers=PLAN)
            ok("⭐ the leaflet page over HTTP with a code: HTML, not cached, the band, the figures, a link to its PDF",
               r.status_code == 200 and r.headers.get("content-type", "").startswith("text/html")
               and r.headers.get("cache-control") == "no-store" and '<header class="band">' in r.text
               and all(f in r.text for f in _figs) and 'href="/api/bulletin/leaflet.pdf?month=10"' in r.text, f"{r.status_code}")
            ok("...a team code is staff too: the page and the PDF open for the North Team code",
               c.get("/bulletin/?month=10", headers=NORTH).status_code == 200
               and c.get("/api/bulletin/month.pdf?month=10", headers=NORTH).status_code == 200)
            ok("🔴 ...a month outside the horizon is refused — 422", c.get("/api/bulletin/month.pdf?month=61", headers=PLAN).status_code == 422
               and c.get("/bulletin/?month=0", headers=PLAN).status_code == 422)
        finally:
            export.urllib.request.urlopen = _real_open
    finally:
        db.execute("DELETE FROM route_geometry WHERE tenant_id = ? AND route_id = ? AND vehicle_profile = ?",
                   ("default", "R001", "Artic Tipper (44t)"))

    # ------------------------------------------- 3c. the carbon target (11 Oct 2026, before the freeze)
    # The demo network is unbaked again here, so the figures are material + waste disposal. The one
    # decision is phase 4's: the surplus cut to the soil-recycling hub instead of the landfill (RE-PINNED 17 Oct 2026: R019's
    # tonnes × 18.53978 kg/t, computed above as _cut19).
    _reuse = [{"driver": "disposal", "scope": {"route_id": "R019"}, "to_route": "R020"}]
    r = c.put("/api/carbon/target", json={"mode": "pct", "levers": _reuse})
    ok("🔴 the carbon target without a code — 401", r.status_code == 401, str(r.status_code))
    r = c.put("/api/carbon/target", headers=NORTH, json={"mode": "pct", "levers": _reuse})
    ok("🔴 a team code cannot set the carbon target — 403", r.status_code == 403, str(r.status_code))
    r = c.get("/api/carbon/target", headers=PLAN)
    # REVERSED 17 Oct 2026 (demo redesign R2, D105): the demo SHIPS a % target — the disposal decision and recycled aggregate
    # in "Large aggregate / ballast" — set by the server's own sums before any bake; status reads it, with the decisions on
    # target. The PUT below replaces it (and the DELETE at the end clears it), as before.
    _sh = r.json() if r.status_code == 200 else {}
    ok("the demo ships its carbon target (D105): a % target from the disposal and recycled-aggregate decisions, on target with them",
       r.status_code == 200 and (_sh.get("target") or {}).get("mode") == "pct"
       and sorted(lv["driver"] for lv in (_sh["target"].get("levers") or [])) == ["disposal", "recycled"]
       and 0 < _sh["target"]["pct_below"] < 100 and (_sh.get("status") or {}).get("on_target_with_decisions") is True
       and _sh["status"].get("invalid") is None, r.text[:300])
    r = c.put("/api/carbon/target", headers=PLAN, json={"mode": "pct", "levers": _reuse, "ui_state": {"disposal": {"R019": {"to_route": "R020"}}},
                                                        "updated_by": "Smoke planner"})
    ct = r.json() if r.status_code == 200 else {}
    tg, st = ct.get("target") or {}, ct.get("status") or {}
    ok("⭐ a planner sets a % target over HTTP: stored with who set it, the decisions in words",
       r.status_code == 200 and tg.get("mode") == "pct" and tg.get("updated_by") == "Smoke planner" and tg.get("levers") == _reuse
       and tg.get("lever_words") and "Soil" in tg["lever_words"][0], r.text[:300])
    ok("⭐ ...the server's own figures read back consistently: the cut is phase 4's (R019's tonnes × 18.53978 kg/t), the % is that over the plan",
       st.get("scope") == "scheme" and abs((st.get("approved_t", 0) - st.get("with_decisions_t", 0)) - _cut19) < 0.002
       and tg.get("set_at_approved_t") == st.get("approved_t") and tg.get("cap_t") == st.get("with_decisions_t")
       # NARROWED 11 Oct 2026 (% held to 6 dp, R1): to 0.00001, was 0.06 (the 1 dp rounding)
       and abs(tg.get("pct_below", 0) - _cut19 * 100 / st.get("approved_t", 1)) < 0.00001, str(st)[:300])
    # NARROWED 11 Oct 2026 (R1): to 0.002, was 0.01; and the plan with the decisions IS the target the moment it is set
    ok("⭐ ...target = approved × (1 − %) to 0.002, the gaps are the differences, and with the decisions it is exactly on target",
       abs(st.get("target_t", 0) - st.get("approved_t", 0) * (1 - tg.get("pct_below", 0) / 100)) < 0.002
       and abs(st.get("gap_as_forecast_t", 0) - (st.get("approved_t", 0) - st.get("target_t", 0))) < 0.002
       and abs(st.get("gap_with_decisions_t", 9) - (st.get("with_decisions_t", 0) - st.get("target_t", 0))) < 0.002
       and abs(st.get("gap_with_decisions_t", 9)) <= 0.001 and st.get("target_t") == st.get("with_decisions_t"), str(st)[:300])
    ok("⭐ R2 over HTTP: the server's own verdicts — as forecast not on target (the cut over), with the decisions on target",
       st.get("on_target_as_forecast") is False and st.get("on_target_with_decisions") is True, str(st)[:300])
    # the one Pending demo line is R013's soil to a compound: soils have no primary factor and the route is
    # unbaked, so it adds 0.0 all-in — counted as 0, not left out
    ok("...the pending line adds 0.0 (soil to a compound, unrouted), transport 0 (nothing routed), months listed, nothing invalid",
       st.get("pending_t") == 0.0 and st.get("approved_transport_t") == 0.0 and st.get("invalid") is None
       and len(st.get("by_month") or []) >= 4
       and abs(sum(m["approved_t"] for m in st["by_month"]) - st.get("approved_t", 0)) < 0.02, str(st)[:300])
    r = c.get("/api/carbon/target", headers=NORTH)
    sn = (r.json() or {}).get("status") or {} if r.status_code == 200 else {}
    # REVERSED 11 Oct 2026 (the review: the scheme's % applied to North's lines, which the decision does not touch, read
    # "253.2 over the target" to every show visitor while the scheme was on target). It asserted North's target_t =
    # its approved plan × (1 − %). Now a team code gets no share of the scheme's target, and so no verdict.
    # REVERSED AGAIN 13 Oct 2026 (the user, 12 Oct: each team gets its own share): the share is the scheme's target split by
    # each team's plan WITH the decisions, so North — whose lines the decision does not touch — gets its own plan as its share
    # and reads ON target (the 11 Oct "253.2 over" came from applying the scheme's % evenly, which this rule does not do).
    # It asserted target_t, the gap and both verdicts null.
    ok("⭐ the North Team code reads it: scope team, its own lines; its share = its own plan (the cut is Central's), on target",
       r.status_code == 200 and sn.get("scope") == "team" and 0 < sn.get("approved_t", 0) < st.get("approved_t", 0)
       and sn.get("with_decisions_t") == sn.get("approved_t")
       and abs((sn.get("target_t") or 0) - sn["approved_t"] * st["target_t"] / st["with_decisions_t"]) <= 0.002
       and sn.get("on_target_with_decisions") is True and 0 < (sn.get("share_pct") or 0) < 100, str(sn)[:300])
    tn = (r.json() or {}).get("target") or {} if r.status_code == 200 else {}
    ok("🔴 R3 over HTTP: North's GET carries no stored decisions and no page state, and no words for the route-scoped decision",
       r.status_code == 200 and tn.get("mode") == "pct" and "levers" not in tn and "ui_state" not in tn
       and tn.get("lever_words") == [] and "R019" not in r.text and "R020" not in r.text, str(tn)[:300])
    r = c.put("/api/carbon/target", headers=PLAN, json={"mode": "cap", "updated_by": "Smoke planner"})
    cc = r.json() if r.status_code == 200 else {}
    ok("⭐ switch to a fixed cap: every number kept, the target is the cap",
       r.status_code == 200 and cc["target"]["mode"] == "cap" and cc["target"]["cap_t"] == tg.get("cap_t")
       and cc["target"]["pct_below"] == tg.get("pct_below") and cc["status"]["target_t"] == tg.get("cap_t"), r.text[:300])
    r = c.get("/api/carbon/target", headers=NORTH)
    # REVERSED 13 Oct 2026 (each team its share): it asserted target_t null for a team under a fixed cap. The cap is split by
    # the same rule; North's share is the cap × its plan ÷ the scheme's plan with the decisions. Still no stored decisions.
    _cn = r.json()["status"] if r.status_code == 200 else {}
    ok("🔴 ...and a team code under a fixed cap sees ITS share of the cap (the same rule), still without the stored decisions",
       r.status_code == 200 and _cn.get("target_t") is not None and 0 < _cn["target_t"] < tg.get("cap_t")
       and _cn.get("share_pct") == sn.get("share_pct") and "levers" not in r.json()["target"], r.text[:200])
    r = c.put("/api/carbon/target", headers=PLAN, json={"mode": "pct", "levers": []})
    ok("🔴 no decisions — 400 with the sentence", r.status_code == 400 and "Choose at least one decision" in r.text, r.text[:160])
    r = c.delete("/api/carbon/target", headers=NORTH)
    ok("🔴 a team code cannot clear it — 403", r.status_code == 403, str(r.status_code))
    r = c.delete("/api/carbon/target", headers=PLAN)
    ok("⭐ a planner clears it: target null, and GET agrees", r.status_code == 200 and r.json()["target"] is None
       and c.get("/api/carbon/target", headers=PLAN).json() == {"target": None, "status": None}, r.text[:160])
    # ---- 11 Oct 2026: Gate arrivals (the staff Dashboard panel) ----------------------------
    # RE-PINNED 17 Oct 2026 (demo redesign R2, §3a): the block moves from Sep 2026 to Apr 2027 (month 16), the first month
    # with all three teams on site; the ranges are set round the regenerated figures (Melton Compound 31.9 a day, the
    # busiest gate; Langton Compound 13.1; Harborough Compound's A6 entrance 3.5), the gaps from the gates' typed hours
    # less the school-run bands (510 open minutes at a compound)
    r = c.get("/api/gates/arrivals?month=16")
    ok("🔴 gate arrivals without a code — 401", r.status_code == 401, str(r.status_code))
    r = c.get("/api/gates/arrivals?month=16", headers=PLAN)
    ga = r.json() if r.status_code == 200 else {}
    ok("⭐ the planner reads month 16 of the demo: scope scheme, 22 days, 10 h, Apr 2027",
       r.status_code == 200 and ga.get("scope") == "scheme" and ga.get("working_days") == 22.0
       and ga.get("shift_hours") == 10.0 and ga.get("label") == "Apr 2027", r.text[:200])
    byname = {(g["location_name"], g["gate_name"]): g for g in ga.get("gates", [])}
    mb = byname.get(("Melton Compound", "Compound gate")) or {}
    hb = byname.get(("Harborough Compound", "A6 entrance (in)")) or {}
    lg = byname.get(("Langton Compound", "Compound gate")) or {}
    ok("⭐ Melton Compound's gate, the month's busiest: 25–40 laden arrivals a day, one every 12–21 min, check-in 3, NOT red (honest: averages alone do not queue)",
       max((g["arrivals_per_day"] for g in ga.get("gates", [])), default=0) == mb.get("arrivals_per_day")
       and 25 <= mb.get("arrivals_per_day", 0) <= 40 and 12 <= mb.get("gap_minutes", 0) <= 21 and mb.get("gate_minutes") == 3.0
       and mb.get("flag") is None and "if arrivals are spread evenly" in mb.get("reason", ""), str(mb)[:300])
    ok("⭐ Langton Compound's gate: 10–17 a day (laden plus the muck-away lorries coming to load, role both), check-in 3, not red",
       10 <= lg.get("arrivals_per_day", 0) <= 17 and lg.get("role") == "both" and lg.get("empty_per_day", 0) > 0
       and lg.get("gate_minutes") == 3.0 and lg.get("flag") is None, str(lg)[:300])
    ok("...Harborough Compound's entrance reads by its new name (the A6, was \"Station Road\"), check-in 5, 2–6 a day",
       2 <= hb.get("arrivals_per_day", 0) <= 6 and hb.get("gate_minutes") == 5.0 and hb.get("flag") is None
       and not any(g["gate_name"].startswith("Station Road") for g in ga.get("gates", [])), str(hb)[:300])
    ok("the quarry gates carry the origin-share note; the compounds do not",
       # NARROWED 11 Oct 2026 (review): the all() over a filter was vacuous with no quarry listed; at least one must be
       any(g["loc_type"] in ("Quarry", "Railhead") for g in ga.get("gates", []))
       and all(g.get("origin_share_note") for g in ga.get("gates", []) if g["loc_type"] in ("Quarry", "Railhead"))
       and not hb.get("origin_share_note") and not lg.get("origin_share_note"))
    ok("the demo's exit gates are not listed: every quarry route is counted at the Weighbridge (in) by fall-through",
       # RE-PINNED 17 Oct 2026 (R2): Northfield's exit is "A6 exit" since R1's rename (it said "A46 exit", on the A6)
       not any(g["gate_name"] in ("Main road exit", "A6 exit", "A46 exit", "Mill Lane exit") for g in ga.get("gates", []))
       and any(g["gate_name"] == "Weighbridge (in)" and g.get("fallthrough_routes") for g in ga.get("gates", [])))
    ok("🔴 nothing in month 16 is red, and the response says so in words rather than inventing a peak",
       all(g.get("flag") is None for g in ga.get("gates", [])) and any("No hourly peak is typed" in n for n in ga.get("notes", [])))
    # 17 Oct 2026 (the user: gate shut times through the day; "default time for gates should be 0800-1700"). The four-year
    # demo (R1) types every gate's hours (compounds 07:30–18:00) and the school-run bands 08:00–09:00 and 15:00–16:00, so
    # the gates rest on their OWN hours less the bands; the 08:00–17:00 default shows on a gate whose hours are cleared.
    ok("⭐ typed hours are the gates' own: Harborough's A6 entrance 07:30–18:00 less the two bands = 510 open minutes",
       hb.get("open_basis") == "gate" and hb.get("open_hours") == "07:30–18:00" and hb.get("open_minutes") == 510.0
       and hb.get("peak_minutes") == 120.0 and hb.get("shut_windows") == [] and ga.get("default_hours") == "08:00–17:00"
       and any(n.startswith("Averages only") and "less the no-delivery bands" in n for n in ga.get("notes", [])), str(hb)[:300])
    _hbid = hb.get("gate_id")
    _gall = (c.get("/api/gates", headers=PLAN).json() or {}).get("gates", [])
    _hbrow = next((g for g in _gall if g["id"] == _hbid), {})
    _sp = next((g for g in _gall if g["name"] == "Stockpile gate" and g.get("open_from")), {})
    _sploc = {l["id"]: l["name"] for l in DEMO["tables"]["locations"]}.get(_sp.get("location_id"))
    rp = c.patch(f"/api/admin/gates/{_sp.get('id')}", params=TOKEN, json={"open_from": "", "open_to": "", "open_days": ""})
    _spg = {g["gate_id"]: g for g in c.get("/api/gates/arrivals?month=16", headers=PLAN).json().get("gates", [])}.get(_sp.get("id")) or {}
    ok("⭐ a gate with its hours cleared is open the default 08:00–17:00, less the bands: 540 − 120 = 420 open minutes, said so",
       rp.status_code == 200 and rp.json().get("open_from") is None
       and (not _spg or (_spg.get("open_basis") == "default" and _spg.get("open_hours") == "08:00–17:00" and _spg.get("open_minutes") == 420.0)),
       str(_spg)[:300])
    _form = {k: _hbrow.get(k) for k in ("name", "direction", "lat", "lon", "safety_minutes", "internal_travel_minutes",
                                         "is_default", "active", "note", "open_from", "open_to", "open_days", "access_zone_id")}
    # FIXED 17 Oct 2026 (the review): the PATCH replies are `rp`, so `r` stays the planner's Gate arrivals read that the
    # banned-words check below is about; the read with a shut time stored is scanned too
    rp = c.patch(f"/api/admin/gates/{_hbid}", params=TOKEN, json={**_form, "open_from": _form["open_from"] or "",
                                                                  "open_to": _form["open_to"] or "", "open_days": _form["open_days"] or "",
                                                                  "shut_windows": [{"from": "12:00", "to": "12:30"}]})
    ok("⭐ PATCH a shut time on Harborough's gate with EVERY field the form sends: stored and read back, and NOTHING to route "
       "again (the review: a field sent is not a field changed)",
       rp.status_code == 200 and rp.json().get("shut_windows") == [{"from": "12:00", "to": "12:30"}]
       and rp.json().get("rebake_needed") is False and rp.json().get("routes_affected") == [], rp.text[:300])
    _ga_s = c.get("/api/gates/arrivals?month=16", headers=PLAN)
    _hb2 = {g["gate_id"]: g for g in _ga_s.json().get("gates", [])}.get(_hbid) or {}
    ok("...and Gate arrivals takes it off: 510 − 30 = 480 open minutes, the shut time in words, and the note says shut times come off",
       _hb2.get("open_minutes") == 480.0 and _hb2.get("shut_words") == "12:00–12:30" and _hb2.get("shut_minutes") == 30.0
       and any("less its shut times" in n for n in _ga_s.json().get("notes", [])), str(_hb2)[:200])
    rp = c.patch(f"/api/admin/gates/{_hbid}", params=TOKEN, json={**_form, "lat": _form["lat"] + 0.0001})
    ok("...while moving the gate does ask to route again, naming the routes", rp.status_code == 200
       and rp.json().get("rebake_needed") is True and len(rp.json().get("routes_affected") or []) > 0, rp.text[:200])
    c.patch(f"/api/admin/gates/{_hbid}", params=TOKEN, json={"lat": _form["lat"]})
    rp = c.patch(f"/api/admin/gates/{_hbid}", params=TOKEN, json={"shut_windows": [{"from": "18:00", "to": "19:00"}]})
    ok("🔴 a shut time outside the gate's hours is refused in words (and the stored one stays)",
       rp.status_code == 200 and rp.json().get("error") == "shut time 18:00–19:00 is outside the open hours 07:30–18:00"
       and ({g["gate_id"]: g for g in c.get("/api/gates/arrivals?month=16", headers=PLAN).json().get("gates", [])}.get(_hbid) or {}).get("shut_words") == "12:00–12:30",
       rp.text[:200])
    rp = c.patch(f"/api/admin/gates/{_hbid}", json={"shut_windows": []})
    ok("🔴 the shut times are an admin write: no token, no change (the shut time is still there)", rp.status_code in (401, 403)
       and ({g["gate_id"]: g for g in c.get("/api/gates/arrivals?month=16", headers=PLAN).json().get("gates", [])}.get(_hbid) or {}).get("shut_words") == "12:00–12:30",
       str(rp.status_code))
    # a Config save from the page that would leave a gate's shut time outside the default gate hours is refused, naming it
    c.patch(f"/api/admin/gates/{_sp.get('id')}", params=TOKEN, json={"shut_windows": [{"from": "09:30", "to": "10:00"}]})
    _cfgdoc = c.get("/api/config/factors", headers=PLAN).json().get("doc") or {}
    _cfgbad = {**_cfgdoc, "planning": {**(_cfgdoc.get("planning") or {}), "default_gate_hours": {"from": "13:00", "to": "17:00"}}}
    rp = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _cfgbad})
    ok("⭐ Config → Planning: default gate hours that would put a gate's shut time outside them are refused (400), naming the gate",
       rp.status_code == 400 and f"{_sploc}, Stockpile gate: shut time 09:30–10:00 is outside the open hours 13:00–17:00 (the default)" in rp.text,
       rp.text[:300])
    c.patch(f"/api/admin/gates/{_sp.get('id')}", params=TOKEN, json={"shut_windows": [], "open_from": _sp.get("open_from"),
                                                                     "open_to": _sp.get("open_to"), "open_days": _sp.get("open_days") or ""})
    rp = c.patch(f"/api/admin/gates/{_hbid}", params=TOKEN, json={"shut_windows": []})
    ok("...an empty list clears them (and the cleared gate has its hours back)", rp.status_code == 200 and rp.json().get("shut_windows") == []
       and next((g for g in (c.get("/api/gates", headers=PLAN).json() or {}).get("gates", []) if g["id"] == _sp.get("id")), {}).get("open_from") == _sp.get("open_from"),
       rp.text[:200])
    rp = c.post("/api/admin/gates", params=TOKEN, json={"location_id": hb.get("location_id"), "name": "Smoke shut gate",
                                                        "lat": _form["lat"] + 0.001, "lon": _form["lon"], "direction": "access",
                                                        "shut_windows": "12:00-12:30,08:30-09:15"})
    ok("a gate POSTed with its shut times as the stored text (not only a list) is created with them, start-sorted",
       rp.status_code == 200 and rp.json().get("shut_windows") == [{"from": "08:30", "to": "09:15"}, {"from": "12:00", "to": "12:30"}],
       rp.text[:200])
    if rp.status_code == 200 and rp.json().get("id"):
        c.delete(f"/api/admin/gates/{rp.json()['id']}", params=TOKEN)
    ok("🔴 the banned words appear nowhere in the planner's response (nor in the read with a shut time stored)",
       not re.search(r"real-time|live tracking|telematics|optimi[sz]|predict|\bAI\b|\btrucks?\b|booking|peak hour", r.text, re.I)
       and r.request.url.path == "/api/gates/arrivals"
       and not re.search(r"real-time|live tracking|telematics|optimi[sz]|predict|\bAI\b|\btrucks?\b|booking|peak hour", _ga_s.text, re.I))
    # RE-PINNED 17 Oct 2026 (R2, §3a, §1.9): the Pending line (R013, soil) sits in Apr 2027, the first earthworks-season
    # month at least two months after the as-of date — month 16 itself
    ok("month 16 carries the Pending line (R013) as a note, never in a figure",
       any(n.startswith("1 pending line in Apr 2027") for n in ga.get("notes", []))
       and [f["month_index"] for f in DEMO["tables"]["forecasts"] if f["status"] == "Pending"] == [16])
    r = c.get("/api/gates/arrivals?month=16", headers=NORTH)
    gn = r.json() if r.status_code == 200 else {}
    ok("🔴 the North Team code: scope team, the SAME totals and team names as the planner, per_team holding only North Team",
       r.status_code == 200 and gn.get("scope") == "team"
       and [(g["gate_id"], g["arrivals_per_day"], g["teams"]) for g in gn.get("gates", [])] == [(g["gate_id"], g["arrivals_per_day"], g["teams"]) for g in ga.get("gates", [])]
       and all(all(t["team"] == "North Team" for t in g.get("per_team", [])) for g in gn.get("gates", []))
       and any(g["per_team"] == [] and "South Team" in g["teams"] for g in gn.get("gates", [])), r.text[:300])
    # the other teams' per-team figures from the planner's view, minus any value that is ALSO a gate total or North's own
    # figure (those may legitimately appear); none of what is left may appear anywhere in North's per_team rows
    _shown = {t["arrivals_per_day"] for g in gn.get("gates", []) for t in g.get("per_team", [])}
    _others = {t["arrivals_per_day"] for g in ga.get("gates", []) for t in g.get("per_team", []) if t["team"] != "North Team"}
    ok("🔴 ...and no other team's per-team figure reaches it", bool(_others) and not (_others & _shown), str(_others & _shown))
    r = c.get("/api/gates/arrivals?month=0", headers=PLAN)
    ok("month 0 is refused — 400", r.status_code == 400, str(r.status_code))
    r = c.get("/api/gates/arrivals", headers=PLAN)
    ok("no month -> today's month, clamped to the horizon, 200", r.status_code == 200 and 1 <= r.json().get("month", 0) <= 60, str(r.status_code))
    # RE-PINNED 17 Oct 2026 (demo redesign R2): month 16 (Apr 2027), as the gate-arrivals block above — Harborough Compound has
    # no line in Sep 2026 any more; in Apr 2027 it does, and the 160 m³ example pour there queues (the Welham Lane
    # underpass base pour of the design, §2)
    # ---- 16 Oct 2026 (S3): the busiest hour on the panel, and the simulation's own read ----------------------------
    ok("⭐ S3: every gate in month 16 carries its busiest hour from the simulation's day, and the reply says where it comes from",
       ga.get("gates") and all((g.get("busiest_hour") or {}).get("lorries", 0) > 0 and re.fullmatch(r"\d\d:00", g["busiest_hour"]["hour"])
                               for g in ga["gates"])
       and "Busiest hour" in (ga.get("busiest_note") or "") and "No pour" in ga["busiest_note"], str(ga.get("busiest_note"))[:200])
    r = c.get("/api/sim/day?month=16")
    ok("🔴 S3: the simulation without a code — 401 (staff only, never on the public map)", r.status_code == 401, str(r.status_code))
    r = c.get("/api/sim/day?month=16", headers=PLAN)
    sd = r.json() if r.status_code == 200 else {}
    ok("⭐ S3: the planner's simulated day of month 16: whole loads, every arrival listed once, the table equals the list",
       r.status_code == 200 and sd.get("loads", 0) > 0 and sd.get("gate_arrivals") == len(sd.get("arrivals", []))
       and sd.get("checks", {}).get("table_equals_arrivals") is True and sd.get("checks", {}).get("hours_add_up") is True
       and sd.get("pour") is None and (sd.get("checks", {}).get("largest_panel_difference") or 0) < 1.0, r.text[:300])
    hb_sd = next((g for g in sd.get("gates", []) if g["location_name"] == "Harborough Compound"), {})
    r = c.get(f"/api/sim/day?month=16&pour_gate={hb_sd.get('gate_id', 'G007')}&pour_start=10:00&pour_m3=160&pour_rate=60", headers=PLAN)
    sp = r.json() if r.status_code == 200 else {}
    hp = next((g for g in sp.get("gates", []) if g["location_name"] == "Harborough Compound"), {})
    ok("⭐ S3: the example pour at Harborough Compound (160 m³, the app's 7.5 m³ mixer, 60 m³ an hour): 22 mixers, a queue",
       r.status_code == 200 and (sp.get("pour") or {}).get("loads") == 22 and hp.get("pour") == 22
       and hp.get("arrivals") == hb_sd.get("arrivals", -99) + 22 and hp.get("queue_start") is not None, r.text[:300])
    # 17 Oct 2026 (found on live): a pour goes to a working site only, and defaults to the one with the most laden lorries
    # RE-PINNED 17 Oct 2026 (demo redesign R3, merging PR #89): on the four-year demo the simulated day is Apr 2027 (month
    # 16), where the compound with the most laden lorries is Melton Compound, not Harborough (Sep 2026 on the old demo):
    # the default is pinned to that rule, read from the Gate arrivals panel of the same month
    _pgn = [x["location_name"] for x in sd.get("pour_gates", [])]
    _most = max((g for g in ga.get("gates", []) if g.get("loc_type") == "Compound"),
                key=lambda g: g.get("laden_per_day") or 0, default={}).get("location_name")
    ok("⭐ S7 fix: on the demo package the pour is offered at the compounds only (no landfill, soil recycling, batching plant, "
       "stockpile, quarry or terminal), and defaults to the compound with the most laden lorries (Melton Compound in Apr 2027)",
       bool(_pgn) and all(n.endswith("Compound") for n in _pgn) and _most == "Melton Compound"
       and next((x["location_name"] for x in sd["pour_gates"] if x["gate_id"] == sd.get("pour_default")), None) == _most,
       str(_pgn) + " " + str(_most))
    _lf = [g for g in c.get("/api/gates").json().get("gates", []) if g.get("location_id") == "L012"]
    r = c.get(f"/api/sim/day?month=16&pour_gate={_lf[0]['id'] if _lf else 'G015'}&pour_start=10:00&pour_m3=160&pour_rate=60", headers=PLAN)
    ok("🔴 S7 fix: a pour at the inert landfill is refused in words — 400",
       bool(_lf) and r.status_code == 400 and "A pour goes to a working site" in r.text and "is a landfill" in r.text, r.text[:200])
    r = c.get("/api/sim/day?month=16&hours=G001=18:00-07:00", headers=PLAN)
    ok("🔴 S3: a window past midnight is refused in words — 400 (D93)", r.status_code == 400 and "past midnight" in r.text, r.text[:200])
    r = c.get("/api/sim/day?month=16&pour_gate=G007&pour_start=23:00&pour_m3=160&pour_rate=60", headers=PLAN)
    ok("🔴 S3: a pour past midnight is refused in words — 400", r.status_code == 400 and "past midnight" in r.text, r.text[:200])
    r = c.get("/api/sim/day?month=16&legs=1", headers=PLAN)
    sl = r.json() if r.status_code == 200 else {}
    # the demo package carries no route lines (they are worked out after an import), so here every route is "not drawn"
    # and its lorries still count at the gates; test_sim.py checks the lines on the live demo's
    ok("S3: legs=1 answers 200; routes with no line worked out are named in not_drawn and still counted at the gates",
       r.status_code == 200 and isinstance(sl.get("legs"), dict) and isinstance(sl.get("not_drawn"), list)
       and set(sl["not_drawn"]) == {l["route_id"] for l in sl.get("lines", []) if not (sl["legs"].get(l.get("laden_leg") or "") or {}).get("drawn")}
       and sl.get("gate_arrivals") == sd.get("gate_arrivals"), r.text[:200])
    r = c.get("/api/sim/day?month=16", headers=NORTH)
    sn = r.json() if r.status_code == 200 else {}
    ok("🔴 S3: the North Team code: scope team, the same gate figures, other teams' lines by route only",
       r.status_code == 200 and sn.get("scope") == "team"
       and [(g["gate_id"], g["arrivals"]) for g in sn.get("gates", [])] == [(g["gate_id"], g["arrivals"]) for g in sd.get("gates", [])]
       and all("loads" not in l for l in sn.get("lines", []) if l.get("team") != "North Team")
       and any("loads" in l for l in sn.get("lines", [])), r.text[:300])
    # the vehicle KEYS (ids such as "Truck mixer 8 m³ (32t)") travel in the reply for the drawing; what a person reads is
    # each *_label field, so the scan drops the keys and keeps the labels
    _sd_words = json.dumps({k: v for k, v in sd.items() if k not in ("legs",)}, ensure_ascii=False)
    _sd_words = re.sub(r'"(vehicle|vehicle_drawn|laden_leg|empty_leg|leg|approach_leg)": "[^"]*"', "", _sd_words)
    ok("🔴 S3: no banned word and no team id in the simulation's reply (vehicle keys aside; their labels kept)",
       not re.search(r"real-time|live tracking|telematics|optimi[sz]|predict|\bAI\b|\btrucks?\b|booking|peak hour", _sd_words, re.I)
       and not re.search(r"\bIPT ?[1-6]\b", json.dumps(sd)) and "Concrete mixer" in _sd_words, _sd_words[:200])
    # the demo figures for the session report (not asserted beyond the ranges above)
    print("  gate arrivals, month 16 (demo package):")
    for g in ga.get("gates", []):
        print(f"    {g['location_name']} · {g['gate_name']}: {g['arrivals_per_day']}/day "
              f"(laden {g['laden_per_day']}, empty {g['empty_per_day']}), gap {g['gap_minutes']} min, "
              f"check-in {g['gate_minutes']}, flag {g['flag']}, role {g['role']}, teams {g['teams']}")

    # 14 Oct 2026 (bundle G, W7): the CLP monitoring read, over the real app and the demo package (which carries no route
    # geometry, so the councils chart has no clipped legs and the carbon is material + disposal only — honest gaps, asserted)
    r = c.get("/api/clp/monitoring?month=10")
    ok("🔴 CLP monitoring without a code — 401", r.status_code == 401, str(r.status_code))
    r = c.get("/api/clp/monitoring?month=10", headers=PLAN)
    cm = r.json() if r.status_code == 200 else {}
    ch = cm.get("charts") or {}
    ok("⭐ the planner reads CLP monitoring for Oct 2026 of the demo: seven charts in order, each set or empty with a sentence",
       r.status_code == 200 and cm.get("label") == "Oct 2026" and cm.get("order") == ["route_caps", "site_caps", "vehicle_types", "euro_class",
                                                                                   "councils", "carbon_rate", "trips_avoided"]
       and all((c_.get("state") == "set") or (c_.get("state") == "empty" and c_.get("missing")) for c_ in ch.values()), r.text[:300])
    ok("the demo's vehicle chart is drawn (approved loads exist), its shares from the server",
       (ch.get("vehicle_types") or {}).get("state") == "set" and (ch.get("vehicle_types") or {}).get("series"), str(ch.get("vehicle_types"))[:200])
    ok("🔴 no team id and no team slot in the reply; the demo notice rides on it",
       "IPT" not in r.text and cm.get("demo_notice"), str(cm.get("demo_notice")))
    r = c.get("/api/clp/monitoring?month=10", headers=NORTH)
    cn = r.json() if r.status_code == 200 else {}
    ok("🔴 the North Team code: scope team; its vehicle chart holds only its own loads (no more than the planner's)",
       r.status_code == 200 and cn.get("scope") == "team"
       and sum(((cn.get("charts") or {}).get("vehicle_types") or {}).get("totals") or [0])
       <= sum((ch.get("vehicle_types") or {}).get("totals") or [0]), r.text[:200])
    r = c.get("/api/clp/monitoring?month=0", headers=PLAN)
    ok("CLP month 0 is refused — 400", r.status_code == 400, str(r.status_code))
    # 14 Oct 2026 (bundle G, W10): the CLP pack over HTTP
    r = c.get("/api/clp/pack.pdf?month=10", headers=PLAN)
    ok("⭐ the CLP pack PDF for Oct 2026: 200, a PDF, the month's file name",
       r.status_code == 200 and r.headers.get("content-type", "").startswith("application/pdf") and r.content[:4] == b"%PDF"
       and 'filename="wayscope-clp-pack-2026-10.pdf"' in r.headers.get("content-disposition", ""), str(r.status_code))
    r = c.get("/api/clp/pack.xlsx?month=10", headers=PLAN)
    ok("⭐ the CLP pack XLSX for Oct 2026: 200, a workbook, the month's file name",
       r.status_code == 200 and r.content[:2] == b"PK"
       and 'filename="wayscope-clp-pack-2026-10.xlsx"' in r.headers.get("content-disposition", ""), str(r.status_code))
    r = c.get("/api/clp/pack.pdf?month=10")
    ok("🔴 the CLP pack without a code — 401", r.status_code == 401, str(r.status_code))
    print("  CLP monitoring, Oct 2026 (demo package):")
    for k in cm.get("order") or []:
        c_ = ch.get(k) or {}
        print(f"    {k}: {c_.get('state')} · {c_.get('headline') or c_.get('missing')}")

    # the Submit forecast preview's all-in carbon, over HTTP: 1,000 t of cut to the landfill = 1000 × 19.55376
    r = c.post("/api/costing/preview", headers=PLAN, json={"route_id": "R019", "vehicle_type": "Artic Tipper (44t)",
                                                          "material_type": "Earthworks / soil", "unit": "t",
                                                          "cells": [{"month_index": 10, "quantity": 1000}]})
    pc = (r.json() or {}).get("carbon") or {} if r.status_code == 200 else {}
    ok("⭐ the preview carries its carbon: 1,000 t to the landfill = 19.554 tCO2e disposal, transport unrouted (never estimated)",
       r.status_code == 200 and pc.get("disposal_t") == 19.554 and pc.get("material_t") is None and pc.get("transport_t") is None
       and pc.get("unrouted") is True and pc.get("total_t") == 19.554 and "landfill" in (pc.get("basis") or ""), str(pc))
    r = c.post("/api/costing/preview", headers=PLAN, json={"route_id": "R002", "vehicle_type": "Artic Tipper (44t)",
                                                          "material_type": "Small aggregate", "unit": "t",
                                                          "cells": [{"month_index": 10, "quantity": 2900}]})
    pc = (r.json() or {}).get("carbon") or {} if r.status_code == 200 else {}
    ok("...and a delivery: 2,900 t of aggregate = 22.629 tCO2e to make (2900 × 7.80307), transport not counted until routed",
       r.status_code == 200 and pc.get("material_t") == 22.629 and pc.get("total_t") == 22.629 and pc.get("unrouted") is True
       and "not counted until the route has a distance" in (pc.get("basis") or ""), str(pc))

    r = c.get("/api/lookahead", headers=PLAN)
    la = r.json() if r.status_code == 200 else {}
    ok("the Look-ahead page reads over HTTP", r.status_code == 200 and "commit" in la, r.text[:200])
    flags = (la.get("clashes") or {}).get("flags") or []
    ok("...its flag codes are the documented ones (the app names them for people)",
       # 13 Oct 2026: LIMIT (lorry limits, limits.py) joins the documented codes
       all(f.get("code") in ("SHORTAGE", "DAYS_NE_WEEK", "STOCKPILE_OVER", "ROUTE_CAP", "IPT_SHARE", "TARK_TEE",
                             "LIMIT", "UNBAKED", "PARENT_CHANGED") for f in flags), str({f.get("code") for f in flags}))
    ok("...the rail reports the stored lorry-limits check as a source (13 Oct 2026)",
       "limits" in ((la.get("clashes") or {}).get("sources") or {}), str(((la.get("clashes") or {}).get("sources") or {}).keys()))
    ok("...and the unbaked demo network is flagged, not estimated",
       any(f.get("code") == "UNBAKED" for f in flags) or (la.get("clashes") or {}).get("by_code", {}).get("UNBAKED"),
       str((la.get("clashes") or {}).get("by_code")))

    r = c.get("/api/forecast-weeks/export", params={"format": "xlsx"}, headers=PLAN)
    ok("⭐ the XLSX export downloads (real openpyxl)", r.status_code == 200
       and r.headers.get("content-type", "").startswith("application/vnd.openxmlformats")
       and r.content[:2] == b"PK" and "attachment" in r.headers.get("content-disposition", ""),
       f"{r.status_code} {r.headers.get('content-type')}")
    xlsx_bytes = r.content
    r = c.get("/api/forecast-weeks/export", params={"format": "pdf"}, headers=PLAN)
    ok("⭐ the PDF export downloads (real reportlab)", r.status_code == 200
       and r.headers.get("content-type") == "application/pdf" and r.content[:5] == b"%PDF-",
       f"{r.status_code} {r.headers.get('content-type')} {r.content[:60]!r}")
    pdf_bytes = r.content

    # ---- 14 Oct 2026 (D9, the mobile shell): the Day sheet, one weekday of the open week as a one-page PDF -------------
    from portable import pdf_text as _pt
    _cw = la.get("commit_week") or {}
    _wds = [d["day_date"] for d in ((la.get("commit") or {}).get("lines") or [{}])[0].get("days", [])
            if __import__("datetime").date.fromisoformat(d["day_date"]).weekday() <= 4]
    _day = _wds[0] if _wds else _cw.get("from")
    r = c.get("/api/forecast-weeks/day-sheet", params={"day": _day})
    ok("🔴 DAY SHEET: without a code — 401", r.status_code == 401, str(r.status_code))
    r = c.get("/api/forecast-weeks/day-sheet", params={"day": _day}, headers=PLAN)
    _dst = " ".join(_pt(r.content).split()) if r.status_code == 200 else ""
    # NARROWED 14 Oct 2026 (civils materials): the concrete pour adds three lines in m³ to the demo's Monday, which brings
    # the PLAN column; 22 lines run on to a second page by the sheet's own rule ("long days run on to a second page
    # rather than shrink the type below 7 pt", export.build_day_sheet_pdf) — A4 portrait, two pages
    # RE-PINNED 17 Oct 2026 (demo redesign R2, D104): the harness pins its date (Thu 26 Nov 2026), so the day is Mon 23 Nov,
    # whose 20 lines (mixer lines among them, so the PLAN column) fit ONE page by the same rule — read from the page itself
    _nl = sum(1 for l in ((la.get("commit") or {}).get("lines") or [])
              if any(d["day_date"] == _day and (d.get("planned_qty") or 0) > 0 for d in l.get("days", [])))
    ok("⭐ DAY SHEET: the planner downloads A4 portrait (real reportlab), named after the day — one page for the pinned Monday's 20 lines",
       _day == "2026-11-23" and _nl == 20 and f"{_nl} lines" in _dst
       and r.status_code == 200 and r.headers.get("content-type") == "application/pdf" and r.content[:5] == b"%PDF-"
       and r.content.count(b"/Type /Page\n") == 1 and f'filename="day-sheet-{_day}.pdf"' in r.headers.get("content-disposition", "")
       and b"/MediaBox [ 0 0 595." in r.content and "day sheet" in _dst and "not a delivery note" in _dst,
       f"{r.status_code} pages={r.content.count(b'/Type /Page\n')} mb={r.content[r.content.find(b'/MediaBox'):r.content.find(b'/MediaBox')+40]!r} {_dst[-400:]}")
    ok("...with the demo notice, team NAMES and no team id; NOT ROUTED counted, never the server's 'baked' working note",
       "DEMO Fictional scheme" in _dst and not re.search(r"\bIPT ?[1-6]\b", _dst)   # NARROWED 15 Oct 2026 (the user)
       and "NOT ROUTED on" in _dst and not re.search(r"\b(un)?baked?\b", _dst, re.I)
       and any(t in _dst for t in ("North Team", "Central Team", "South Team")), _dst[:300])
    r = c.get("/api/forecast-weeks/day-sheet", params={"day": _day}, headers=NORTH)
    _dsn = " ".join(_pt(r.content).split()) if r.status_code == 200 else ""
    ok("⭐ DAY SHEET: a team code gets ITS lines only (the page's own scope)",
       r.status_code == 200 and ("North Team" in _dsn or "Nothing is planned on this day" in _dsn)
       and "Central Team" not in _dsn and "South Team" not in _dsn, _dsn[:300])
    _sat = (__import__("datetime").date.fromisoformat(_cw["from"]) + __import__("datetime").timedelta(days=5)).isoformat() if _cw.get("from") else "2000-01-01"
    r1 = c.get("/api/forecast-weeks/day-sheet", params={"day": _sat}, headers=PLAN)
    r2 = c.get("/api/forecast-weeks/day-sheet", params={"day": "2000-01-03"}, headers=PLAN)
    r3 = c.get("/api/forecast-weeks/day-sheet", params={"day": "not-a-day"}, headers=PLAN)
    ok("🔴 DAY SHEET: a weekend, a day outside the week and a non-date are each a 400 with the reason",
       r1.status_code == 400 and r2.status_code == 400 and r3.status_code == 400 and "not a weekday of the week" in r1.text
       and "YYYY-MM-DD" in r3.text, f"{r1.status_code} {r2.status_code} {r3.status_code} {r1.text[:120]}")

    try:
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb[wb.sheetnames[0]]
        head = [cell.value for cell in ws[1]]
        teams_col = head.index("Team") if "Team" in head else None
        vals = [row[teams_col].value for row in ws.iter_rows(min_row=2)] if teams_col is not None else []
        ok("the XLSX header uses the tenant's words — 'Team' and £", "Team" in head and "£" in head and "IPT" not in head,
           str(head[:20]))
        ok("⭐ ...and its Team column carries team NAMES, never ids",
           vals and all(v in ("North Team", "Central Team", "South Team") for v in vals), str(sorted(set(map(str, vals)))[:5]))
        cl = wb["Clashes"] if "Clashes" in wb.sheetnames else None
        cl_txt = " ".join(str(c.value) for row in cl.iter_rows() for c in row if c.value) if cl else ""
        ok("⭐ the Clashes sheet shows flag NAMES and team names — no code with IPT in it, no team id",
           cl is not None and "UNBAKED" in cl_txt and not re.search(r"IPT_|\bIPT ?[1-6]\b", cl_txt)
           and [c.value for c in cl[1]][2] == "Team", cl_txt[:200])
        about = wb["About"] if "About" in wb.sheetnames else None
        about_txt = " ".join(str(c.value) for row in about.iter_rows() for c in row if c.value) if about else ""
        ok("the About sheet names the tenant", "Wolds Link" in about_txt, about_txt[:120])
    except Exception as e:
        ok("the XLSX opens in openpyxl", False, repr(e))

    try:
        with open(os.path.join(TMP, "week.pdf"), "wb") as f:
            f.write(pdf_bytes)
        txt = ""
        try:  # pdftotext, when the machine has it: poppler's (CI installs poppler-utils) or
            # Xpdf's (Git for Windows ships one). Ask both for UTF-8 — Xpdf defaults to Latin-1,
            # which fails a UTF-8 decode and silently emptied this block on Windows (7 Oct 2026).
            import subprocess
            txt = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", "-eol", "unix",
                                  os.path.join(TMP, "week.pdf"), "-"],
                                 capture_output=True, timeout=60).stdout.decode("utf-8", errors="replace")
        except Exception:
            txt = ""
        if txt:
            ok("the PDF names the tenant and never a team id", "Wolds Link" in txt and not re.search(r"\bIPT ?[1-6]\b", txt),
               txt[:160])
            flat = re.sub(r"\s+", " ", txt)
            ok("...its team column is headed with the tenant's word, not IPT", "TEAM /" in txt and "IPT /" not in txt,
               txt[:400])
            ok("...it shows team names (the cell may wrap) and £, and says the routes are not baked",
               re.search(r"\b(North|Central|South)\b", txt) and "Team" in txt and "£ WEEK" in txt
               and "not baked" in flat.lower(), flat[:300])
            ok("...and its footer is country-neutral", "Road-user charges are not on this sheet" in txt
               and "vignette" not in txt.lower() and "Estonia" not in txt, txt[-300:])
        else:
            ok("the PDF is a non-trivial document (text extraction unavailable here)", len(pdf_bytes) > 5000, str(len(pdf_bytes)))
    except Exception as e:
        ok("the PDF could be written", False, repr(e))

    r = c.get("/api/public/map-data")
    ok("the public map data reads (staff cookie)", r.status_code == 200, str(r.status_code))

    # ------------------------------------------------- 3a. 13 Oct 2026: lorry limits (limits.py)
    # No import has run on the demo: the endpoint answers an empty collection that still says
    # it is advisory and names both sources; the admin import is refused without the token.
    r = c.get("/api/public/limits")
    lm = r.json() if r.status_code == 200 else {}
    ok("⭐ /api/public/limits answers 200 with an empty, advisory collection before any import",
       r.status_code == 200 and lm.get("features") == [] and str(lm.get("note", "")).startswith("Advisory")
       and set((lm.get("sources") or {}).keys()) == {"osm", "here"}, f"{r.status_code} {r.text[:160]}")
    ok("...a bad bbox is a 400", c.get("/api/public/limits", params={"bbox": "x"}).status_code == 400)
    ok("🔴 the limits import is refused without the admin token",
       c.post("/api/admin/limits/import", params={"source": "osm"}).status_code == 403)
    ok("🔴 ...and so is the limits diagnostic", c.get("/api/admin/diagnostics/limits").status_code == 403)

    # ------------------------------------------------- 3b. 12 Oct 2026 (§2.2d): shared road sections
    # The demo package ships 0 geometry rows, so the index is empty here: the endpoint must
    # still answer, say so, and never invent a section.
    r = c.get("/api/public/road-load", params={"month": 10})
    ok("⭐ /api/public/road-load answers 200 on the unbaked demo", r.status_code == 200, f"{r.status_code} {r.text[:120]}")
    rl = r.json() if r.status_code == 200 else {}
    # REVERSED 12 Oct 2026 (the review and fix pass): the feed counts what the public map bar counts — lines on routes
    # the map DRAWS. On the unbaked demo that is none, so the totals are empty and every October line is NAMED in
    # routes_not_drawn (until then they were counted in the totals, which the bar beside the map never shows)
    ok("...empty runs, the month, the loaded leg, the planning block's working days and a note naming the undrawn lines",
       rl.get("runs") == [] and rl.get("month") == 10 and rl.get("leg") == "loaded" and (rl.get("working_days") or 0) > 0
       and rl.get("totals", {}).get("runs") == 0 and rl.get("totals", {}).get("lines") == 0
       # NARROWED 22 Oct 2026 (words sweep): the note reads on the map, so it names no field ("routes_not_drawn" left it)
       and any("18 lines not routed yet, not counted" in n for n in rl.get("notes", []))
       and not any("routes_not_drawn" in n for n in rl.get("notes", [])), str(rl.get("notes")))
    # NARROWED 14 Oct 2026: 22 October lines on 21 routes (the concrete pour). RE-PINNED 17 Oct 2026 (R2, §3a): 18 on 13
    ok("...and the October demo volume is NOT in the totals (the map draws none of it): all 18 lines named in routes_not_drawn",
       rl.get("totals", {}).get("routes_with_volume") == 0 and rl.get("totals", {}).get("trips") == 0
       and rl.get("routes_not_drawn", {}).get("lines") == 18 and len(rl.get("routes_not_drawn", {}).get("routes", [])) == 13
       and rl.get("totals", {}).get("co2_t") == 0.0, str(rl.get("totals")))
    _ret = c.get("/api/public/road-load", params={"month": 10, "leg": "return"})
    # NARROWED 12 Oct 2026: the status is checked too — called as a function, month-kpis' `unit` default was
    # FastAPI's Query object and this answered 400 "unit must be one of [...]"; the note is the LAST line
    ok("...the return leg answers 200 too, with its 'never added' note",
       _ret.status_code == 200 and (_ret.json().get("notes") or [""])[-1].startswith("Return leg"), _ret.text[:160])
    # NARROWED 12 Oct 2026 (the tenancy / wording review): the 400 reads as words, not a Python list
    _bad = c.get("/api/public/road-load", params={"month": 10, "leg": "x"})
    ok("an unknown leg is a 400 whose reason reads \"leg must be one of 'loaded' or 'return'\"",
       _bad.status_code == 400 and _bad.json().get("detail") == "leg must be one of 'loaded' or 'return'", _bad.text[:120])
    ok("⭐ the feed carries the Carbon page's planning-estimate stamp beside its CO2e figures (12 Oct 2026 review)",
       rl.get("carbon", {}).get("stamp") == "Wayscope planning estimate, DESNZ 2026 basis, not Carbon Tool figures"
       and "approved plan only" in rl.get("carbon", {}).get("rule", ""), str(rl.get("carbon")))
    ok("a month out of range is a 422", c.get("/api/public/road-load", params={"month": 99}).status_code == 422)
    r = c.post("/api/admin/road-sections/rebuild")
    ok("🔴 the admin rebuild is 403 without the token", r.status_code == 403, str(r.status_code))
    ok("🔴 ...so are status and the diagnostic", c.get("/api/admin/road-sections/status").status_code == 403
       and c.get("/api/admin/diagnostics/road-sections").status_code == 403)
    r = c.post("/api/admin/road-sections/rebuild", params=TOKEN)
    ok("with the token the rebuild answers: 0 routes indexed, 0 runs, not stale", r.status_code == 200
       and r.json().get("routes_indexed") == 0 and r.json().get("runs") == 0 and r.json().get("stale") is False, r.text[:200])
    r = c.get("/api/admin/diagnostics/road-sections", params=TOKEN)
    ok("the diagnostic prints the status text", r.status_code == 200 and "0 routes indexed" in r.json().get("text", ""), r.text[:200])
    # 12 Oct 2026 (the tests-and-regressions review): status had only its 403 here; the 200 and the
    # diagnostic's leg check over the real transport
    r = c.get("/api/admin/road-sections/status", params=TOKEN)
    ok("with the token status answers 200: 0 routes indexed, both legs listed with 0 runs, not stale",
       r.status_code == 200 and r.json().get("routes_indexed") == 0 and r.json().get("stale") is False
       and set(r.json().get("legs", {})) == {"loaded", "return"} and r.json()["legs"]["return"]["runs"] == 0, r.text[:200])
    ok("the diagnostic refuses an unknown leg with 400 (token accepted)",
       c.get("/api/admin/diagnostics/road-sections", params=dict(TOKEN, leg="x")).status_code == 400)
    # ------------------------------------------------- 3c. 12 Oct 2026 (§2.2b, D27): movements and carbon per council
    # REVERSED 17 Oct 2026 (demo redesign R2, D102): the shipped package carries the nine council areas, so the By council
    # view works straight after an import; the admin import of the seed file (backend/data/areas_demo.geojson) below still
    # REPLACES them over the real transport. Nothing is baked here, so every area reads 0 movements and the clip is empty —
    # the live clip runs on Render after a bake.
    AREAS_SEED = json.load(open(os.path.join(BACKEND, "data", "areas_demo.geojson"), encoding="utf-8"))
    r = c.get("/api/public/areas", params={"month": 10})
    ar = r.json() if r.status_code == 200 else {}
    ok("⭐ /api/public/areas answers 200 straight after the demo import: the nine shipped areas, 0 movements each (nothing drawn), the licence lines",
       r.status_code == 200 and len(ar.get("areas") or []) == 9 and (ar.get("working_days") or 0) > 0
       and all(a["movements_per_day"] == 0 for a in ar["areas"]) and ar.get("attribution") == AREAS_SEED["attribution"]
       and "No council areas are loaded for this tenant." not in ar.get("notes", []), r.text[:200])
    ok("a month out of range or no month is a 422", c.get("/api/public/areas", params={"month": 99}).status_code == 422
       and c.get("/api/public/areas").status_code == 422)
    ok("🔴 the areas import is 403 without the token, and so is the list",
       c.post("/api/admin/areas/import", json=AREAS_SEED).status_code == 403 and c.get("/api/admin/areas").status_code == 403)
    _bad_fc = {"type": "FeatureCollection", "features": AREAS_SEED["features"][:1]}          # no licence lines
    r = c.post("/api/admin/areas/import", params=TOKEN, json=_bad_fc)
    ok("an import with no licence lines is a 400 that says so (token accepted)",
       r.status_code == 400 and "licence lines" in r.json().get("detail", ""), r.text[:200])
    r = c.post("/api/admin/areas/import", params=TOKEN, json=AREAS_SEED)
    ok("⭐ with the token the seed imports: 9 areas (8 districts, 1 county), the two ONS lines, the clip empty (nothing baked)",
       r.status_code == 200 and r.json().get("imported") == 9 and r.json().get("tiers") == {"district": 8, "county": 1}
       and r.json().get("attribution") == sorted(AREAS_SEED["attribution"]) and r.json()["clip"]["rows"] == 0, r.text[:300])
    r = c.get("/api/admin/areas", params=TOKEN)
    ok("the admin list answers 200: nine areas without geometry, no route entering any (unbaked)",
       r.status_code == 200 and len(r.json().get("areas", [])) == 9 and all("geometry" not in a for a in r.json()["areas"])
       and all(a["routes_entering"] == [] for a in r.json()["areas"]), r.text[:200])
    r = c.get("/api/public/areas", params={"month": 10})
    ar = r.json() if r.status_code == 200 else {}
    _by = {a["name"]: a for a in ar.get("areas", [])}
    ok("⭐ the feed now serves the nine polygons with the two licence lines, 0 movements each (nothing drawn), 0.0 t",
       r.status_code == 200 and len(_by) == 9 and ar.get("attribution") == AREAS_SEED["attribution"]
       and all(a["movements_per_day"] == 0 and a["co2_t"] == 0.0 and a["geometry"]["type"] in ("Polygon", "MultiPolygon")
               for a in ar["areas"]), r.text[:200])
    # NARROWED 14 Oct 2026 (civils materials): the batching plant is a fourteenth site, in Melton
    # RE-PINNED 17 Oct 2026 (R2, §3a): the asphalt plant and Welland Ready-Mix are both in Harborough (10 of 16)
    # NARROWED 22 Oct 2026 (abnormal loads, phase 3, D75): the crane depot in Charnwood, the precast works in Hinckley and Bosworth (18)
    ok("...the demo's sites where the package coordinates put them (Harborough 10, Melton 3, Blaby 1, Charnwood 2, Rutland 1, Hinckley and Bosworth 1)",
       [_by.get(n, {}).get("sites_count") for n in ("Harborough", "Melton", "Blaby", "Charnwood", "Rutland", "Hinckley and Bosworth")] == [10, 3, 1, 2, 1, 1])
    ok("...of the 18 the totals carry (the card's '10 of 18'), and one label point per area for the map's one label",
       ar.get("totals", {}).get("sites") == 18
       and all(isinstance(a.get("label_point"), list) and len(a["label_point"]) == 2 for a in ar.get("areas", [])), r.text[:200])
    ok("...and the October lines are named as undrawn, never counted in an area",
       len(ar.get("routes_not_drawn", [])) == 13 and ar.get("totals", {}).get("loads") == 0)   # NARROWED 14 Oct: 21 routes; RE-PINNED 17 Oct: 13
    # 12 Oct 2026 (Track C review): the polygons once (no month roll-up), the month's figures without them
    r = c.get("/api/public/areas/shapes")
    sh = r.json() if r.status_code == 200 else {}
    ok("⭐ Track C review: /api/public/areas/shapes serves the nine polygons, their label points and the two licence lines, and no month figure",
       r.status_code == 200 and len(sh.get("areas", [])) == 9 and sh.get("attribution") == AREAS_SEED["attribution"]
       and all(a["geometry"]["type"] in ("Polygon", "MultiPolygon") and len(a.get("label_point") or []) == 2 for a in sh["areas"])
       and all("movements_per_day" not in a for a in sh["areas"]), r.text[:200])
    r = c.get("/api/public/areas", params={"month": 10, "geometry": 0})
    ok("Track C review: /api/public/areas?geometry=0 gives the same figures without the polygons",
       r.status_code == 200 and len(r.json().get("areas", [])) == 9 and all("geometry" not in a and "label_point" not in a for a in r.json()["areas"])
       and [(a["area_id"], a["movements_per_day"]) for a in r.json()["areas"]] == [(a["area_id"], a["movements_per_day"]) for a in ar["areas"]], r.text[:200])
    # 13 Oct 2026 (W12, bundle B): the council colour bands per scheme, saved through Config, served with the areas
    # NARROWED 13 Oct 2026 (bundle C regenerated the demo package once with carbon 20 / 50 / 100): the demo's own bands
    # RE-PINNED 18 Oct 2026 (demo redesign R3b): the package now carries both measures, re-chosen from the live four-year
    # demo's district figures — movements 25 / 50 / 100 a day, carbon 10 / 25 / 50 t in the month
    ok("⭐ W12: both area payloads carry the scheme's bands — the demo package's (movements 25 / 50 / 100, carbon 10 / 25 / 50)",
       sh.get("stops") == {"movements": [25, 50, 100], "carbon": [10, 25, 50]} and ar.get("stops") == sh.get("stops"),
       str((sh.get("stops"), ar.get("stops"))))
    _cfg = c.get("/api/config/factors", headers=PLAN).json()["doc"]
    _bad = json.loads(json.dumps(_cfg)); _bad.setdefault("tenant", {})["council_breaks"] = {"carbon": [100, 50, 20]}
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _bad, "updated_by": "smoke"})
    ok("🔴 W12: a Config save with falling bands is refused (400) with the plain-words reason, nothing stored",
       r.status_code == 400 and "tenant.council_breaks.carbon must rise from left to right" in r.text
       and c.get("/api/public/areas/shapes").json().get("stops", {}).get("carbon") == [10, 25, 50], r.text[:200])   # RE-PINNED 18 Oct 2026 (R3b)
    _good = json.loads(json.dumps(_cfg)); _good.setdefault("tenant", {})["council_breaks"] = {"carbon": [25, 60, 120]}
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _good, "updated_by": "smoke"})
    import config as _cfgmod
    _cfgmod.invalidate()
    _s2 = c.get("/api/public/areas/shapes").json().get("stops")
    _a2 = c.get("/api/public/areas", params={"month": 10, "geometry": 0}).json().get("stops")
    ok("⭐ W12: saved bands (carbon 25 / 60 / 120) reach both payloads the map reads; movements keep the default",
       r.status_code == 200 and _s2 == _a2 == {"movements": [50, 100, 200], "carbon": [25, 60, 120]}, str((r.status_code, _s2, _a2)))
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _cfg, "updated_by": "smoke"})
    _cfgmod.invalidate()
    # RE-PINNED 18 Oct 2026 (R3b): the demo's bands are movements 25 / 50 / 100 and carbon 10 / 25 / 50, both put back
    ok("W12: ...and the document put back reads the demo's bands again", r.status_code == 200
       and c.get("/api/public/areas/shapes").json().get("stops") == {"movements": [25, 50, 100], "carbon": [10, 25, 50]})
    # 22 Oct 2026 (read_cache): the every-month council read is kept until the next write — here through the REAL
    # middleware and endpoints, which test_areas.py's stubbed HTTP layer cannot run
    import read_cache as _rc
    _m1 = c.get("/api/public/areas/months")
    _bm = _rc.STATS["builds"]
    _m2 = c.get("/api/public/areas/months")
    ok("⭐ read cache: a second /api/public/areas/months is a hit — the same bytes, no build, application/json",
       _m1.status_code == _m2.status_code == 200 and _m2.content == _m1.content and _rc.STATS["builds"] == _bm
       and _m2.headers.get("content-type", "").startswith("application/json")
       and len(_m1.json()["months"]) == main.MONTH_COUNT and len(_m1.json()["area_fixed"]) == 9, str(_rc.STATS))
    _v = db.data_version()
    c.get("/api/public/areas/shapes"); c.get("/api/health"); c.get("/api/public/areas", params={"month": 10, "geometry": 0})
    ok("read cache: reads (GET) never move the data version", db.data_version() == _v)
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _bad, "updated_by": "smoke"})
    ok("⭐ read cache: a write that is REFUSED (400, nothing stored) still moves the version — the middleware's second net",
       r.status_code == 400 and db.data_version() > _v, f"{r.status_code} {db.data_version()} {_v}")
    _good = json.loads(json.dumps(_cfg)); _good.setdefault("tenant", {})["council_breaks"] = {"carbon": [11, 22, 33]}
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _good, "updated_by": "smoke"})
    _bm = _rc.STATS["builds"]
    _m3 = c.get("/api/public/areas/months")
    ok("⭐ read cache: a Config save, then the very next read carries the new bands (rebuilt once, never the cached body)",
       r.status_code == 200 and _m3.json()["common"]["stops"]["carbon"] == [11, 22, 33] and _rc.STATS["builds"] == _bm + 1
       and _m3.content != _m1.content, str((r.status_code, _m3.json()["common"].get("stops"))))
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _cfg, "updated_by": "smoke"})
    _m4 = c.get("/api/public/areas/months")
    ok("read cache: ...and the document put back, the read is the first one again, byte for byte",
       r.status_code == 200 and _m4.content == _m1.content)
    # 14 Oct 2026 (HU7): the Config page reads the file's per-country fair-price sets from the same call
    _fj = c.get("/api/config/factors", headers=PLAN).json()
    _gbf = ((_fj.get("file") or {}).get("fair_price_by_country") or {}).get("GB") or {}
    ok("⭐ HU7: GET /api/config/factors hands the page the file's GB fair-price set (£24/h driver, rigid 31, no thaw)",
       _gbf.get("driver_eur_per_h") == 24.0 and _gbf["consumption"]["rigid"]["l_per_100km_empty"] == 31.0
       and _gbf["season"]["thaw_months"] == [] and (_fj.get("file") or {}).get("fair_price", {}).get("driver_eur_per_h") == 20.0,
       str(_gbf)[:200])
    _bad = json.loads(json.dumps(_cfg)); _bad["fair_price_by_country"] = {"GB": dict(_gbf, driver_eur_per_h=-1)}
    r = c.put("/api/admin/config/factors", params=TOKEN, json={"doc": _bad, "updated_by": "smoke"})
    ok("🔴 HU7: a Config save with a negative GB driver rate in fair_price_by_country is refused (400), naming the field",
       r.status_code == 400 and "fair_price_by_country.GB.driver_eur_per_h" in r.text, r.text[:200])
    # ------------------------------------------- 3d. the network behind the gate (12 Oct 2026, map review D18)
    # The twelve reads that needed no credential until today, read three ways: the staff COOKIE (this client, no
    # header), the planner CODE in the header with no cookie, and — at the end of this file — the MAP cookie and
    # nobody at all. The two staff writes under the same prefixes keep their own checks behind the gate.
    NETWORK_READS = ("/api/routes/status", "/api/routes/summary", "/api/routes/geojson", "/api/routes/R001/geometries",
                     "/api/routes/R001/analysis", "/api/routes/analysis-batch?route_ids=R001,R002", "/api/locations/geojson",
                     "/api/gates", "/api/haul-roads", "/api/routes/R001/haul-roads", "/api/routes/R001/restrictions",
                     "/api/routes/restrictions")
    import restrictions as _rx
    # the National Highways feed is never fetched from a harness: two closures, one current and far from every route,
    # one ENDED and drawn on the stretch R016 is given below
    _now_ms = int(__import__("time").time() * 1000)
    _day_ms = 86400000
    def _fake_nh_page(offset, count=_rx.NH_PAGE, timeout=None):
        if offset:
            return {"type": "FeatureCollection", "features": []}
        return {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[-2.5, 53.5], [-2.49, 53.51]]},
             "properties": {"OBJECTID": 2, "description": "Lane closure for barrier works", "road_number": "M6",
                            "eventtype": "Lane closure", "natureofworks": "Barrier works", "formattedeventnumber": "NH-0002",
                            "scheduledplannedstartdate": _now_ms - 2 * _day_ms, "scheduledplannedenddate": _now_ms + 400 * _day_ms}},
            {"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[-0.9, 52.7], [-0.895, 52.705]]},
             "properties": {"OBJECTID": 1, "description": "Carriageway closure for resurfacing, finished", "road_number": "A607",
                            "eventtype": "Carriageway closure", "natureofworks": "Resurfacing", "formattedeventnumber": "NH-0001",
                            "scheduledplannedstartdate": _now_ms - 40 * _day_ms, "scheduledplannedenddate": _now_ms - 10 * _day_ms}}]}
    _real_nh_page = _rx._nh_page
    _rx._nh_page = _fake_nh_page
    _rx.clear_cache()
    # R016 (cap 8 in the demo since 17 Oct 2026, the Burrough lane agreement) baked by hand along the ended closure, both legs
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
               "duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", ("default", "R016", "Artic Tipper (44t)", "loaded", 0, "[[-0.9,52.7],[-0.89,52.71]]", 1.3, 0.05))
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
               "duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", ("default", "R016", "Artic Tipper (44t)", "return", 0, "[[-0.89,52.71],[-0.9,52.7]]", 1.3, 0.05))
    try:
        _codes = {p: c.get(p).status_code for p in NETWORK_READS}
        ok("⭐ D18: all twelve network reads answer 200 to the staff COOKIE alone (the app's iframe sends no header)",
           all(v == 200 for v in _codes.values()), str(_codes))
        _jar = c.cookies
        c.cookies = type(_jar)()
        try:
            _codes_hdr = {p: c.get(p, headers=PLAN).status_code for p in NETWORK_READS}
            _codes_none = {p: c.get(p).status_code for p in NETWORK_READS}
            _bodies_none = {p: c.get(p).json().get("detail") for p in NETWORK_READS}
            r_put_none = c.put("/api/routes/R999/status", json={"status": "Approved"})
            r_put_hdr = c.put("/api/routes/R999/status", json={"status": "Approved"}, headers=PLAN)
            r_put_north = c.put("/api/routes/R999/status", json={"status": "Approved"}, headers=NORTH)
            r_cap_none = c.put("/api/locations/L001/capacity", json={"capacity_qty": 1, "capacity_unit": "t"})
            r_cap_hdr = c.put("/api/locations/L001/capacity", json={"capacity_qty": 1, "capacity_unit": "t"}, headers=PLAN)
            r_who_plan = c.get("/api/public/whoami", headers=PLAN)
            r_who_north = c.get("/api/public/whoami", headers=NORTH)
            r_who_none = c.get("/api/public/whoami")
        finally:
            c.cookies = _jar
        ok("⭐ D18: ...and 200 to the planner CODE in the header with no cookie (the staff app's own fetches)",
           all(v == 200 for v in _codes_hdr.values()), str(_codes_hdr))
        ok("🔴 D18: ...and 401 to NOBODY — every one of the twelve, with the map's wording",
           all(v == 401 for v in _codes_none.values()) and all(d == "the map is password protected" for d in _bodies_none.values()),
           str(_codes_none) + str(_bodies_none)[:200])
        ok("🔴 D18: the staff WRITE PUT /api/routes/{id}/status with no credential is refused at the gate (401, the map's wording) — before its own approver check",
           r_put_none.status_code == 401 and r_put_none.json().get("detail") == "the map is password protected", r_put_none.text[:120])
        ok("D18: ...with the planner code it passes the gate unchanged and reaches its body (404: no forecast for R999)",
           r_put_hdr.status_code == 404 and "No forecast" in r_put_hdr.text, f"{r_put_hdr.status_code} {r_put_hdr.text[:80]}")
        ok("D18: ...and a team code passes the gate but is refused by the body's own approver check (403)",
           r_put_north.status_code == 403 and "cannot approve" in r_put_north.text, f"{r_put_north.status_code} {r_put_north.text[:80]}")
        ok("D18: PUT /api/locations/{id}/capacity: 401 at the gate with nothing; with a code the body's admin check answers (403, no token)",
           r_cap_none.status_code == 401 and r_cap_hdr.status_code == 403 and "admin token" in r_cap_hdr.text,
           f"{r_cap_none.status_code} {r_cap_hdr.status_code} {r_cap_hdr.text[:80]}")
        ok("⭐ whoami: the planner code is staff with no team; the North Team code is staff with its team's NAME",
           r_who_plan.status_code == 200 and r_who_plan.json() == {"level": "staff", "team_label": None}
           and r_who_north.status_code == 200 and r_who_north.json() == {"level": "staff", "team_label": "North Team"},
           f"{r_who_plan.text} {r_who_north.text}")
        ok("🔴 whoami: never a code, never a team id", "IPT" not in r_who_north.text and "north-code-xyz" not in r_who_north.text
           and "plan-code-xyz" not in r_who_plan.text and r_who_none.status_code == 401)
        # §2.5 row 2: the cap on the public map's route features
        _md = c.get("/api/public/map-data").json()
        _lines = {(f["properties"]["route_id"], f["properties"]["leg"]): f["properties"] for f in _md["features"]
                  if f["geometry"]["type"] == "LineString" and f["properties"].get("route_id")}
        # RE-PINNED 17 Oct 2026 (R2, §3a): the typed cap is 8 a day (was 20)
        ok("⭐ §2.5: the hand-baked R016 draws both legs on the public map and BOTH carry max_vehicles_per_day 8 (its typed cap)",
           set(_lines) == {("R016", "loaded"), ("R016", "return")} and all(p["max_vehicles_per_day"] == 8 for p in _lines.values()),
           str({k: v.get("max_vehicles_per_day") for k, v in _lines.items()}))
        # §2.2c (c) and §2.3: the restriction check and the drawn layer, over HTTP, against the stubbed feed
        r = c.get("/api/routes/restrictions")
        ok("§2.2c: GET /api/routes/restrictions by default sees no hit — the only closure on R016 has ENDED",
           r.status_code == 200 and r.json()["routes"] == {} and r.json()["current_only"] is True
           and r.json()["expired_or_future_excluded"] == 1, r.text[:200])
        r = c.get("/api/routes/restrictions?include_expired=1")
        _h = (r.json().get("routes") or {}).get("R016", {}).get("hits") or []
        ok("⭐ §2.2c: ...and with include_expired=1 R016 lists the ended closure NH-0001 with its dates, so a past month can show it",
           r.status_code == 200 and r.json()["current_only"] is False and sorted({x["event_number"] for x in _h}) == ["NH-0001"]
           and all(x["from"] and x["to"] for x in _h), r.text[:300])
        r = c.get("/api/restrictions")
        _fc0 = r.json()
        ok("§2.3: GET /api/restrictions by default carries the one current closure and reports the ended one dropped",
           r.status_code == 200 and [f["properties"]["event_number"] for f in _fc0["features"]] == ["NH-0002"]
           and _fc0["expired_or_future"] == 1 and _fc0["total_records"] == 2 and "clipped" not in _fc0, r.text[:200])
        r = c.get("/api/restrictions?include_expired=1&bbox=-1.0,52.6,-0.8,52.8")
        _fc1 = r.json()
        ok("⭐ §2.3: ...bbox=<the corridor>&include_expired=1 keeps only the closure inside the window and says it clipped 1",
           r.status_code == 200 and [f["properties"]["event_number"] for f in _fc1["features"]] == ["NH-0001"]
           and _fc1["clipped"] == 1 and _fc1["bbox"] == [-1.0, 52.6, -0.8, 52.8] and _fc1["counts"] == {"nh_closures": 1}, r.text[:200])
        r = c.get("/api/restrictions?bbox=10,10,11,11")
        ok("§2.3: a window with nothing in it is an empty collection, clipped 1 (the ended one was already dropped by the default read)",
           r.status_code == 200 and r.json()["features"] == [] and r.json()["clipped"] == 1, r.text[:200])
        r = c.get("/api/restrictions?bbox=junk")
        ok("🔴 §2.3: a bad bbox is a 400 that says the shape", r.status_code == 400 and "west,south,east,north" in r.text, r.text[:120])
    finally:
        _rx._nh_page = _real_nh_page
        _rx.clear_cache()
        db.execute("DELETE FROM route_geometry WHERE tenant_id = ? AND route_id = ?", ("default", "R016"))

    # ------------------------------------------------------------ 4. export and replace
    r = c.get("/api/admin/tenant/export", params=TOKEN)
    ok("⭐ the tenant exports as a JSON attachment", r.status_code == 200
       and r.headers.get("content-type", "").startswith("application/json")
       and "attachment" in r.headers.get("content-disposition", ""), str(r.status_code))
    pkg = r.json() if r.status_code == 200 else {}
    # NARROWED 14 Oct 2026 (civils materials): the demo's new size
    # RE-PINNED 17 Oct 2026 (R2, §3a): read from the package file (exact). Weeks may only GROW on the round trip: reads above
    # materialise the weeks of later months on demand (§1.9: future weeks materialise on read), never drop one
    ok("...a valid package of the same size", pkg.get("modus_package") == 1
       and len(pkg["tables"]["routes"]) == len(DEMO["tables"]["routes"]) == 43
       and len(pkg["tables"]["forecasts"]) == len(DEMO["tables"]["forecasts"]) == 907
       and len(pkg["tables"]["forecast_weeks"]) >= len(DEMO["tables"]["forecast_weeks"])
       and len(pkg["tables"]["location_gates"]) == len(DEMO["tables"]["location_gates"]) == 21      # NARROWED 22 Oct 2026: the two abnormal-load sites
       and len(pkg["tables"]["abnormal_loads"]) == len(DEMO["tables"]["abnormal_loads"]) == 3)
    ok("...and it carries the route_sections table (empty: nothing is baked here)",
       "route_sections" in pkg.get("tables", {}) and pkg["tables"]["route_sections"] == [])
    ok("...and the nine council areas imported above, with an empty clip table (§2.2b)",
       len(pkg.get("tables", {}).get("areas", [])) == 9 and pkg["tables"].get("route_area_km") == [])
    r = c.post("/api/admin/tenant/import", params=dict(TOKEN, replace=1), json=pkg)
    ok("⭐ replace=1 re-imports the exported tenant over itself", r.status_code == 200 and r.json().get("replaced") is True,
       r.text[:160])
    m2 = c.get("/api/meta").json()
    ok("...and nothing is lost or doubled", len(m2["routes"]) == 43 and len(m2["ipts"]) == 3   # NARROWED 14 Oct: 22 / 88; RE-PINNED 17 Oct: 43 / 907
       and len(c.get("/api/forecasts", headers=PLAN).json()) == 907)
    ok("...and the council areas came back with the package (9)",
       len(c.get("/api/admin/areas", params=TOKEN).json().get("areas", [])) == 9)

    # ------------------------------------------------------------ 5. H9: demo requests (8 Oct)
    # Signed out, no cookie, no code: the landing page's visitor. SMTP stubbed at the module.
    r = c.post("/api/gate-signout")
    c.cookies.clear()
    import demo_request as _dr
    _sent = []
    class _SMTPStub:
        def __init__(self, *a, **k): pass
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def ehlo(self): pass
        def starttls(self): pass
        def login(self, u, p): pass
        def send_message(self, m): _sent.append(m)
        ehlo_resp = b"250-AUTH LOGIN PLAIN"
        def auth_plain(self, challenge=None): return ""
        def auth_login(self, challenge=None): return ""
        def auth(self, mech, authobject, initial_response_ok=True): return (235, b"Accepted")
    _real = _dr.smtplib.SMTP
    _dr.smtplib.SMTP = _SMTPStub
    try:
        for v in ("SMTP_USER", "SMTP_PASSWORD", "DEMO_FROM", "DEMO_NOTIFY"):
            os.environ.pop(v, None)
        form = {"name": "Ada Example", "email": "ada@example.org", "organisation": "Example Civils",
                "role": "Principal contractor", "wants": "the look-ahead", "updates": "yes", "website": ""}
        os.environ.pop("DEMO_LEADS", None)
        r = c.post("/api/public/demo-request", follow_redirects=False, data=form)
        ok("🔴 F3 a: with DEMO_LEADS unset the form is off: a 303 to error.html and no login made",
           r.status_code == 303 and r.headers["location"].endswith("/error.html") and "#demo=" not in r.text)
        os.environ["DEMO_LEADS"] = "1"
        r = c.post("/api/public/demo-request", follow_redirects=False, data=form)
        # REVERSED 14 Oct 2026 (G3 / HU14): the access page carried both shared codes; it now carries the lead's OWN link
        ok("⭐ no sender configured: the form POST (signed out) answers 200 with the access page — the lead's own link, no code",
           r.status_code == 200 and "https://app.wayscope.co.uk/#demo=" in r.text and "fictional" in r.text
           and "map-pass-xyz" not in r.text and "north-code-xyz" not in r.text, str(r.status_code))
        _screen_link = re.search(r"https://app\.wayscope\.co\.uk/#demo=([A-Za-z0-9_-]+)", r.text).group(1)
        ok("...no-store, and it never carries the planner or admin code", r.headers.get("cache-control") == "no-store"
           and "plan-code-xyz" not in r.text and "admin-code-xyz" not in r.text)
        r = c.post("/api/public/demo-request", follow_redirects=False, data=dict(form, website="bot"))
        ok("a honeypot hit is a 303 to thanks.html", r.status_code == 303 and r.headers["location"] == "https://wayscope.co.uk/thanks.html")
        r = c.post("/api/public/demo-request", follow_redirects=False, data=dict(form, email="nope"))
        ok("an invalid form is a 303 to error.html", r.status_code == 303 and r.headers["location"] == "https://wayscope.co.uk/error.html")
        os.environ.update({"SMTP_USER": "sender@wayscope.co.uk", "SMTP_PASSWORD": "app-pass", "DEMO_FROM": "Wayscope <sender@wayscope.co.uk>",
                           "DEMO_NOTIFY": "notify@wayscope.co.uk"})
        r = c.post("/api/public/demo-request", follow_redirects=False, data=dict(form, email="bea@example.org"))
        ok("⭐ sender configured: the POST is a 303 to thanks.html and the background task sent two emails",
           r.status_code == 303 and r.headers["location"] == "https://wayscope.co.uk/thanks.html" and len(_sent) == 2
           and _sent[0]["To"] == "bea@example.org" and _sent[1]["To"] == "notify@wayscope.co.uk", f"{r.status_code} {len(_sent)}")
        # REVERSED 14 Oct 2026 (G3 / HU14): the email carried both shared codes; it now carries the lead's own link
        _bea = _sent[0].get_body(preferencelist=("plain",)).get_content()
        ok("...the email carries the lead's own link and NEITHER code", "#demo=" in _bea
           and "map-pass-xyz" not in _bea and "north-code-xyz" not in _bea)
        _bea_link = re.search(r"#demo=([A-Za-z0-9_-]+)", _bea).group(1)
        ok("🔴 /api/public/route-forecasts is still closed to the same signed-out visitor",
           c.get("/api/public/route-forecasts").status_code == 401)
        ok("🔴 ...and so is /api/public/road-load (12 Oct 2026, §2.2d)",
           c.get("/api/public/road-load", params={"month": 10}).status_code == 401)
        ok("🔴 ...and so is /api/public/areas (12 Oct 2026, §2.2b)",
           c.get("/api/public/areas", params={"month": 10}).status_code == 401)
        ok("🔴 ...and so is /api/public/areas/shapes (Track C review, 12 Oct 2026)",
           c.get("/api/public/areas/shapes").status_code == 401)
        ok("🔴 ...and so is /api/public/limits (13 Oct 2026)", c.get("/api/public/limits").status_code == 401)
        r = c.get("/api/admin/diagnostics/smtp")
        ok("the SMTP probe is refused without the admin token", r.status_code == 403, str(r.status_code))
        r = c.get("/api/admin/diagnostics/smtp", params=TOKEN)
        ok("...and with it reports a login (stubbed) as ok at step 'done'", r.status_code == 200 and r.json().get("ok") is True
           and r.json().get("step") == "done" and r.json().get("user") == "sender@wayscope.co.uk", r.text[:200])
        r = c.get("/api/admin/demo-requests.csv")
        ok("the CSV is refused without the admin token", r.status_code == 403, str(r.status_code))
        r = c.get("/api/admin/demo-requests.csv", params=TOKEN)
        ok("...and served with it: text/csv, three stored rows (the honeypot and the invalid one are not rows)",
           r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
           and r.text.count(chr(10)) == 3 and "ada@example.org" in r.text and "bea@example.org" in r.text
           and "shown on screen" in r.text and ",sent," in r.text, r.text[:300])
    finally:
        _dr.smtplib.SMTP = _real
        for v in ("SMTP_USER", "SMTP_PASSWORD", "DEMO_FROM", "DEMO_NOTIFY"):
            os.environ.pop(v, None)

    r = c.get("/map/")
    ok("after sign-out the map is closed again", r.status_code == 401, str(r.status_code))
    ok("…and so are the public map's abnormal loads (phase 3, D68): 401 signed out", c.get("/api/public/abnormal").status_code == 401)
    ok("🔴 D18: ...and so is the network — the twelve reads answer 401 to the signed-out visitor",
       all(c.get(p).status_code == 401 for p in NETWORK_READS))
    r = c.post("/api/map-auth", json={"password": "map-pass-xyz"})
    ok("the map password opens the map…", r.status_code == 200 and c.get("/map/").status_code == 200)
    _pub = c.get("/api/public/abnormal")
    ok("…and (phase 3, D68) the map cookie reads the public abnormal loads: a collection of the approved-or-later movements only — none yet — "
       "with no grades, no candidate words and no structures words (the staff map's), only the indicative line words",
       _pub.status_code == 200 and _pub.json().get("type") == "FeatureCollection" and _pub.json().get("features") == []
       and _pub.json().get("grades") is None and _pub.json().get("words", {}).get("candidates") is None
       and _pub.json().get("words", {}).get("structures") is None and "Indicative" in (_pub.json().get("words", {}).get("indicative") or ""),
       _pub.text[:200])
    ok("…but 🔴 not the guide", c.get("/help/").status_code == 401)
    ok("…and 🔴 not the monthly bulletin (staff only, D10)", c.get("/bulletin/?month=10").status_code == 401
       and c.get("/api/bulletin/month.pdf?month=10").status_code == 401)
    ok("…and 🔴 not the route map export (staff only, the user 14 Oct): the map password's cookie gets 401, and whoami says 'map'",
       c.get("/api/routes/map-export.pdf?month=10").status_code == 401 and c.get("/api/public/whoami").json().get("level") == "map"
       and c.get("/api/routes/map-export.png?month=10").status_code == 401
       and c.get("/api/routes/map-export/places?month=10").status_code == 401)
    # 14 Oct 2026 (night): the community considerations layer is never on the password view
    ok("…and 🔴 not the community considerations layer (staff only, the user 14 Oct): the map password's cookie gets 401",
       c.get("/api/community").status_code == 401 and c.get("/api/community/status").status_code == 401)
    _cm = c.get("/api/community", headers=PLAN)
    _cj = _cm.json() if _cm.status_code == 200 else {}
    ok("⭐ …a planner's code reads it: 200, a FeatureCollection, the nine kinds, nothing fetched yet on this deployment",
       _cm.status_code == 200 and _cj.get("type") == "FeatureCollection" and len(_cj.get("kinds") or []) == 9
       and _cj.get("state") == "never" and _cj.get("features") == [] and _cj.get("fill_alpha") == 0.22, _cm.text[:200])
    ok("🔴 …a team code may read it but not fetch it: 403 (no refresh is started from this harness: it would go online)",
       c.get("/api/community", headers=NORTH).status_code == 200
       and c.post("/api/community/refresh", headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403)
    # 16 Oct 2026 (CLOCS v5 C1): the collision history is never on the password view either (R6)
    ok("…and 🔴 not the collision history (staff only, R6): the map password's cookie gets 401, the evidence too",
       c.get("/api/collisions").status_code == 401 and c.get("/api/collisions/status").status_code == 401
       and c.get("/api/routes/R016/collision-evidence.pdf").status_code == 401)
    _co = c.get("/api/collisions", headers=PLAN)
    _coj = _co.json() if _co.status_code == 200 else {}
    ok("⭐ …a planner's code reads it: 200, a FeatureCollection, nothing fetched yet on this deployment",
       _co.status_code == 200 and _coj.get("type") == "FeatureCollection" and _coj.get("state") == "never"
       and _coj.get("features") == [] and _coj.get("strip") == {"route_m": 250, "site_m": 500, "on_route_m": 25}, _co.text[:200])
    _ce = c.get("/api/routes/R016/collision-evidence.pdf", headers=PLAN)
    ok("⭐ …and a route's evidence PDF before any fetch: 200, a PDF that says it has not been fetched",
       _ce.status_code == 200 and _ce.content[:5] == b"%PDF-" and "no-store" in _ce.headers.get("cache-control", ""),
       str(_ce.status_code))
    ok("🔴 …a team code may read it but not fetch it: 403 (no fetch is started from this harness: it would go online)",
       c.get("/api/collisions", headers=NORTH).status_code == 200
       and c.post("/api/collisions/refresh", headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403)
    # 22 Oct 2026 (S16 / B5): the abnormal loads — the real endpoints end to end on the demo package
    ok("🔴 abnormal loads: the page, the rules preview and the settings need a sign-in (the map cookie gets 401)",
       c.get("/api/abnormal").status_code == 401 and c.get("/api/abnormal/rules?width_m=5").status_code == 401
       and c.get("/api/abnormal/settings").status_code == 401)
    _ab = c.get("/api/abnormal", headers=PLAN)
    _abj = _ab.json() if _ab.status_code == 200 else {}
    # NARROWED 22 Oct 2026 (phase 3, D75): the demo ships three loads — the units approved, the crane and the beams pending — and 18 sites
    ok("⭐ …a planner reads the page: 200, the demo's three loads (D75: AL001 the units approved, AL002 the crane and AL003 the beams pending), the shipped bank holidays, no time bans, the pickers",
       _ab.status_code == 200 and [m["id"] for m in _abj.get("movements") or []] == ["AL001", "AL002", "AL003"]
       and [m["status"] for m in _abj["movements"]] == ["approved", "pending", "pending"] and all(m.get("geometry") is None for m in _abj["movements"])
       and _abj.get("settings", {}).get("holidays", {}).get("source") == "seed"
       and _abj.get("settings", {}).get("time_bans") == [] and len(_abj.get("choices", {}).get("locations") or []) == 18
       and "the haulier, esdal, the police" in (_abj.get("words", {}).get("advice") or "").lower(), _ab.text[:200])   # NARROWED 27 Oct 2026: the one advice line
    _pv = c.get("/api/abnormal/rules", params={"width_m": 4.6, "height_m": 3.9, "length_m": 16.5, "gross_t": 40, "max_axle_t": 9,
                                                "axle_count": 6, "indivisible": "true", "date_from": "2027-03-08"}, headers=NORTH)
    _pvj = _pv.json() if _pv.status_code == 200 else {}
    ok("⭐ …the preview for the 4.6 m modular units: STGO, police 2 clear working days by Wed 3 Mar 2027, attendant, escort expected "
       "off motorways; nothing saved",
       _pv.status_code == 200 and _pvj.get("regime") == "stgo" and _pvj.get("notices", {}).get("police_days") == 2
       and _pvj.get("notify_by") == "2027-03-03" and _pvj.get("attendant") is True and _pvj.get("escort_expected") == "other_roads"
       and c.get("/api/abnormal", headers=PLAN).json().get("count") == 3, _pv.text[:200])
    _mk = c.post("/api/abnormal", json={"name": "Site office units", "origin_id": "L004", "dest_id": "L006", "dest_gate_id": "G009",
                                        "date_from": "2027-03-08", "date_to": "2027-03-12", "movements": 3, "width_m": 4.6, "height_m": 3.9,
                                        "length_m": 16.5, "gross_t": 40, "max_axle_t": 9, "axle_count": 6, "indivisible": True,
                                        "vehicle_type": "low_loader"}, headers=dict(NORTH, **{"X-Wayscope": "1"}))
    _mkj = _mk.json() if _mk.status_code == 200 else {}
    ok("🔴 …a team code plans one: 200, AL004 (after the demo's three), pending, ITS team (the body named none), the team's name never the id, the advice on it",
       _mk.status_code == 200 and _mkj.get("movement", {}).get("id") == "AL004" and _mkj["movement"]["status"] == "pending"
       and _mkj["movement"]["ipt"] == "IPT1" and _mkj["movement"]["team_name"] == "North Team"
       and _mkj["movement"]["advice"]["summary"].startswith("STGO Cat 1"), _mk.text[:300])
    ok("🔴 …a bad body is a 400 in sentences (a slipped decimal, a missing figure)",
       c.post("/api/abnormal", json={"name": "x", "origin_id": "L004", "dest_id": "L006", "width_m": 46, "height_m": 3.9,
                                     "date_from": "2027-03-08"}, headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 400)
    ok("🔴 …approving is a planner's: the team's code gets 403, the planner's 200; a wrong step a 400",
       c.post("/api/abnormal/AL004/status", json={"status": "approved"}, headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.post("/api/abnormal/AL004/status", json={"status": "approved"}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.post("/api/abnormal/AL004/status", json={"status": "pending"}, headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.post("/api/abnormal/AL004/status", json={"status": "moved"}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.post("/api/abnormal/AL004/status", json={"status": "pending"}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 400)
    ok("…the time bans and the holidays fetch are a planner's: 403 to a team code; a bad ban list a 400; the fetch is NOT pressed here (it would go online)",
       c.put("/api/abnormal/time-bans", json={"time_bans": []}, headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.post("/api/abnormal/bank-holidays/refresh", headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.put("/api/abnormal/time-bans", json={"time_bans": [{"force": "", "from": "7", "to": "9"}]}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 400
       and c.put("/api/abnormal/time-bans", json={"time_bans": [{"force": "Leicestershire", "from": "07:00", "to": "09:30"}]},
                 headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.get("/api/abnormal/settings", headers=NORTH).json().get("time_bans", [{}])[0].get("force") == "Leicestershire")
    ok("🔴 …another team's code cannot see it (404, never 403) and a planner deletes it; the demo's three stay",
       c.get("/api/abnormal/AL004", headers={"X-Access-Code": "central-code-xyz"}).status_code in (401, 404)
       and c.delete("/api/abnormal/AL004", headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.get("/api/abnormal", headers=PLAN).json().get("count") == 3)
    # 22 Oct 2026 (phase 2): the route at the typed size — no HERE key on this harness, so the bake says so; via points; the staff map's collection
    ok("🔴 abnormal loads, the route: the map cookie gets 401 on the staff collection; a team code may not route (403); a planner's bake with no HERE key is a 400 in words",
       c.get("/api/abnormal/geojson").status_code == 401 and c.get("/api/abnormal/geojson", headers=NORTH).status_code == 200
       and c.post("/api/abnormal", json={"name": "Beams", "ipt": "IPT1", "origin_id": "L002", "dest_id": "L008", "date_from": "2027-04-12", "width_m": 2.8,
                                         "height_m": 4.2, "length_m": 36, "gross_t": 78, "max_axle_t": 11, "axle_count": 8}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.post("/api/abnormal/AL001/route", headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.post("/api/abnormal/AL001/route", headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 400
       and "no routing key" in c.post("/api/abnormal/AL001/route", headers=dict(PLAN, **{"X-Wayscope": "1"})).json().get("detail", "")
       and c.put("/api/abnormal/AL001/via", json={"via": [[52.6, -0.9]]}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200
       and c.put("/api/abnormal/AL001/via", json={"via": [[95, 0]]}, headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 400
       and c.get("/api/abnormal/geojson", headers=PLAN).json().get("count") == 0
       and c.delete("/api/abnormal/AL001", headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200)
    # 22 Oct 2026 (seen on a local copy): the map page sends the STAFF gate cookie and no header, so access.current() is None
    # and the staff collection came back EMPTY for a code sign-in. Now the cookie alone reads the scheme's staff view.
    _ck = TestClient(main.app, base_url="https://testserver")
    _ck.post("/api/auth", json={"code": "plan-code-xyz"})
    _cm = c.post("/api/abnormal", json={"name": "Cookie check", "ipt": "IPT1", "origin_id": "L002", "dest_id": "L008", "date_from": "2027-04-12",
                                        "width_m": 3.2, "height_m": 4.2, "length_m": 18, "gross_t": 48, "max_axle_t": 10, "axle_count": 6},
                 headers=dict(PLAN, **{"X-Wayscope": "1"})).json().get("movement", {}).get("id")
    ok("🔴 abnormal loads, the staff map's collection: the STAFF gate cookie alone (the map page of a code sign-in, no header) reads the same "
       "collection a planner's header does — the words, the grades and the movement's candidates — not an empty one; the map cookie still gets 401",
       gate.COOKIE in _ck.cookies and _ck.get("/api/abnormal/geojson").status_code == 200
       and _ck.get("/api/abnormal/geojson").json().get("grades") == c.get("/api/abnormal/geojson", headers=PLAN).json().get("grades")
       and _ck.get("/api/abnormal/geojson").json().get("words", {}).get("candidates")
       and c.get("/api/abnormal", headers=PLAN).json().get("count") == 4      # AL002, AL003, the Beams, the Cookie check
       and _ck.get("/api/abnormal").status_code == 401, str(_cm))
    ok("...the collection with the cookie carries every movement a planner sees once drawn: none drawn here (no HERE key), so count 0 both ways, and the cookie's read is byte-identical to the header's",
       _ck.get("/api/abnormal/geojson").json() == c.get("/api/abnormal/geojson", headers=PLAN).json())
    _cal = c.get("/api/abnormal/calendar.xlsx", headers=PLAN)
    _calp = c.get("/api/abnormal/calendar.pdf", headers=PLAN)
    _es = c.get(f"/api/abnormal/{_cm}/esdal.pdf", headers=PLAN)
    # REVERSED 27 Oct 2026 (the final UI pass, card 2): the ESDAL details sheet is gone — its path answers nothing now
    ok("abnormal loads, phase 3 (D73): the notice calendar as XLSX and PDF — any sign-in (401 signed out), the right media types and file names; the ESDAL sheet's old path is no route (404)",
       c.get("/api/abnormal/calendar.xlsx").status_code == 401 and c.get(f"/api/abnormal/{_cm}/esdal.pdf").status_code in (401, 404)
       and _cal.status_code == 200 and "spreadsheetml" in _cal.headers.get("content-type", "") and "notice-calendar" in _cal.headers.get("content-disposition", "")
       and _calp.status_code == 200 and _calp.headers.get("content-type") == "application/pdf" and _calp.content[:4] == b"%PDF"
       and _es.status_code == 404 and c.get("/api/abnormal/AL999/esdal.pdf", headers=PLAN).status_code == 404, f"{_cal.status_code} {_calp.status_code} {_es.status_code}")
    _chk = c.post(f"/api/abnormal/{_cm}/structures", headers=dict(PLAN, **{"X-Wayscope": "1"}))
    ok("abnormal loads, phase 3: Check the roads again — a team code may not (403), a missing movement is a 404, an unrouted one a 400 in words (draw first), no HERE call needed",
       c.post(f"/api/abnormal/{_cm}/structures", headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403
       and c.post("/api/abnormal/AL999/structures", headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 404
       and _chk.status_code == 400 and "draw the route first" in _chk.json().get("detail", "")
       and c.delete(f"/api/abnormal/{_cm}", headers=dict(PLAN, **{"X-Wayscope": "1"})).status_code == 200, _chk.text[:160])
    ok("abnormal loads, phase 2: the Dashboard's summary (any sign-in; 401 signed out) and the pinch rows (a planner's)",
       c.get("/api/abnormal/summary").status_code == 401 and c.get("/api/abnormal/summary", headers=NORTH).status_code == 200
       and c.get("/api/abnormal/summary", headers=NORTH).json().get("planned") == 3      # the demo's crane and beams, the Beams (all North Team)
       and c.put("/api/abnormal/AL001/pinch", json={"pinch": []}, headers=dict(NORTH, **{"X-Wayscope": "1"})).status_code == 403)
    # 12 Oct 2026 (D18 / D31 / whoami): the MAP cookie — the outside viewer with the shared password
    _codes_map = {p: c.get(p).status_code for p in NETWORK_READS}
    ok("⭐ D18: the map cookie opens all twelve network reads (the public-safe reads the map draws from)",
       all(v == 200 for v in _codes_map.values()), str(_codes_map))
    r = c.get("/api/public/areas", params={"month": 10})
    ok("⭐ §2.2b: the map cookie reads the council areas (the outside viewer is the audience, D27) — the nine polygons",
       r.status_code == 200 and len(r.json().get("areas", [])) == 9, f"{r.status_code} {r.text[:120]}")
    ok("🔴 ...but not the admin list or import", c.get("/api/admin/areas").status_code == 403
       and c.post("/api/admin/areas/import", json={"type": "FeatureCollection", "features": []}).status_code == 403)
    r = c.get("/api/public/whoami")
    ok("⭐ whoami: the map cookie is level `map`, no team", r.status_code == 200 and r.json() == {"level": "map", "team_label": None}, r.text)
    _mm = c.get("/api/meta").json()
    ok("🔴 D31: the map cookie is NOT staff — /api/meta stays the public shape (no routes, no teams, no sections) while the map still gets its factors and seasonal windows",
       _mm["scope"] == "public" and _mm["routes"] == [] and _mm["ipts"] == [] and _mm["work_sections"] == []
       and isinstance(_mm["factors"], dict) and isinstance(_mm["seasonal_restrictions"], list), str(_mm.get("scope")))
    r = c.put("/api/routes/R001/status", json={"status": "Approved"})
    ok("🔴 D18: the map cookie passes the gate on the staff write but the body's own check refuses it — 401 with the STAFF wording (sign in), nothing approved",
       r.status_code == 401 and "access code" in r.text and "password" not in r.text, f"{r.status_code} {r.text[:100]}")
    r = c.put("/api/locations/L001/capacity", json={"capacity_qty": 1, "capacity_unit": "t"})
    ok("🔴 D18: ...and the capacity write is refused by its admin check (403)", r.status_code == 403, f"{r.status_code} {r.text[:80]}")

    # 12 Oct 2026 (the review and fix pass): the app's Sign out on a shared machine. It used to clear the code in the
    # tab only, so the 12-hour STAFF cookie kept the staff shape of /api/meta, whoami "staff" and /help/ open to the
    # next person. Now it POSTs /api/gate-signout?keep=map: the staff cookie becomes a map cookie (expiring no later).
    r = c.post("/api/auth", json={"code": "plan-code-xyz"})
    ok("⭐ the STAFF cookie alone opens the monthly bulletin page (the Dashboard's link opens it in a new tab with no header)",
       r.status_code == 200 and c.get("/bulletin/?month=10").status_code == 200)
    ok("sign-out check: a planner sign-in sets the STAFF cookie (whoami staff, /help/ open, the staff /api/meta, all without a header)",
       r.status_code == 200 and c.get("/api/public/whoami").json().get("level") == "staff"
       and c.get("/help/").status_code == 200 and c.get("/api/meta").json().get("scope") == "staff")
    r = c.post("/api/gate-signout", params={"keep": "map"})
    ok("⭐ Sign out (keep=map): 200, and the same machine is now MAP level — whoami says map", r.status_code == 200
       and c.get("/api/public/whoami").json() == {"level": "map", "team_label": None}, c.get("/api/public/whoami").text)
    _mo = c.get("/api/meta").json()
    ok("🔴 ...the staff surface is closed to the next person: /api/meta is the public shape and /help/ is refused",
       _mo.get("scope") == "public" and _mo.get("routes") == [] and _mo.get("ipts") == [] and c.get("/help/").status_code == 401,
       str(_mo.get("scope")))
    ok("🔴 ...and the monthly bulletin is closed to the next person too", c.get("/bulletin/?month=10").status_code == 401)
    ok("⭐ ...while the public map still opens on that machine", c.get("/map/").status_code == 200)
    r = c.post("/api/gate-signout", params={"keep": "map"})
    ok("...a second keep=map sign-out on a map cookie leaves it alone (still map)", r.status_code == 200
       and c.get("/api/public/whoami").json().get("level") == "map")
    r = c.get("/api/admin/road-sections/status", params={"token": "é-not-the-token"})
    ok("🔴 a non-ASCII admin token is a 403, not a 500 (hmac.compare_digest on str raised TypeError)", r.status_code == 403,
       f"{r.status_code} {r.text[:80]}")

    # ------------------------------------------------------------ 5b. pilot readiness phase 2 (27 Oct 2026): the headers
    # on the real app's answers — a page, the API, the gate's refusal page, the guide, the sign-in calls, the vendored file.
    import security_headers as _sh
    _anon2 = TestClient(main.app, base_url="https://testserver")
    for _p in ("/", "/api/health", "/map/", "/help/", "/api/auth/me", "/api/meta"):
        _r = _anon2.get(_p)
        ok(f"⭐ HEADERS: {_p} ({_r.status_code}) carries every security header, the CSP whole",
           all(_r.headers.get(k) == v for k, v in _sh.HEADERS.items()), str({k: _r.headers.get(k) for k in _sh.HEADERS})[:300])
    ok("🔴 HEADERS: the gate's refusal page (a 401 HTML page) carries them too",
       _anon2.get("/map/").status_code == 401 and _anon2.get("/map/").headers.get("x-frame-options") == "SAMEORIGIN")
    ok("HEADERS: /api/auth/me is never cached; the app shell keeps its own no-cache (revalidated, not refetched)",
       _anon2.get("/api/auth/me").headers.get("cache-control") == "no-store"
       and _anon2.get("/").headers.get("cache-control") == "no-cache")
    _r = _anon2.post("/api/auth/signin", json={"email": "x@example.org", "password": "y"}, headers={"X-Wayscope": "1"})
    ok("HEADERS: a refused sign-in is no-store as well", _r.status_code == 401 and _r.headers.get("cache-control") == "no-store")
    _r = _anon2.get("/vendor/tailwind-play-3.4.17.js")
    ok("⭐ VENDOR: the app serves Tailwind 3.4.17 itself, as JavaScript, unchanged",
       _r.status_code == 200 and "javascript" in (_r.headers.get("content-type") or "")
       and len(_r.content) == os.path.getsize(os.path.join(ROOT, "frontend", "vendor", "tailwind-play-3.4.17.js")), str(_r.status_code))
    ok("🔴 VENDOR: nothing else under /vendor/ (a missing file is a 404)", _anon2.get("/vendor/nope.js").status_code == 404)

    # ------------------------------------------------------------ 6. G3: per-person logins (14 Oct 2026)
    # The real app: the cookie the browser gets, the middleware's guard header and audit row, CORS, the package
    # re-import leaving the logins alone, HU14's link. users.py's every branch is in test_users.py.
    import users as _users
    c.cookies.clear()
    W = {"X-Wayscope": "1"}
    r = c.post("/api/auth/signin", json={"email": "nobody@example.org", "password": "whatever pass"})
    ok("🔴 G3: an /api/auth/ call without the app's header is refused (403) before anything is checked",
       r.status_code == 403 and "Wayscope app" in r.text, f"{r.status_code} {r.text[:80]}")
    r = c.post("/api/auth/signin", json={"email": "nobody@example.org", "password": "whatever pass"}, headers=W)
    ok("G3: an unknown address is the one refusal, 401", r.status_code == 401 and "don't match" in r.text, r.text[:120])
    r = c.get("/api/auth/me")
    ok("G3: /api/auth/me signed out says so, and that the shared codes are still on (they are set here)",
       r.status_code == 200 and r.json() == {"signed_in": False, "access": None, "codes": True}, r.text[:200])
    # the first admin, D47: the admin token on the Users endpoints
    r = c.post("/api/admin/users", params=TOKEN, json={"email": "Ada@Example.org", "name": "Ada Lovelace", "role": "admin"})
    ok("⭐ G3 D47: the admin token adds the first admin; with no sender the invite link comes back",
       r.status_code == 200 and r.json()["email"] == "unsent" and "#invite=" in r.json().get("link", ""), r.text[:200])
    ok("G3: Ada had asked for a demo above (a demo login): adding her as staff CONVERTS that login — no end date, Admin",
       r.json().get("converted") is True and r.json()["user"]["kind"] == "staff" and r.json()["user"]["expires_at"] is None
       and r.json()["user"]["role"] == "admin")
    ok("🔴 ...and her old personal demo link no longer signs anyone in",
       c.post("/api/auth/link", json={"token": _screen_link}, headers={"X-Wayscope": "1"}).status_code == 410)
    _inv = r.json()["link"].split("#invite=", 1)[1]
    r = c.post("/api/admin/users", json={"email": "x@example.org", "name": "X", "role": "planner"}, headers=W)
    ok("🔴 G3: without the token or an admin login, adding a login is refused (403)", r.status_code == 403)
    r = c.post("/api/auth/check", json={"token": _inv}, headers=W)
    ok("G3: the invite greets Ada", r.status_code == 200 and r.json()["name"] == "Ada Lovelace", r.text[:120])
    r = c.post("/api/auth/reset", json={"token": _inv, "password": "correct horse battery"}, headers=W)
    # 27 Oct 2026 (F3 b): the password is set, but a staff login is signed in only after two-step sign-in is set up
    ok("⭐ F3 b: the invite's password answers the SET-UP step: a ticket, the key, a QR code — and no session cookie",
       r.status_code == 200 and r.json().get("step") == "setup" and r.json().get("ticket") and "<svg" in r.json().get("qr", "")
       and "wayscope_session=" not in r.headers.get("set-cookie", ""), r.text[:200])
    _ada_key = r.json()["key"].replace(" ", "")
    _t = r.json()["ticket"]
    r = c.post("/api/auth/setup", json={"ticket": _t, "code": "000000"}, headers=W)
    ok("🔴 F3 b: a wrong set-up code is refused in words (401), no cookie",
       r.status_code == 401 and "didn't work" in r.text and "wayscope_session=" not in r.headers.get("set-cookie", ""), r.text[:120])
    r = c.post("/api/auth/setup", json={"ticket": _t, "code": "123456"})
    ok("🔴 F3 b: the set-up step needs the app's header like every /api/auth/ call", r.status_code == 403)
    r = c.post("/api/auth/setup", json={"ticket": _t, "code": _users.totp_at(_ada_key, _users.totp_step())}, headers=W)
    _ada_codes = r.json().get("recovery_codes") or []
    ok("⭐ F3 b: the right code turns two-step on and answers ten recovery codes",
       r.status_code == 200 and len(_ada_codes) == 10 and r.json().get("two_step") is True, r.text[:200])
    _sc = r.headers.get("set-cookie", "")
    ok("⭐ G3: setting the password signs Ada in; the session cookie is HttpOnly, Secure, SameSite=Lax, path /, 12 hours",
       r.status_code == 200 and r.json()["via"] == "login" and r.json()["role"] == "admin"
       and "wayscope_session=" in _sc and "HttpOnly" in _sc and "Secure" in _sc and "samesite=lax" in _sc.lower()
       and "Path=/" in _sc and "Max-Age=43200" in _sc, _sc[:300])
    ok("🔴 ...and the raw session value is in no table", _sc.split("wayscope_session=", 1)[1].split(";", 1)[0]
       not in json.dumps(db.query("SELECT * FROM user_sessions")))
    r = c.get("/api/auth/me")
    ok("G3: the cookie alone (no header) is who you are", r.json()["signed_in"] and r.json()["access"]["email"] == "ada@example.org")
    ok("⭐ G3: the session opens what the staff cookie opened: /help/, /bulletin/, /map/, the staff /api/meta",
       c.get("/help/").status_code == 200 and c.get("/bulletin/?month=10").status_code == 200
       and c.get("/map/").status_code == 200 and c.get("/api/meta").json()["scope"] == "staff")
    ok("G3: whoami says staff for a login", c.get("/api/public/whoami").json()["level"] == "staff")
    # D53: an admin LOGIN saves Config with no token; the audit row names Ada
    _doc = c.get("/api/config/factors").json()["doc"]
    r = c.put("/api/admin/config/factors", json={"doc": _doc, "updated_by": "Mallory"})
    ok("🔴 G3: a cookie write WITHOUT the app's header is refused (403), nothing saved",
       r.status_code == 403 and "Wayscope app" in r.text, f"{r.status_code} {r.text[:80]}")
    r = c.put("/api/admin/config/factors", json={"doc": _doc, "updated_by": "Mallory"}, headers=W)
    ok("⭐ G3 D53: an admin LOGIN saves Config with no admin token", r.status_code == 200, f"{r.status_code} {r.text[:120]}")
    ok("⭐ G3: ...recorded as Ada, whatever the browser typed (main._by)",
       c.get("/api/config/factors").json()["status"].get("updated_by") == "Ada Lovelace")
    _a = c.get("/api/admin/audit", params={"kind": "write"}).json()["rows"]
    ok("⭐ G3 D55: the middleware wrote ONE audit row for that write: Ada, the path, the body; the refused one is not there",
       len([x for x in _a if x["path"] == "/api/admin/config/factors" and x["actor_kind"] == "user"]) == 1
       and _a[0]["actor_label"] == "Ada Lovelace <ada@example.org>" and _a[0]["method"] == "PUT" and _a[0]["status"] == 200
       and _a[0]["detail"].startswith('{"body":{"doc":{') and len(_a[0]["detail"]) <= 4000, json.dumps(_a[:2])[:400])
    # a team login: the line filter works on a session
    r = c.post("/api/admin/users", json={"email": "bob@example.org", "name": "Bob North", "role": "team", "team": "IPT1"}, headers=W)
    _bob_inv = r.json()["link"].split("#invite=", 1)[1]
    r = c.post("/api/admin/users", json={"email": "pat@example.org", "name": "Pat Planner", "role": "planner"}, headers=W)
    ok("G3: Ada (a login) adds a team login and a planner, no token", r.status_code == 200)
    _ul = c.get("/api/admin/users").json()
    ok("G3: the Users list carries the team's NAME beside the slot, and no hash",
       [u for u in _ul["users"] if u["email"] == "bob@example.org"][0]["team_label"] not in (None, "", "IPT1")
       and "password_hash" not in json.dumps(_ul) and "scrypt$" not in json.dumps(_ul))
    # 23 Oct 2026 — Users -> Teams through the real app (admins only; at most six; John)
    _tl = c.get("/api/admin/teams").json()
    ok("TEAMS: Ada (an admin login, no token) reads the demo's three teams by name, three of six used",
       [t["label"] for t in _tl["teams"]] == ["North Team", "Central Team", "South Team"] and _tl["used"] == 3 and _tl["max"] == 6, json.dumps(_tl)[:200])
    r = c.post("/api/admin/teams", json={"label": "Drainage package"}, headers=W)
    _new = r.json().get("team") or {}
    ok("⭐ TEAMS: Ada adds a fourth team by its name; it takes a slot nothing has carried",
       r.status_code == 200 and _new.get("label") == "Drainage package" and _new.get("id") not in ("IPT1", "IPT2", "IPT3"), r.text[:200])
    ok("TEAMS: ...it is in the staff team pickers at once (by name)", "Drainage package" in [i["label"] for i in c.get("/api/meta").json()["ipts"]])
    _a = c.get("/api/admin/audit", params={"kind": "write"}).json()["rows"]
    ok("⭐ TEAMS: the middleware recorded the add: Ada, the path, the name",
       any(x["path"] == "/api/admin/teams" and x["method"] == "POST" and x["actor_label"] == "Ada Lovelace <ada@example.org>"
           and "Drainage package" in (x["detail"] or "") for x in _a), json.dumps(_a[:2])[:300])
    r = c.post("/api/admin/teams", json={"label": "north team"}, headers=W)
    ok("TEAMS: a second team with a name already used is refused, in words", r.status_code == 400 and "already a team with that name" in r.text, r.text[:120])
    r = c.patch(f"/api/admin/teams/{_new.get('id')}", json={"label": "Drainage and ducts"}, headers=W)
    ok("TEAMS: rename", r.status_code == 200 and r.json()["team"]["label"] == "Drainage and ducts", r.text[:160])
    r = c.patch("/api/admin/teams/IPT1", json={"active": False}, headers=W)
    ok("🔴 TEAMS: North Team cannot be stopped while its logins can sign in", r.status_code == 400 and "first" in r.text, r.text[:160])
    r = c.patch(f"/api/admin/teams/{_new.get('id')}", json={"active": False}, headers=W)
    ok("TEAMS: the new, empty team can be stopped; it leaves the pickers", r.status_code == 200
       and "Drainage and ducts" not in [i["label"] for i in c.get("/api/meta").json()["ipts"]])
    r = c.post("/api/admin/users", json={"email": "dee@example.org", "name": "Dee Drain", "role": "team", "team": _new.get("id")}, headers=W)
    ok("🔴 TEAMS: nobody can be invited into a stopped team", r.status_code == 400, f"{r.status_code} {r.text[:120]}")
    ok("TEAMS: the rest of the page reads the demo's three teams again", len(c.get("/api/meta").json()["ipts"]) == 3)
    _ada_jar = dict(c.cookies)
    c.cookies.clear()
    r = c.post("/api/auth/reset", json={"token": _bob_inv, "password": "bob's long password"}, headers=W)
    _bob_key = r.json()["key"].replace(" ", "")
    r = c.post("/api/auth/setup", json={"ticket": r.json()["ticket"], "code": _users.totp_at(_bob_key, _users.totp_step())}, headers=W)
    _bob_codes = r.json().get("recovery_codes") or []
    ok("G3: Bob sets his password and is signed in as his team", r.status_code == 200 and r.json()["ipt"] == "IPT1"
       and r.json()["ipt_locked"] is True)
    _bob_lines = c.get("/api/forecasts").json()
    ok("⭐ G3: a team LOGIN sees only its own team's lines, exactly as the team code does",
       len(_bob_lines) > 0 and all(x.get("ipt") == "IPT1" for x in _bob_lines)
       and len(_bob_lines) == len(c.get("/api/forecasts", headers={"X-Access-Code": "north-code-xyz"}).json()))
    r = c.put("/api/routes/R001/status", json={"status": "Approved"}, headers=W)
    ok("🔴 G3: a team login cannot approve (403)", r.status_code == 403, f"{r.status_code} {r.text[:80]}")
    ok("🔴 G3: ...nor reach Users", c.get("/api/admin/users").status_code == 403)
    r = c.post("/api/auth/signout", headers=W)
    ok("⭐ G3 sign out: the session row is gone, the cookie dropped, and the map still opens on this machine (map cookie)",
       r.status_code == 200 and c.get("/api/auth/me").json()["signed_in"] is False
       and c.get("/api/public/whoami").json()["level"] == "map" and c.get("/help/").status_code == 401)
    # 27 Oct 2026 (F3 b): the code step through the real app — password, then a recovery code (the app's own code
    # was used seconds ago at set-up; a code works once), and an admin's reset of a lost phone
    c.cookies.clear()
    r = c.post("/api/auth/signin", json={"email": "bob@example.org", "password": "bob's long password"}, headers=W)
    ok("⭐ F3 b: with two-step on, the password answers the CODE step and no session cookie",
       r.status_code == 200 and r.json().get("step") == "code" and "key" not in r.json()
       and "wayscope_session=" not in r.headers.get("set-cookie", ""), r.text[:200])
    _t = r.json()["ticket"]
    ok("🔴 F3 b: the ticket alone opens nothing", c.get("/api/auth/me").json()["signed_in"] is False
       and c.get("/api/forecasts", cookies={"wayscope_session": _t}).status_code in (401, 403))
    r = c.post("/api/auth/code", json={"ticket": _t, "code": _users.totp_at(_bob_key, _users.totp_step())}, headers=W)
    ok("🔴 F3 b: the code used seconds ago at set-up is refused (a code works once)", r.status_code == 401, r.text[:120])
    r = c.post("/api/auth/code", json={"ticket": _t, "code": _bob_codes[0]}, headers=W)
    ok("⭐ F3 b: a recovery code signs Bob in; nine left", r.status_code == 200 and r.json()["recovery_left"] == 9
       and "wayscope_session=" in r.headers.get("set-cookie", ""), r.text[:200])
    _bob_jar = dict(c.cookies)
    c.cookies.clear(); c.cookies.update(_ada_jar)
    r = c.post(f"/api/admin/users/{[u for u in _ul['users'] if u['email'] == 'bob@example.org'][0]['id']}/two-step-reset", headers=W)
    ok("⭐ F3 b: Ada resets Bob's two-step (a lost phone): it is off on the Users list",
       r.status_code == 200 and r.json()["user"]["two_step"] is False, r.text[:200])
    _rr = [x for x in c.get("/api/admin/audit").json()["rows"] if (x.get("path") or "").endswith("/two-step-reset")]
    ok("F3 b: ...recorded, naming Ada", _rr and "Ada Lovelace" in _rr[0]["actor_label"], json.dumps(_rr[:1])[:200])
    c.cookies.clear(); c.cookies.update(_bob_jar)
    ok("🔴 F3 b: ...and Bob's session ended at once", c.get("/api/auth/me").json()["signed_in"] is False)
    # rate limit: five failures lock the address
    c.cookies.clear()
    for _i in range(5):
        c.post("/api/auth/signin", json={"email": "bob@example.org", "password": "not his password"}, headers=W)
    r = c.post("/api/auth/signin", json={"email": "bob@example.org", "password": "bob's long password"}, headers=W)
    ok("⭐ G3: after five failures the right password is refused for 15 minutes — 429", r.status_code == 429, r.text[:120])
    # the shared codes keep working beside logins (D49)
    ok("G3 D49: the planner code still signs in while it is set",
       c.post("/api/auth", json={"code": "plan-code-xyz"}).status_code == 200)
    _sg = c.get("/api/admin/audit", params=dict(TOKEN, kind="signin-code")).json()["rows"]
    ok("G3 D49: ...and that sign-in is recorded, as the code", _sg and _sg[0]["actor_label"] == "Planner code (shared)", json.dumps(_sg[:1]))
    r = c.put("/api/routes/R001/status", json={"status": "Approved"}, headers=PLAN)
    _cr = c.get("/api/admin/audit", params=dict(TOKEN, kind="write")).json()["rows"][0]
    ok("G3: a write on a shared code needs no extra header, and is recorded as the code",
       r.status_code == 200 and _cr["actor_label"] == "Planner code (shared)" and _cr["path"] == "/api/routes/R001/status",
       json.dumps(_cr)[:200])
    _tok_rows = [x for x in c.get("/api/admin/audit", params=dict(TOKEN, kind="write")).json()["rows"] if x["path"] == "/api/admin/users"]
    ok("G3: the first-admin add (by token) is recorded as the admin token, without the token's value",
       any(x["actor_label"] == "admin token" for x in _tok_rows) and "adm-token-xyz" not in json.dumps(_tok_rows))
    # CORS
    r = c.get("/api/meta", headers={"Origin": "https://evil.example", "Cookie": "x=1"})
    ok("🔴 G3: CORS no longer answers another origin (no Access-Control-Allow-Origin, no credentials)",
       "access-control-allow-origin" not in {k.lower() for k in r.headers.keys()}
       and "access-control-allow-credentials" not in {k.lower() for k in r.headers.keys()}, str(dict(r.headers))[:300])
    r = c.options("/api/forecasts/bulk", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST",
                                                    "Access-Control-Request-Headers": "x-wayscope,content-type"})
    ok("🔴 G3: ...and a preflight from another origin is not granted", r.headers.get("access-control-allow-origin") is None,
       f"{r.status_code} {dict(r.headers)}")
    # the package re-import leaves the logins
    _n_users = len(db.query("SELECT * FROM users"))
    _pkg = c.get("/api/admin/tenant/export", params=TOKEN).json()
    r = c.post("/api/admin/tenant/import", params=dict(TOKEN, replace=1), json=_pkg)
    ok("⭐ G3: re-importing the tenant (replace) leaves every login and the audit trail in place",
       r.status_code == 200 and len(db.query("SELECT * FROM users")) == _n_users >= 3
       and len(db.query("SELECT * FROM audit_log")) > 0, f"{r.status_code} {_n_users}")
    # HU14: the lead's link from the form above signs them in, and their page opens are recorded
    c.cookies.clear()
    r = c.post("/api/auth/link", json={"token": _bea_link}, headers=W)
    ok("⭐ G3 HU14: the emailed link signs the lead in as a demo login on the team slot, for its 14 days",
       r.status_code == 200 and r.json()["kind"] == "demo" and r.json()["ipt"] == "IPT1"
       and "Max-Age=1209" in r.headers.get("set-cookie", ""), f"{r.status_code} {r.headers.get('set-cookie', '')[:200]}")
    ok("G3 HU14: the lead's page open is recorded", c.post("/api/activity", json={"page": "lookahead"}, headers=W).json()["recorded"] is True)
    c.cookies.clear()
    r = c.get("/api/admin/demo-requests.csv", params=TOKEN)
    _bea_line = [l for l in r.text.split("\n") if "bea@example.org" in l][0]
    ok("⭐ G3 HU14: the leads CSV says who looked at what — login, expiry, sign-ins, pages",
       r.text.split("\n")[0].endswith(",login,login_expires_at,sign_ins,first_seen,last_seen,pages")
       and ",active," in _bea_line and _bea_line.endswith(",lookahead 1"), _bea_line)

    # ---- 28 Oct 2026 (pilot readiness phase 3): the outside hand-over check against the REAL app ----
    # backend/tools/check_pilot.py asks a deployment what anyone can ask. Here its requests go to this app (Render's
    # http -> https redirect and the address lookup cannot be had locally, so those two rows are not asserted).
    sys.path.insert(0, os.path.join(BACKEND, "tools"))
    import check_pilot  # noqa: E402

    def _cp_fetch(method, url, data=None, headers=None, timeout=30):
        if url.startswith("http://"):
            return 301, {"location": "https://" + url[len("http://"):]}, ""
        path = url[len("https://pilot.example.test"):] or "/"
        rr = c.request(method, path, content=data, headers=dict(headers or {}), follow_redirects=False)
        return rr.status_code, {k.lower(): v for k, v in rr.headers.items()}, rr.text

    c.cookies.clear()
    _cp_saved = {k: os.environ.pop(k, None) for k in ("ADMIN_TOKEN", "DEMO_LEADS")}
    try:
        _cp = dict((w.split(":")[0].split(" (")[0], r) for r, w in check_pilot.checks("https://pilot.example.test", fetch=_cp_fetch))
        _n_req = len(db.query("SELECT * FROM demo_requests"))
        os.environ["ADMIN_TOKEN"] = "adm-token-xyz"
        _cp_tok = dict((w.split(":")[0].split(" (")[0], r) for r, w in check_pilot.checks("https://pilot.example.test", fetch=_cp_fetch))
        os.environ["DEMO_LEADS"] = "1"
        _cp_demo = dict((w.split(":")[0].split(" (")[0], r) for r, w in
                        check_pilot.checks("https://pilot.example.test", fetch=_cp_fetch, demo=True))
        _n_req2 = len(db.query("SELECT * FROM demo_requests"))
    finally:
        for k, v in _cp_saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
    _want = ("Health", "Sign-in page", "Security headers on the app page", "The page refuses to be framed by another site",
             "Sign-in answers are never cached", "Admin calls refused without an admin login", "The landing form is off here",
             "The public map asks for its password", "The guides ask for a staff sign-in")
    ok("⭐ phase 3: check_pilot passes every outside check on an organisation's settings (no token, form off)",
       all(_cp.get(w) == "pass" for w in _want), json.dumps(_cp))
    ok("🔴 phase 3: ...and FAILS the admin row when ADMIN_TOKEN is set (the 403 words tell a set token from none)",
       _cp_tok.get("Admin calls refused without an admin login") == "FAIL", json.dumps(_cp_tok))
    ok("phase 3: ...on the demo (--demo) a set token is a note and the form must be ON",
       _cp_demo.get("Admin calls refused without an admin login") == "note"
       and _cp_demo.get("The landing form is ON here") == "pass", json.dumps(_cp_demo))
    ok("🔴 phase 3: the form check (its trap field only) stores no request, form on or off", _n_req2 == _n_req,
       f"{_n_req} -> {_n_req2}")

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\n{PASS} passed, {len(FAIL)} failed")
sys.exit(1 if FAIL else 0)
