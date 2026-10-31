"""
Notes and tasks on each CLP, 31 Oct 2026 (task 4a, PR 3b) — backend/clp_notes.py, the steps and kept lines written by
clp_records, and the /api/clp/documents/{id}/notes|tasks endpoints. Mockup notes/clp-notes-mockup-1031/ (Q1–Q12).

What is asserted: who may write a note, add a task and tick it (everyone who can see the CLP; a team its own team's only);
who deletes a note or changes a task (its writer, a planner, an admin); what a task can be For (the CLP's team, the planners
and admins, a person who can see it — never a demo lead's login, never one who cannot see it); the one email (a task given
to one named person other than yourself); the Tasks column's counts (open, overdue, given to me); the approval steps as
grey lines, written as they happen and read back from the CLP and its issues for the steps before 3b, never twice; the
kept "Still to complete" lines and a task started from one, shown as gone when a later build drops it, never ticked by
itself; the approve card's open tasks; Ask for changes adding its note as a task; Delete of a + New CLP removing them,
clearing a listed row keeping them; the table in the organisation file and kept by a demo refresh; a note ending no kept
read (NO_DATA_TABLES and the middleware's path rule).

WHAT THIS DOES NOT PROVE
------------------------
  * The HTTP layer is stubbed (http_smoke.py drives the real one); the page is checked in Chrome by hand.
  * SQLite only. The email is captured, never sent: a real send from Render is not tested.

Run (PowerShell):  $env:PYTHONUTF8 = "1"; .venv-tests\\Scripts\\python backend/tests/test_clp_notes.py
"""
import copy
import json
import io
import math
import os
import re
import sys
import tempfile
import types
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, as test_gate_arrivals.py ----------------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)


TMP = tempfile.mkdtemp(prefix="rbe_clpnotes_")
from portable import remove_scratch_db, offline   # noqa: E402,F401  tests/portable.py
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import costing  # noqa: E402
import config  # noqa: E402
import gates  # noqa: E402
import gate_arrivals  # noqa: E402
import gate_hours  # noqa: E402
import network  # noqa: E402
import site_areas  # noqa: E402
import taxonomy  # noqa: E402
import carbon_target  # noqa: E402
import clp_doc  # noqa: E402
import export  # noqa: E402
offline(export)            # PR 2 draws plans: Mapbox is never called from a test (the schematic branch)
from docx import Document  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


ART = network.DEFAULT_PROFILE          # "Artic Tipper (44t)", 29 t a load
FLAT = "Artic Flatbed (44t)"
BAL = "Large aggregate / ballast"
SPOIL = "Earthworks / soil"
PRE = "Precast units"

# =========================================================================== #
#  0. A small scheme by hand                                                   #
# =========================================================================== #
remove_scratch_db(db, TMP)
db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db(); db.init_gates_db()
db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db(); db.init_costing_db()
db.init_tenant()
taxonomy.seed_taxonomy()
config.invalidate()
config.seed_from_file(conversions)
config.invalidate()
costing.invalidate()
carbon_target.invalidate()
access.set_current("planner123")
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc["planning"]["peak_bands"] = [{"from": "08:00", "to": "09:00"}, {"from": "15:00", "to": "16:00"}]
_doc.setdefault("tenant", {})["demo_notice"] = "Fictional scheme"
_doc["tenant"]["name"] = "Test scheme"
for _k, _v in _doc["vehicles"].items():
    if isinstance(_v, dict):
        _v["euro_class"] = "Euro VI"
ok("a scheme with two peak bands, a demo notice and every lorry Euro VI passes Config", config.validate(_doc) == [],
   str(config.validate(_doc)))
config.save(_doc, by="tester")
config.invalidate()
factors = conversions.load_factors()
WD = 22.0

PIT = network.create_location("Pit West", "origin", lat=52.55, lon=-0.95, loc_type="Quarry", supplies=[BAL])["id"]
YARD = network.create_location("Yard East", "destination", lat=52.6447, lon=-0.8822, loc_type="Compound", receives=[BAL, PRE], supplies=[BAL, SPOIL])["id"]
PILE = network.create_location("Pile North", "destination", lat=52.58, lon=-0.92, loc_type="Stockpile", receives=[BAL])["id"]
TIP = network.create_location("Tip South", "destination", lat=52.50, lon=-0.90, loc_type="Landfill", receives=[SPOIL])["id"]
RAIL = network.create_location("Rail Head", "origin", lat=52.48, lon=-0.92, loc_type="Railhead", supplies=[PRE])["id"]
R_IN = network.create_route(PIT, YARD, material_category=BAL, route_id="R001", ipt="IPT1")["id"]
R_OUT = network.create_route(YARD, TIP, material_category=SPOIL, route_id="R002", ipt="IPT1")["id"]
R_BTW = network.create_route(YARD, PILE, material_category=BAL, route_id="R003", ipt="IPT2")["id"]
R_RAIL = network.create_route(RAIL, YARD, material_category=PRE, route_id="R004", ipt="IPT2")["id"]
R_ELSE = network.create_route(PIT, PILE, material_category=BAL, route_id="R005", ipt="IPT1")["id"]


def add_geom(rid, km, profile):
    for leg in ("loaded", "return"):
        db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
                   "duration_hr, computed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                   (db.current_tenant(), rid, profile, leg, 0, "[]", km, 0.5, "2026-10-30T00:00:00Z"))


add_geom(R_IN, 32.5, ART)
add_geom(R_OUT, 20.0, ART)
add_geom(R_BTW, 6.0, ART)
add_geom(R_RAIL, 22.1, FLAT)          # R005 is never drawn


def _line(rid, disc, sect, month, veh, ipt, mat, approve=True):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline=disc, section_id=sect, material_type=mat, material_description=None,
        vehicle_type=veh, submitted_by="tester", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=q) for m, q in month], ipt=ipt))
    if approve:
        main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline=disc, section_id=sect)


# month 10 = Oct 2026, 11 = Nov 2026, 13 = Jan 2027 (a new quarter and a new year)
_line(R_IN, "substructure", "WS1", [(10, 2900.0), (11, 5800.0), (13, 1450.0)], ART, "IPT1", BAL)   # 100, 200, 50 loads
_line(R_OUT, "earthworks", "WS2", [(11, 2900.0)], ART, "IPT1", SPOIL)                             # 100 loads out
_line(R_BTW, "substructure", "WS3", [(10, 580.0)], ART, "IPT2", BAL)                              # 20 loads yard -> pile
_line(R_RAIL, "structures", "WS4", [(13, 1100.0)], FLAT, "IPT2", PRE)                             # precast off the railhead
_line(R_ELSE, "earthworks", "WS5", [(10, 2900.0)], ART, "IPT1", BAL)                              # pit -> pile only
_line(R_IN, "earthworks", "WS6", [(12, 290.0)], ART, "IPT1", BAL, approve=False)                  # waiting for approval
config.invalidate()
factors = conversions.load_factors()

G_IN = gates.create_gate(YARD, "Yard gate in", 52.6447, -0.8822, direction="access", safety_minutes=3,
                         is_default=True, open_from="07:30", open_to="18:00", open_days="1111100",
                         note="Check in here")
G_OUT = gates.create_gate(YARD, "Yard gate out", 52.6449, -0.8820, direction="both", open_from="07:30",
                          open_to="18:00", closed_from="2027-03-01", closed_to="2027-05-31")
ok("the two gates are made", "id" in G_IN and "id" in G_OUT, f"{G_IN} {G_OUT}")
G_IN, G_OUT = G_IN.get("id"), G_OUT.get("id")
_ar = site_areas.create_area(YARD, "Pier bases", activity="unload", service_minutes=5, bays=1, level="river terrace")
ok("the work area is made", "error" not in _ar, str(_ar))
db.execute("INSERT INTO forecast_weeks (tenant_id, route_id, month_index, discipline, section_id, week_index, planned_qty, "
           "unit, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
           (db.current_tenant(), R_IN, 10, "substructure", "WS1", 40, 725.0, "t", "confirmed"))


def trips(rid, disc, sect, m):
    row = db.query("SELECT * FROM forecasts WHERE tenant_id = ? AND route_id = ? AND discipline = ? AND section_id = ? "
                   "AND month_index = ?", (db.current_tenant(), rid, disc, sect, m))[0]
    return gate_arrivals.trips_for(row, factors)



import base64  # noqa: E402
import clp_records  # noqa: E402
import tenant_package  # noqa: E402
for _tid, _lab in (("IPT1", "North Team"), ("IPT2", "Central Team")):
    db.execute("INSERT INTO ipts (tenant_id, id, label, manager, active, merged_into) VALUES (?, ?, ?, ?, ?, ?)",
               (db.current_tenant(), _tid, _lab, None, True, None))
PLANNER = {"role": "planner", "ipt": None, "label": "Planner", "name": "Pat Planner"}
NORTH = {"role": "ipt", "ipt": "IPT1", "label": "North Team"}
CENTRAL = {"role": "ipt", "ipt": "IPT2", "label": "Central Team"}
READER = {"role": "submitter", "ipt": None, "label": "Submitter"}


def as_(acc):
    access._ACCESS.set(acc)


def text_of(body):
    d = Document(io.BytesIO(body))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    head = " ".join(p.text for s in d.sections for p in s.header.paragraphs)
    foot = " ".join(p.text for s in d.sections for p in s.footer.paragraphs)
    return "\n".join(parts), head, foot


def build_fast(data, rec):
    return clp_doc.build(data, impact=rec["impact"], pictures=False)


def count_rows():
    return db.query("SELECT COUNT(*) AS n FROM clp_documents WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]


import clp_notes  # noqa: E402
import users  # noqa: E402
import demo_request  # noqa: E402
db.init_users_db()

def login(email, name, role, team=None, kind="staff"):
    u, err = users.create_user(email, name, role, team=team, kind=kind)
    assert u, err
    return u


U_NORA = login("nora@example.com", "Nora North", "team", "IPT1")
U_NED = login("ned@example.com", "Ned North", "team", "IPT1")
U_CAL = login("cal@example.com", "Cal Central", "team", "IPT2")
U_PAT = login("pat@example.com", "Pat Planner", "planner")
U_ADA = login("ada@example.com", "Ada Admin", "admin")
U_LEAD = login("lead@example.com", "Lee Lead", "team", "IPT1", kind="demo")
U_GONE = login("gone@example.com", "Gus Gone", "team", "IPT1")
db.execute("UPDATE users SET disabled = 1 WHERE tenant_id = ? AND id = ?", (db.current_tenant(), U_GONE["id"]))
NORA = users.access_for(users.get_user(U_NORA["id"]))
NED = users.access_for(users.get_user(U_NED["id"]))
CAL = users.access_for(users.get_user(U_CAL["id"]))
PAT = users.access_for(users.get_user(U_PAT["id"]))
LEAD = users.access_for(users.get_user(U_LEAD["id"]))
NORTH_CODE = {"role": "ipt", "ipt": "IPT1", "label": "North Team"}
READER = {"role": "submitter", "ipt": None, "label": "Submitter"}

SENT = []
users.mail_settings = lambda: {"DEMO_FROM": "Wayscope <hello@example.com>", "reply_to": "hello@example.com"}
demo_request.send = lambda msg, settings: SENT.append(msg)

Y1 = f"auto-{YARD}-IPT1"          # the North Team's CLP for the yard
Y2 = f"auto-{YARD}-IPT2"          # the Central Team's
YA = f"auto-{YARD}-all"           # every team's


def th(clp_id, acc):
    return clp_notes.thread(clp_id, acc, factors)


def row_of(clp_id, acc):
    return next(c for c in clp_records.list_for(acc, factors)["clps"] if c["id"] == clp_id)


# =========================================================================== #
#  1. The table: tenanted, in the organisation file, never a kept read's       #
# =========================================================================== #
ok("the table is tenanted, in the organisation file after clp_issues, kept by a demo refresh",
   "clp_notes" in db.TENANTED_TABLES
   and tenant_package.ORDER.index("clp_notes") > tenant_package.ORDER.index("clp_issues")
   and '"clp_notes"' in open(os.path.join(BACKEND, "tools", "make_demo_tenant.py"), encoding="utf-8").read()
   .split("LIVE_TABLES = (")[1].split(")")[0])
ok("🔴 no kept read is built from it (NO_DATA_TABLES), so a note never ends a kept CLP draft",
   "clp_notes" in db.NO_DATA_TABLES)
ok("...and the middleware's second net skips the notes and tasks paths (not Send, Changes or Approve)",
   main._CLP_NOTES_PATH.match(f"/api/clp/documents/{Y1}/notes") and main._CLP_NOTES_PATH.match(f"/api/clp/documents/{Y1}/tasks/n-1")
   and main._CLP_NOTES_PATH.match(f"/api/clp/documents/{Y1}/notes/n-1")
   and not main._CLP_NOTES_PATH.match(f"/api/clp/documents/{Y1}/changes")
   and not main._CLP_NOTES_PATH.match(f"/api/clp/documents/{Y1}") and not main._CLP_NOTES_PATH.match("/api/clp/documents"))
ok("...while every notes and tasks write is still in the audit trail", main._audited(f"/api/clp/documents/{Y1}/notes")
   and main._audited(f"/api/clp/documents/{Y1}/tasks/n-1"))

# =========================================================================== #
#  2. For whom (Q4): the team, the planners and admins, a person who can see it #
# =========================================================================== #
as_(PAT)
F = {o["key"]: o["word"] for o in th(Y1, PAT)["for_options"]}
ok("⭐ a team CLP's task can be for its team (by name), the planners and admins, or a person who can see it",
   F.get("team") == "North Team" and F.get("approvers") == "Planners and admins"
   and {F.get(f"user:{u['id']}") for u in (U_NORA, U_NED, U_PAT, U_ADA)} == {"Nora North", "Ned North", "Pat Planner", "Ada Admin"},
   str(F))
ok("🔴 ...never a person who cannot see it (another team), a demo lead's login or a disabled one",
   not any(f"user:{u['id']}" in F for u in (U_CAL, U_LEAD, U_GONE)), str(F))
FA = {o["key"] for o in th(YA, PAT)["for_options"]}
ok("...an every-team CLP has no 'team' choice, and only the planners and admins see it among the logins",
   "team" not in FA and f"user:{U_PAT['id']}" in FA and f"user:{U_NORA['id']}" not in FA and f"user:{U_CAL['id']}" not in FA, str(FA))
FL = [o["key"] for o in clp_notes.thread(Y1, LEAD, factors)["for_options"]]
ok("🔴 a demo lead is offered no person's name", FL == ["team", "approvers"], str(FL))

# =========================================================================== #
#  3. Notes (Q3, Q5)                                                           #
# =========================================================================== #
as_(NORA)
r = clp_notes.add_note(Y1, "  The council asked for the Saturday hours.  ", NORA, factors)
ok("⭐ a team member writes a note on the team's CLP; it reads back with who and when",
   r.get("thread") and r["thread"]["notes"][-1]["words"] == "The council asked for the Saturday hours."
   and r["thread"]["notes"][-1]["by"] == "Nora North" and r["thread"]["notes"][-1]["at"], str(r)[:300])
N_NORA = r.get("id")
ok("🔴 not on another team's CLP, nor the every-team one (not found, like one that does not exist)",
   clp_notes.add_note(Y2, "x", NORA, factors).get("status") == 404 and clp_notes.add_note(YA, "x", NORA, factors).get("status") == 404
   and clp_notes.thread(Y2, NORA, factors) is None and clp_notes.thread(YA, NORA, factors) is None)
ok("🔴 an empty note and a banned word are refused in words",
   clp_notes.add_note(Y1, "  ", NORA, factors).get("error") == "Write a note first."
   and "secure" in clp_notes.add_note(Y1, "Make the gate secure", NORA, factors).get("error", ""))
r = clp_notes.add_note(Y1, "y" * 1500, NORA, factors)
ok("...a long one is cut at 1,000 characters", len(r["thread"]["notes"][-1]["words"]) == 1000)
N_LONG = r["id"]
ok("a sign-in with a shared team code writes on its team's CLP too (by the team's name)",
   clp_notes.add_note(Y1, "Gate B shut on Friday.", NORTH_CODE, factors)["thread"]["notes"][-1]["by"] == "North Team")
r = clp_notes.add_note(Y1, "Thanks, will add them.", PAT, factors)
N_PAT = r["id"]
ok("...and a planner on any CLP", r["thread"]["notes"][-1]["by"] == "Pat Planner")
ok("🔴 a team member cannot delete a planner's note, nor another member's",
   clp_notes.remove(Y1, N_PAT, NORA, factors).get("status") == 403 and clp_notes.remove(Y1, N_NORA, NED, factors).get("status") == 403)
_t = th(Y1, NORA)
ok("...the page is told which notes it may offer Delete on", {n["id"]: n["can_delete"] for n in _t["notes"] if n["kind"] == "note"}
   .get(N_NORA) is True and {n["id"]: n["can_delete"] for n in _t["notes"] if n["kind"] == "note"}.get(N_PAT) is False
   and _t["can_moderate"] is False)
ok("⭐ its writer deletes her own note", N_LONG not in [n.get("id") for n in clp_notes.remove(Y1, N_LONG, NORA, factors)["thread"]["notes"]])
ok("...a planner deletes anyone's", N_NORA not in [n.get("id") for n in clp_notes.remove(Y1, N_NORA, PAT, factors)["thread"]["notes"]])
r = clp_notes.add_note(Y1, "The council asked for the Saturday hours.", NORA, factors)
N_NORA = r["id"]

# =========================================================================== #
#  4. Tasks (Q3, Q4, Q10): For, due, the one email                             #
# =========================================================================== #
SENT.clear()
r = clp_notes.add_task(Y1, {"words": "Ask Melton BC for their code of construction practice", "for": f"user:{U_PAT['id']}",
                            "due": "2026-11-06"}, NORA, factors)
T_PAT = r.get("id")
ok("⭐ a team member gives a task to a planner by name, with a due date", T_PAT and r["thread"]["tasks"][0]["for_word"] == "Pat Planner"
   and r["thread"]["tasks"][0]["due"] == "2026-11-06" and not r["thread"]["tasks"][0]["done"], str(r)[:400])
ok("⭐ ...and that person gets ONE email (Q10) with the task, the CLP and the date, and the page is told it went",
   r.get("email") == "sent" and len(SENT) == 1 and SENT[0]["To"] == "pat@example.com"
   and "Ask Melton BC" in SENT[0].get_content() and "Yard East" in SENT[0].get_content()
   and "Due 6 Nov 2026" in SENT[0].get_content() and "Nora North gave you a task" in SENT[0].get_content()
   and "#clp" in SENT[0].get_content(), str(SENT[0] if SENT else None)[:500])
ok("🔴 the email names no team id", not re.search(r"IPT ?[1-6]", SENT[0].as_string() if SENT else "IPT1"))
SENT.clear()
r = clp_notes.add_task(Y1, {"words": "Type the Saturday hours on the gate", "for": "team"}, NORA, factors)
T_TEAM = r["id"]
r2 = clp_notes.add_task(Y1, {"words": "Check the swept path at Gate A", "for": f"user:{U_NORA['id']}", "due": "2026-10-01"}, NORA, factors)
T_SELF = r2["id"]
r3 = clp_notes.add_task(Y1, {"words": "Approve the gate dates", "for": "approvers", "due": "2026-10-02"}, NORA, factors)
T_APP = r3["id"]
ok("🔴 nothing else emails: a task for the team, for yourself or for the planners and admins", SENT == []
   and r.get("email") is None and r2.get("email") is None and r3.get("email") is None)
ok("🔴 a task for a person who cannot see the CLP is refused in words",
   clp_notes.add_task(Y1, {"words": "x", "for": f"user:{U_CAL['id']}"}, PAT, factors).get("error") == "That person cannot see this CLP.")
ok("🔴 ...nor a demo lead, a disabled login, or one that does not exist",
   all(clp_notes.add_task(Y1, {"words": "x", "for": f"user:{u}"}, PAT, factors).get("status") == 422
       for u in (U_LEAD["id"], U_GONE["id"], "nobody")))
ok("🔴 a demo lead cannot give a task to a person", clp_notes.add_task(Y1, {"words": "x", "for": f"user:{U_PAT['id']}"}, LEAD,
                                                                     factors).get("status") == 422)
ok("🔴 an every-team CLP's task cannot be for 'the team'",
   "covers every team" in clp_notes.add_task(YA, {"words": "x", "for": "team"}, PAT, factors).get("error", ""))
ok("🔴 no For, a bad date or no words is refused in words",
   clp_notes.add_task(Y1, {"words": "x"}, NORA, factors).get("error") == "Say who the task is for."
   and clp_notes.add_task(Y1, {"words": "x", "for": "team", "due": "6/11/2026"}, NORA, factors).get("error") == "The due date must be a date."
   and clp_notes.add_task(Y1, {"words": "x", "for": "team", "due": "2026-02-30"}, NORA, factors).get("status") == 422
   and clp_notes.add_task(Y1, {"words": "", "for": "team"}, NORA, factors).get("error") == "Write the task first.")
_t = th(Y1, NORA)
_by = {t["id"]: t for t in _t["tasks"]}
ok("⭐ open tasks come first, by due date (none last)", [t["id"] for t in _t["tasks"]][:4] == [T_SELF, T_APP, T_PAT, T_TEAM],
   str([(t["words"][:12], t["due"]) for t in _t["tasks"]]))
ok("...a past due date reads overdue (the page's red)", _by[T_SELF]["overdue"] and _by[T_APP]["overdue"]
   and not _by[T_TEAM]["overdue"])
ok("...'mine' = by name, my team, or the planners and admins when I am one",
   _by[T_SELF]["mine"] and _by[T_TEAM]["mine"] and not _by[T_PAT]["mine"] and not _by[T_APP]["mine"]
   and {t["id"]: t["mine"] for t in th(Y1, PAT)["tasks"]}[T_APP] and {t["id"]: t["mine"] for t in th(Y1, PAT)["tasks"]}[T_PAT])

# =========================================================================== #
#  5. The Tasks column and filter (Q1, Q2)                                     #
# =========================================================================== #
_rn = row_of(Y1, NORA)["tasks"]
ok("⭐ the list carries each CLP's open, overdue and next due tasks, its notes and the caller's own",
   _rn["open"] == 4 and _rn["overdue"] == 2 and _rn["next_due"] == "2026-11-06" and _rn["notes"] == 3
   and _rn["mine_open"] == 2 and _rn["done"] == 0, str(_rn))
ok("...the planner's own count differs (the task by name and the one for the planners and admins)",
   row_of(Y1, PAT)["tasks"]["mine_open"] == 2)
ok("...a CLP with none reads zeros", row_of(Y2, PAT)["tasks"] == {"open": 0, "overdue": 0, "notes": 0, "done": 0,
                                                                     "mine_open": 0, "next_due": None}, str(row_of(Y2, PAT)["tasks"]))

# =========================================================================== #
#  6. Tick, untick, change (Q3, Q5)                                            #
# =========================================================================== #
r = clp_notes.change_task(Y1, T_PAT, {"done": True}, NED, factors)
_d = {t["id"]: t for t in r["thread"]["tasks"]}[T_PAT]
ok("⭐ anyone who can see the CLP ticks a task done: who and when", _d["done"] and _d["done_by"] == "Ned North" and _d["done_at"])
ok("...done ones go last", r["thread"]["tasks"][-1]["id"] == T_PAT)
ok("...and leave the open count", row_of(Y1, NORA)["tasks"]["open"] == 3 and row_of(Y1, NORA)["tasks"]["done"] == 1)
r = clp_notes.change_task(Y1, T_PAT, {"done": True}, PAT, factors)
ok("...ticking it again keeps who did it first", {t["id"]: t for t in r["thread"]["tasks"]}[T_PAT]["done_by"] == "Ned North")
r = clp_notes.change_task(Y1, T_PAT, {"done": False}, READER, factors)
ok("a sign-in with no team that only reads ticks it back (it can see every CLP)",
   not {t["id"]: t for t in r["thread"]["tasks"]}[T_PAT]["done"] and not {t["id"]: t for t in r["thread"]["tasks"]}[T_PAT]["done_by"])
ok("🔴 another team cannot tick it (not found)", clp_notes.change_task(Y1, T_PAT, {"done": True}, CAL, factors).get("status") == 404)
ok("🔴 Ned cannot change the words Nora wrote; a planner can",
   clp_notes.change_task(Y1, T_TEAM, {"words": "Mine now"}, NED, factors).get("status") == 403
   and {t["id"]: t for t in clp_notes.change_task(Y1, T_TEAM, {"words": "Type the Saturday hours on both gates"}, PAT,
                                                      factors)["thread"]["tasks"]}[T_TEAM]["words"] == "Type the Saturday hours on both gates")
SENT.clear()
r = clp_notes.change_task(Y1, T_TEAM, {"for": f"user:{U_NED['id']}", "due": "2026-11-13"}, NORA, factors)
ok("⭐ given to someone else by name: they get the one email; the due date moves",
   r.get("email") == "sent" and len(SENT) == 1 and SENT[0]["To"] == "ned@example.com"
   and {t["id"]: t for t in r["thread"]["tasks"]}[T_TEAM]["for_word"] == "Ned North"
   and {t["id"]: t for t in r["thread"]["tasks"]}[T_TEAM]["due"] == "2026-11-13", str(r.get("email")))
SENT.clear()
r = clp_notes.change_task(Y1, T_TEAM, {"due": ""}, NORA, factors)
ok("...changing only the date sends nothing, and an empty date clears it", SENT == [] and r.get("email") is None
   and {t["id"]: t for t in r["thread"]["tasks"]}[T_TEAM]["due"] is None)
ok("🔴 a task that does not exist: not found", clp_notes.change_task(Y1, "n-none", {"done": True}, NORA, factors).get("status") == 404)
users.mail_settings = lambda: None
ok("🔴 no sender set up: the task is kept and the page is told the email did not go",
   clp_notes.add_task(Y1, {"words": "Ring the haulier", "for": f"user:{U_ADA['id']}"}, PAT, factors).get("email") == "unsent")
users.mail_settings = lambda: {"DEMO_FROM": "Wayscope <hello@example.com>", "reply_to": "hello@example.com"}
_old_send = demo_request.send
demo_request.send = lambda msg, settings: (_ for _ in ()).throw(OSError("down"))
_r = clp_notes.add_task(Y1, {"words": "Ring the council", "for": f"user:{U_ADA['id']}"}, PAT, factors)
ok("🔴 a failed send never loses the task", _r.get("email") == "failed" and _r.get("id"))
demo_request.send = _old_send
for _x in (_r["id"],):
    clp_notes.remove(Y1, _x, PAT, factors)

# =========================================================================== #
#  7. The chain's steps as grey lines (Q6); Ask for changes as a task (Q9)     #
# =========================================================================== #
clp_records.send(Y1, "Gate dates updated", NORTH_CODE, factors)
ok("🔴 Ask for changes with a bad date as a task is refused before anything moves",
   clp_records.ask_changes(Y1, "Add the Saturday hours", PAT, factors, task={"due": "soon"}).get("status") == 422
   and clp_records.get(Y1, PAT)["status"] == "sent")
r = clp_records.ask_changes(Y1, "Add the Saturday hours to section 7", PAT, factors, task={"due": "2026-11-09"})
_t = th(Y1, NORA)
_steps = [n for n in _t["notes"] if n["kind"] == "step"]
ok("⭐ Send and Ask for changes are grey lines in the thread, with their notes",
   [s["step_word"] for s in _steps] == ["Sent for approval", "Changes asked"]
   and _steps[0]["by"] == "North Team" and _steps[0]["words"] == "Gate dates updated"
   and _steps[1]["by"] == "Pat Planner" and _steps[1]["words"] == "Add the Saturday hours to section 7", str(_steps))
_ct = [t for t in _t["tasks"] if t["words"] == "Add the Saturday hours to section 7"]
ok("⭐ ...and the change asked for is a task for the team, with its date (Q9)",
   len(_ct) == 1 and _ct[0]["for_word"] == "North Team" and _ct[0]["due"] == "2026-11-09" and _ct[0]["mine"], str(_ct))
clp_records.send(Y1, "Saturday hours added", NORTH_CODE, factors)
P = clp_records.preview(Y1, PAT, factors, WD)
ok("⭐ the approve card is told the open tasks and the first overdue one (never blocking)",
   P["tasks"]["open"] >= 4 and P["tasks"]["overdue"] == 2 and P["tasks"]["first_overdue"] == "Check the swept path at Gate A",
   str(P.get("tasks")))
r = clp_records.approve(Y1, "Approved with the gate dates", PAT, factors, WD, build_fast)
ok("approving with open tasks works", r.get("last_issue", {}).get("no") == 1, str(r)[:200])
_t = th(Y1, NORA)
_steps = [n for n in _t["notes"] if n["kind"] == "step"]
ok("⭐ ...and reads 'Approved as Issue 1' with the approver's note, last in time",
   _steps[-1]["step_word"] == "Approved as Issue 1" and _steps[-1]["words"] == "Approved with the gate dates"
   and _t["notes"][-1]["kind"] == "step", str(_steps[-1]))
ok("...every step once (the CLP's own sent / changes fields are not shown again)",
   [s["step_word"] for s in _steps] == ["Sent for approval", "Changes asked", "Sent for approval", "Approved as Issue 1"],
   str([s["step_word"] for s in _steps]))
db.execute("DELETE FROM clp_notes WHERE tenant_id = ? AND kind = 'step'", (db.current_tenant(),))
_steps = [n for n in th(Y1, NORA)["notes"] if n["kind"] == "step"]
ok("⭐ steps from before 3b are read back from the CLP and its issues (the issue, the last send)",
   "Approved as Issue 1" in [s["step_word"] for s in _steps] and "Sent for approval" in [s["step_word"] for s in _steps]
   and next(s for s in _steps if s["step"] == "approved")["by"] == "Pat Planner", str(_steps))
ok("🔴 a step cannot be deleted", clp_notes.remove(Y1, "step", PAT, factors).get("status") == 404
   and not [r for r in clp_notes._rows(Y1, ("step",))])

# =========================================================================== #
#  8. The kept "Still to complete" lines (Q7, Q8)                              #
# =========================================================================== #
L1 = th(Y1, NORA)["lines"]
ok("⭐ approving kept the issue's 'Still to complete' list, with each line's section",
   L1 and L1["at"] and len(L1["items"]) >= 5 and all(i["section"] and i["words"] for i in L1["items"]),
   str(L1)[:300])
_cop = next((i for i in L1["items"] if i["section"].startswith("7")), L1["items"][0])
r = clp_notes.add_task(Y1, {"words": "Type the CLP cap", "for": "team", "from_line": _cop["line"]}, NORA, factors)
T_LINE = r["id"]
_tl = {t["id"]: t for t in r["thread"]["tasks"]}[T_LINE]
ok("⭐ a task starts from a line: it keeps the line and its section", _tl["from_line"] == _cop["line"]
   and _tl["from_section"] == _cop["section"] and _tl["line_gone_at"] is None, str(_tl))
ok("...an unchanged rebuild writes nothing", clp_notes.keep_lines(Y1, [i["line"] for i in L1["items"]]) is False)
clp_notes.keep_lines(Y1, [i["line"] for i in L1["items"] if i["line"] != _cop["line"]])
_tl = {t["id"]: t for t in th(Y1, NORA)["tasks"]}[T_LINE]
ok("⭐ a later build without the line: the task says so, and is NOT ticked by itself (Q8)",
   _tl["line_gone_at"] and not _tl["done"], str(_tl))
ok("...one kept list a CLP", len(clp_notes._rows(Y1, ("lines",))) == 1)
ok("split_line names the section", clp_notes.split_line("2.4.1: Melton BC: the code") == ("2.4.1", "Melton BC: the code")
   and clp_notes.split_line("no section") == ("", "no section"))
_doc_text = "\n".join(p.text for p in __import__("docx").Document(io.BytesIO(clp_records.issue_file(Y1, 1, PAT)[1])).paragraphs)
ok("🔴 nothing of a note or a task prints in the Word file (Q11)",
   "Saturday hours to section 7" not in _doc_text and "Melton BC for their code" not in _doc_text
   and "Check the swept path" not in _doc_text)

# =========================================================================== #
#  9. Delete (Q12): + New CLP takes its notes, a listed row keeps them         #
# =========================================================================== #
as_(PAT)
NEW = clp_records.create({"name": "Pile and pit", "location_ids": [PILE], "team": None}, PAT, factors)["id"]
clp_notes.add_note(NEW, "One for the pile", PAT, factors)
clp_notes.add_task(NEW, {"words": "Count the pile", "for": "approvers"}, PAT, factors)
ok("a + New CLP carries its notes and tasks in the list", row_of(NEW, PAT)["tasks"]["open"] == 1 and row_of(NEW, PAT)["tasks"]["notes"] == 1)
clp_records.delete(NEW, PAT, factors)
ok("⭐ deleting it removes its notes and tasks", clp_notes._rows(NEW) == [])
clp_records.update(Y2, {"cover": {"address": "2 Yard Lane"}}, PAT, factors)
clp_notes.add_note(Y2, "Central note", PAT, factors)
r = clp_records.delete(Y2, PAT, factors)
ok("⭐ clearing a listed row's typed details keeps its notes and tasks", r.get("cleared") and len(clp_notes._rows(Y2, ("note",))) == 1)

# =========================================================================== #
#  10. The endpoints                                                           #
# =========================================================================== #
as_(NORA)
EP = main.clp_notes_read(Y1)
ok("GET notes: the thread, the tasks, the lines and the choices", EP["tasks"] and EP["notes"] and EP["lines"] and EP["for_options"])
_e = _raises(lambda: main.clp_notes_read(Y2))
ok("🔴 GET another team's notes: 404", getattr(_e, "status_code", None) == 404, str(_e))
_n = main.clp_notes_add(Y1, main.ClpNoteIn(note="From the page"))
ok("POST a note", _n["thread"]["notes"][-1]["words"] == "From the page")
_e = _raises(lambda: main.clp_notes_add(Y1, main.ClpNoteIn(note="")))
ok("🔴 POST an empty note: 422 in words", getattr(_e, "status_code", None) == 422 and "Write a note" in str(_e), str(_e))
_tk = main.clp_tasks_add(Y1, main.ClpTaskIn(**{"words": "From the page too", "for": "approvers", "due": "2026-12-01"}))
_new = {t["id"]: t for t in _tk["thread"]["tasks"]}[_tk["id"]]
ok("⭐ POST a task, 'for' and all", _new["for_word"] == "Planners and admins" and _new["due"] == "2026-12-01", str(_new))
_p = main.clp_tasks_change(Y1, _tk["id"], main.ClpTaskIn(done=True))
ok("PATCH ticks it, sending only 'done' (the words are kept)", {t["id"]: t for t in _p["thread"]["tasks"]}[_tk["id"]]["done"]
   and {t["id"]: t for t in _p["thread"]["tasks"]}[_tk["id"]]["words"] == "From the page too")
as_(NED)
_e = _raises(lambda: main.clp_tasks_change(Y1, _tk["id"], main.ClpTaskIn(words="x")))
ok("🔴 PATCH another's words as a team member: 403", getattr(_e, "status_code", None) == 403, str(_e))
_e = _raises(lambda: main.clp_tasks_delete(Y1, _tk["id"]))
ok("🔴 DELETE another's task: 403", getattr(_e, "status_code", None) == 403, str(_e))
as_(NORA)
ok("DELETE her own task and note", _tk["id"] not in [t["id"] for t in main.clp_tasks_delete(Y1, _tk["id"])["thread"]["tasks"]]
   and _n["id"] not in [n.get("id") for n in main.clp_notes_delete(Y1, _n["id"])["thread"]["notes"]])
as_(PAT)
_v = db.data_version()
main.clp_notes_add(Y1, main.ClpNoteIn(note="Version check"))
ok("⭐ a note moves no data version (the kept CLP drafts stay)", db.data_version() == _v)
_ch = main.clp_documents_changes(Y1, main.ClpNoteIn(note="One more change", task={"due": "2026-12-02"}))
ok("POST changes with a task (the endpoint passes it on)", any(t["words"] == "One more change" and t["due"] == "2026-12-02"
                                                               for t in th(Y1, PAT)["tasks"]))
_ch = main.clp_documents_changes(Y1, main.ClpNoteIn(note="Without a task"))
ok("...and without one adds none", not any(t["words"] == "Without a task" for t in th(Y1, PAT)["tasks"]))
db.execute("DELETE FROM clp_notes WHERE tenant_id = ? AND kind = 'lines'", (db.current_tenant(),))
_dr = main.clp_documents_draft(Y1)
ok("⭐ the draft download keeps its 'Still to complete' list", _dr.body[:2] == b"PK" and th(Y1, PAT)["lines"]
   and th(Y1, PAT)["lines"]["items"])
_pkg = tenant_package.export_tenant()
ok("⭐ the notes and tasks travel in the organisation file", any(
   r.get("clp_id") == Y1 and r.get("kind") == "task" for r in (_pkg.get("tables") or {}).get("clp_notes") or []))

# =========================================================================== #
#  11. The source                                                              #
# =========================================================================== #
SRC = open(os.path.join(BACKEND, "clp_notes.py"), encoding="utf-8").read()
_calls = [SRC[m.end():m.end() + 300] for m in re.finditer(r"db\.(?:query|execute)\(", SRC)]
ok("🔴 every statement in clp_notes names tenant_id", len(_calls) >= 5
   and all("tenant_id" in c.split("))")[0] for c in _calls), str(len(_calls)))
ok("🔴 no BOOLEAN compared with a number", not re.search(r"\b(active|disabled)\s*=\s*[01]\b", SRC))
ok("🔴 no team slot id in the module's words", re.search(r"\bIPT ?[1-6]\b", SRC) is None)
ok("🔴 no reply word carries a team id", not re.search(r"IPT ?[1-6]", json.dumps(
    {k: v for k, v in th(Y1, PAT).items() if k != "id"}, ensure_ascii=False)))
_REC = open(os.path.join(BACKEND, "clp_records.py"), encoding="utf-8").read()
ok("the chain writes its steps (send, changes, approve) and approve keeps the lines",
   _REC.count("clp_notes.add_step(") == 3 and "clp_notes.keep_lines(rec[\"id\"], data.get(\"todos\"))" in _REC)

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\ntest_clp_notes: {PASS} passed, {len(FAIL)} failed")
remove_scratch_db(db, TMP)
sys.exit(1 if FAIL else 0)
