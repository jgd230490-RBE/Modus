"""
H1, 2026-09-29 — UK localisation: a GB tenant sees no Estonian source, name or unit.

WHAT IS ASSERTED
----------------
  1. Providers per country: GB → DESNZ diesel + National Highways closures; EE → the EU
     bulletin + Tark Tee; NO country → typed diesel, no restriction provider (the
     pre-H1 fallback to Estonia is gone from every module).
  2. The demo tenant (GB, GBP, miles) imported into an empty tenant: every JSON value a
     GB page reads — /api/meta, /api/costing, /api/fuel-index, /api/restrictions/layers,
     /api/restrictions, /api/lookahead, /api/forecast-weeks/restrictions,
     /api/public/alignment — contains NO Estonian term, NO EU-bulletin term, NO "€" and
     NO "EE" country value. Same for the XLSX (read back with openpyxl) and the PDF
     (pdftotext) of the commit week.
  3. Source level: "Tark Tee" and "Maa-amet" reach a GB page only through code that is
     gated on the tenant's country or on the API's provider block (the guide, which
     documents both countries, is the deliberate exception).
  4. km/miles (HU5, per tenant): stored km, shown in the tenant's unit — the helpers in
     the app and the map, and the XLSX/PDF headers and figures.
  5. The GB vehicle set: lead names, hidden EU names, the two truck mixers, DESNZ 2026
     CO₂ factors with their source on every >17 t rigid and >33 t artic entry.
  6. The Mapbox token default in config.py equals the one in map/config.js; GB → Europe/London.
  7. The probes exist and are shaped for the deployment: /api/admin/diagnostics/fuel-index
     and the provider-aware /api/admin/diagnostics/restrictions.

WHAT THIS DOES NOT PROVE
------------------------
  * HTTP is stubbed; endpoints are called as functions (http_smoke.py has the real layer).
  * gov.uk (DESNZ) and the National Highways FeatureServer are NEVER called: both fetches
    run against fixtures shaped from their published schemas. The probes on Render are
    the first real run.
  * Nothing renders. The frontend and map assertions are source-level.

Run:  python3 backend/tests/test_h1_uk.py
"""
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import types

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# --------------------------------------------------------------------------- #
#  Stubs (the same shape every other harness uses)                             #
# --------------------------------------------------------------------------- #
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


class _HTTPException(Exception):
    def __init__(self, status_code, detail=None):
        super().__init__(detail)
        self.status_code, self.detail = status_code, detail


class _App:
    def __init__(self, *a, **k):
        pass

    def _deco(self, *a, **k):
        return lambda f: f
    get = post = put = delete = patch = _deco

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


def _Query(default=None, **k):
    return default


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query
_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None
_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)

TMP = tempfile.mkdtemp(prefix="wayscope_h1_")
os.environ.pop("DATABASE_URL", None)
os.environ.pop("TENANT_ID", None)
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE", "ADMIN_TOKEN"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import config  # noqa: E402
import taxonomy  # noqa: E402
import network  # noqa: E402
import restrictions  # noqa: E402
import fuel  # noqa: E402
import costing  # noqa: E402
import tenant_package  # noqa: E402
import lookahead  # noqa: E402
import export  # noqa: E402
import here_routing  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
access.set_current("planner123")

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def read(rel):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return f.read()


def reset_db():
    if os.path.exists(db._SQLITE_PATH):
        os.remove(db._SQLITE_PATH)
    db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db()
    db.init_gates_db(); db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db()
    db.init_costing_db(); db.init_tenant()
    config.invalidate(); costing.invalidate()


def boot():
    reset_db()
    taxonomy.seed_taxonomy()
    config.seed_from_file(conversions)
    config.invalidate(); costing.invalidate()


def values_of(o, out=None):
    """Every string VALUE in a JSON-shaped object (keys are internal names, not shown)."""
    out = out if out is not None else []
    if isinstance(o, dict):
        for v in o.values():
            values_of(v, out)
    elif isinstance(o, (list, tuple)):
        for v in o:
            values_of(v, out)
    elif isinstance(o, str):
        out.append(o)
    return out


def js(o):
    """Round-trip through JSON, as the HTTP layer would."""
    return json.loads(json.dumps(o, default=str))


# Estonian names, EU-bulletin names and the euro must not reach a GB page. Case-insensitive.
FORBIDDEN = [r"tark\s?tee", r"tarktee", r"maa-?amet", r"estonia", r"eesti", r"tallinn",
             r"euro\s?oil\s?watch", r"oil bulletin", r"weekly oil", r"transpordiamet", r"€"]
FORBIDDEN_RX = re.compile("|".join(FORBIDDEN), re.IGNORECASE)


def leaks(o):
    vals = values_of(js(o))
    hits = [v[:80] for v in vals if FORBIDDEN_RX.search(v)]
    # a country VALUE of "EE" (never the lower-case 'ee' vehicle-label language key, which the
    # app hides for a non-Estonian tenant) or the Estonian time zone
    ee = [v for v in vals if v.strip() == "EE" or v.startswith("Europe/Tallinn")]
    return hits + ee


# =========================================================================== #
#  1. Providers per country — and the EE fallback is gone                       #
# =========================================================================== #
ok("fuel: GB → DESNZ (GBP/L), EE → the EU bulletin (EUR/L), NO → nothing, no country → nothing",
   fuel.provider_for("GB")["key"] == "desnz" and fuel.provider_for("EE")["key"] == "eu_bulletin"
   and fuel.provider_for("NO") is None and fuel.provider_for(None) is None
   and fuel.norm_country(None) == fuel.NO_COUNTRY)
ok("restrictions: GB → National Highways (100 m, dated, ROADWORKS), EE → Tark Tee (30 m, undated, TARK TEE)",
   restrictions.provider_for("GB")["key"] == "nh_closures" and restrictions.provider_for("GB")["match_m"] == 100.0
   and restrictions.provider_for("GB")["dated"] is True and restrictions.provider_for("GB")["flag_label"] == "ROADWORKS"
   and restrictions.provider_for("EE")["key"] == "tark_tee" and restrictions.provider_for("EE")["match_m"] == 30.0
   and restrictions.provider_for("EE")["dated"] is False and restrictions.provider_for("EE")["flag_label"] == "TARK TEE"
   and restrictions.provider_for("NO") is None and restrictions.provider_for(None) is None)
ok("config.RESTRICTION_PROVIDERS is read from restrictions.PROVIDERS (the two cannot drift)",
   config.RESTRICTION_PROVIDERS == set(restrictions.PROVIDERS) == {"EE", "GB"})
_srcs = {f: read(f) for f in ("backend/fuel.py", "backend/costing.py", "backend/derived.py", "backend/main.py",
                              "backend/export.py", "backend/restrictions.py", "backend/config.py")}
ok("🔴 no DEFAULT_COUNTRY anywhere in the backend — nothing falls back to Estonia",
   all("DEFAULT_COUNTRY" not in s for s in _srcs.values()))
ok("🔴 the frontend's fuel widget has no country fallback either", 'fuelS.country || "EE"' not in read("frontend/index.html"))

boot()
ok("a tenant with no country: no restriction provider, no automatic fuel index, typed only",
   restrictions.provider() is None and fuel.state(costing.settings()["fuel"]["country"])["no_country"] is True
   and costing.settings()["fuel"]["country"] == fuel.NO_COUNTRY
   and fuel.refresh(costing.settings()["fuel"]["country"], sync=True)["status"] == "manual_only")
_m0 = main.meta()["tenant"]
ok("...and /api/meta says so: no providers, unit km", _m0["fuel_provider"] is None and _m0["restrictions"] is None
   and _m0["restrictions_provider"] is False and _m0["distance_unit"] == "km")

# =========================================================================== #
#  2. The demo tenant: nothing Estonian reaches a GB page                       #
# =========================================================================== #
DEMO = json.load(open(os.path.join(ROOT, "demo", "uk-corridor.package.json"), encoding="utf-8"))
_fac = json.loads(next(r["value"] for r in DEMO["tables"]["config"] if r["key"] == "factors"))
ok("the demo package is GB / GBP / miles and its fuel country is GB",
   _fac["tenant"]["country"] == "GB" and _fac["tenant"]["currency"] == "GBP"
   and _fac["tenant"].get("distance_unit") == "mi"
   and json.loads(next(r["value"] for r in DEMO["tables"]["config"] if r["key"] == "costing"))["fuel"]["country"] == "GB",
   str(_fac["tenant"]))
ok("the demo package names no Estonian source anywhere in its values", not leaks(DEMO), str(leaks(DEMO))[:300])
boot()
res = tenant_package.import_tenant(DEMO, replace=True)
ok("the demo imports into an empty tenant", res.get("ok") is True, str(res)[:200])
config.invalidate(); costing.invalidate(); restrictions.clear_cache()
tn = main.meta()["tenant"]
ok("GB tenant: DESNZ and National Highways on /api/meta, £, miles",
   tn["country"] == "GB" and tn["fuel_provider"]["key"] == "desnz" and tn["restrictions"]["key"] == "nh_closures"
   and tn["currency_symbol"] == "£" and tn["distance_unit"] == "mi", str(tn))

# stub the two live sources with fixtures shaped from their published schemas
_CSV = ("Date,ULSP:  Pump price in pence/litre,ULSD:  Pump price in pence/litre,ULSP:  VAT (%) rate applicable,ULSD:  VAT (%) rate applicable\n"
        "14/09/2026,135.20,194.80,20,20\n21/09/2026,135.60,195.50,20,20\n")
fuel.fetch_desnz = lambda timeout=None: {"csv": _CSV, "csv_url": "https://assets.publishing.service.gov.uk/x.csv", "csv_title": "Weekly road fuel prices (CSV)"}
import time as _t
_now_ms = int(_t.time() * 1000)
_day = 86400000
_r1 = db.query("SELECT id FROM routes WHERE tenant_id = ? ORDER BY id", (db.current_tenant(),))[0]["id"]
# bake the first route for every vehicle its forecast lines use, along the closure fixture below
_veh = [r["vehicle_type"] for r in db.query("SELECT DISTINCT vehicle_type FROM forecasts WHERE tenant_id = ? AND route_id = ?",
                                            (db.current_tenant(), _r1))] or ["Rigid 8-wheeler (32t)"]
for _v in _veh:
    network._upsert_geom(_r1, _v, "[[-0.905,52.499],[-0.8955,52.5052],[-0.885,52.515]]", 2.0, 0.1, None, leg="loaded", alt_index=0)
    network._upsert_geom(_r1, _v, "[[-0.885,52.515],[-0.8955,52.5052],[-0.905,52.499]]", 2.0, 0.1, None, leg="return", alt_index=0)


def _fake_nh_page(offset, count=restrictions.NH_PAGE, timeout=None):
    if offset:
        return {"type": "FeatureCollection", "features": []}
    return {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[-0.9, 52.5], [-0.89, 52.51]]},
         "properties": {"OBJECTID": 2, "description": "Lane closure for barrier works", "road_number": "A46",
                        "eventtype": "Lane closure", "natureofworks": "Barrier works", "formattedeventnumber": "NH-0002",
                        "scheduledplannedstartdate": _now_ms - 2 * _day, "scheduledplannedenddate": _now_ms + 400 * _day}}]}


# --- 5 Oct 2026, H2: demo v2 — what the stand needs is IN the data ------------------------
_T = DEMO["tables"]
ok("H2: the scheme is a road — tenant name, overlay note, discipline words; no railway term reaches a reader",
   "dualling" in _T["config"][0]["value"] or any("dualling" in c["value"] for c in _T["config"])
   and not re.search(r"track ballast|rail corridor|Melton South station|Rail and sleepers", json.dumps(_T, ensure_ascii=False))
   and [d["label"] for d in _T["disciplines"] if d["id"] == "superstructure"] == ["Pavement"]
   and [d["in_scope"] for d in _T["disciplines"] if d["id"] == "stations"] == [0])
ok("H2: every site has a gate at its own coordinates; the quarries and the town compound enter and leave by different gates",
   len(_T["location_gates"]) == 14 and {g["location_id"] for g in _T["location_gates"]} == {l["id"] for l in _T["locations"]}
   and all(g["lat"] and g["lon"] for g in _T["location_gates"])
   and sorted(g["direction"] for g in _T["location_gates"] if g["location_id"] == "L001") == ["access", "egress"]
   and sorted(g["direction"] for g in _T["location_gates"] if g["location_id"] == "L005") == ["access", "egress"])
ok("H2: every route names its origin and destination gate",
   all(r["origin_gate_id"] and r["dest_gate_id"] for r in _T["routes"]))
ok("H2: two haul roads, spliced (HERE does not know them), attached to the routes that use them",
   [(z["id"], z["haul_mode"]) for z in _T["zones"] if z["kind"] == "haul_road"] == [("Z001", "splice"), ("Z002", "splice")]
   and sorted((l["route_id"], l["zone_id"]) for l in _T["route_haul_roads"]) == [("R004", "Z001"), ("R005", "Z001"), ("R006", "Z001"), ("R014", "Z001"), ("R016", "Z002"), ("R018", "Z002")])
ok("H2: three geofences — a dated closure that re-routes (Tilton, from the show week), an advisory 7.5 t village limit, an advisory works area",
   [(z["kind"], z["affects_routing"], z["starts_on"]) for z in _T["zones"] if z["kind"] != "haul_road"]
   == [("closure", 1, "2026-10-12"), ("weight_limit", 0, "2026-09-01"), ("works", 0, "2026-10-01")])
ok("H2: daily vehicle caps on the two village routes, so ROUTE_CAP has something to fire on",
   {r["id"]: r["max_vehicles_per_day"] for r in _T["routes"] if r["max_vehicles_per_day"]} == {"R016": 20, "R018": 30})
ok("H2: the show week (12–18 Oct = Oct W3) and the week after are populated for every approved line, and October W1/W2 carry typed actuals",
   sum(1 for w in _T["forecast_weeks"] if (w["month_index"], w["week_index"]) == (10, 3)) == 18
   and sum(1 for w in _T["forecast_weeks"] if (w["month_index"], w["week_index"]) == (10, 4)) == 18
   and sum(1 for w in _T["forecast_weeks"] if w["month_index"] == 10 and w["week_index"] in (1, 2) and w.get("actual_qty") is not None) == 36)
ok("H2: the tenant carries the demo notice, /api/meta serves it, and it validates as short text or null",
   main.meta()["tenant"]["demo_notice"] == "Fictional demo scheme · nothing here is real"
   and config.validate(dict(json.loads(DEMO["tables"]["config"][[c["key"] for c in DEMO["tables"]["config"]].index("factors")]["value"]), tenant={"name": "x", "demo_notice": "y" * 121})) != []
   and "demo_notice" in config.TENANT_DEFAULTS)
ok("H2: the notice is on both exports (XLSX About row, PDF footer) and the app's header, and the map hides an empty rail layer",
   '("Notice", _demo_notice() or "")' in read("backend/export.py") and 'ft.insert(0, P("<b>" + E(_demo_notice()) + "</b>", st_grey))' in read("backend/export.py")
   and 'className="demo-badge' in read("frontend/index.html") and "demo_notice: (tn.demo_notice || null)" in read("frontend/index.html")
   and "railItem.hidden = !hasRail" in read("map/index.html") and "dn.textContent = TENANT.demo_notice" in read("map/index.html"))
# --- 30 Sep 2026: the national feed against baked routes must be pre-filtered by box -----
# The first GB run with baked routes put the live service at 100 % CPU for hours: 2,920
# closures × 170,000 vertices against 18 routes × 4 legs, vertex by segment. These pin the
# bounding-box short cut in _min_distance_km and that it never hides a real hit.
import random as _random
import time as _time
_random.seed(30)
_line = [[-0.905 + i * 0.0001, 52.499 + i * 0.00004] for i in range(500)]          # a 500-vertex HERE-like leg
_far = [{"type": "Feature", "geometry": {"type": "MultiLineString",
                                          "coordinates": [[[-0.23 + _random.random() * 0.01, 51.68 + _random.random() * 0.01] for _ in range(58)]]},
         "properties": {}} for _ in range(3000)]                                       # M25-sized, 90 km south
_near = {"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[-0.9, 52.5], [-0.89, 52.51]]}, "properties": {}}
_calls = {"n": 0}
_orig_pts = restrictions._point_to_segment_km
def _counting(p, a, b):
    _calls["n"] += 1
    return _orig_pts(p, a, b)
restrictions._point_to_segment_km = _counting
try:
    _lb = restrictions._bbox(_line)
    _t0 = _time.time()
    _skipped = sum(1 for f in _far if restrictions._min_distance_km(f["geometry"], _line, max_km=0.1, line_bbox=_lb) is None)
    _dt = _time.time() - _t0
    ok("🔴 a closure whose box is further than the match distance is skipped without a single segment test",
       _skipped == 3000 and _calls["n"] == 0, f"skipped {_skipped}, segment tests {_calls['n']}")
    ok("🔴 …and 3,000 of them against a 500-vertex leg take well under a second, not hours", _dt < 1.0, f"{_dt:.2f}s")
    _calls["n"] = 0
    _d = restrictions._min_distance_km(_near["geometry"], [[-0.905, 52.499], [-0.8955, 52.5052], [-0.885, 52.515]], max_km=0.1)
    ok("🔴 a closure inside the box still goes through the exact test and is a hit", _d is not None and _d * 1000 <= 100 and _calls["n"] > 0, str(_d))
    ok("the box test is a lower bound on the exact distance (never hides a pair): 2,000 random pairs",
       all(restrictions._bbox_apart_km(restrictions._bbox([p]), restrictions._bbox([q])) <= _orig_pts(p, q, q) + 1e-9
           for p, q in (([-1 + _random.random() * 2, 51 + _random.random() * 4], [-1 + _random.random() * 2, 51 + _random.random() * 4]) for _ in range(2000))))
    ok("without max_km the old exact behaviour is unchanged (a far feature still gets a distance)",
       restrictions._min_distance_km(_far[0]["geometry"], _line) is not None)
    ok("check_route passes the match distance and the leg's box (source)",
       "max_km=max_km, line_bbox=line_bbox" in read("backend/restrictions.py") and "max_km = mm / 1000.0" in read("backend/restrictions.py"))
finally:
    restrictions._point_to_segment_km = _orig_pts

restrictions._nh_page = _fake_nh_page
fuel.refresh("GB", sync=True)
restrictions.refresh_async(sync=True)
ok("GB: the DESNZ row is stored (£1.955/L incl. VAT, £1.629 ex-VAT) and the NH check stored one hit on the baked route",
   (fuel.get_index("GB") or {}).get("eur_per_l") == 1.955 and fuel.state("GB")["ex_vat_per_l"] == 1.629
   and restrictions.stored_checks([_r1])[_r1]["hits"] and restrictions.stored_checks([_r1])[_r1]["status"] == "ok")

PAGES = {
    "/api/meta": main.meta(),
    "/api/costing": main.get_costing(),
    "/api/fuel-index": main.get_fuel_index(lazy=0),
    "/api/restrictions/layers": main.restriction_layers(),
    "/api/restrictions": json.loads(main.get_restrictions().body),     # NARROWED 30 Sep: a pre-serialised Response
    "/api/lookahead": main.lookahead_page(bucket="commit"),
    "/api/forecast-weeks/restrictions": main.forecast_week_restrictions(bucket="commit"),
    "/api/forecast-weeks/clashes": main.forecast_week_clashes(bucket="commit", restrictions_on=1),
    "/api/public/alignment": json.loads(main.public_alignment(_Request()).body),
    "/api/routes/{id}/restrictions": main.route_restrictions(_r1),
    "/api/admin/diagnostics/fuel-index": main.diagnostics_fuel_index(probe=False),
    "/api/admin/diagnostics/restrictions (no probe)": restrictions.diagnostics(probe=False),
}
for path, payload in PAGES.items():
    lk = leaks(payload)
    ok(f"🔴 GB page {path}: no Estonian term, no EU-bulletin term, no €, no EE value", not lk, str(lk)[:240])
_la = PAGES["/api/lookahead"]
ok("GB look-ahead: the restriction source is National Highways with the ROADWORKS word and the week window",
   _la["clashes"]["sources"]["restrictions_provider"] == "nh_closures" and _la["clashes"]["sources"]["restrictions_label"] == "ROADWORKS"
   and _la["clashes"]["sources"]["restrictions_ui_label"] == "Road closures (National Highways)"
   and _la["clashes"]["sources"]["restrictions_week"] and len(_la["clashes"]["sources"]["restrictions_week"]) == 2)
_fi = PAGES["/api/fuel-index"]["index"]
ok("GB fuel index: provider DESNZ, unit GBP/L, attribution names gov.uk, VAT 20, ex-VAT derived",
   _fi["provider"]["key"] == "desnz" and _fi["unit"] == "GBP/L" and "gov.uk" in _fi["attribution"]
   and _fi["vat_pct"] == 20.0 and _fi["ex_vat_per_l"] == 1.629 and _fi["source"] == fuel.SOURCE_DESNZ, str(_fi)[:200])
ok("GB restrictions catalogue: one layer, planned road closures, 100 m, OGL attribution",
   [l["key"] for l in PAGES["/api/restrictions/layers"]["layers"]] == ["nh_closures"]
   and PAGES["/api/restrictions/layers"]["match_m"] == 100.0 and "Open Government Licence" in PAGES["/api/restrictions/layers"]["attribution"])
_r_ = main.get_restrictions()
_fc_ = restrictions.fetch_all(None, current_only=True)
ok("🔴 GET /api/restrictions is a pre-serialised JSON Response (not walked by FastAPI's encoder), cached per collection",
   getattr(_r_, "media_type", "") == "application/json" and isinstance(_r_.body, str)
   and json.loads(_r_.body)["type"] == "FeatureCollection"
   and restrictions.serialised(_fc_) is restrictions.serialised(_fc_))
_rr = PAGES["/api/routes/{id}/restrictions"]
ok("GB route check: the hit carries its dates, event number and the ROADWORKS word; verdict unknown (a closure, not a limit)",
   _rr["hits"] and _rr["hits"][0]["event_number"] == "NH-0002" and _rr["hits"][0]["from"] and _rr["hits"][0]["to"]
   and _rr["flag_label"] == "ROADWORKS" and _rr["hits"][0]["verdict"] == "unknown" and _rr["match_m"] == 100.0)

# the exports, read back for real
try:
    import openpyxl
    xb = export.build_xlsx(_la)
    wb = openpyxl.load_workbook(io.BytesIO(xb), read_only=True)
    _xtext = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            _xtext.extend(str(c) for c in row if c is not None)
    _xall = "\n".join(_xtext)
    ok("GB XLSX: no Estonian term, no EU-bulletin term, no € anywhere on any sheet",
       not FORBIDDEN_RX.search(_xall) and not re.search(r"\bEE\b", _xall), str([t for t in _xtext if FORBIDDEN_RX.search(t)])[:200])
    _hdr = [c for c in _xtext]
    ok("GB XLSX (HU5): the day sheet's distance headers read mi/trip, mi/day and t·mi — not km",
       "mi/trip" in _hdr and "mi/day" in _hdr and "t·mi" in _hdr and "km/trip" not in _hdr and "t·km" not in _hdr)
    ok("GB XLSX: the Fuel sheet names DESNZ and the running rate per mile",
       any("DESNZ" in t for t in _xtext) and any("running £/mi" in t for t in _xtext))
except ImportError:
    for _ in range(3):
        ok("(XLSX read-back skipped — openpyxl not installed here)", True)
try:
    import reportlab  # noqa: F401
    pb = export.build_pdf(_la)
    _ptxt = (subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb, capture_output=True).stdout.decode("utf-8")
             if shutil.which("pdftotext") else "")
    if _ptxt:
        ok("GB PDF: no Estonian term, no EU-bulletin term, no € in the rendered text",
           not FORBIDDEN_RX.search(_ptxt), str(FORBIDDEN_RX.findall(_ptxt))[:200])
        ok("GB PDF (HU5): the week column is headed MI/TRIP and totals read t·mi",
           "MI/TRIP" in _ptxt and "t·mi" in _ptxt and "KM/TRIP" not in _ptxt)
    else:
        for _ in range(2):
            ok("(PDF text skipped — pdftotext missing)", True)
except ImportError:
    for _ in range(2):
        ok("(PDF skipped — reportlab not installed here)", True)

# 29 Sep 2026 first look on the LIVE demo tenant (miles, DESNZ): three wording leaks in the
# exports and one misleading caption. "km omitted" on every UNBAKED line, "Km from the baked
# HERE route" in the footer, "Bulletin date" on the Fuel sheet of a DESNZ tenant, and a PDF
# with nothing baked said "Mapbox could not be reached". Pinned here on the same GB page.
_KM_OK = re.compile(r"L/100\s?km")          # consumption stays per 100 km on purpose (README)
try:
    import openpyxl
    _xt = []
    for ws in openpyxl.load_workbook(io.BytesIO(export.build_xlsx(_la)), read_only=True).worksheets:
        for row in ws.iter_rows(values_only=True):
            _xt.extend(str(c) for c in row if c is not None)
    _km_leaks = [t for t in _xt if re.search(r"\b[Kk]m\b", _KM_OK.sub("", t))]
    ok("🔴 GB XLSX (miles): no 'km' word reaches a sheet except the L/100 km consumption line",
       not _km_leaks, str(_km_leaks)[:240])
    ok("GB XLSX (miles): an UNBAKED line says 'mi omitted', the footer says 'Miles from the baked HERE route'",
       any("mi omitted on that line" in t for t in _xt) and any("Miles from the baked HERE route" in t for t in _xt))
    ok("🔴 GB XLSX (DESNZ): the Fuel sheet never says 'bulletin' — the date row reads 'Published (week ending)'",
       not any(re.search(r"bulletin", t, re.IGNORECASE) for t in _xt)
       and "Published (week ending)" in _xt and "BAF base published (week ending)" in _xt, str([t for t in _xt if "ulletin" in t])[:200])
    _ee_cost = {"index_source": "eu_weekly_oil_bulletin", "fuel": {"country": "EE"}}
    ok("…while an Estonian tenant's sheet keeps 'Bulletin date' (the series IS a bulletin), a typed one 'Index date'",
       export._index_date_label(_ee_cost) == "Bulletin date"
       and export._index_date_label({"index_source": "manual", "fuel": {"country": "--"}}) == "Index date"
       and export._index_date_label({"index_source": "manual", "fuel": {"country": "GB"}}) == "Published (week ending)")
except ImportError:
    for _ in range(4):
        ok("(XLSX wording read-back skipped — openpyxl not installed here)", True)
try:
    import reportlab  # noqa: F401
    if shutil.which("pdftotext"):
        _pt = subprocess.run(["pdftotext", "-layout", "-", "-"], input=export.build_pdf(_la), capture_output=True).stdout.decode("utf-8")
        ok("🔴 GB PDF (miles): no 'km' word in the rendered text except the L/100 km line",
           not re.search(r"\b[Kk]m\b", _KM_OK.sub("", _pt)), str(re.findall(r".{0,40}\b[Kk]m\b.{0,40}", _KM_OK.sub("", _pt)))[:240])
        _rg = export.route_geometries
        try:
            export.route_geometries = lambda page: []          # nothing baked at all
            _pt0 = subprocess.run(["pdftotext", "-layout", "-", "-"], input=export.build_pdf(_la), capture_output=True).stdout.decode("utf-8")
        finally:
            export.route_geometries = _rg
        ok("🔴 PDF with nothing baked: the caption says so and does not blame Mapbox",
           "no baked route to draw" in _pt0 and "Mapbox could not be reached" not in _pt0, _pt0[:300])
        ok("…and with a baked route but no Mapbox, the schematic caption still names Mapbox (the branch the sandbox takes)",
           "schematic from the baked geometry" in _pt and "Mapbox could not be reached" in _pt)
    else:
        for _ in range(3):
            ok("(PDF wording skipped — pdftotext missing)", True)
except ImportError:
    for _ in range(3):
        ok("(PDF wording skipped — reportlab not installed here)", True)
_mp_src = read("map/index.html")
ok("🔴 the map's legend carries no first-tenant figure ('~60 km' unsurveyed, '7 package edges') — the words are generic",
   "~60" not in _mp_src and "7 package edges" not in _mp_src
   and "Where the alignment file has no surveyed main track the line is interpolated" in _mp_src
   and "Package edges · ticks from zoom 11" in _mp_src)

# =========================================================================== #
#  3. Source level: Estonian things are gated                                   #
# =========================================================================== #
fe = read("frontend/index.html")
fe_code = "\n".join(l.split("//", 1)[0] if not re.search(r"https?://", l) else l for l in fe.splitlines())
ok("the staff app names Tark Tee only in comments — every label comes from the API's provider block",
   "Tark Tee" not in fe_code and "Estonian Transport Administration" not in fe
   and "TENANT.restrictions.short" in fe and "ttSrc.restrictions_ui_label" in fe)
ok("the staff app offers the Estonian orthophoto to an Estonian tenant only",
   'val !== "maaamet-ortho" || c === "EE"' in fe and "basemapsFor(TENANT.country).map" in fe and "{BASEMAPS.map(" not in fe)
ok("the staff app hides the tenant's hidden vehicles in its pickers and the EU label option for GB",
   "const vehShown = (v, keep) => !HIDDEN_VEHICLES.includes(v) || v === keep;" in fe
   and fe.count(".filter(v => vehShown(v") >= 4 and '(k === "eu" && String(TENANT.country || "").toUpperCase() !== "GB")' in fe)
# 29 Sep 2026 first look on the LIVE GB demo: the staff app still said "bulletin" (Dashboard
# status line, Config fuel widget), "km" on a miles tenant (Dashboard status line, hover hint,
# the model notes, the route form's help text), "the usual Estonian quote" on the route form,
# an "Estonian" label column on Config → Vehicles, a literal "Fair $£" in the Dashboard note,
# and the Submit form DEFAULTED to a hidden EU vehicle because the material categories list
# the N-category entries first. Pinned at source; the render harness covers the JSX.
import re as _re2
_fits = [m.start() for m in _re2.finditer(r"\.fitBounds\(", fe)]
ok("🔴 every fitBounds in the app is inside a try (30 Sep: a NaN bounds on the Look-ahead map white-screened the whole app)",
   len(_fits) == 3 and all("try {" in fe[max(0, i - 60):i] for i in _fits), str([fe[max(0, i - 60):i] for i in _fits if "try {" not in fe[max(0, i - 60):i]])[:200])
# --- 5 Oct 2026: Appearance (Ink / Light / Dark) ------------------------------------------
_PAL = re.compile(r'(?:(?:hover|focus|disabled|group-hover|sm|md|lg):)?(?:bg|text|border|divide|ring|placeholder)-(?:white|black|transparent|current|\[color:[^\]]+\]|\[#[0-9a-fA-F]{3,8}\]|\[rgba?\([^\]]+\)\]|(?:slate|gray|red|amber|blue|emerald|green)-\d{2,3})(?:/\d{1,3})?')
_SKIP = re.compile(r'^(?:\w+:)?(?:text-white|bg-white/\d+|border-white/\d+|text-white/\d+|bg-black/\d+|\w+-transparent|\w+-current|(?:bg|text|border|ring)-\[color:var|ring-\[)')
_dark = fe[fe.index("/* DARK-REMAP-START"):fe.index("/* DARK-REMAP-END */")]
def _esc(u):
    for a, b in ((":", "\\:"), ("/", "\\/"), ("[", "\\["), ("]", "\\]"), ("(", "\\("), (")", "\\)"), (",", "\\,"), (".", "\\.")):
        u = u.replace(a, b)
    return u
_fe_wo = fe.replace(_dark, "")          # the utilities in the app markup, not the remap block itself
_uncovered = sorted({u for u in set(_PAL.findall(_fe_wo)) if not _SKIP.match(u) and ("." + _esc(u)) not in _dark})
ok("🔴 APPEARANCE: every Tailwind colour utility the app uses has a dark remap (generated block, html[data-theme=\"dark\"])",
   not _uncovered and _dark.count('html[data-theme="dark"]') >= 40, str(_uncovered)[:300])
ok("APPEARANCE: three looks as tokens — ink on :root, light = header tokens only, dark = a full palette with orange controls and color-scheme dark",
   '--hdr-logo:url("/brand/logo-header-on-dark.svg")' in fe and 'html[data-theme="light"]{' in fe and '--hdr-logo:url("/brand/logo-header-on-light.svg")' in fe
   and 'html[data-theme="dark"]{' in fe and "--navy:#FF8C14; --navy-deep:#F3F4F6; --blue:#FFA10A;" in fe and fe.count("color-scheme:dark;") == 1)
ok("APPEARANCE: the header reads header tokens, never a white literal — bar, lockup, rule, muted text, buttons, the language segment",
   ".brand-bar{background:var(--hdr-bg);color:var(--hdr-fg);" in fe and "content:var(--hdr-logo);" in fe and ".brand-rule{width:1px;height:28px;background:var(--hdr-rule);}" in fe
   and 'className="brand-bar shrink-0"' in fe and "bg-white/10" not in fe and "text-white/70" not in fe and "border-white/20" not in fe and "bg-white/25" not in fe)
ok("APPEARANCE: the picker is a 32 px button at the right of the header with a popover of menuitemradio items, closing on outside click and Escape",
   "function AppearancePicker()" in fe and "<AppearancePicker />" in fe and 'role="menuitemradio"' in fe and 'aria-haspopup="menu"' in fe
   and 'e.key === "Escape"' in fe and '.appearance-btn{width:32px;height:32px;' in fe and fe.index("<AppearancePicker />") < fe.index("Sign out\n"))
ok("APPEARANCE: the choice is per browser (localStorage modus_theme), applied before React mounts, an unknown value falls back to ink",
   'key: "modus_theme"' in fe and 'applyTheme(localStorage.getItem(THEME.key) || "ink", false)' in fe
   and 'if(!THEMES.some(([k]) => k === name)) name = "ink";' in fe and fe.index("applyTheme(localStorage.getItem") < fe.index("ReactDOM.createRoot"))
ok("APPEARANCE: charts rebuild on a change and read token colours — no literal blue left in the Dashboard's series",
   "}, [JSON.stringify(config), theme]);" in fe and 'window.Chart.defaults.color = tok("--muted"' in fe
   and 'backgroundColor: "#2563EB"' not in fe and 'backgroundColor: "#3B82F6"' not in fe and 'borderColor: "#2563EB"' not in fe and 'tok("--chart-1"' in fe)
ok("APPEARANCE: cards, the big-screen bar and the Look-ahead's white buttons read surface tokens — no literal white surface left in component CSS",
   ".card{background:var(--surface);" in fe and ".dash-wallbar{display:flex;flex-wrap:wrap;align-items:center;gap:7px;background:var(--surface);" in fe
   and 'background: "white"' not in fe and "background:#fff;" not in fe)
ok("APPEARANCE: the map's dark basemap and the exports are untouched — the exports keep the print palette (export.py has no theme code)",
   "data-theme" not in read("backend/export.py") and '["Dark", "mapbox://styles/mapbox/dark-v11"]' in fe)
ok("🔴 the Dashboard status line reads the tenant's unit word and dates the index 'published' / 'typed', never 'bulletin'",
   "· {DU()}, CO₂e and fair {CUR()} exclude the" in fe and "· km, CO₂e and fair" not in fe
   and '(${cost.index_source === "manual" ? "typed" : "published"} ${longDate(' in fe and "(bulletin ${" not in fe)
ok("🔴 the Config fuel widget dates the index 'published' / 'typed' like the compact card — no 'bulletin' word on a DESNZ screen",
   "`bulletin ${fuelDate(" not in fe and fe.count('? "published" : "typed"') >= 2)
ok("🔴 no 'km' in the Dashboard's hover hint or model notes, no '$' before the currency symbol",
   "{CUR()}/{DU()} on hover" in fe and "{CUR()}/km on hover" not in fe
   and "{DU()}, cycle and vehicles from the baked HERE route" in fe
   and "Fair {CUR()} = the cost model on Config" in fe and "Fair ${CUR()} = the cost model on Config" not in fe
   and "+ {DU()} × running" in fe and "+ km × running" not in fe)
ok("🔴 the route form's rate help names no country and reads the tenant's unit",
   "usual Estonian quote" not in fe and "per-{DU()} rate is a common haulage quote" in fe)
ok("🔴 Config → Vehicles shows the Estonian label column to an Estonian tenant only",
   '.filter(h => h !== "Estonian" || vehLangShown("ee"))' in fe and '{vehLangShown("ee") && <td' in fe)
ok("🔴 the Submit form never defaults to a hidden vehicle (first SHOWN of the category's list)",
   "const firstShown = (list) => (list || []).find(v => !HIDDEN_VEHICLES.includes(v))" in fe
   and "useState(firstShown(vehicles))" in fe and "(firstShown(suggestedVehicles) || v)" in fe
   and "useState(vehicles[0] || \"\")" not in fe and "(suggestedVehicles[0] || v)" not in fe)
ok("🔴 the Look-ahead's UNBAKED text reads the tenant's unit (clashes._dist_word), never a fixed 'km'",
   "km omitted on that line" not in read("backend/clashes.py") and "{_dist_word()} omitted on that line" in read("backend/clashes.py"))
mp = read("map/index.html")
ok("the map hides the Estonian orthophoto unless the tenant is Estonian and labels Satellite as Aerial for GB",
   "bm.hidden = (TENANT.country || '').toUpperCase() !== 'EE';" in mp
   and "'Aerial (Mapbox Satellite)' : 'Satellite (Mapbox)'" in mp)
ok("the map's restriction panel heading and source come from the API's provider",
   "cat.provider.ui_label" in mp and "restriction-heading" in mp and "Estonian Transport Administration" not in mp.split("<script>")[0])
hp = read("frontend/help/index.html")
ok("the guide documents BOTH providers and both diesel series (the one page that may name Estonia)",
   "National Highways" in hp and "Tark Tee" in hp and "DESNZ" in hp and "ROADWORKS" in hp)

# =========================================================================== #
#  4. km / miles (HU5)                                                          #
# =========================================================================== #
ok("config: distance_unit validated (km|mi), default km, stored on the tenant block",
   config.DISTANCE_UNITS == ("km", "mi") and abs(config.KM_PER_MILE - 1.609344) < 1e-9
   and config.TENANT_DEFAULTS["distance_unit"] == "km")
_doc = json.loads(json.dumps(config.load(conversions, use_cache=False)))
_doc["tenant"]["distance_unit"] = "furlongs"
ok("...a unit that is not km or mi is refused", not config.save(_doc, by="test")["ok"])
ok("export: km → mi at the sheet boundary only (60 km → 37.28 mi; £2/km → £3.22/mi); km tenant unchanged",
   abs(export._dist(60.0) - 37.282) < 0.01 and abs(export._dist_rate(2.0) - 3.219) < 0.01 and export._du() == "mi"
   and export._tkm() == "t·mi")
ok("the app's helpers convert at the display boundary (dist, distRate, distInRate, speed) and the map's (distU, fmtDist)",
   all(x in fe for x in ("const KM_PER_MI = 1.609344;", "const dist = (km) =>", "const distRate = (perKm) =>",
                          "const distInRate = (perUnit) =>", "const speed = (kph) =>", "const fmtDist = (km, dp) =>"))
   and all(x in mp for x in ("const KM_PER_MI = 1.609344;", "const distU = (km) =>", "const fmtDist = (km, dp) =>")))
ok("the app types per-distance rates in the tenant's unit and stores per km (route form, target rate, running coefficient)",
   fe.count("distInRate(") >= 3 and "rate_eur_per_km: route.rate_eur_per_km == null" in fe)

# =========================================================================== #
#  5. The GB vehicle set and the 2026 CO₂ factors                               #
# =========================================================================== #
F = conversions.load_factors()
lead = conversions.planning_vehicle_names(F)
hidden = conversions.hidden_vehicle_names(F)
ok("GB lead vehicles are the UK trade names, the EU N-category names are hidden, nothing is removed",
   lead[:3] == ["Rigid 8-wheeler (32t)", "Rigid 6-wheeler (26t)", "Artic Tipper (44t)"]
   and "Truck mixer 8 m³ (32t)" in lead and all(n.startswith("N3 ") for n in hidden if n != "Rigid 7.5t")
   and len(hidden) == 5 and all(h in conversions.vehicle_names(F) for h in hidden)
   and set(lead).isdisjoint(hidden), str((lead, hidden)))
ok("/api/meta carries hidden_vehicles and the GB lead list", main.meta()["hidden_vehicles"] == hidden
   and main.meta()["planning_vehicles"] == lead)
V = F["vehicles"]
ok("the two GB truck mixers carry a stated payload basis (ASSUMPTION), drum m³ and the rigid CO₂ factor",
   V["Truck mixer 8 m³ (32t)"]["payload_t"] == 18.0 and V["Truck mixer 8 m³ (32t)"]["payload_m3"] == 7.5
   and "ASSUMPTION" in V["Truck mixer 8 m³ (32t)"]["_payload_basis"]
   and V["Truck mixer 6 m³ (26t)"]["payload_t"] == 14.4 and V["Truck mixer 6 m³ (26t)"]["gvw_t"] == 26)
_rigid = ("Rigid 4-wheeler (18t)", "Rigid 6-wheeler (26t)", "Rigid 8-wheeler (32t)", "Truck mixer 8 m³ (32t)", "Truck mixer 6 m³ (26t)")
_artic = ("Artic Tipper (44t)", "Artic Flatbed (44t)")
ok("DESNZ 2026: rigid >17 t 0.99773 and artic >33 t 0.93939 kg CO₂e/km, average laden, with the source on every entry",
   all(V[k]["emissions_kg_co2e_per_km"] == 0.99773 and "DESNZ/DEFRA 2026" in V[k]["_emissions_basis"] for k in _rigid)
   and all(V[k]["emissions_kg_co2e_per_km"] == 0.93939 and "DESNZ/DEFRA 2026" in V[k]["_emissions_basis"] for k in _artic))
ok("...the laden/empty pair is stored with the figure and says it is not read by code yet",
   V["Rigid 8-wheeler (32t)"]["emissions_kg_co2e_per_km_empty"] == 0.77910 and V["Rigid 8-wheeler (32t)"]["emissions_kg_co2e_per_km_laden"] == 1.11639
   and V["Artic Tipper (44t)"]["emissions_kg_co2e_per_km_empty"] == 0.65383 and "NOT read by code yet" in V["Artic Tipper (44t)"]["_emissions_basis"])
ok("...and the file says the figures were read from a secondary copy and must be verified against gov.uk",
   "VERIFY against the gov.uk flat file" in V["Rigid 8-wheeler (32t)"]["_emissions_basis"])
# config.validate() returns the PROBLEMS (a list of sentences; empty = fine)
ok("a vehicle_sets entry naming an unknown vehicle, or a name both led and hidden, is refused",
   any("No such lorry" in x for x in config.validate(dict(_doc, tenant=dict(_doc["tenant"], distance_unit="mi"),
                                                          vehicle_sets={"GB": {"lead": ["No such lorry"], "hide": []}})))
   and any("both lead and hidden" in x for x in config.validate(dict(_doc, tenant=dict(_doc["tenant"], distance_unit="mi"),
                                                                     vehicle_sets={"GB": {"lead": ["Artic Tipper (44t)"], "hide": ["Artic Tipper (44t)"]}})))
   and not config.validate(dict(_doc, tenant=dict(_doc["tenant"], distance_unit="mi"))))

# =========================================================================== #
#  6. Token and time zone                                                       #
# =========================================================================== #
_js_tok = re.search(r'MAPBOX_TOKEN:\s*"([^"]+)"', read("map/config.js")).group(1)
ok("🔴 config.MAPBOX_TOKEN_DEFAULT equals the token in map/config.js (the revoked one is gone)",
   config.MAPBOX_TOKEN_DEFAULT == _js_tok and _js_tok.startswith("pk.")
   and "Y21xbnJzaTRrMDYyOTJxcXowczRxNTlxdyJ9" not in read("backend/config.py"))
ok("GB → Europe/London for the departure-time probe", here_routing.COUNTRY_TZ.get("GB") == "Europe/London"
   and "Europe/London" in here_routing.default_departure_times(country="GB")[1])

# =========================================================================== #
#  7. The probes                                                                #
# =========================================================================== #
_pf = main.diagnostics_fuel_index(probe=True)
ok("the fuel-index probe runs the DESNZ path inline and shows the resolved CSV and the parsed row",
   _pf["provider"]["key"] == "desnz" and _pf["probe"]["ok"] is True and _pf["probe"]["parsed"]["eur_per_l"] == 1.955
   and _pf["probe"]["detail"]["csv_url"].endswith(".csv") and _pf["urls"] == [fuel.DESNZ_CONTENT_API])
_pr = restrictions.diagnostics(probe=True)
ok("the restrictions probe on a GB tenant shows the NH field map, the raw property names and the mapped dates",
   _pr["provider"]["key"] == "nh_closures" and "field_map" in _pr and _pr["probe"]["sample_dates"]["from"]
   and "scheduledplannedstartdate" in _pr["probe"]["sample_property_names"])
ok("the probes are behind the admin check and shipped in main.py",
   '@app.get("/api/admin/diagnostics/fuel-index")' in _srcs["backend/main.py"]
   and "def diagnostics_fuel_index(" in _srcs["backend/main.py"] and "_nh_diagnostics(probe)" in _srcs["backend/restrictions.py"])

print(f"\n{PASS} passed, {len(FAIL)} failed")
for f in FAIL:
    print("  FAIL:", f)
sys.exit(1 if FAIL else 0)
