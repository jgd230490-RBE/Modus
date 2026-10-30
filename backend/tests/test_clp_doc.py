"""
The CLP document, 30 Oct 2026 (task 4a, PR 1) — backend/clp_doc.py and GET /api/clp/document.docx.

What is asserted: a CLP's lorries are the approved lines with an end at one of its locations (a route between two of its
own locations counted once); every figure equals the module that owns it (trips = gate_arrivals.trips_for, the hourly
profile = Gate arrivals' arrivals a day spread over gate_hours.open_minutes); the CLOCS tool's rules (busiest month,
÷ working days, quarters, weight bands); the stages from planning.clp, and a discipline no stage names kept as its own
stage; §5's ticks only on evidence; then the .docx READ BACK with python-docx: the template's headings in order, no
CLOCS guidance sentence, every placeholder in the "To complete" style and in Still to complete, the demo line and the
snapshot, the word rule; the OS grid reference; Config's check; the endpoint's edges (planner or admin, 422s).

WHAT THIS DOES NOT PROVE
------------------------
  * The HTTP layer is stubbed (http_smoke.py drives the real one).
  * How Word lays the file out (looked at by hand in Word on this laptop, 30 Oct 2026).
  * SQLite only; hand-made routes, no HERE geometry.

Run (PowerShell):  $env:PYTHONUTF8 = "1"; .venv-tests\\Scripts\\python backend/tests/test_clp_doc.py
"""
import copy
import json
import io
import math
import os
import re
import sys
import tempfile
import types
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, as test_gate_arrivals.py ----------------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)


TMP = tempfile.mkdtemp(prefix="rbe_clpdoc_")
from portable import remove_scratch_db, offline   # noqa: E402,F401  tests/portable.py
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import costing  # noqa: E402
import config  # noqa: E402
import gates  # noqa: E402
import gate_arrivals  # noqa: E402
import gate_hours  # noqa: E402
import network  # noqa: E402
import site_areas  # noqa: E402
import taxonomy  # noqa: E402
import carbon_target  # noqa: E402
import clp_doc  # noqa: E402
from docx import Document  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


ART = network.DEFAULT_PROFILE          # "Artic Tipper (44t)", 29 t a load
FLAT = "Artic Flatbed (44t)"
BAL = "Large aggregate / ballast"
SPOIL = "Earthworks / soil"
PRE = "Precast units"

# =========================================================================== #
#  0. A small scheme by hand                                                   #
# =========================================================================== #
remove_scratch_db(db, TMP)
db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db(); db.init_gates_db()
db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db(); db.init_costing_db()
db.init_tenant()
taxonomy.seed_taxonomy()
config.invalidate()
config.seed_from_file(conversions)
config.invalidate()
costing.invalidate()
carbon_target.invalidate()
access.set_current("planner123")
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc["planning"]["peak_bands"] = [{"from": "08:00", "to": "09:00"}, {"from": "15:00", "to": "16:00"}]
_doc.setdefault("tenant", {})["demo_notice"] = "Fictional scheme"
_doc["tenant"]["name"] = "Test scheme"
for _k, _v in _doc["vehicles"].items():
    if isinstance(_v, dict):
        _v["euro_class"] = "Euro VI"
ok("a scheme with two peak bands, a demo notice and every lorry Euro VI passes Config", config.validate(_doc) == [],
   str(config.validate(_doc)))
config.save(_doc, by="tester")
config.invalidate()
factors = conversions.load_factors()
WD = 22.0

PIT = network.create_location("Pit West", "origin", lat=52.55, lon=-0.95, loc_type="Quarry", supplies=[BAL])["id"]
YARD = network.create_location("Yard East", "destination", lat=52.6447, lon=-0.8822, loc_type="Compound", receives=[BAL, PRE], supplies=[BAL, SPOIL])["id"]
PILE = network.create_location("Pile North", "destination", lat=52.58, lon=-0.92, loc_type="Stockpile", receives=[BAL])["id"]
TIP = network.create_location("Tip South", "destination", lat=52.50, lon=-0.90, loc_type="Landfill", receives=[SPOIL])["id"]
RAIL = network.create_location("Rail Head", "origin", lat=52.48, lon=-0.92, loc_type="Railhead", supplies=[PRE])["id"]
R_IN = network.create_route(PIT, YARD, material_category=BAL, route_id="R001", ipt="IPT1")["id"]
R_OUT = network.create_route(YARD, TIP, material_category=SPOIL, route_id="R002", ipt="IPT1")["id"]
R_BTW = network.create_route(YARD, PILE, material_category=BAL, route_id="R003", ipt="IPT2")["id"]
R_RAIL = network.create_route(RAIL, YARD, material_category=PRE, route_id="R004", ipt="IPT2")["id"]
R_ELSE = network.create_route(PIT, PILE, material_category=BAL, route_id="R005", ipt="IPT1")["id"]


def add_geom(rid, km, profile):
    for leg in ("loaded", "return"):
        db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
                   "duration_hr, computed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                   (db.current_tenant(), rid, profile, leg, 0, "[]", km, 0.5, "2026-10-30T00:00:00Z"))


add_geom(R_IN, 32.5, ART)
add_geom(R_OUT, 20.0, ART)
add_geom(R_BTW, 6.0, ART)
add_geom(R_RAIL, 22.1, FLAT)          # R005 is never drawn


def _line(rid, disc, sect, month, veh, ipt, mat, approve=True):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline=disc, section_id=sect, material_type=mat, material_description=None,
        vehicle_type=veh, submitted_by="tester", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=q) for m, q in month], ipt=ipt))
    if approve:
        main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline=disc, section_id=sect)


# month 10 = Oct 2026, 11 = Nov 2026, 13 = Jan 2027 (a new quarter and a new year)
_line(R_IN, "substructure", "WS1", [(10, 2900.0), (11, 5800.0), (13, 1450.0)], ART, "IPT1", BAL)   # 100, 200, 50 loads
_line(R_OUT, "earthworks", "WS2", [(11, 2900.0)], ART, "IPT1", SPOIL)                             # 100 loads out
_line(R_BTW, "substructure", "WS3", [(10, 580.0)], ART, "IPT2", BAL)                              # 20 loads yard -> pile
_line(R_RAIL, "structures", "WS4", [(13, 1100.0)], FLAT, "IPT2", PRE)                             # precast off the railhead
_line(R_ELSE, "earthworks", "WS5", [(10, 2900.0)], ART, "IPT1", BAL)                              # pit -> pile only
_line(R_IN, "earthworks", "WS6", [(12, 290.0)], ART, "IPT1", BAL, approve=False)                  # waiting for approval
config.invalidate()
factors = conversions.load_factors()

G_IN = gates.create_gate(YARD, "Yard gate in", 52.6447, -0.8822, direction="access", safety_minutes=3,
                         is_default=True, open_from="07:30", open_to="18:00", open_days="1111100",
                         note="Check in here")
G_OUT = gates.create_gate(YARD, "Yard gate out", 52.6449, -0.8820, direction="both", open_from="07:30",
                          open_to="18:00", closed_from="2027-03-01", closed_to="2027-05-31")
ok("the two gates are made", "id" in G_IN and "id" in G_OUT, f"{G_IN} {G_OUT}")
G_IN, G_OUT = G_IN.get("id"), G_OUT.get("id")
_ar = site_areas.create_area(YARD, "Pier bases", activity="unload", service_minutes=5, bays=1, level="river terrace")
ok("the work area is made", "error" not in _ar, str(_ar))
db.execute("INSERT INTO forecast_weeks (tenant_id, route_id, month_index, discipline, section_id, week_index, planned_qty, "
           "unit, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
           (db.current_tenant(), R_IN, 10, "substructure", "WS1", 40, 725.0, "t", "confirmed"))


def trips(rid, disc, sect, m):
    row = db.query("SELECT * FROM forecasts WHERE tenant_id = ? AND route_id = ? AND discipline = ? AND section_id = ? "
                   "AND month_index = ?", (db.current_tenant(), rid, disc, sect, m))[0]
    return gate_arrivals.trips_for(row, factors)


# =========================================================================== #
#  1. Which lorries                                                            #
# =========================================================================== #
D1 = clp_doc.gather([YARD], factors=factors, working_days=WD)
T1 = D1["totals"]
exp_in = (trips(R_IN, "substructure", "WS1", 10) + trips(R_IN, "substructure", "WS1", 11)
          + trips(R_IN, "substructure", "WS1", 13) + trips(R_RAIL, "structures", "WS4", 13))
exp_out = trips(R_OUT, "earthworks", "WS2", 11) + trips(R_BTW, "substructure", "WS3", 10)
ok("⭐ one site: its lorries = approved lines with an end there; in and out by the trips rule (gate_arrivals.trips_for)",
   T1["in"] == exp_in and T1["out"] == exp_out and T1["between"] == 0 and T1["trips"] == exp_in + exp_out, str(T1))
ok("...the hand sums: 100 + 200 + 50 ballast in, 100 spoil + 20 ballast out at 29 t a load",
   trips(R_IN, "substructure", "WS1", 11) == 200 and exp_out == 120, f"{exp_in} {exp_out}")
ok("...the pit → pile route (no end at the yard) is left out; the waiting line is a note, never a figure",
   "R005" not in {r["route_id"] for r in D1["routes"]} and T1["pending_lines"] == 1
   and all(l["month"] != 12 for l in D1["lines"]), str(T1["pending_lines"]))
D2 = clp_doc.gather([YARD, PILE], factors=factors, working_days=WD)
T2 = D2["totals"]
ok("⭐ two sites in one CLP: the yard → pile route is BETWEEN them and counted once",
   T2["between"] == trips(R_BTW, "substructure", "WS3", 10) == 20
   and T2["trips"] == T1["trips"] + trips(R_ELSE, "earthworks", "WS5", 10)
   and T2["out"] == trips(R_OUT, "earthworks", "WS2", 11), str(T2))
DP = clp_doc.gather([PILE], factors=factors, working_days=WD)
ok("...and the same route is 'in' in the pile's own CLP (a location can sit in more than one CLP)",
   DP["totals"]["in"] == 20 + 100 and DP["totals"]["between"] == 0, str(DP["totals"]))
DQ = clp_doc.gather([PIT], factors=factors, working_days=WD)
ok("any location can have one (the user, 30 Oct): a quarry's CLP holds its lorries out, and no gate lists none",
   DQ["totals"]["out"] == 350 + 100 and DQ["totals"]["in"] == 0 and DQ["sites"][0]["gates"] == []
   and DQ["hourly"] == [], str(DQ["totals"]))
ok("the same location twice, or with spaces, is one location",
   clp_doc.gather([YARD, f" {YARD} "], factors=factors, working_days=WD)["location_ids"] == [YARD])
e0 = _raises(lambda: clp_doc.gather([], factors=factors))
e1 = _raises(lambda: clp_doc.gather(["L999"], factors=factors))
ok("no location, or an unknown one, is a ValueError in words",
   isinstance(e0, ValueError) and str(e0) == "Pick at least one location."
   and isinstance(e1, ValueError) and str(e1) == "No such location: L999.", f"{e0!r} {e1!r}")
ok("tonnes: the lines' own tonnes", abs(T1["tonnes"] - sum(l["tonnes"] for l in D1["lines"])) < 0.05
   and abs(sum(l["tonnes"] for l in D1["lines"] if l["route_id"] == R_IN) - 10150.0) < 0.05, str(T1["tonnes"]))
ok("materials by direction", {(m["material"], m["direction"]) for m in D1["materials"]}
   == {(BAL, "in"), (SPOIL, "out"), (BAL, "out"), (PRE, "in")}, str(D1["materials"]))
ok("the routes' other ends, by name", D1["other_ends"] == ["Pile North", "Pit West", "Rail Head", "Tip South"],
   str(D1["other_ends"]))

# =========================================================================== #
#  2. Stages and the CLOCS tool's rules                                        #
# =========================================================================== #
S = {s["number"]: s for s in D1["stages"]}
ok("⭐ the infrastructure stages by default: 3.2.1–3.2.7, Pavement and surfacing as 3.2.7 (the user's Q3)",
   [s["number"] for s in D1["stages"]] == ["3.2.1", "3.2.2", "3.2.3", "3.2.4", "3.2.5", "3.2.6", "3.2.7"]
   and S["3.2.7"]["name"] == "Pavement and surfacing" and S["3.2.2"]["disciplines"] == ["earthworks"]
   and D1["programme"] == "infrastructure" and D1["prefix"] == "3.2")
sub = S["3.2.3"]
ok("⭐ a stage's trips (monthly) = its busiest month; daily = ÷ working days, one decimal (the CLOCS tool's rule)",
   sub["peak_month"] == 11 and sub["peak_trips"] == 200 and sub["peak_daily"] == round(200 / 22, 1)
   and sub["first"] == 10 and sub["last"] == 13 and sub["trips_total"] == 100 + 200 + 50 + 20
   and sub["routes"] == ["R001", "R003"], str({k: sub[k] for k in ("peak_month", "peak_trips", "peak_daily", "first", "last")}))
ok("...an empty stage has no dates and no zero figure",
   S["3.2.1"]["first"] is None and S["3.2.1"]["peak_daily"] is None and S["3.2.1"]["trips_total"] == 0)
ok("...a tie goes to the earlier month", clp_doc.gather([PIT], factors=factors, working_days=WD)["totals"]["peak_month"] == 10)
ok("...the site's peak period = its busiest month over every stage",
   T1["peak_month"] == 11 and T1["peak_trips"] == 300 and T1["peak_daily"] == round(300 / 22, 1), str(T1))
ok("periods in quarters, as the template prints them", clp_doc.period_words(10, 13, quarters=True) == "Q4 2026 – Q1 2027"
   and clp_doc.period_words(10, 12, quarters=True) == "Q4 2026" and clp_doc.quarter_label(13) == "Q1 2027"
   and clp_doc.period_words(10, 13) == "Oct 2026 – Jan 2027", clp_doc.period_words(10, 13))
ok("weight bands from the gross weight: 3.5 / 7.5 / over, and 'Weight not typed' when none",
   clp_doc.band_of(3.5) == "3.5 t and under" and clp_doc.band_of(7.5) == "3.5 t to 7.5 t"
   and clp_doc.band_of(44) == "Over 7.5 t" and clp_doc.band_of(None) == "Weight not typed")
_gvw = (factors["vehicles"].get(ART) or {}).get("gvw_t")
ok("⭐ a stage's vehicles in its busiest month, banded by the vehicle's own gross weight",
   sub["vehicles"] == [{"vehicle": ART, "label": conversions.vehicle_label(factors, ART), "gvw_t": _gvw,
                        "band": clp_doc.band_of(_gvw), "trips": 200}], str(sub["vehicles"]))
_bad = copy.deepcopy(_doc)
_bad["planning"]["clp"] = {"programme": "infrastructure", "stages": [{"name": "Dig", "disciplines": ["earthworks"]},
                                                                     {"name": "Build", "disciplines": ["earthworks"]}]}
ok("🔴 Config refuses a discipline in two stages",
   "planning.clp.stages: earthworks is in two stages (1 and 2)" in config.validate(_bad), str(config.validate(_bad)))
_bad["planning"]["clp"] = {"programme": "bridges"}
ok("...an unknown programme", "planning.clp.programme must be infrastructure or building" in config.validate(_bad))
_bad["planning"]["clp"] = {"stages": [{"name": f"S{i}", "disciplines": []} for i in range(9)]}
ok("...more than 8 stages", "planning.clp.stages: at most 8 stages" in config.validate(_bad))
_bad["planning"]["clp"] = {"stages": [{"name": " ", "disciplines": []}]}
ok("...a stage with no name", "planning.clp.stages[0] needs a name" in config.validate(_bad))
ok("...and passes none (the defaults) or a programme alone", clp_doc.settings_problem(None) is None
   and clp_doc.settings_problem({"programme": "building"}) is None)
_own = copy.deepcopy(_doc)
_own["planning"]["clp"] = {"programme": "building", "stages": [{"number": "3.1.1", "name": "Groundworks",
                                                                "disciplines": ["earthworks", "substructure"]}]}
ok("a scheme's own stages pass Config", config.validate(_own) == [], str(config.validate(_own)))
DB_ = clp_doc.gather([YARD], factors=_own, working_days=WD)
_lab = (taxonomy.get_discipline("structures") or {}).get("label") if hasattr(taxonomy, "get_discipline") else None
ok("⭐ a discipline no stage names becomes a stage of its own after the last (no load lost); building = 3.1",
   [s["number"] for s in DB_["stages"]] == ["3.1.1", "3.1.2"] and DB_["stages"][1]["added"]
   and DB_["stages"][1]["disciplines"] == ["structures"] and DB_["prefix"] == "3.1"
   and sum(s["trips_total"] for s in DB_["stages"]) == DB_["totals"]["trips"],
   str([(s["number"], s["name"]) for s in DB_["stages"]]))

# =========================================================================== #
#  3. Gates and the hourly profile                                             #
# =========================================================================== #
site = D1["sites"][0]
gin = next(g for g in site["gates"] if g["id"] == G_IN)
om = gate_hours.open_minutes(gates.get_gate(G_IN), factors)
ok("⭐ a gate: its use, days, hours, check-in and capacity = gate_hours.capacity(open minutes, check-in)",
   gin["direction"] == "in" and gin["days"] == "Mon–Fri" and gin["hours"] == om["hours"] and gin["check_in"] == 3
   and gin["capacity"] == round(gate_hours.capacity(om["minutes"], 3), 1) and gin["note"] == "Check in here", str(gin))
gout = next(g for g in site["gates"] if g["id"] == G_OUT)
ok("...the other gate: in and out, its closed period in words",
   gout["direction"] == "in and out" and gout["closed"] == gates.closed_words(gates.get_gate(G_OUT))
   and gout["closed"].startswith("closed 1 Mar"), str(gout))
ok("the work areas", site["areas"] == [{"name": "Pier bases", "activity": "unload", "bays": 1, "minutes": 5,
                                        "level": "river terrace"}], str(site["areas"]))
H = D1["hourly"]
ar = gate_arrivals.arrivals(11, {"role": "planner"}, factors=factors, working_days=WD)
arg = {g["gate_id"]: g for g in ar["gates"] if g["location_id"] == YARD and g["arrivals_per_day"]}
ok("⭐ the hourly profile: the site's busiest month, one table per gate with arrivals, Gate arrivals' own figure",
   sorted(h["gate"] for h in H) == sorted(g["gate_name"] for g in arg.values())
   and all(h["month"] == 11 for h in H)
   and all(h["arrivals_per_day"] == next(g for g in arg.values() if g["gate_name"] == h["gate"])["arrivals_per_day"]
           for h in H), str([(h["gate"], h["arrivals_per_day"]) for h in H]))
hin = next((h for h in H if h["gate"] == "Yard gate in"), None)
per = {x["hour"]: x for x in (hin or {}).get("per_hour", [])}
ok("...spread over the open minutes: none in the bands (08, 15) or outside the hours; half an hour at 07:30",
   hin is not None and per[8]["lorries"] == 0 and per[15]["lorries"] == 0 and per[6]["lorries"] == 0
   and per[18]["lorries"] == 0 and per[7]["open_minutes"] == 30 and per[9]["open_minutes"] == 60
   and per[8]["minutes_between"] is None, str([per.get(h) for h in (7, 8, 9)]))
ok("...the hours add back to the day (to rounding), and minutes between = open minutes ÷ lorries",
   hin is not None and abs(sum(x["lorries"] for x in hin["per_hour"]) - hin["arrivals_per_day"]) < 0.3
   and abs(per[9]["minutes_between"] - 60 / (hin["arrivals_per_day"] * 60 / om["minutes"])) < 0.2, str(per.get(9)))
H2 = clp_doc.gather([YARD, PILE], factors=factors, working_days=WD)["hourly"]
ok("two sites: profiles only where a site has a gate (the pile has none)", {h["site"] for h in H2} == {"Yard East"})

# =========================================================================== #
#  4. §5's evidence                                                             #
# =========================================================================== #
M = D1["measures"]
ok("⭐ ticks only where the plan already does it: scheduling (a confirmed week), out of peak (bands), vehicle choice",
   M["Delivery scheduling"]["tick"] and M["Re-timing for out of peak deliveries"]["tick"] and M["Vehicle choice"]["tick"]
   and M["Re-timing for out of peak deliveries"]["evidence"] == f"No deliveries {gate_hours.peak_words(D1['peak_bands'])} at any gate.")
ok("...every route drawn: designated routes ticked", M["Adherence to designated routes"]["tick"]
   and M["Adherence to designated routes"]["evidence"].startswith("4 approved routes"), str(M["Adherence to designated routes"]))
ok("...rail and precast: the evidence, never a tick; Euro VI said, not ticked",
   M["Freight by rail*"]["tick"] is False and "Rail Head" in M["Freight by rail*"]["evidence"]
   and M["DfMA and offsite manufacture"]["tick"] is False
   and M["Safety and environmental standards and programmes"]["tick"] is False
   and M["Safety and environmental standards and programmes"]["evidence"] == "Every lorry type planned is Euro VI.")
ok("...no evidence, no entry (out of hours, holding areas, waterways, consolidation, re-use, staff travel)",
   not any(k in M for k in ("Re-timing for out of hours deliveries", "Use of holding areas and vehicle call off areas",
                            "Freight by water*", "Use of logistics and consolidation centres", "Re-use of material on site",
                            "Implement a staff travel plan")))
MP = clp_doc.gather([PILE], factors=factors, working_days=WD)["measures"]
ok("🔴 a route not drawn: designated routes NOT ticked, and the route named",
   MP["Adherence to designated routes"]["tick"] is False
   and MP["Adherence to designated routes"]["evidence"] == "1 of 2 routes drawn as lorry routes; not drawn yet: R005.",
   str(MP.get("Adherence to designated routes")))
ok("...and no confirmed week on its routes: no scheduling tick", "Delivery scheduling" not in MP)

# =========================================================================== #
#  5. The Word file, read back                                                 #
# =========================================================================== #
B1 = clp_doc.build(D1)
ok("a .docx (a zip)", B1[:2] == b"PK")
W = Document(io.BytesIO(B1))
paras = W.paragraphs
heads = [p.text for p in paras if p.style.name.startswith("Heading")]
want = ["Still to complete", "Contents", "1 Introduction", "1.1 CLP objectives", "1.2 Site context",
        "1.3 Development proposal", "1.4 CLP structure", "2 Context, considerations and challenges", "2.1 Policy context",
        "2.2 Three plans at different scales",
        "2.3 Local access including highway, public transport, cycling, walking and waterways",
        "2.3.1 Highways, carriageways and footways", "2.3.2 Railway / underground", "2.3.3 Bus routes", "2.3.4 Cycling",
        "2.3.5 Waterways", "2.4 Community considerations", "2.4.1 Local policy",
        "2.4.2 Schools, hospitals and other sensitive places", "3 Construction programme and methodology",
        "3.2 Infrastructure programme phases", "3.2.1 Site establishment, clearance and alterations",
        "3.2.2 Excavation and foundations", "3.2.3 Sub-structure", "3.2.4 Super-structure",
        "3.2.5 Services and systems installation", "3.2.6 Fit-out, testing and commissioning", "3.2.7 Pavement and surfacing",
        "4 Vehicle routing and site access", "5 Strategies to reduce impacts",
        "5.1.1 Measures influencing construction vehicles and deliveries", "5.1.2 Measures to encourage sustainable freight",
        "5.1.3 Material procurement measures", "5.1.4 Other measures", "6 Estimated vehicle movements",
        "7 Implementing, monitoring and updating"]
pos = [heads.index(h) if h in heads else -1 for h in want]
ok("⭐ the template's section numbers and headings, all present and in order", -1 not in pos and pos == sorted(pos),
   str([h for h, p in zip(want, pos) if p == -1]))
ok("...3.1 (a building programme's stages) left out of an infrastructure programme, and said so",
   not any(h.startswith("3.1") for h in heads)
   and any(p.text == "3.1 (stages of a building programme) is left out: this programme is set as infrastructure."
           for p in paras))


def all_text(doc):
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    for s in doc.sections:
        out += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return out


JOINED = "\n".join(all_text(W))
GUIDE = ("outline measures that will be undertaken", "The charts in this section are to be made using",
         "Delete Medium/Higher Impact", "outline the system that will be implemented", "outline proposals for",
         "identify suppliers who have been recognised", "outline the feasibility of delivering")
ok("🔴 no CLOCS guidance sentence copied (the template's own instructions, a negative list)",
   not any(g.lower() in JOINED.lower() for g in GUIDE), str([g for g in GUIDE if g.lower() in JOINED.lower()]))
ok("🔴 the word rule over every string in the file (the landing list, truck, no back-end words)",
   clp_doc.banned_words(JOINED) == [], str(clp_doc.banned_words(JOINED)[:5]))
ok("...and the rule catches them", clp_doc.banned_words("a Secure API on the server, real-time, AI, Truck, said") ==
   ["secure", "api", "server", "real-time", "truck", "AI"], str(clp_doc.banned_words("a Secure API on the server, real-time, AI, Truck, said")))
ok("🔴 no team slot id and no team word: a CLP is the site's", re.search(r"\bIPT ?[1-6]\b", JOINED) is None
   and "team" not in JOINED.lower())
ok("⭐ not a CLOCS publication, said once on the cover", JOINED.count("this is not a CLOCS publication") == 1
   and "The section numbers follow the CLOCS outline CLP template" in JOINED and "CLOCS logo" not in JOINED)
_snap = D1["snapshot"]
_snap_short = f"{_snap.day} {clp_doc.MONTHS[_snap.month - 1]} {_snap.year}"
_foot = W.sections[0].footer.paragraphs[0].text
ok("⭐ a dated snapshot on the cover; the footer: demo line, CLP, Draft, snapshot date, page",
   f"Draft · a snapshot of the approved plan in Wayscope on {_snap_short} at {_snap:%H:%M}." in JOINED
   and _foot.startswith(f"DEMO · Fictional scheme · CLP Yard East · Draft · snapshot {_snap_short} · page "), _foot)
ok("...the demo line heads every page and the cover says the figures are fictional",
   W.sections[0].header.paragraphs[0].text == "DEMO · Fictional scheme"
   and "DEMO · Fictional scheme: every site, route and figure in this document is fictional." in JOINED)
todo_paras = [p for p in paras if p.text.startswith("To complete: ")]
todo_cells = [c for t in W.tables for row in t.rows for c in row.cells if c.text.startswith("To complete: ")]
ok("⭐ every placeholder is in the 'To complete' style (paragraphs and the cover's cells)",
   len(todo_paras) > 20 and all(p.style.name == clp_doc.TODO_STYLE for p in todo_paras)
   and len(todo_cells) == 7 and all(c.paragraphs[0].style.name == clp_doc.TODO_STYLE for c in todo_cells),
   f"{len(todo_paras)} {len(todo_cells)}")
ok("...and the style itself is the yellow box with the left bar",
   'w:fill="FFF6C8"' in W.styles[clp_doc.TODO_STYLE].element.xml and 'w:color="C89500"' in W.styles[clp_doc.TODO_STYLE].element.xml)
listed = [p.text for p in paras if p.style.name == "List Number"]
ok("⭐ ...each one listed once under Still to complete, which comes before the contents and section 1",
   len(listed) == len(todo_paras) + len(todo_cells)
   and all(any(x.endswith(": " + p.text[len("To complete: "):]) for x in listed) for p in todo_paras)
   and all(any(x.endswith(": " + c.text[len("To complete: "):]) for x in listed) for c in todo_cells)
   and paras.index(next(p for p in paras if p.style.name == "List Number"))
   < paras.index(next(p for p in paras if p.text == "Contents")),
   f"{len(listed)} vs {len(todo_paras) + len(todo_cells)}")
ok("...the list says where each one stands", "Cover, landowner: who owns the land the site sits on." in listed
   and "5.1.1 Vehicle choice: how this measure will be carried out here, or why it is not used." in listed
   and "3.2.3 Sub-structure: the plant used and the works in this stage." in listed
   and "5.1.2 Freight by rail: how this measure will be carried out here, or why it is not used." in listed)
ok("no other placeholder anywhere (no empty brackets, no TBC, no XXX, no lorem)",
   not re.search(r"\[\s*\]|\bTBC\b|XXX|lorem", JOINED, re.IGNORECASE))
_z = zipfile.ZipFile(io.BytesIO(B1))
ok("the contents and page numbers are Word fields, updated when the file opens",
   b"TOC \\o" in _z.read("word/document.xml") and b"PAGE" in _z.read("word/footer1.xml")
   and b"updateFields" in _z.read("word/settings.xml"))
ok("the cover's position: the OS grid reference", "Yard East: 52.6447° N, 0.8822° W · OS grid reference SK 757 058" in JOINED)


def table_with(doc, first_head):
    return next((t for t in doc.tables if t.rows[0].cells[0].text == first_head), None)


t6 = table_with(W, "Construction stage")
rows6 = {r.cells[0].text: [c.text for c in r.cells] for r in t6.rows[1:]}
ok("⭐ §6: the stage's quarters, busiest month and daily figure, as gathered",
   rows6.get("Sub-structure") == ["Sub-structure", "Q4 2026 – Q1 2027", "200 (Nov 2026)", f"{round(200 / 22, 1):g}"]
   and rows6.get("Peak period of construction") == ["Peak period of construction", "Nov 2026", "300", f"{round(300 / 22, 1):g}"]
   and rows6.get("Site establishment, clearance and alterations", ["", ""])[1] == "No lorry loads in Wayscope", str(rows6))
ok("...and the waiting line said under it", "1 line waiting for approval is not in these figures." in JOINED)
tv = table_with(W, "Stage (busiest month)")
ok("§6: number and type of vehicle in each stage's busiest month, with the band",
   tv is not None and [c.text for c in tv.rows[1].cells] == [
       "Excavation and foundations (Nov 2026)", conversions.vehicle_label(factors, ART),
       (f"{_gvw:g} t" if _gvw is not None else "not typed"), clp_doc.band_of(_gvw), "100"],
   str([c.text for c in tv.rows[1].cells]) if tv else "no table")
tr = table_with(W, "Route")
rr = {r.cells[0].text: [c.text for c in r.cells] for r in tr.rows[1:]}
_u = D1["unit"]
ok("§4: each route from → to, distance in the scheme's unit, lorry, months, loads",
   rr["R001"][1] == "Pit West → Yard East" and rr["R001"][2] == clp_doc.n(D1["routes"][0]["distance"])
   and rr["R001"][4] == "Oct 2026 – Jan 2027" and rr["R001"][5] == "350" and rr["R002"][5] == "100"
   and rr["R004"][1] == "Rail Head → Yard East" and "R005" not in rr
   and tr.rows[0].cells[2].text == ("Miles" if _u == "mi" else "Km"), str(rr))
tg = table_with(W, "Gate")
_g1 = [c.text for c in tg.rows[1].cells] if tg else []
ok("§4: the gates' hours with the no-delivery bands, check-in, capacity, closed period and note",
   _g1[:2] == ["Yard gate in", "in"] and _g1[2].startswith(f"Mon–Fri {om['hours']}, not ")
   and _g1[3] == "3 min" and _g1[4] == clp_doc.n(gin["capacity"]) and _g1[5] == "Check in here"
   and tg.rows[2].cells[5].text.startswith("closed 1 Mar"), str(_g1))
tc = table_with(W, "Higher impact site planned measures")
ticks = {r.cells[0].paragraphs[0].text: r.cells[1].text for r in tc.rows[1:]}
ok("⭐ §5's checklist: the template's rows in order, ticks only where gathered",
   list(ticks) == [x for grp, rows in clp_doc.MEASURES for x in (grp,) + rows]
   and {k for k, v in ticks.items() if v == "✔"} == {k for k, v in M.items() if v["tick"]}, str(ticks))
th = next((t for t in W.tables if t.rows[0].cells[0].text == "Hour"), None)
_hr = [[c.text for c in r.cells] for r in th.rows[1:]] if th else []
ok("§6's hourly table: open hours only, from the gathered profile",
   th is not None and _hr[0][0] == "07:00–08:00" and not any(r[0].startswith("08:00") for r in _hr)
   and not any(r[0].startswith("06:00") for r in _hr) and _hr[0][1] == f"{per[7]['lorries']:g}"
   and _hr[0][2] == f"{per[7]['minutes_between']:g}", str(_hr[:3]))
ok("...said to be the plan's spread, not a count", "It is the plan's spread, not a count." in JOINED)
ok("§7: a site with no CLP cap is a placeholder, never a zero",
   any(p.text.startswith("To complete: a CLP cap (lorry visits a day) for Yard East") for p in todo_paras)
   and "CLP caps typed in Wayscope" not in JOINED)
WM = Document(io.BytesIO(clp_doc.build(D1, impact="medium")))
ok("medium impact picks the Medium checklist", table_with(WM, "Medium impact site planned measures") is not None
   and table_with(WM, "Higher impact site planned measures") is None)
ok("an unknown impact is refused", isinstance(_raises(lambda: clp_doc.build(D1, impact="low")), ValueError))
W2 = Document(io.BytesIO(clp_doc.build(D2)))
J2 = "\n".join(all_text(W2))
ok("two sites: one gates heading each, the loads between them said once, the cover's title both names",
   [p.text for p in W2.paragraphs if p.text.startswith("Gates at ")] == ["Gates at Yard East", "Gates at Pile North"]
   and "20 loads run between sites in this plan; they are counted once." in J2
   and "No gates recorded for this location." in J2 and "Yard East, Pile North" in J2)
ok("the file name: CLP <site> <date>.docx; a set of sites by their names",
   clp_doc.file_name(D1) == f"CLP Yard East {D1['snapshot']:%Y-%m-%d}.docx"
   and clp_doc.file_name(D2) == f"CLP Yard East, Pile North {D2['snapshot']:%Y-%m-%d}.docx")
WE = Document(io.BytesIO(clp_doc.build(clp_doc.gather([TIP], factors=factors, working_days=WD))))
ok("a location with lorries only one way still builds, every stage it lacks a placeholder",
   "Tip South" in "\n".join(all_text(WE)))
_plain = copy.deepcopy(_doc)
_plain["tenant"]["demo_notice"] = ""
config.save(_plain, by="tester")
config.invalidate()
WP = Document(io.BytesIO(clp_doc.build(clp_doc.gather([YARD], factors=conversions.load_factors(), working_days=WD))))
ok("a tenant that is not a demo: no DEMO line anywhere", "DEMO" not in "\n".join(all_text(WP)))
config.save(_doc, by="tester")
config.invalidate()

# =========================================================================== #
#  6. The OS grid reference                                                    #
# =========================================================================== #
ok("⭐ the OS grid reference: Big Ben TQ 302 796, the yard SK 757 058, Paris none",
   clp_doc.grid_ref(51.50072, -0.12462) == "TQ 302 796" and clp_doc.grid_ref(52.6447, -0.8822) == "SK 757 058"
   and clp_doc.grid_ref(48.8566, 2.3522) is None and clp_doc.grid_ref(None, None) is None,
   f"{clp_doc.grid_ref(51.50072, -0.12462)} {clp_doc.grid_ref(52.6447, -0.8822)}")

# =========================================================================== #
#  7. The endpoint                                                             #
# =========================================================================== #
R = main.clp_document_docx(locations=f"{YARD},{PILE}")
ok("⭐ GET /api/clp/document.docx?locations=…: a Word file with its name",
   R.body[:2] == b"PK" and R.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
   and R.headers["Content-Disposition"].startswith('attachment; filename="CLP Yard East, Pile North '), str(R.headers))
for args, word in ((dict(locations=""), "Pick at least one location."), (dict(locations="L999"), "No such location: L999."),
                   (dict(locations=YARD, impact="low"), "impact must be higher or medium")):
    e = _raises(lambda: main.clp_document_docx(**args))
    ok(f"422 in words: {word}", getattr(e, "status_code", None) == 422 and getattr(e, "detail", "") == word, repr(e))
ok("impact=medium answers too", main.clp_document_docx(locations=YARD, impact="Medium").body[:2] == b"PK")
os.environ["IPT2_CODE"] = "ipt2-secret"
try:
    access.set_current("ipt2-secret")
    e403 = _raises(lambda: main.clp_document_docx(locations=YARD))
finally:
    os.environ.pop("IPT2_CODE", None)
    access.set_current("planner123")
ok("🔴 a team code cannot download it (planner or admin only)", getattr(e403, "status_code", None) == 403, repr(e403))
access.set_current("")
e401 = _raises(lambda: main.clp_document_docx(locations=YARD))
access.set_current("planner123")
ok("🔴 no sign-in: 401", getattr(e401, "status_code", None) == 401, repr(e401))
MAIN = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
ok("kept in the staff cache until the next write, keyed by the locations and the checklist",
   '_staff_bytes("clp/document.docx", {"locations": ids, "impact": imp}' in MAIN)

# =========================================================================== #
#  8. The source                                                               #
# =========================================================================== #
SRC = open(os.path.join(BACKEND, "clp_doc.py"), encoding="utf-8").read()
ok("🔴 clp_doc reads only: no INSERT, UPDATE or DELETE", not re.search(r"\b(INSERT|UPDATE|DELETE)\b", SRC))
sqls = re.findall(r'"(SELECT[^"]*)"', SRC)
ok("🔴 every SQL literal carries the tenant predicate", sqls and all("tenant_id = ?" in q for q in sqls), str(sqls))
ok("🔴 no BOOLEAN compared with a number", not re.search(r"\bactive\s*=\s*1\b", SRC))
ok("🔴 no team slot id in the module", re.search(r"\bIPT ?[1-6]\b", SRC) is None)
REQ = open(os.path.join(BACKEND, "requirements.txt"), encoding="utf-8").read()
ok("python-docx pinned (with lxml) in requirements.txt", "python-docx==1.2.0" in REQ and "lxml==6.1.3" in REQ)

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\ntest_clp_doc: {PASS} passed, {len(FAIL)} failed")
remove_scratch_db(db, TMP)
sys.exit(1 if FAIL else 0)
