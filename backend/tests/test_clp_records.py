"""
The CLPs as records, 31 Oct 2026 (task 4a, PR 3a) — backend/clp_records.py, the team cut in clp_doc.gather and the
/api/clp/documents endpoints.

What is asserted: a team's CLP counts its own lorries only (trips, routes, waiting lines, the gate's hourly profile) and
says so on its cover; the rows Wayscope lists by itself (an every-team CLP for each compound, site and stockpile, one per
team with lorries there) and that listing writes nothing; a team sees ITS OWN CLPs only and another team's is "not found";
a team edits its cover details, a planner the rest; + New CLP; the chain draft -> sent -> changes -> approved as Issue N,
"What changed" worked out from the plan, the file kept exactly as approved; a draft marked DRAFT; delete only without an
issue; the endpoints' 403 / 404 / 422; the tables in the organisation file and kept by a demo refresh.

WHAT THIS DOES NOT PROVE
------------------------
  * The HTTP layer is stubbed (http_smoke.py drives the real one); the page is checked in Chrome by hand.
  * SQLite only; hand-made routes, no HERE geometry; the plans are the offline schematic.

Run (PowerShell):  $env:PYTHONUTF8 = "1"; .venv-tests\\Scripts\\python backend/tests/test_clp_records.py
"""
import copy
import json
import io
import math
import os
import re
import sys
import tempfile
import types
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, as test_gate_arrivals.py ----------------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)


TMP = tempfile.mkdtemp(prefix="rbe_clprec_")
from portable import remove_scratch_db, offline   # noqa: E402,F401  tests/portable.py
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import costing  # noqa: E402
import config  # noqa: E402
import gates  # noqa: E402
import gate_arrivals  # noqa: E402
import gate_hours  # noqa: E402
import network  # noqa: E402
import site_areas  # noqa: E402
import taxonomy  # noqa: E402
import carbon_target  # noqa: E402
import clp_doc  # noqa: E402
import export  # noqa: E402
offline(export)            # PR 2 draws plans: Mapbox is never called from a test (the schematic branch)
from docx import Document  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


ART = network.DEFAULT_PROFILE          # "Artic Tipper (44t)", 29 t a load
FLAT = "Artic Flatbed (44t)"
BAL = "Large aggregate / ballast"
SPOIL = "Earthworks / soil"
PRE = "Precast units"

# =========================================================================== #
#  0. A small scheme by hand                                                   #
# =========================================================================== #
remove_scratch_db(db, TMP)
db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db(); db.init_gates_db()
db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db(); db.init_costing_db()
db.init_tenant()
taxonomy.seed_taxonomy()
config.invalidate()
config.seed_from_file(conversions)
config.invalidate()
costing.invalidate()
carbon_target.invalidate()
access.set_current("planner123")
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc["planning"]["peak_bands"] = [{"from": "08:00", "to": "09:00"}, {"from": "15:00", "to": "16:00"}]
_doc.setdefault("tenant", {})["demo_notice"] = "Fictional scheme"
_doc["tenant"]["name"] = "Test scheme"
for _k, _v in _doc["vehicles"].items():
    if isinstance(_v, dict):
        _v["euro_class"] = "Euro VI"
ok("a scheme with two peak bands, a demo notice and every lorry Euro VI passes Config", config.validate(_doc) == [],
   str(config.validate(_doc)))
config.save(_doc, by="tester")
config.invalidate()
factors = conversions.load_factors()
WD = 22.0

PIT = network.create_location("Pit West", "origin", lat=52.55, lon=-0.95, loc_type="Quarry", supplies=[BAL])["id"]
YARD = network.create_location("Yard East", "destination", lat=52.6447, lon=-0.8822, loc_type="Compound", receives=[BAL, PRE], supplies=[BAL, SPOIL])["id"]
PILE = network.create_location("Pile North", "destination", lat=52.58, lon=-0.92, loc_type="Stockpile", receives=[BAL])["id"]
TIP = network.create_location("Tip South", "destination", lat=52.50, lon=-0.90, loc_type="Landfill", receives=[SPOIL])["id"]
RAIL = network.create_location("Rail Head", "origin", lat=52.48, lon=-0.92, loc_type="Railhead", supplies=[PRE])["id"]
R_IN = network.create_route(PIT, YARD, material_category=BAL, route_id="R001", ipt="IPT1")["id"]
R_OUT = network.create_route(YARD, TIP, material_category=SPOIL, route_id="R002", ipt="IPT1")["id"]
R_BTW = network.create_route(YARD, PILE, material_category=BAL, route_id="R003", ipt="IPT2")["id"]
R_RAIL = network.create_route(RAIL, YARD, material_category=PRE, route_id="R004", ipt="IPT2")["id"]
R_ELSE = network.create_route(PIT, PILE, material_category=BAL, route_id="R005", ipt="IPT1")["id"]


def add_geom(rid, km, profile):
    for leg in ("loaded", "return"):
        db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
                   "duration_hr, computed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                   (db.current_tenant(), rid, profile, leg, 0, "[]", km, 0.5, "2026-10-30T00:00:00Z"))


add_geom(R_IN, 32.5, ART)
add_geom(R_OUT, 20.0, ART)
add_geom(R_BTW, 6.0, ART)
add_geom(R_RAIL, 22.1, FLAT)          # R005 is never drawn


def _line(rid, disc, sect, month, veh, ipt, mat, approve=True):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline=disc, section_id=sect, material_type=mat, material_description=None,
        vehicle_type=veh, submitted_by="tester", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=q) for m, q in month], ipt=ipt))
    if approve:
        main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline=disc, section_id=sect)


# month 10 = Oct 2026, 11 = Nov 2026, 13 = Jan 2027 (a new quarter and a new year)
_line(R_IN, "substructure", "WS1", [(10, 2900.0), (11, 5800.0), (13, 1450.0)], ART, "IPT1", BAL)   # 100, 200, 50 loads
_line(R_OUT, "earthworks", "WS2", [(11, 2900.0)], ART, "IPT1", SPOIL)                             # 100 loads out
_line(R_BTW, "substructure", "WS3", [(10, 580.0)], ART, "IPT2", BAL)                              # 20 loads yard -> pile
_line(R_RAIL, "structures", "WS4", [(13, 1100.0)], FLAT, "IPT2", PRE)                             # precast off the railhead
_line(R_ELSE, "earthworks", "WS5", [(10, 2900.0)], ART, "IPT1", BAL)                              # pit -> pile only
_line(R_IN, "earthworks", "WS6", [(12, 290.0)], ART, "IPT1", BAL, approve=False)                  # waiting for approval
config.invalidate()
factors = conversions.load_factors()

G_IN = gates.create_gate(YARD, "Yard gate in", 52.6447, -0.8822, direction="access", safety_minutes=3,
                         is_default=True, open_from="07:30", open_to="18:00", open_days="1111100",
                         note="Check in here")
G_OUT = gates.create_gate(YARD, "Yard gate out", 52.6449, -0.8820, direction="both", open_from="07:30",
                          open_to="18:00", closed_from="2027-03-01", closed_to="2027-05-31")
ok("the two gates are made", "id" in G_IN and "id" in G_OUT, f"{G_IN} {G_OUT}")
G_IN, G_OUT = G_IN.get("id"), G_OUT.get("id")
_ar = site_areas.create_area(YARD, "Pier bases", activity="unload", service_minutes=5, bays=1, level="river terrace")
ok("the work area is made", "error" not in _ar, str(_ar))
db.execute("INSERT INTO forecast_weeks (tenant_id, route_id, month_index, discipline, section_id, week_index, planned_qty, "
           "unit, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
           (db.current_tenant(), R_IN, 10, "substructure", "WS1", 40, 725.0, "t", "confirmed"))


def trips(rid, disc, sect, m):
    row = db.query("SELECT * FROM forecasts WHERE tenant_id = ? AND route_id = ? AND discipline = ? AND section_id = ? "
                   "AND month_index = ?", (db.current_tenant(), rid, disc, sect, m))[0]
    return gate_arrivals.trips_for(row, factors)



import base64  # noqa: E402
import clp_records  # noqa: E402
import tenant_package  # noqa: E402
for _tid, _lab in (("IPT1", "North Team"), ("IPT2", "Central Team")):
    db.execute("INSERT INTO ipts (tenant_id, id, label, manager, active, merged_into) VALUES (?, ?, ?, ?, ?, ?)",
               (db.current_tenant(), _tid, _lab, None, True, None))
PLANNER = {"role": "planner", "ipt": None, "label": "Planner", "name": "Pat Planner"}
NORTH = {"role": "ipt", "ipt": "IPT1", "label": "North Team"}
CENTRAL = {"role": "ipt", "ipt": "IPT2", "label": "Central Team"}
READER = {"role": "submitter", "ipt": None, "label": "Submitter"}


def as_(acc):
    access._ACCESS.set(acc)


def text_of(body):
    d = Document(io.BytesIO(body))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    head = " ".join(p.text for s in d.sections for p in s.header.paragraphs)
    foot = " ".join(p.text for s in d.sections for p in s.footer.paragraphs)
    return "\n".join(parts), head, foot


def build_fast(data, rec):
    return clp_doc.build(data, impact=rec["impact"], pictures=False)


def count_rows():
    return db.query("SELECT COUNT(*) AS n FROM clp_documents WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]


# =========================================================================== #
#  1. A team's document counts the team's lorries only (John's Q4)             #
# =========================================================================== #
ALL = clp_doc.gather([YARD], factors=factors, working_days=WD, layers=False)
N1 = clp_doc.gather([YARD], factors=factors, working_days=WD, layers=False, team="IPT1")
C2 = clp_doc.gather([YARD], factors=factors, working_days=WD, layers=False, team="ipt 2")
own1 = trips(R_IN, "substructure", "WS1", 10) + trips(R_IN, "substructure", "WS1", 11) + trips(R_IN, "substructure", "WS1", 13) \
    + trips(R_OUT, "earthworks", "WS2", 11)
ok("⭐ the North Team's CLP counts only its own lorries at the yard", N1["totals"]["trips"] == own1,
   f"{N1['totals']['trips']} vs {own1}")
ok("...the Central Team's only its own (R003 between the yard and the pile, R004 off the railhead)",
   C2["totals"]["trips"] == trips(R_BTW, "substructure", "WS3", 10) + trips(R_RAIL, "structures", "WS4", 13),
   str(C2["totals"]["trips"]))
ok("...and the two together are the every-team CLP", N1["totals"]["trips"] + C2["totals"]["trips"] == ALL["totals"]["trips"])
ok("...the team's routes only", {r["route_id"] for r in N1["routes"]} == {R_IN, R_OUT}
   and {r["route_id"] for r in C2["routes"]} == {R_BTW, R_RAIL}, str([r["route_id"] for r in N1["routes"]]))
ok("...a team's line waiting for approval is the team's note (the North Team's one), not the Central Team's",
   N1["totals"]["pending_lines"] == 1 and C2["totals"]["pending_lines"] == 0 and ALL["totals"]["pending_lines"] == 1)
ok("...the team's name, never its id ('ipt 2' is read as the Central Team)", N1["team_name"] == "North Team"
   and C2["team_name"] == "Central Team" and C2["team"] == "IPT2" and ALL["team"] is None and ALL["team_name"] is None)
_h_n1 = {h["gate"]: h for h in N1["hourly"]}
_ga = gate_arrivals.arrivals(N1["sites"][0]["peak_month"], NORTH, factors=factors, working_days=WD)
_mine = [g for g in _ga["gates"] if g["location_id"] == YARD and g["arrivals_per_day"] and g.get("per_team")]
ok("⭐ the team's hourly profile is the team's arrivals at the gate (gate_arrivals' per-team figure)",
   _h_n1 and _mine and all(abs(_h_n1[g["gate_name"]]["arrivals_per_day"] - g["per_team"][0]["arrivals_per_day"]) < 1e-9
                           for g in _mine), str(list(_h_n1)))
ok("...its % of capacity is every team's at that gate, and marked so", all(h["all_teams_pct"] for h in N1["hourly"])
   and not any(h["all_teams_pct"] for h in ALL["hourly"]))
_bn = clp_doc.build(N1, pictures=False)
_tn, _hn, _fn = text_of(_bn)
ok("⭐ the team's file says on its cover that it covers the team's lorries only",
   "This plan covers the North Team's lorries only." in _tn and "North Team" in _fn, _fn)
ok("...and words the gate's figures as the team's lorries / every team's together",
   ("of this team's lorries a day" in _tn) and ("every team's lorries together" in _tn
                                                 or not any(h["capacity_pct"] is not None for h in N1["hourly"])))
ok("🔴 the every-team file names no team", "This plan covers" not in text_of(clp_doc.build(ALL, pictures=False))[0])

# =========================================================================== #
#  2. The rows Wayscope lists by itself (Q2), and who sees which (Q3, Q4)      #
# =========================================================================== #
as_(PLANNER)
L = clp_records.list_for(PLANNER, factors=factors)
ids = {c["id"] for c in L["clps"]}
ok("⭐ a compound and a stockpile each get an every-team CLP and one per team with lorries there",
   {f"auto-{YARD}-all", f"auto-{YARD}-IPT1", f"auto-{YARD}-IPT2", f"auto-{PILE}-all", f"auto-{PILE}-IPT1",
    f"auto-{PILE}-IPT2"} <= ids, str(sorted(ids)))
ok("...not a quarry, a railhead or a landfill", not any(x in ids for x in (f"auto-{PIT}-all", f"auto-{RAIL}-all", f"auto-{TIP}-all")))
ok("🔴 opening the page writes nothing", count_rows() == 0)
_y1 = next(c for c in L["clps"] if c["id"] == f"auto-{YARD}-IPT1")
ok("...a row carries its team's name, its place, a status in words and the cover count",
   _y1["team_name"] == "North Team" and _y1["places"][0]["name"] == "Yard East" and _y1["status_word"] == "Draft"
   and _y1["cover_typed"] == 0 and _y1["cover_total"] == 11 and _y1["auto"] and not _y1["saved"], str(_y1))
ok("...its busiest month is the team's own (Nov 2026)", _y1["busiest_label"] == "Nov 2026", str(_y1["busiest_label"]))
LN = clp_records.list_for(NORTH, factors=factors)
ok("⭐ a team login sees its own team's CLPs only (no every-team CLP, no other team's)",
   LN["clps"] and all(c["team"] == "IPT1" for c in LN["clps"]) and not LN["can_approve"] and LN["team_name"] == "North Team",
   str([c["id"] for c in LN["clps"]]))
ok("🔴 another team's CLP, or the every-team one, is 'not found' to a team, like one that does not exist",
   clp_records.get(f"auto-{YARD}-IPT2", NORTH) is None and clp_records.get(f"auto-{YARD}-all", NORTH) is None
   and clp_records.get("auto-nothing-all", NORTH) is None and clp_records.get(f"auto-{YARD}-IPT1", NORTH) is not None)
ok("...a planner sees every CLP", clp_records.get(f"auto-{YARD}-IPT2", PLANNER) is not None and L["can_approve"])
ok("...a sign-in with no team that cannot approve reads every CLP",
   len(clp_records.list_for(READER, factors)["clps"]) == len(L["clps"]))
ok("🔴 no reply word carries a team id", not re.search(r"IPT ?[1-6]", json.dumps(
    [{k: v for k, v in c.items() if k not in ("id", "team")} for c in L["clps"]])))

# =========================================================================== #
#  3. Edit (Q5): a team its cover details; a planner the rest                  #
# =========================================================================== #
Y1 = f"auto-{YARD}-IPT1"
r = clp_records.update(Y1, {"cover": {"address": "1 Yard Lane", "postcode": "LE7 9XX", "manager_name": "A. Example",
                                      "manager_phone": "07700 900123"}}, NORTH, factors)
ok("⭐ a team edits the cover details of its own CLP, which writes the row then", r.get("cover_typed") == 4
   and r.get("saved") and count_rows() == 1, str(r))
ok("🔴 a team cannot rename it, or change its places, team or impact",
   clp_records.update(Y1, {"name": "Mine"}, NORTH, factors).get("status") == 403
   and clp_records.update(Y1, {"impact": "medium"}, NORTH, factors).get("status") == 403)
ok("🔴 a team cannot edit another team's CLP (not found)",
   clp_records.update(f"auto-{YARD}-IPT2", {"cover": {"address": "x"}}, NORTH, factors).get("status") == 404)
ok("🔴 a sign-in that only reads changes nothing",
   clp_records.update(Y1, {"cover": {"address": "x"}}, READER, factors).get("status") == 403)
ok("🔴 an unknown cover field, an overlong one and a banned word are refused in words",
   clp_records.update(Y1, {"cover": {"colour": "red"}}, NORTH, factors).get("status") == 422
   and clp_records.update(Y1, {"cover": {"address": "x" * 500}}, NORTH, factors).get("status") == 422
   and "secure" in clp_records.update(Y1, {"cover": {"summary": "A secure site"}}, NORTH, factors).get("error", ""))
r = clp_records.update(Y1, {"name": "Yard East (North)", "impact": "medium"}, PLANNER, factors)
ok("a planner renames it and sets the impact", r.get("name") == "Yard East (North)" and r.get("impact") == "medium", str(r))
ok("🔴 a row Wayscope lists keeps its places and team (+ New CLP for another set)",
   clp_records.update(Y1, {"location_ids": [YARD, PILE]}, PLANNER, factors).get("status") == 422
   and clp_records.update(Y1, {"team": "IPT2"}, PLANNER, factors).get("status") == 422)
ok("...and the typed cover survives the planner's edit", (clp_records.get(Y1, PLANNER)["cover"] or {}).get("postcode") == "LE7 9XX")

r = clp_records.create({"name": "Yard and pile", "location_ids": [YARD, PILE, YARD], "team": None}, PLANNER, factors)
NEW = r.get("id")
ok("⭐ a planner makes a CLP for any set of places, for every team", bool(NEW) and NEW.startswith("clp-")
   and r["location_ids"] == [YARD, PILE] and r["team"] is None and not r["auto"], str(r))
ok("🔴 a team cannot make one", clp_records.create({"location_ids": [YARD]}, NORTH, factors).get("status") == 403)
ok("🔴 no place, an unknown place, or a team not set up here is refused in words",
   clp_records.create({"location_ids": []}, PLANNER, factors).get("status") == 422
   and clp_records.create({"location_ids": ["L999"]}, PLANNER, factors).get("error") == "That place is not set up here."
   and clp_records.create({"location_ids": [YARD], "team": "IPT5"}, PLANNER, factors).get("error") == "That team is not set up here.")
r = clp_records.update(NEW, {"location_ids": [PILE], "team": "IPT2"}, PLANNER, factors)
ok("...and changes its places and team later", r.get("location_ids") == [PILE] and r.get("team") == "IPT2", str(r))
ok("...which the Central Team now sees and the North Team does not", clp_records.get(NEW, CENTRAL) is not None
   and clp_records.get(NEW, NORTH) is None)

# =========================================================================== #
#  4. The chain (Q6): send, ask for changes, approve as Issue N                #
# =========================================================================== #
ok("🔴 a planner does not send (they approve)", clp_records.send(Y1, "x", PLANNER, factors).get("status") == 403)
r = clp_records.send(Y1, "Gate dates updated after the council's comments", NORTH, factors)
ok("⭐ the team sends its CLP for approval with a note", r.get("status") == "sent" and r.get("sent_by") == "North Team"
   and r.get("sent_note", "").startswith("Gate dates"), str(r))
L = clp_records.list_for(PLANNER, factors=factors)
ok("...and the planner's list shows it first", L["clps"][0]["id"] == Y1, L["clps"][0]["id"])
ok("🔴 asking for changes needs a note, and a planner", clp_records.ask_changes(Y1, "", PLANNER, factors).get("status") == 422
   and clp_records.ask_changes(Y1, "x", NORTH, factors).get("status") == 403)
r = clp_records.ask_changes(Y1, "Add the Saturday hours", PLANNER, factors)
ok("⭐ changes asked: back to the team with the planner's note", r.get("status") == "changes"
   and r.get("changes_note") == "Add the Saturday hours" and r.get("changes_by") == "Pat Planner", str(r))
ok("...which the team reads on its own list",
   next(c for c in clp_records.list_for(NORTH, factors)["clps"] if c["id"] == Y1)["status_word"] == "Changes asked")
clp_records.send(Y1, "Saturday hours added", NORTH, factors)
P = clp_records.preview(Y1, PLANNER, factors, WD)
ok("⭐ the approve card: Issue 1, the team's note, 'First issue.'", P.get("next_issue") == 1 and P.get("sent")
   and P.get("sent_note") == "Saturday hours added" and P.get("changes") == ["First issue."], str(P))
ok("🔴 a team cannot approve, nor see the approve card", clp_records.preview(Y1, NORTH, factors, WD).get("status") == 403
   and clp_records.approve(Y1, "", NORTH, factors, WD, build_fast).get("status") == 403)
r = clp_records.approve(Y1, "Approved with the gate dates", PLANNER, factors, WD, build_fast)
ok("⭐ approved: Issue 1 kept, by the planner", r.get("status") == "approved" and r.get("issues") == 1
   and (r.get("last_issue") or {}).get("no") == 1 and r["last_issue"]["by"] == "Pat Planner", str(r))
_f1 = clp_records.issue_file(Y1, 1, PLANNER, factors)
ok("...the file is kept", bool(_f1) and _f1[0].endswith("Issue 1.docx") and _f1[1][:2] == b"PK", str(_f1 and _f1[0]))
_t1, _h1, _ft1 = text_of(_f1[1])
ok("⭐ the file says Issue 1 · approved on every page, not DRAFT", "Issue 1 · approved" in _h1 and "DRAFT" not in _h1
   and "Issue 1" in _ft1, _h1)
ok("...the typed cover details fill the cover (no 'To complete' for them)", "1 Yard Lane" in _t1 and "LE7 9XX" in _t1
   and "A. Example · 07700 900123" in _t1 and "Cover, site address" not in _t1 and "Cover, landowner" in _t1)
ok("...its Document history: Issue 1, First issue, the notes", "Document history" in _t1
   and re.search(r"1 \| \d+ \w{3} 20\d\d \| Pat Planner \| First issue\. Saturday hours added\. Approved with the gate dates\.", _t1)
   is not None, _t1[_t1.find("Document history"):][:300])
ok("...the title is the CLP's name", "Yard East (North)" in _t1)
H = clp_records.history(Y1, NORTH, factors)
ok("⭐ the team reads the history", bool(H) and H["issues"][0]["no"] == 1 and H["issues"][0]["approved_by"] == "Pat Planner")
ok("🔴 another team cannot read the history or the file", clp_records.history(Y1, CENTRAL, factors) is None
   and clp_records.issue_file(Y1, 1, CENTRAL, factors) is None)

# the plan changes: Issue 2
_line(R_IN, "substructure", "WS7", [(11, 1450.0)], ART, "IPT1", BAL)
_line(R_OUT, "earthworks", "WS8", [(16, 290.0)], ART, "IPT1", SPOIL)
config.invalidate()
factors = conversions.load_factors()
P2 = clp_records.preview(Y1, PLANNER, factors, WD)
ok("⭐ Issue 2's 'What changed' is worked out from the plan: more loads, a stage that ends later",
   P2["next_issue"] == 2 and not P2["sent"] and any("loads" in x and "→" in x for x in P2["changes"])
   and any("now ends" in x for x in P2["changes"]), str(P2["changes"]))
r = clp_records.approve(Y1, "", PLANNER, factors, WD, build_fast)
ok("⭐ a planner approves a CLP nobody sent (Q6)", (r.get("last_issue") or {}).get("no") == 2 and r.get("status") == "approved")
H = clp_records.history(Y1, PLANNER, factors)
ok("...the history is newest first", [i["no"] for i in H["issues"]] == [2, 1] and H["issues"][0]["sent_by"] is None)
ok("⭐ Issue 1's file is still exactly as approved", clp_records.issue_file(Y1, 1, PLANNER, factors)[1] == _f1[1])
_t2 = text_of(clp_records.issue_file(Y1, 2, PLANNER, factors)[1])[0]
ok("...Issue 2's file lists both issues, newest first", re.search(r"\n2 \| [^\n]*\n1 \| ", _t2) is not None)
r = clp_records.update(Y1, {"cover": {"landowner": "The county council"}}, NORTH, factors)
ok("⭐ editing an approved CLP makes it a draft again, its issues kept", r.get("status") == "draft" and r.get("issues") == 2)
_rec, _dd = clp_records.draft_data(Y1, NORTH, factors, WD)
_td, _hd, _fd = text_of(build_fast(_dd, _rec))
ok("⭐ a draft says DRAFT · not issued on every page, with the history so far", "DRAFT · not issued" in _hd
   and "Draft" in _fd and "Document history" in _td and "The county council" in _td, _hd)

# =========================================================================== #
#  5. Delete (Q7)                                                              #
# =========================================================================== #
ok("🔴 a CLP with an issue is kept", clp_records.delete(Y1, PLANNER, factors).get("status") == 422)
ok("🔴 a team cannot delete", clp_records.delete(NEW, CENTRAL, factors).get("status") == 403)
Y2 = f"auto-{YARD}-IPT2"
clp_records.update(Y2, {"cover": {"address": "2 Yard Lane"}}, CENTRAL, factors)
_n = count_rows()
r = clp_records.delete(Y2, PLANNER, factors)
ok("⭐ deleting a row Wayscope lists only clears what was typed (it is listed again, blank)", r.get("cleared")
   and count_rows() == _n - 1 and clp_records.get(Y2, PLANNER)["cover_typed"] == 0)
r = clp_records.delete(NEW, PLANNER, factors)
ok("...a CLP made with + New CLP is gone", r.get("deleted") == NEW and clp_records.get(NEW, PLANNER) is None)

# =========================================================================== #
#  6. The endpoints, the package, the source                                   #
# =========================================================================== #
as_(NORTH)
EL = main.clp_documents_list()
ok("GET /api/clp/documents for a team: its own", EL["clps"] and all(c["team"] == "IPT1" for c in EL["clps"]))
_e = _raises(lambda: main.clp_documents_get(f"auto-{YARD}-IPT2"))
ok("🔴 GET another team's CLP: 404", getattr(_e, "status_code", None) == 404, str(_e))
_e = _raises(lambda: main.clp_documents_create(main.ClpIn(location_ids=[YARD])))
ok("🔴 POST as a team: 403", getattr(_e, "status_code", None) == 403, str(_e))
_e = _raises(lambda: main.clp_documents_approve(Y1, main.ClpNoteIn(note="")))
ok("🔴 approve as a team: 403", getattr(_e, "status_code", None) == 403, str(_e))
_e = _raises(lambda: main.clp_documents_update(Y1, main.ClpIn(name="x")))
ok("🔴 PATCH a name as a team: 403", getattr(_e, "status_code", None) == 403, str(_e))
_ok_patch = main.clp_documents_update(Y1, main.ClpIn(cover={"site_use": "Farmland"}))
ok("PATCH sends only the fields named (the name and the other cover details are kept)",
   _ok_patch["name"] == "Yard East (North)" and _ok_patch["cover"].get("landowner") == "The county council"
   and _ok_patch["cover"].get("site_use") == "Farmland", str(_ok_patch))
_iss = main.clp_documents_issue(Y1, 1)
ok("⭐ GET an issue's file: the bytes as approved", _iss.body == _f1[1] and _iss.media_type == main.DOCX_TYPE
   and "Issue 1.docx" in _iss.headers["Content-Disposition"])
_dr = main.clp_documents_draft(Y1)
ok("⭐ GET the draft: a Word file marked DRAFT, built with the plans (offline)", _dr.body[:2] == b"PK"
   and "DRAFT · not issued" in text_of(_dr.body)[1] and "DRAFT.docx" in _dr.headers["Content-Disposition"])
as_(PLANNER)
_ap = main.clp_documents_approve(f"auto-{PILE}-all", main.ClpNoteIn(note="First look"))
ok("⭐ POST approve as a planner builds and keeps Issue 1 (through the one-at-a-time build)",
   _ap["last_issue"]["no"] == 1 and _ap["status"] == "approved")
MAIN = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
ok("🔴 approve and the draft build through _clp_build (alone, outside the public slots)",
   'lambda data, rec: _clp_build(data, rec["impact"], acc)' in MAIN
   # NARROWED 31 Oct 2026 (PR 3b): the draft builds in a small _build() that keeps its Still to complete lines
   # NARROWED 31 Oct 2026 (PR 3c): the kept draft's key carries the client groups' decisions (printed in its history)
   and 'out = _clp_build(data, rec["impact"], acc)' in MAIN
   and '"clp/draft.docx", {"id": rec["id"], "reviews": clp_reviews.marker(rec["id"])}, _build, acc=acc' in MAIN
   and 'with heavy_reads.solo("clp-document"' in MAIN)
_e = _raises(lambda: main.clp_documents_changes(Y1, main.ClpNoteIn(note="")))
ok("🔴 changes with no note: 422 in words", getattr(_e, "status_code", None) == 422 and "Say what" in str(_e), str(_e))

ok("the two tables are tenanted and in the organisation file, after the locations",
   {"clp_documents", "clp_issues"} <= db.TENANTED_TABLES
   and tenant_package.ORDER.index("clp_documents") > tenant_package.ORDER.index("locations")
   and tenant_package.ORDER.index("clp_issues") > tenant_package.ORDER.index("clp_documents"))
_pkg = tenant_package.export_tenant()
_iss_rows = (_pkg.get("tables") or {}).get("clp_issues") or []
ok("⭐ an issue's Word file travels in the organisation file", any(
   r.get("clp_id") == Y1 and r.get("issue_no") == 1 and base64.b64decode(r.get("file_b64") or "") == _f1[1] for r in _iss_rows),
   str(len(_iss_rows)))
_GEN = open(os.path.join(BACKEND, "tools", "make_demo_tenant.py"), encoding="utf-8").read()
_LT = _GEN.split("LIVE_TABLES = (")[1].split(")")[0]
ok("...and a demo refresh keeps live's CLPs", '"clp_documents"' in _LT and '"clp_issues"' in _LT)
SRC = open(os.path.join(BACKEND, "clp_records.py"), encoding="utf-8").read()
_calls = [SRC[m.end():m.end() + 420] for m in re.finditer(r"db\.(?:query|execute)\(", SRC)]
ok("🔴 every statement in clp_records names tenant_id (test_tenant_audit checks the binding)",
   len(_calls) >= 8 and all("tenant_id" in c.split("))")[0] for c in _calls), str(len(_calls)))
ok("🔴 no BOOLEAN compared with a number", not re.search(r"\bactive\s*=\s*1\b", SRC))
ok("🔴 no team slot id in the module's words", re.search(r"\bIPT ?[1-6]\b", SRC) is None)

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\ntest_clp_records: {PASS} passed, {len(FAIL)} failed")
remove_scratch_db(db, TMP)
sys.exit(1 if FAIL else 0)
