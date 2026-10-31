"""
The CLP document's pictures and layers, 30 Oct 2026 (task 4a, PR 2) — backend/clp_doc_layers.py, clp_doc_pictures.py and
what clp_doc.build() does with them.

What is asserted: §2.3 / §2.4 from what Wayscope stores, measured on THIS plan's routes and sites only (a school by
another site's route is not this plan's): the councils (a route's own ends counted even before a road-sections rebuild),
the stored lorry-limit and closure checks, the dated diversions, the community places within 250 m of a route / 500 m of
a site with the nearest one and how far, the collisions on and near each route; "not fetched" is never printed as a
zero. The charts (programme, monthly with a day scale at round numbers, hourly per gate) and the plans (a regional
plan, then each site's local and site plan in the Streets style) as pictures with their alternative text. The .docx
read back: the pictures inline with their descriptions, the tables, the word rule. The endpoint: a heavy slot, kept
until the next write only when every plan had its basemap.

WHAT THIS DOES NOT PROVE
------------------------
  * The councils, the lorry-limit and closure checks and the route choices are STUBBED here (areas.area_load,
    limits / restrictions.stored_checks, route_choice.periods): their own harnesses own them; this one asserts what
    the CLP does with their answers.
  * Mapbox is never called (portable.offline): the plans are the schematic branch. The real basemap was looked at
    on the demo copy on this laptop and on live after the merge.
  * How Word lays the pictures out (looked at in Word, 30 Oct 2026).
  * SQLite only; hand-made straight-line routes.

Run (PowerShell):  $env:PYTHONUTF8 = "1"; .venv-tests\\Scripts\\python backend/tests/test_clp_doc_maps.py
"""
import copy
import json
import io
import math
import os
import re
import sys
import tempfile
import types
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, as test_gate_arrivals.py ----------------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query


class _Request:
    def __init__(self, headers=None):
        self.headers = {k.lower(): v for k, v in (headers or {}).items()}


_fa.Request = _Request
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None


class _Response:
    def __init__(self, content=None, status_code=200, headers=None, media_type=None):
        self.body = content
        self.status_code = status_code
        self.headers = dict(headers or {})
        self.media_type = media_type


_resp.Response = _Response
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)


TMP = tempfile.mkdtemp(prefix="rbe_clpdocmaps_")
from portable import remove_scratch_db, offline   # noqa: E402,F401  tests/portable.py
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import costing  # noqa: E402
import config  # noqa: E402
import gates  # noqa: E402
import gate_arrivals  # noqa: E402
import gate_hours  # noqa: E402
import network  # noqa: E402
import site_areas  # noqa: E402
import taxonomy  # noqa: E402
import carbon_target  # noqa: E402
import clp_doc  # noqa: E402
import export  # noqa: E402
offline(export)            # PR 2 draws plans: Mapbox is never called from a test (the schematic branch)
from docx import Document  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
os.environ["ALLOW_DEMO_CODES"] = "1"
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


ART = network.DEFAULT_PROFILE          # "Artic Tipper (44t)", 29 t a load
FLAT = "Artic Flatbed (44t)"
BAL = "Large aggregate / ballast"
SPOIL = "Earthworks / soil"
PRE = "Precast units"


def all_text(doc):
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    for s in doc.sections:
        out += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return out


import contextlib  # noqa: E402
import areas  # noqa: E402
import basemap  # noqa: E402
import collisions  # noqa: E402
import community  # noqa: E402
import limits  # noqa: E402
import read_cache  # noqa: E402
import restrictions  # noqa: E402
import route_choice  # noqa: E402
import clp_doc_layers  # noqa: E402
import clp_doc_pictures as PIC  # noqa: E402
from reportlab import rl_config  # noqa: E402
from PIL import Image  # noqa: E402

# =========================================================================== #
#  0. A small scheme by hand: a yard, a pit, a tip; R005 is not this plan's    #
# =========================================================================== #
remove_scratch_db(db, TMP)
db.init_db(); db.init_network_db(); db.init_taxonomy_db(); db.init_zones_db(); db.init_gates_db()
db.init_weeks_db(); db.init_lookahead_db(); db.init_config_db(); db.init_costing_db()
db.init_tenant()
taxonomy.seed_taxonomy()
config.invalidate()
config.seed_from_file(conversions)
config.invalidate()
costing.invalidate()
carbon_target.invalidate()
access.set_current("planner123")
_doc = copy.deepcopy(config.load(conversions, use_cache=False))
_doc.setdefault("tenant", {})["demo_notice"] = "Fictional scheme"
config.save(_doc, by="tester")
config.invalidate()
factors = conversions.load_factors()
WD = 22.0
TEN = db.current_tenant()

Y_AT, P_AT, T_AT = (-0.8822, 52.6447), (-0.95, 52.55), (-0.90, 52.50)
YARD = network.create_location("Yard East", "destination", lat=Y_AT[1], lon=Y_AT[0], loc_type="Compound",
                               receives=[BAL], supplies=[SPOIL])["id"]
PIT = network.create_location("Pit West", "origin", lat=P_AT[1], lon=P_AT[0], loc_type="Quarry", supplies=[BAL])["id"]
TIP = network.create_location("Tip South", "destination", lat=T_AT[1], lon=T_AT[0], loc_type="Landfill",
                              receives=[SPOIL, BAL])["id"]
R1 = network.create_route(PIT, YARD, material_category=BAL, route_id="R001", ipt="IPT1")["id"]
R2 = network.create_route(YARD, TIP, material_category=SPOIL, route_id="R002", ipt="IPT1")["id"]
R5 = network.create_route(PIT, TIP, material_category=BAL, route_id="R005", ipt="IPT1")["id"]


def line(a, b, n=12):
    return [[a[0] + (b[0] - a[0]) * i / n, a[1] + (b[1] - a[1]) * i / n] for i in range(n + 1)]


def off(a, b, t, m):
    """The point a fraction t along a -> b, m metres to its right (east of a northbound line)."""
    lon, lat = a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t
    kx, ky = 111320.0 * math.cos(math.radians(lat)), 110574.0
    dx, dy = (b[0] - a[0]) * kx, (b[1] - a[1]) * ky
    L = math.hypot(dx, dy)
    return [lon + dy / L * m / kx, lat - dx / L * m / ky]


def add_geom(rid, a, b):
    for leg, pts in (("loaded", line(a, b)), ("return", line(b, a))):
        db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, geometry, distance_km, "
                   "duration_hr, computed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                   (TEN, rid, ART, leg, 0, json.dumps(pts), 12.0, 0.4, "2026-10-30T00:00:00Z"))


add_geom(R1, P_AT, Y_AT)
add_geom(R2, Y_AT, T_AT)
add_geom(R5, P_AT, T_AT)


def _line(rid, disc, sect, month, mat):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline=disc, section_id=sect, material_type=mat, material_description=None,
        vehicle_type=ART, submitted_by="tester", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=q) for m, q in month], ipt="IPT1"))
    main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline=disc, section_id=sect)


_line(R1, "substructure", "WS1", [(10, 2900.0), (11, 5800.0)], BAL)
_line(R2, "earthworks", "WS2", [(11, 2900.0), (13, 1450.0)], SPOIL)
_line(R5, "earthworks", "WS5", [(10, 2900.0)], BAL)
G_IN = gates.create_gate(YARD, "Yard gate", Y_AT[1], Y_AT[0], direction="both", is_default=True,
                         open_from="07:30", open_to="18:00", open_days="1111100")
ok("the gate is made", "id" in G_IN, str(G_IN))
config.invalidate()
factors = conversions.load_factors()

# =========================================================================== #
#  1. Nothing fetched yet: no zero printed                                     #
# =========================================================================== #
D0 = clp_doc.gather([YARD], factors=factors, working_days=WD)
L0 = D0["layers"]
ok("⭐ the layers ride on gather(): this plan's routes only (R005 has no end at the yard)",
   L0["route_ids"] == ["R001", "R002"], str(L0["route_ids"]))
ok("...gather(layers=False) leaves them for the endpoint to fill in its heavy slot",
   clp_doc.gather([YARD], factors=factors, working_days=WD, layers=False)["layers"] is None)
ok("🔴 nothing fetched: community 'not stored' and collisions 'not fetched', never a zero count",
   L0["community"]["stored"] is False and L0["collisions"]["fetched"] is False and L0["collisions"]["rows"] == []
   and L0["limits"]["checked"] is False and L0["closures"]["checked"] is False, str(L0["community"]["stored"]))
W0 = Document(io.BytesIO(clp_doc.build(D0, pictures=False)))
T0 = "\n".join(p.text for p in W0.paragraphs)
ok("...and the document says nothing it does not hold: no bus stop count, no collision table, no limit line",
   "bus stop" not in T0 and "Injury collisions" not in T0 and "lorry limits" not in T0
   and "No National Highways closures" not in T0, "")
ok("...the 2.2 plans placeholder stays when there are no pictures",
   any(p.text.startswith("To complete: a regional plan (1:15,000 or smaller)") for p in W0.paragraphs))

# =========================================================================== #
#  2. The layers                                                               #
# =========================================================================== #
_real = {"area_load": areas.area_load, "lim": limits.stored_checks, "res": restrictions.stored_checks,
         "per": route_choice.periods}
areas.area_load = lambda mi, working_days=None, geometry=True: {"areas": [
    {"name": "Harborough", "tier_words": "local authority district", "highway_authority": "Leicestershire County Council",
     "routes": ["R001", "R005"], "sites": ["Yard East", "Pit West"]},
    {"name": "Rutland", "tier_words": "county or unitary authority", "highway_authority": "Rutland County Council",
     "routes": [], "sites": ["Tip South"]},
    {"name": "Elsewhere", "tier_words": "", "highway_authority": "", "routes": ["R005"], "sites": []}]}
limits.stored_checks = lambda ids: {
    "R001": {"checked_at": "2026-10-29T10:00:00Z", "hits": [
        {"headline": "Weak bridge 7.5 t", "verdict": "exceeds", "vehicle_label": ART},
        {"headline": "Height 5.0 m", "verdict": "within", "vehicle_label": ART}]},
    "R002": {"checked_at": "2026-10-29T10:00:00Z", "hits": [
        {"headline": "Weak bridge 7.5 t", "verdict": "exceeds", "vehicle_label": ART},
        {"headline": "Width 2.0 m", "verdict": "unknown", "vehicle_label": ""}]}}
restrictions.stored_checks = lambda ids: {
    "R001": {"checked_at": "2026-10-29T06:00:00Z", "hits": [{"description": "A47 closed overnight", "severity": "high"}]},
    "R002": {"checked_at": "2026-10-28T06:00:00Z", "hits": []}}
route_choice.periods = lambda *a, **k: [
    {"route_id": "R002", "reason_kind": "zone", "status": "in_use", "zone_names": ["Bridge works"],
     "starts_on": "2027-03-01", "ends_on": "2027-05-31"},
    {"route_id": "R005", "reason_kind": "zone", "status": "in_use", "zone_names": ["Not ours"],
     "starts_on": "2027-03-01", "ends_on": "2027-05-31"},
    {"route_id": "R001", "reason_kind": "peak", "status": "offered", "zone_names": [], "starts_on": None, "ends_on": None},
    {"route_id": "R001", "reason_kind": "zone", "status": "retired", "zone_names": ["Old works"],
     "starts_on": "2026-01-01", "ends_on": "2026-02-01"}]


def feat(fid, kind, source, name, geom):
    db.execute("INSERT INTO community_features (tenant_id, feature_id, kind, source, source_ref, name, detail, geometry_json, "
               "dist_m, fetched_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
               (TEN, fid, kind, source, fid, name, None, json.dumps(geom), 0.0, "2026-10-29"))


pt = lambda c: {"type": "Point", "coordinates": c}                   # noqa: E731
feat("g1", "schools", "gias", "Near School", pt(off(P_AT, Y_AT, 0.5, 100)))         # 100 m off R001
feat("g2", "schools", "gias", "Far School", pt(off(P_AT, Y_AT, 0.5, 400)))          # 400 m: outside the strip
feat("g3", "schools", "gias", "Other Route School", pt(off(P_AT, T_AT, 0.5, 50)))   # by R005 only: not this plan's
feat("h1", "hospitals", "osm", "Wolds Hospital", pt(off(Y_AT, T_AT, 0.4, 60)))       # 60 m off R002
feat("s1", "stations", "naptan", "Yard Halt", pt(off(P_AT, Y_AT, 0.3, 200)))        # 200 m off R001
feat("b1", "bus_stops", "naptan", "Yard Lane", pt([Y_AT[0], Y_AT[1] + 300 / 110574.0]))   # 300 m north of the yard
feat("b2", "bus_stops", "naptan", "Pit Road", pt(off(P_AT, Y_AT, 0.2, 30)))
feat("c1", "cycle", "osm", "NCN 64", {"type": "LineString", "coordinates": [off(Y_AT, T_AT, 0.6, -300),
                                                                            off(Y_AT, T_AT, 0.6, 300)]})
feat("c2", "cycle", "osm", "", {"type": "LineString", "coordinates": [off(P_AT, Y_AT, 0.7, -100),
                                                                      off(P_AT, Y_AT, 0.7, 100)]})
_ring = [off(Y_AT, T_AT, 0.2, 80), off(Y_AT, T_AT, 0.25, 80), off(Y_AT, T_AT, 0.25, 180), off(Y_AT, T_AT, 0.2, 80)]
feat("r1", "residential", "osm", "Southfield estate", {"type": "Polygon", "coordinates": [_ring]})


def coll(cid, sev, routes, sites, vru=(), goods=0):
    db.execute("INSERT INTO collision_records (tenant_id, " + ", ".join(collisions._COLS) + ") VALUES ("
               + ", ".join("?" for _ in range(len(collisions._COLS) + 1)) + ")",
               (TEN, cid, 2023, "2023-05", Y_AT[0], Y_AT[1], sev, "A6003", "", "", "", 60, "", "", 1, 2,
                json.dumps(list(vru)), "[]", goods, 10.0, json.dumps(routes), json.dumps(sites), "2026-10-29", "{}"))


coll("k1", 2, {"R001": 10.0}, {}, vru=("cyclist",), goods=1)          # on R001, serious, a cyclist, a lorry
coll("k2", 3, {"R001": 200.0}, {})                                    # near R001 only
coll("k3", 3, {"R002": 5.0}, {YARD: 300.0})                            # on R002 and round the yard
coll("k4", 1, {"R005": 5.0}, {})                                      # R005: not this plan's
collisions._write_meta("done", {"years": [2020, 2021, 2022, 2023, 2024]}, fetched_at="2026-10-29T09:00:00Z")

D1 = clp_doc.gather([YARD], factors=factors, working_days=WD)
L1 = D1["layers"]
C1 = L1["community"]
ok("⭐ councils: a council with this plan's route or site; a route's own end counted even with no clip ('Rutland' holds "
   "Tip South, the end of R002); another plan's council left out",
   [(c["name"], c["routes"], c["sites"]) for c in L1["councils"]]
   == [("Harborough", ["R001", "R002"], ["Yard East"]), ("Rutland", ["R002"], [])]
   and L1["councils_month"] == D1["totals"]["peak_month"], str(L1["councils"]))
ok("⭐ lorry limits: only over / at / to-check rows, merged across routes, the rest counted as within",
   L1["limits"]["checked"] is True and L1["limits"]["within"] == 1
   and [(e["what"], e["verdict"], e["routes"]) for e in L1["limits"]["rows"]]
   == [("Weak bridge 7.5 t", "exceeds", ["R001", "R002"]), ("Width 2.0 m", "unknown", ["R002"])]
   and L1["limits"]["rows"][0]["lorries"] == [ART], str(L1["limits"]))
ok("closures: the newest check's time, and each closure with its routes",
   L1["closures"]["checked_at"] == "2026-10-29T06:00:00Z"
   and L1["closures"]["rows"] == [{"what": "A47 closed overnight", "severity": "high", "routes": ["R001"]}],
   str(L1["closures"]))
ok("diversions: this plan's dated zone diversions only (no peak offer, no retired one, no other plan's route)",
   L1["diversions"] == [{"route_id": "R002", "zones": ["Bridge works"], "from": "2027-03-01", "to": "2027-05-31",
                         "status": "in_use", "kept": False}], str(L1["diversions"]))
places = {x["name"]: x for x in C1["places"]}
ok("⭐ the strip on THIS plan's routes: the school 100 m off R001 in, the one 400 m off out, the one by R005 out",
   set(places) == {"Near School", "Wolds Hospital"}, str(sorted(places)))
ok("...each with the route it is nearest and how far (metres, ±3)",
   places["Near School"]["where"] == "R001" and abs(places["Near School"]["metres"] - 100) <= 3
   and places["Wolds Hospital"]["where"] == "R002" and abs(places["Wolds Hospital"]["metres"] - 60) <= 3,
   str(places))
ok("...nearest first", [x["name"] for x in C1["places"]] == ["Wolds Hospital", "Near School"])
ok("stations: Yard Halt, 200 m from R001", [(x["name"], x["where"], abs(x["metres"] - 200) <= 3) for x in C1["stations"]]
   == [("Yard Halt", "R001", True)], str(C1["stations"]))
ok("bus stops: two, one of them 300 m from the yard and nearest it (outside 250 m of every route)",
   C1["bus_stops"] == 2 and C1["bus_stops_near_sites"] == 1, f"{C1['bus_stops']} {C1['bus_stops_near_sites']}")
ok("cycle routes: the named one by name, the unnamed one counted", C1["cycle"] == ["NCN 64"] and C1["cycle_unnamed"] == 1,
   f"{C1['cycle']} {C1['cycle_unnamed']}")
ok("residential areas counted", C1["residential"] == 1)
ok("the credit names the sources the plan's features came from (OpenStreetMap, and the government ones under the OGL)",
   "OpenStreetMap" in C1["credit"] and "Open Government Licence" in C1["credit"], C1["credit"])
_mp = clp_doc_layers.MAX_PLACES
clp_doc_layers.MAX_PLACES = 1
C1b = clp_doc_layers.gather(D1)["community"]
clp_doc_layers.MAX_PLACES = _mp
ok("past MAX_PLACES the nearest are listed and the rest counted", [x["name"] for x in C1b["places"]] == ["Wolds Hospital"]
   and C1b["places_more"] == 1)
CO = L1["collisions"]
rows = {r["route_id"]: r for r in CO["rows"]}
ok("⭐ collisions per route: on the route (within 25 m) and within 250 m, the killed or seriously injured, a person on "
   "foot / cycle / motorcycle hurt, a goods vehicle over 3.5 t; R005's not counted",
   (rows["R001"]["on_all"], rows["R001"]["on_ksi"], rows["R001"]["near_all"], rows["R001"]["near_vru"],
    rows["R001"]["near_goods"]) == (1, 1, 2, 1, 1)
   and (rows["R002"]["on_all"], rows["R002"]["near_all"], rows["R002"]["near_ksi"]) == (1, 1, 0)
   and set(rows) == {"R001", "R002"}, str(CO["rows"]))
ok("...round the yard: one, none killed or seriously injured; the period and the credit",
   CO["sites"] == [{"name": "Yard East", "all": 1, "ksi": 0}] and CO["period"] == "2020–2024"
   and CO["credit"] == collisions.credit_line(), str(CO))

# =========================================================================== #
#  3. The pictures                                                             #
# =========================================================================== #


def png_size(b):
    return Image.open(io.BytesIO(b)).size


ch_p, ch_m = PIC.programme_chart(D1), PIC.monthly_chart(D1)
ch_h = PIC.hourly_chart(D1["hourly"][0])
ok("⭐ the three charts are PNGs 170 mm wide at 200 dpi, each with its alternative text",
   all(c and c["png"][:8] == b"\x89PNG\r\n\x1a\n" and c["alt"] for c in (ch_p, ch_m, ch_h))
   and all(abs(png_size(c["png"])[0] - round(170 / 25.4 * 200)) <= 2 for c in (ch_p, ch_m, ch_h)),
   str([png_size(c["png"]) for c in (ch_p, ch_m, ch_h) if c]))
ok("...the monthly chart's words carry the busiest month and its trips",
   clp_doc.month_label(D1["totals"]["peak_month"]) in ch_m["alt"] and f"{D1['totals']['peak_trips']:,}" in ch_m["alt"],
   ch_m["alt"])
ok("...the hourly chart's words carry the gate, the month and the lorries a day",
   D1["hourly"][0]["gate"] in ch_h["alt"] and f"{D1['hourly'][0]['arrivals_per_day']:g} lorries a day" in ch_h["alt"],
   ch_h["alt"])
ok("the day scale at round numbers: 400 a month over 22 days = 0, 5, 10, 15; nothing to draw = 0",
   PIC.day_ticks(400, 22) == [0, 5, 10, 15] and PIC.day_ticks(0, 22) == [0] and PIC.day_ticks(100, 0) == [0],
   str(PIC.day_ticks(400, 22)))
DT = clp_doc.gather([TIP], factors=factors, working_days=WD, layers=False)
ok("no stage with loads: no programme or monthly chart (None, never an empty picture)",
   PIC.programme_chart(dict(D1, stages=[dict(s, first=None) for s in D1["stages"]])) is None
   and PIC.monthly_chart(dict(D1, totals=dict(D1["totals"], first=None))) is None)
_styles = []
_real_use = basemap.use


@contextlib.contextmanager
def _rec_use(key):
    _styles.append(key)
    with _real_use(key):
        yield


basemap.use = _rec_use
_a85 = rl_config.useA85
PL = PIC.plans(D1)
basemap.use = _real_use
ok("⭐ the plans: one regional, then the yard's local and site plan (1 + 2 a site with a position)",
   [(p["kind"], p["site"]) for p in PL] == [("regional", None), ("local", "Yard East"), ("site", "Yard East")],
   str([(p["kind"], p["site"]) for p in PL]))
ok("...JPEGs with alternative text and the plan's own scale words",
   all(p["jpg"][:3] == b"\xff\xd8\xff" and p["alt"] and p["scale"] in p["alt"] for p in PL)
   and PL[1]["scale"].startswith("1:2,") and PL[2]["scale"].startswith("1:"), str([p["scale"] for p in PL]))
ok("...the local and site plans in the Streets style, the regional plan in the export's own",
   _styles == ["streets", "streets"] and PIC.PLAN_STYLE == {"regional": None, "local": "streets", "site": "streets"},
   str(_styles))
ok("...the PDF writer's ASCII85 setting put back afterwards", rl_config.useA85 == _a85)
ok("...each site's plans at that site's own busiest month", PL[1]["month"] == PIC._site_month(D1, YARD))
ok("no drawn route: no plans", PIC.plans(dict(D1, routes=[dict(r, routed=False) for r in D1["routes"]])) == [])

# =========================================================================== #
#  4. The Word file, read back                                                 #
# =========================================================================== #
B1 = clp_doc.build(D1)
W1 = Document(io.BytesIO(B1))
pics = W1.inline_shapes
n_hour = len(D1["hourly"])
ok("⭐ the pictures are in the file: 3 plans, the programme, the monthly chart and one hourly chart a gate",
   len(pics) == 3 + 2 + n_hour, str(len(pics)))
descr = [s._inline.docPr.get("descr") for s in pics]
ok("...every picture carries its description (a screen reader reads it)", all(d for d in descr), str(descr))
ok("...the plans first, in 2.2, then the programme chart, the monthly chart, the hourly ones",
   descr[0].startswith("Regional plan") and descr[3].startswith("Programme:") and descr[4] == ch_m["alt"],
   str([d[:20] for d in descr]))
P1 = [p.text for p in W1.paragraphs]
TX = "\n".join(all_text(W1))
ok("2.2's plan placeholder gone once the plans are drawn; the consolidation one stays",
   not any(t.startswith("To complete: a regional plan") for t in P1)
   and any(t.startswith("To complete: consolidation centres") for t in P1))
ok("...each plan captioned with its scale and month, and the line on the scale bar",
   any(t.startswith("Local plan of Yard East · 1:2,") for t in P1)
   and any("use it, not the stated ratio" in t for t in P1), "")
ok("2.3.1: the councils table, the limits, the closure and the diversion with its dates in words",
   "Leicestershire County Council" in TX and "Rutland County Council" in TX and "Weak bridge 7.5 t" in TX
   and "over the limit" in TX and "to be checked" in TX and "1 other lorry limit on the routes is within every lorry's size and weight." in TX
   and "A47 closed overnight (R001)" in TX and "Bridge works" in TX and "1 Mar 2027" in TX and "31 May 2027" in TX
   and "last checked 2026-10-29" in TX, "")
ok("2.3.2–2.3.4: the station, the bus stops, the cycle routes, from the strip",
   "Yard Halt (200 m from R001)" in TX.replace("199 m", "200 m").replace("201 m", "200 m")
   and "2 bus stops within 250 m of this plan's routes or 500 m of this site, 1 of them nearest this site." in TX
   and "NCN 64, 1 unnamed cycle path" in TX, "")
ok("2.4.1: one line a council, by name", sum(1 for t in P1 if t.startswith("To complete: Harborough: the code")) == 1
   and sum(1 for t in P1 if t.startswith("To complete: Rutland: the code")) == 1)
ok("2.4.2: the places table and the collisions table, with their credits",
   "Near School" in TX and "Wolds Hospital" in TX and "Far School" not in TX and "Other Route School" not in TX
   and "Injury collisions reported to the police, 2020–2024:" in TX and "Counts, not a rate" in TX
   and collisions.credit_line() in TX and C1["credit"] in TX
   and "Within 500 m of Yard East: 1 collision, 0 killed or seriously injured." in TX, "")
ok("🔴 the word rule over the whole file (no Mapbox, no back-end words)", not clp_doc._BAN_RE.search(TX),
   str(clp_doc._BAN_RE.findall(TX)))
ok("§4 points to 2.2 for the local and site plans", "The local and site plans of each site are in 2.2." in P1)
ok("pictures=False: the same text, no pictures (the fast path the tests and a future preview use)",
   len(Document(io.BytesIO(clp_doc.build(D1, pictures=False))).inline_shapes) == 0)

# =========================================================================== #
#  5. The endpoint: a heavy slot; kept only when every plan had its basemap    #
# =========================================================================== #
_builds = [0]
_real_build = clp_doc.build


def _count_build(*a, **k):
    _builds[0] += 1
    return _real_build(*a, **k)


clp_doc.build = _count_build
_tok = config.mapbox_token
config.mapbox_token = lambda: ""            # no token: no basemap asked for, nothing missed
r1 = main.clp_document_docx(locations=YARD)
r2 = main.clp_document_docx(locations=YARD)
ok("⭐ no basemap missed: built once, then kept until the next write",
   _builds[0] == 1 and r1.body == r2.body and r1.body[:2] == b"PK", str(_builds[0]))
config.mapbox_token = lambda: "pk.harness"   # a token, but offline: every plan misses its basemap
_builds[0] = 0
main.clp_document_docx(locations=YARD, impact="medium")
main.clp_document_docx(locations=YARD, impact="medium")
ok("🔴 a plan drawn without its basemap: not kept, so the next download asks again", _builds[0] == 2, str(_builds[0]))
config.mapbox_token = _tok
clp_doc.build = _real_build
MAIN = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
_ep = MAIN[MAIN.index('@app.get("/api/clp/document.docx")'):MAIN.index('@app.post("/api/admin/gates")')]
ok("the build runs in a heavy slot of its own, busy in words", 'heavy_reads.slot("clp-document"' in _ep
   and '"Wayscope is busy; try again in a minute."' in _ep and "keep=lambda: basemap.MISSES[0] == misses" in _ep)

areas.area_load, limits.stored_checks = _real["area_load"], _real["lim"]
restrictions.stored_checks, route_choice.periods = _real["res"], _real["per"]

# =========================================================================== #
#  6. The source                                                               #
# =========================================================================== #
for fn in ("clp_doc_layers.py", "clp_doc_pictures.py"):
    S = open(os.path.join(BACKEND, fn), encoding="utf-8").read()
    ok(f"🔴 {fn} reads only: no INSERT, UPDATE or DELETE", not re.search(r"\b(INSERT|UPDATE|DELETE)\b", S))
    ok(f"🔴 {fn}: no BOOLEAN compared with a number, no team slot id",
       not re.search(r"\bactive\s*=\s*1\b", S) and re.search(r"\bIPT ?[1-6]\b", S) is None)
SL = open(os.path.join(BACKEND, "clp_doc_layers.py"), encoding="utf-8").read()
sqls = re.findall(r'"(SELECT[^"]*)"', SL)
ok("🔴 every SQL literal in the layers carries the tenant predicate", sqls and all("tenant_id = ?" in q for q in sqls),
   str(sqls))

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\ntest_clp_doc_maps: {PASS} passed, {len(FAIL)} failed")
remove_scratch_db(db, TMP)
sys.exit(1 if FAIL else 0)
